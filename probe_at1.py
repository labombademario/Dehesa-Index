import json, urllib.request, urllib.parse, os, re, time, io, subprocess, sys, site
subprocess.run([sys.executable,'-m','pip','install','-q','--user','--break-system-packages','xlrd','pypdf','openpyxl'])
sys.path.append(site.getusersitepackages())
os.makedirs('probe',exist_ok=True); S=open('probe/summary_at1.txt','w')
def log(*a): s=' '.join(str(x) for x in a); print(s,flush=True); S.write(s+'\n'); S.flush()
UA={'User-Agent':'Mozilla/5.0 (compatible; DehesaIndex-research/1.0; +https://dehesaindex.com)','Accept':'*/*'}
def get(u,t=90):
    for i in range(2):
        try:
            with urllib.request.urlopen(urllib.request.Request(u,headers=UA),timeout=t) as r: return r.status,r.read(),r.geturl()
        except urllib.error.HTTPError as e:
            if e.code in (502,503,504,429): time.sleep(15); continue
            return e.code,e.read()[:800],u
        except Exception as e: time.sleep(5); err=repr(e)
    return 0,b'fail',u
def text(b): t=re.sub(rb'<script.*?</script>',b' ',b,flags=re.S); t=re.sub(rb'<[^>]+>',b' ',t); return re.sub(r'\s+',' ',t.decode('utf-8','replace'))
def links(b,base,rx=r'\.(xlsx?|csv|zip|pdf)'):
    out=[]
    for h in re.findall(rb'href="([^"]+)"',b):
        h=h.decode('utf-8','replace')
        if re.search(rx,h,re.I): out.append(urllib.parse.urljoin(base,h))
    return sorted(set(out))
def pdftext(b):
    from pypdf import PdfReader
    return re.sub(r'\s+',' ',' '.join(p.extract_text() or '' for p in PdfReader(io.BytesIO(b)).pages))
def xls_dump(b,nm,maxr=8):
    try:
        if b[:4]==b'PK\x03\x04':
            import openpyxl; wb=openpyxl.load_workbook(io.BytesIO(b),read_only=True,data_only=True)
            for ws in wb.worksheets[:25]:
                rows=list(ws.iter_rows(values_only=True)); log('   SHEET',nm,'|',ws.title,len(rows))
                for rr in rows[:maxr]+rows[-3:]: 
                    v=[str(x)[:22] for x in rr[:14] if x is not None]
                    if v: log('      ',v)
        else:
            import xlrd; wb=xlrd.open_workbook(file_contents=b)
            for sh in wb.sheets()[:25]:
                log('   SHEET',nm,'|',sh.name,sh.nrows,sh.ncols)
                for r in list(range(min(sh.nrows,maxr)))+list(range(max(0,sh.nrows-3),sh.nrows)):
                    v=[str(sh.cell_value(r,c))[:22] for c in range(min(sh.ncols,14)) if str(sh.cell_value(r,c)).strip()]
                    if v: log('      ',v)
    except Exception as e: log('   xls err',repr(e)[:200])

import openpyxl
os.makedirs('probe/at',exist_ok=True)
P='https://www.statistik.at/statistiken/land-und-forstwirtschaft/landwirtschaftliche-bilanzen/versorgungsbilanzen'
st,b,fu=get(P,60); t=text(b); log('PAGE',st,len(b),fu)
for m in list(re.finditer(r'.{0,250}(Creative|CC BY|Lizenz|Nutzungsbedingung|Weiterverwend|Quellenangabe|Urheber|Copyright|\u00a9).{0,250}',t,re.I))[:8]: log('   LIC',m.group(0))
hs=sorted(set(urllib.parse.urljoin(fu,h) for h in re.findall(r'href=["\']([^"\']+)["\']',b.decode('utf-8','replace'))))
for h in hs:
    if re.search(r'\.(xlsx?|ods|csv|pdf)($|\?)',h,re.I) or re.search(r'nutzung|impressum|lizenz|rechtlich|legal|terms|copyright|open-data',h,re.I): log('   L',h)
files=[h for h in hs if re.search(r'\.(xlsx|ods|xls|csv)($|\?)',h,re.I)]
for h in files[:40]:
    st2,b2,_=get(urllib.parse.quote(h,safe=':/%?=&'),120); nm=h.rsplit('/',1)[-1].split('?')[0]; log('  >> FILE',nm,st2,len(b2))
    if st2==200:
        open('probe/at/'+nm,'wb').write(b2)
        if nm.lower().endswith('xlsx'):
            try:
                wb=openpyxl.load_workbook(io.BytesIO(b2),read_only=True,data_only=True)
                for ws in wb.worksheets[:3]:
                    rows=list(ws.iter_rows(values_only=True)); log('   SHEET',ws.title,len(rows))
                    for r in rows[:45]:
                        v=[str(x)[:16] for x in r if x is not None]
                        if v: log('     ',v[:9])
            except Exception as e: log('   err',repr(e)[:150])
for u in ('https://www.statistik.at/fileadmin/announcement/2026/08/20260901Versorgungsbilanzen2025.pdf','https://www.statistik.at/fileadmin/publications/SB_1-27_Versorgungsbilanzen-pflanzliche-Produkte-2024-25.pdf'):
    st,b,_=get(u,120); log('PDF',u,st,len(b))
    if st==200:
        open('probe/at/'+u.rsplit('/',1)[-1],'wb').write(b)
        from pypdf import PdfReader
        R=PdfReader(io.BytesIO(b)); log('  pages',len(R.pages))
        tt=' '.join(p.extract_text() or '' for p in R.pages)
        for m in list(re.finditer(r'.{0,200}(Creative|CC BY|Nachdruck|Quellenangabe|Copyright|\u00a9).{0,200}',tt,re.I))[:4]: log('   PLIC',re.sub(r'\s+',' ',m.group(0)))
        for i,pg in enumerate(R.pages[:4]): log('  ---page',i); log((pg.extract_text() or '')[:1800])
for u in ('https://www.statistik.at/ueber-uns/impressum','https://www.statistik.at/impressum','https://www.statistik.at/en/about-us/imprint','https://www.statistik.at/ueber-uns/rechtliche-hinweise','https://www.statistik.at/services/tools/services/rechtliches'):
    st,b,fu=get(u,40); t=text(b); log('TERMS',u,st,fu)
    for m in list(re.finditer(r'.{0,300}(Creative|CC BY|Weiterverwend|Nutzung der|Quellenangabe|kommerziell).{0,300}',t,re.I))[:5]: log('   T',m.group(0))
S.close()
