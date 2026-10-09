import urllib.request, re, io, os
LOG=[]
def log(*a):
    LOG.append(' '.join(str(x) for x in a)); open('probe/probe-kc.txt','w').write('\n'.join(LOG)+'\n')
H={'User-Agent':'Mozilla/5.0 (dehesaindex probe)'}
def get(u):
    r=urllib.request.urlopen(urllib.request.Request(u,headers=H),timeout=90); return r.read()
t=get('https://www.kansascityfed.org/agriculture/ag-credit-survey/').decode('utf8','replace')
links=sorted(set(re.findall(r'href="(/documents/\d+/[^"]+\.xlsx)"',t)))
log('links',links)
import openpyxl
for l in links:
    try:
        b=get('https://www.kansascityfed.org'+l); open('probe/'+os.path.basename(l),'wb').write(b)
        wb=openpyxl.load_workbook(io.BytesIO(b),data_only=True)
        log('==',l,len(b),wb.sheetnames)
        for ws in wb.worksheets[:3]:
            log('  sheet',ws.title,ws.max_row,ws.max_column)
            for row in ws.iter_rows(min_row=1,max_row=12,values_only=True): log('   ',[c for c in row[:10]])
    except Exception as e: log('ERR',l,repr(e)[:150])
# sitemap of historical files
hist=sorted(set(re.findall(r'href="([^"]+)"[^>]*>[^<]*(?:Historical|historical|Data)[^<]*<',t)))
log('hist',hist[:20])
for m in re.finditer(r'.{200}(?:©|Copyright|copyright).{200}',re.sub(r'<[^>]+>',' ',t)): log('COPY',re.sub(r'\s+',' ',m.group(0))); break
