import os, sys, io, re, urllib.request, html
sys.path.insert(0, 'scripts')
from lib_ods import read_ods
import openpyxl
OUT = []
def p(*a):
    s = ' '.join(str(x) for x in a); OUT.append(s); print(s, flush=True)
UA = {'User-Agent': 'Mozilla/5.0 (X11; Linux x86_64) DehesaIndex/1.0 (+https://dehesaindex.com)'}
def get(u):
    with urllib.request.urlopen(urllib.request.Request(u, headers=UA), timeout=120) as r: return r.read()
wb = openpyxl.load_workbook(io.BytesIO(get('https://www.daera-ni.gov.uk/sites/default/files/2026-10/Vol%2089%20Quarter%204%20Reports.xlsx')), read_only=True, data_only=True)
p('SHEETS', wb.sheetnames)
ws = wb[wb.sheetnames[-1]]
for i, r in enumerate(ws.iter_rows(values_only=True)):
    if i < 38: continue
    if any(c not in (None, '') for c in r): p(i, [str(c)[:26] if c is not None else '' for c in r][:17])
t = html.unescape(get('https://www.gov.scot/publications/results-from-the-scottish-agricultural-census-june-2026/documents/').decode('utf8', 'ignore'))
u = [x for x in re.findall(r'href="([^"]+)"', t) if 'main-tables/' in x and 'ods' not in x.split('main-tables')[1][:6]][0]
u = u if u.startswith('http') else 'https://www.gov.scot' + u
wb = openpyxl.load_workbook(io.BytesIO(get(u)), read_only=True, data_only=True)
for n in ['Table_3', 'Table_8', 'Table_13']:
    if n not in wb.sheetnames: continue
    p('--- ', n)
    for i, r in enumerate(wb[n].iter_rows(values_only=True)):
        if i > 40: break
        p(i, [str(c)[:30] if c is not None else '' for c in r][:5])
S = read_ods(get('https://assets.publishing.service.gov.uk/media/6aad019ece3f006bd4346c7d/ITL3_Regional_Accounts_Current_Prices-20260924.ods'), max_cols=30)
rows = S['North_East']
for i, r in enumerate(rows[:30]): p(i, [str(c)[:30] if c is not None else '' for c in r][:3], len(r), r[-2:] if r else '')
os.makedirs('probe/out', exist_ok=True); open('probe/out/uk7.txt', 'w').write('\n'.join(OUT))
