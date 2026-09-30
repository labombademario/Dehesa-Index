#!/usr/bin/env python3
"""Gas natural Europa (TTF) — Banco Mundial, Pink Sheet mensual, USD/MMBtu."""
import json, re, urllib.request
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SNAP_DIR = ROOT / "data" / "snapshots"
LOG = ROOT / "data" / "energy-markets-log-eu.txt"
LANDING = "https://www.worldbank.org/en/research/commodity-markets"

def fetch(url):
    req = urllib.request.Request(url, headers={"User-Agent": "Dehesa-Index-data-bot/1.0"})
    with urllib.request.urlopen(req, timeout=60) as r:
        return r.read()

html = fetch(LANDING).decode("utf-8", "ignore")
matches = re.findall(r'https://thedocs\.worldbank\.org/[^"\']+CMO-Historical-Data-Monthly\.xlsx', html)
if not matches:
    raise RuntimeError("World Bank monthly XLSX link not found")
xlsx_url = matches[0].replace("&amp;", "&")
pub = re.search(r'\((\d{2}/\d{2}/\d{4})\)', re.sub(r'<[^>]+>', ' ', html))
publication_date = datetime.strptime(pub.group(1), "%m/%d/%Y").date().isoformat() if pub else None

tmp = ROOT / ".worldbank-monthly-gas.xlsx"
tmp.write_bytes(fetch(xlsx_url))
wb = load_workbook(tmp, read_only=True, data_only=True)
rows = list(wb["Monthly Prices"].iter_rows(values_only=True))
header_row = col = None
for idx, row in enumerate(rows[:100]):
    for c, val in enumerate(row):
        n = re.sub(r"\s+", " ", str(val or "").strip().lower())
        if n == "natural gas, europe" or n == "natural gas europe":
            header_row, col = idx, c
            break
    if col is not None:
        break
if col is None:
    heads = [str(v) for v in rows[4] if v] if len(rows) > 4 else []
    LOG.write_text("Natural gas, Europe column not found. Header sample: " + " | ".join(heads[:40]) + "\n")
    raise RuntimeError("World Bank 'Natural gas, Europe' series not found")

points = []
for row in rows[header_row + 1:]:
    if not row:
        continue
    period = str(row[0] or "").strip()
    if not re.fullmatch(r"\d{4}M\d{2}", period, re.I):
        continue
    v = row[col]
    if isinstance(v, (int, float)):
        points.append((period.upper(), float(v)))
if len(points) < 12:
    raise RuntimeError("Insufficient natural gas history")
latest_period, latest_value = points[-1]
year, month = int(latest_period[:4]), int(latest_period[5:])
obs_date = f"{year:04d}-{month:02d}-01"
change = round((latest_value / points[-2][1] - 1) * 100, 4) if points[-2][1] else None
now = datetime.utcnow().isoformat(timespec="seconds") + "Z"
observation = {
    "id": "di_gas_natural_eu", "product": "gas_natural", "region": "eu", "sourceId": "world_bank",
    "observationDate": obs_date, "publicationDate": publication_date or obs_date,
    "value": latest_value, "currency": "USD", "unit": "mmbtu", "frequency": "monthly",
    "status": "verified", "verifiedAt": now, "comparability": "directional", "changePct": change,
    "history": [{"period": int(p[5:]), "year": int(p[:4]), "value": v} for p, v in points[-240:]]
}
SNAP_DIR.mkdir(parents=True, exist_ok=True)
snap = SNAP_DIR / f"{datetime.utcnow().date().isoformat()}.json"
if snap.exists():
    old = json.loads(snap.read_text(encoding="utf-8"))
    old["observations"] = [o for o in old.get("observations", []) if not (o.get("product") == "gas_natural" and o.get("region") == "eu")] + [observation]
else:
    old = {"schemaVersion": "1.0", "generatedAt": now, "observations": [observation]}
snap.write_text(json.dumps(old, indent=2) + "\n", encoding="utf-8")
LOG.write_text(f"{now}\nOK Natural gas, Europe: {latest_period} = {latest_value} USD/MMBtu ({len(points)} pts)\n")
tmp.unlink(missing_ok=True)
print("World Bank gas EU:", latest_period, latest_value)
