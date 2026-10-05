#!/usr/bin/env python3
"""Tasas de ARC/PLC de la FSA (USDA Farm Service Agency) -> data/us-arcplc/<ESTADO>.json, data/us-arcplc/national.json, data/us-arcplc/index.json y log.txt
  ARC-CO (Agriculture Risk Coverage, County Option): por condado y cultivo, para cada campana publicada, el rendimiento y el precio de referencia
    (benchmark), el ingreso de referencia y el garantizado, el rendimiento real del condado, el precio nacional, el ingreso real y la TASA DE PAGO
    ARC-CO en USD por acre base. Son las tasas que publica la FSA, no los dolares pagados.
  PLC (Price Loss Coverage): la tasa de pago es nacional, por unidad de producto (USD por bushel o por libra), con el precio de referencia efectivo,
    el precio medio de campana (MYA, proyectado P o final F), la tasa de prestamo y la tasa maxima.
Ficheros: pagina oficial https://www.fsa.usda.gov/resources/programs/arc-plc/program-data (los enlaces se descubren cada vez; no se inventa ninguna URL).
Lo que la FSA deja en blanco (campana aun abierta) queda como hueco, nunca como cero. Si un fichero falla, se conserva lo anterior de esa campana.
Atribucion: "U. S. Department of Agriculture, Farm Service Agency."
Uso: update-us-arcplc.py [--outdir DIR] [--years N] [--from-dir DIR]
  --from-dir: lee la pagina (program-data.html) y los xlsx de una carpeta en vez de la red (mismo procesado; para una carga inicial si la FSA no responde al servidor)."""
import datetime, io, json, re, sys, time, urllib.error, urllib.parse, urllib.request
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
PAGE = "https://www.fsa.usda.gov/resources/programs/arc-plc/program-data"
PAGE_ALT = "https://www.fsa.usda.gov/programs-and-services/arcplc_program/arcplc-program-data/index"
BASE = "https://www.fsa.usda.gov/"
LOG = []
FROM = None  # carpeta local con arcco_AAAA.xlsx y plc_AAAA.xlsx (opcion --from-dir)
NAMES = {"Alabama": "AL", "Alaska": "AK", "Arizona": "AZ", "Arkansas": "AR", "California": "CA", "Colorado": "CO", "Connecticut": "CT", "Delaware": "DE", "Florida": "FL", "Georgia": "GA",
         "Hawaii": "HI", "Idaho": "ID", "Illinois": "IL", "Indiana": "IN", "Iowa": "IA", "Kansas": "KS", "Kentucky": "KY", "Louisiana": "LA", "Maine": "ME", "Maryland": "MD", "Massachusetts": "MA",
         "Michigan": "MI", "Minnesota": "MN", "Mississippi": "MS", "Missouri": "MO", "Montana": "MT", "Nebraska": "NE", "Nevada": "NV", "New Hampshire": "NH", "New Jersey": "NJ", "New Mexico": "NM",
         "New York": "NY", "North Carolina": "NC", "North Dakota": "ND", "Ohio": "OH", "Oklahoma": "OK", "Oregon": "OR", "Pennsylvania": "PA", "Rhode Island": "RI", "South Carolina": "SC",
         "South Dakota": "SD", "Tennessee": "TN", "Texas": "TX", "Utah": "UT", "Vermont": "VT", "Virginia": "VA", "Washington": "WA", "West Virginia": "WV", "Wisconsin": "WI", "Wyoming": "WY"}
def log(*a): s = " ".join(str(x) for x in a); LOG.append(s); print(s, flush=True)
def now(): return datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
def fetch(u, tries=3, deadline=240):
    """Descarga con limite total de tiempo por intento (un servidor que manda el contenido gota a gota no cuelga el proceso)."""
    if FROM and not u.startswith("http"): return Path(u).read_bytes()
    last = None
    for i in range(tries):
        t0 = time.time()
        try:
            with urllib.request.urlopen(urllib.request.Request(u, headers={"User-Agent": "Mozilla/5.0 (X11; Linux x86_64) dehesaindex.com-pipeline (+https://dehesaindex.com)", "Accept": "*/*"}), timeout=60) as r:
                buf = bytearray()
                while True:
                    ch = r.read(1 << 16)
                    if not ch: return bytes(buf)
                    buf += ch
                    if time.time() - t0 > deadline: raise TimeoutError("mas de %d s descargando" % deadline)
        except urllib.error.HTTPError as e:
            last = "HTTP %s" % e.code
            if e.code in (403, 404): break
        except Exception as e: last = repr(e)[:200]
        time.sleep(10 * (i + 1))
    raise RuntimeError("%s: %s" % (u, last))
def xlsx_of(href):
    """Un enlace de la pagina puede ser el xlsx o una pagina /documents/... que enlaza al xlsx."""
    if FROM and Path(href).exists(): return href
    u = urllib.parse.urljoin(BASE, href)
    if re.search(r"\.xlsx?$", u, re.I): return u
    t = fetch(u).decode("utf-8", "replace"); m = re.findall(r'href="([^"]+\.xlsx)"', t)
    if not m: raise RuntimeError("sin xlsx en " + u)
    return urllib.parse.urljoin(BASE, m[0])
def num(v):
    if v is None or isinstance(v, bool): return None
    if isinstance(v, (int, float)): return round(float(v), 4)
    s = str(v).strip().replace(",", "").replace("$", "")
    try: return round(float(s), 4)
    except Exception: return None
def rows_of(b):
    import openpyxl
    try: wb = openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True); ws = wb.worksheets[0]; return [list(r) for r in ws.iter_rows(values_only=True)]
    except TypeError:  # algunas celdas con fechas mal tipadas rompen el modo de solo lectura
        wb = openpyxl.load_workbook(io.BytesIO(b), data_only=True); ws = wb.worksheets[0]; return [list(r) for r in ws.iter_rows(values_only=True)]
def links():
    if FROM:
        arc, plc = {}, {}
        for f in sorted(Path(FROM).glob("*.xlsx")):
            m = re.match(r"(arcco|plc)_(20\d\d)\.xlsx$", f.name)
            if m: (arc if m.group(1) == "arcco" else plc)[int(m.group(2))] = str(f)
        return arc, plc
    try: t = fetch(PAGE).decode("utf-8", "replace")
    except Exception as e: log("pagina principal", e, "-> direccion alternativa"); t = fetch(PAGE_ALT).decode("utf-8", "replace")
    hs = sorted(set(re.findall(r'href="([^"]+)"', t)))
    arc, plc = {}, {}
    for h in hs:
        hl = urllib.parse.unquote(h).lower()
        m = re.search(r"arcco[-_ ](20\d\d)[-_ ]data", hl)
        if m and "source" not in hl and not hl.endswith(".pdf"): arc.setdefault(int(m.group(1)), h)
        m = re.search(r"/(20\d\d)[-_ ]plc(?:[-_ ]excel|\.?xlsx)", hl)
        if m: plc.setdefault(int(m.group(1)), h)
    return arc, plc
COLS = [("bmYield", r"bench ?mark \(\d{4}-\d{2} olympic"), ("bmPrice", r"bench ?mark price"), ("bmRev", r"benchmark revenue"), ("guarRev", r"guarantee revenue"),
        ("maxRate", r"maximum payment rate"), ("actYield", r"actual yield"), ("natPrice", r"national price"), ("actRev", r"actual revenue"), ("formula", r"formula payment rate"), ("rate", r"arc-co payment rate")]
def arcco(year, href, out):
    u = xlsx_of(href); rows = rows_of(fetch(u))
    hi = next(i for i, r in enumerate(rows[:12]) if r and str(r[0] or "").strip().lower() in ("st_cty", "st_cty "))
    head = [str(x or "").strip().lower() for x in rows[hi]]; col = {}
    for k, rx in COLS:
        for j, h in enumerate(head):
            if re.search(rx, h) and k not in col and not (k == "bmYield" and "price" in h): col[k] = j
    if "rate" not in col or "bmRev" not in col: raise RuntimeError("cabecera inesperada %s: %s" % (year, head))
    n = 0
    for r in rows[hi + 1:]:
        if not r or not r[0]: continue
        fips = str(r[0]).strip().zfill(5); st = NAMES.get(str(r[1] or "").strip())
        if not st or not re.match(r"^\d{5}$", fips): continue
        crop = str(r[4] or "").strip(); unit = str(r[5] or "").strip(); sub = str(r[3] or "").strip(); des = str(r[6] or "").strip() or "All"
        if not crop: continue
        c = out.setdefault(st, {}).setdefault(fips, {"name": str(r[2] or "").strip(), "crops": {}})
        key = crop + ("|" + des if des != "All" else "") + ("|" + sub if sub else "")
        e = c["crops"].setdefault(key, {"crop": crop, "unit": unit, "type": des, "sub": sub or None, "years": {}})
        e["years"][str(year)] = {k: num(r[j]) if j < len(r) else None for k, j in col.items()}; n += 1
    log("ARC-CO", year, n, "filas", u.rsplit("/", 1)[-1])
    return u if u.startswith("http") else Path(u).name
def plc(year, href):
    u = xlsx_of(href); rows = rows_of(fetch(u))
    hi = next(i for i, r in enumerate(rows[:12]) if r and str(r[0] or "").strip().lower() == "commodity")
    out = []
    for r in rows[hi + 1:]:
        name = str(r[0] or "").strip()
        if not name or name.startswith(("1/", "2/", "3/", "4/", "5/", "Note", "Source")) or len(r) < 13: continue
        ref, mya, loan, eff, rate, mx = num(r[4]), num(r[5]), num(r[7]), num(r[8]), num(r[10]), num(r[12])
        if ref is None and rate is None: continue
        out.append({"crop": re.sub(r"\s+\d/$", "", name), "unit": str(r[3] or "").strip(), "my": str(r[1] or "").strip(), "refPrice": ref, "mya": mya, "myaFlag": (str(r[6] or "").strip() or None),
                    "loanRate": loan, "effPrice": eff, "rate": rate, "rateFlag": (str(r[11] or "").strip() or None), "maxRate": mx})
    if len(out) < 10: raise RuntimeError("PLC %s: solo %d cultivos" % (year, len(out)))
    log("PLC", year, len(out), "cultivos", u.rsplit("/", 1)[-1]); return out, (u if u.startswith("http") else Path(u).name)
def now_s(): return __import__('datetime').datetime.now(__import__('datetime').timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ')
def main():
    a = sys.argv[1:]; outdir = Path(a[a.index("--outdir") + 1]) if "--outdir" in a else ROOT / "data" / "us-arcplc"
    years = int(a[a.index("--years") + 1]) if "--years" in a else 5
    global FROM
    if "--from-dir" in a: FROM = a[a.index("--from-dir") + 1]
    outdir.mkdir(parents=True, exist_ok=True)
    try: arc_l, plc_l = links()
    except Exception as e:
        log("la web de la FSA no responde (%s): no se toca ningun dato; se reintenta en la proxima ejecucion" % str(e)[:160]); (outdir / "log.txt").write_text(now_s() + "\n" + "\n".join(LOG) + "\n"); return 0
    log("enlaces ARC-CO", sorted(arc_l), "PLC", sorted(plc_l))
    last = max(list(arc_l) + list(plc_l) + [0])
    if not last: log("la pagina de la FSA no tiene enlaces reconocibles"); (outdir / "log.txt").write_text("\n".join(LOG) + "\n"); return 1
    want = [y for y in range(last - years + 1, last + 1)]
    out, files = {}, {}
    for y in want:
        if y in arc_l:
            try: files["arcco_%d" % y] = arcco(y, arc_l[y], out)
            except Exception as e: log("ARC-CO", y, "FALLO", e)
    oldus = None
    try: oldus = json.loads((outdir / "national.json").read_text())
    except Exception: pass
    plcs = dict((oldus or {}).get("plc") or {}); pfiles = dict((oldus or {}).get("files") or {})
    for y in want:
        if y in plc_l:
            try: plcs[str(y)], files["plc_%d" % y] = plc(y, plc_l[y])
            except Exception as e: log("PLC", y, "FALLO", e)
    plcs = {k: v for k, v in plcs.items() if int(k) >= want[0]}
    idx = {}
    for st in sorted(NAMES.values()):
        p = outdir / (st + ".json"); old = None
        try: old = json.loads(p.read_text())
        except Exception: pass
        cur = out.get(st, {})
        if old:  # una campana que no se pudo leer conserva sus datos anteriores
            for fips, c in (old.get("counties") or {}).items():
                for key, e in c["crops"].items():
                    for y, v in e["years"].items():
                        if int(y) < want[0]: continue
                        ne = cur.setdefault(fips, {"name": c["name"], "crops": {}})["crops"].setdefault(key, {k: e[k] for k in ("crop", "unit", "type", "sub")} | {"years": {}})
                        ne["years"].setdefault(y, v)
        if not cur: continue
        ys = sorted({y for c in cur.values() for e in c["crops"].values() for y in e["years"]})
        doc = {"schemaVersion": 1, "generatedAt": now(), "state": st, "sourceId": "usda_fsa", "unit": "USD per base acre", "years": ys, "counties": cur}
        p.write_text(json.dumps(doc, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
        idx[st] = {"counties": len(cur), "years": ys, "bytes": p.stat().st_size}
    if len(idx) < (int(a[a.index("--min-states") + 1]) if "--min-states" in a else 40): log("demasiados pocos estados", len(idx)); (outdir / "log.txt").write_text("\n".join(LOG) + "\n"); return 1
    pfiles.update(files)
    (outdir / "national.json").write_text(json.dumps({"schemaVersion": 1, "generatedAt": now(), "sourceId": "usda_fsa", "plc": plcs, "files": pfiles}, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    (outdir / "index.json").write_text(json.dumps({"schemaVersion": 1, "generatedAt": now(), "page": PAGE, "states": idx}, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    (outdir / "log.txt").write_text(now() + "\n" + "\n".join(LOG[-300:]) + "\n", encoding="utf-8"); log("estados", len(idx)); return 0
if __name__ == "__main__": sys.exit(main())
