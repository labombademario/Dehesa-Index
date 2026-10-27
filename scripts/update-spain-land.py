#!/usr/bin/env python3
"""España: precio medio anual de las tierras de uso agrario y cánones de arrendamiento rústico por comunidad autónoma, tipo de tierra y provincia
(MAPA, Encuesta de Precios de la Tierra y Encuesta de Cánones de Arrendamiento Rústico, libros .xlsx publicados en la web del Ministerio) -> data/spain-land/land.json.

Fuentes (se buscan en la página del MAPA, el año más reciente publicado):
  * https://www.mapa.gob.es/es/estadistica/temas/estadisticas-agrarias/economia/tierra-precios  («Resultados definitivos/provisionales AAAA»)
      Cuadro 1 = precio medio general; Cuadros 2-19 = por tipo de tierra (ARZ arrozal, CBP bajo plástico, CIT cítricos, FCR/FCS frutal carnoso regadío/secano, FSR/FSS frutos secos,
      HUE huerta, OLR/OLS olivar, PLT platanera, PRD prados, PST pastizal, SBT subtropicales, TAR/TAS tierras arables, VIR/VIS viñedo) por CCAA y año (€/ha, 5 años);
      Cuadro 20 = precios provinciales del último año por tipo.
  * https://www.mapa.gob.es/es/estadistica/temas/estadisticas-agrarias/economia/canones-anuales-arrendamientos-rusticos  («Tablas Cánones de arrendamiento AAAA Base 2020»)
      Cuadro 1 = canon general (€/ha y año); Cuadros 2-7 = TAS, TAR, OLS, VIS, PRD, PST por CCAA y año; Cuadro 8 = otros cultivos (solo el último año).
Licencia: aviso legal del MAPA, reutilizable citando la fuente (registro de licencias, fuente mapa_es). Valores tal cual; «..» (dato no significativo) y celdas vacías son huecos, nunca cero.
Cada libro cubre cinco años: al llegar uno nuevo el fichero ACUMULA historia (los años que ya no vienen en el libro se conservan; si el MAPA revisa un año gana el libro nuevo y queda en el log).
Validaciones (si falla alguna no se escribe nada): 17 comunidades + Total Nacional, 50 provincias, precios > 0, el total nacional del Cuadro 1 coincide con el del Cuadro 0, el canon general también, y
el total nacional es la media de las comunidades ponderada por las superficies base 2020 que trae el propio libro (tolerancia 0,5 %).
Uso: update-spain-land.py [--out DIR] [--price-file X.xlsx] [--rent-file Y.xlsx]   (con los ficheros no se toca la red; así lo usan los tests)"""
import datetime, io, json, os, re, ssl, sys, time, unicodedata, urllib.request
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
SITE = "https://www.mapa.gob.es"
PAGE_PRICE = SITE + "/es/estadistica/temas/estadisticas-agrarias/economia/tierra-precios"
PAGE_RENT = SITE + "/es/estadistica/temas/estadisticas-agrarias/economia/canones-anuales-arrendamientos-rusticos"
UA = {"User-Agent": "Mozilla/5.0 (compatible; DehesaIndex/1.0; +https://dehesaindex.com)"}
CCAA = {"galicia": "GA", "principado de asturias": "AS", "cantabria": "CB", "pais vasco": "PV", "comunidad foral de navarra": "NC", "la rioja": "RI", "aragon": "AR", "cataluna": "CT",
        "illes balears": "IB", "castilla y leon": "CL", "comunidad de madrid": "MD", "castilla-la mancha": "CM", "comunitat valenciana": "VC", "region de murcia": "MC", "extremadura": "EX",
        "andalucia": "AN", "canarias": "CN"}
TYPES = {"PT": "General", "ARZ": "Arrozal", "CBP": "Cultivos bajo plástico", "CIT": "Cítricos", "FCR": "Frutal carnoso regadío", "FCS": "Frutal carnoso secano", "FSR": "Frutos secos regadío",
         "FSS": "Frutos secos secano", "HUE": "Huerta", "OLR": "Olivar regadío", "OLS": "Olivar secano", "PLT": "Platanera", "PRD": "Prados", "PST": "Pastizal", "SBT": "Cultivos subtropicales y tropicales",
         "TAR": "Tierras arables regadío", "TAS": "Tierras arables secano", "VIR": "Viñedo regadío", "VIS": "Viñedo secano"}
RENT_TYPES = ["GEN", "TAS", "TAR", "OLS", "VIS", "PRD", "PST"]
RENT_SHEETS = {"1": "GEN", "2": "TAS", "3": "TAR", "4": "OLS", "5": "VIS", "6": "PRD", "7": "PST"}
OTHER_CROPS = ["ARZ", "CBP", "CIT", "FCR", "FCS", "FSR", "FSS", "HUE", "OLR", "PLT", "SBT", "VIR"]
LOG = []
def log(*a):
    s = " ".join(str(x) for x in a); LOG.append(s); print(s, flush=True)
def norm(s):
    s = unicodedata.normalize("NFD", str(s).lower()); s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", s).strip()
def http(url, tries=3, timeout=180):
    last = None
    for i in range(tries):
        try:
            with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=timeout) as r: return r.read()
        except Exception as e: last = e; time.sleep(5 * (i + 1))
    raise RuntimeError("%s -> %s" % (url, str(last)[:160]))
def links(html):
    out = []
    for h, txt in re.findall(r'<a[^>]+href="([^"#]+\.xlsx)[^"]*"[^>]*>(.*?)</a>', html, re.S | re.I):
        t = re.sub(r"<[^>]+>|\s+", " ", txt).replace("&nbsp;", " ").strip()
        out.append((h if h.startswith("http") else SITE + h, t))
    return out
def pick_price(html):
    """Libro de precios de la tierra mas reciente: texto «... Resultados [definitivos|provisionales] AAAA» (no los de rango AAAA-AAAA)."""
    best = None
    for u, t in links(html):
        m = re.search(r"Resultados\s+(definitivos|provisionales)?\s*(\d{4})\s*$", t, re.I)
        if m:
            k = (int(m.group(2)), 1 if (m.group(1) or "").lower().startswith("def") else 0)
            if best is None or k > best[0]: best = (k, u, t)
    if not best: raise RuntimeError("no se encontro el libro de precios de la tierra en la pagina del MAPA")
    return best[1], best[0][0], best[0][1] == 1, best[2]
def pick_rent(html):
    """Libro de canones mas reciente: «Tablas Canones de arrendamiento AAAA Base 2020». Si cambia la base se para: los pesos y la definicion ya no serian los mismos."""
    best = None
    for u, t in links(html):
        m = re.search(r"Tablas\s+C[aá]nones de arrendamiento\s+(\d{4})\s+Base\s+(\d{4})\s*$", t, re.I)
        if m and (best is None or int(m.group(1)) > best[0]): best = (int(m.group(1)), int(m.group(2)), u, t)
    if not best: raise RuntimeError("no se encontro el libro de canones en la pagina del MAPA")
    if best[1] != 2020: raise RuntimeError("el libro de canones %d esta en base %d (se esperaba 2020): revisar el parser" % (best[0], best[1]))
    return best[2], best[0], best[3]
def num(v):
    """Numero de una celda del MAPA: vacio o «..» (no significativo) = None; cualquier otro texto es un error."""
    if v is None or (isinstance(v, str) and v.strip() in ("", "..", "-", "--")): return None
    if isinstance(v, bool) or not isinstance(v, (int, float)): raise ValueError("celda no numerica: %r" % (v,))
    return float(v)
def header(ws, group_start):
    """Fila de cabecera (la de la celda «CCAA»): devuelve (fila, columna de CCAA, [(columna, año)]) con los años del bloque cuyo titulo empieza por group_start."""
    for row in ws.iter_rows(min_row=1, max_row=12):
        g = next((c for c in row if isinstance(c.value, str) and norm(c.value).startswith(norm(group_start))), None)
        if g is None: continue
        cc = max((c for c in row if isinstance(c.value, str) and c.value.strip() == "CCAA" and c.column < g.column), key=lambda c: c.column, default=None)   # la CCAA de la izquierda del bloque, no la de la tabla de indices
        if cc is None: continue
        yrs = []; col = g.column
        while True:
            v = ws.cell(cc.row + 1, col).value
            if isinstance(v, int) and 1990 < v < 2100: yrs.append((col, v)); col += 1
            else: break
        if yrs: return cc.row, cc.column, yrs
    raise ValueError("hoja %r: no se encontro la cabecera (CCAA / %s)" % (ws.title, group_start))
def table(ws, group_start, extra_weight=False):
    """{codigo CCAA o 'ES': {año: valor}} de una hoja de precios/canones por CCAA; filas sin ningun dato se omiten. Con extra_weight tambien {codigo: superficie base}."""
    hr, cc, yrs = header(ws, group_start); out = {}; wts = {}
    for r in range(hr + 2, hr + 2 + 30):
        nm = ws.cell(r, cc).value
        if nm is None: continue
        k = norm(nm)
        code = "ES" if k == "total nacional" else CCAA.get(k)
        if code is None: raise ValueError("hoja %r fila %d: comunidad desconocida %r" % (ws.title, r, nm))
        vals = {y: num(ws.cell(r, c).value) for c, y in yrs}
        if any(v is not None for v in vals.values()): out[code] = {y: v for y, v in vals.items() if v is not None}
        if extra_weight:
            w = ws.cell(r, cc + 1).value; wts[code] = float(w) if isinstance(w, (int, float)) else None
        if code == "ES": break
    else: raise ValueError("hoja %r: no aparece la fila Total Nacional" % ws.title)
    if len(out) < 2 or "ES" not in out: raise ValueError("hoja %r: sin total nacional o sin comunidades con datos" % ws.title)
    return (out, wts) if extra_weight else out
def sheets_by_key(wb, pat):
    d = {}
    for ws in wb.worksheets:
        m = re.match(pat, ws.title.strip())
        if m: d[m.group(1)] = ws
    return d
def parse_price(wb):
    sh = sheets_by_key(wb, r"^Cuadro_\d+_([A-Z]{2,3})$"); price = {}; weights = {}
    if "PT" not in sh: raise ValueError("falta el Cuadro 1 (PT, precio general)")
    for t in TYPES:
        if t not in sh: raise ValueError("falta la hoja del tipo %s" % t)
        tb, wt = table(sh[t], "Precios medios", True)
        for code, ys in tb.items(): price.setdefault(code, {})[t] = ys
        weights[t] = wt
    g0 = next((w for w in wb.worksheets if w.title.strip().startswith("Cuadro_0")), None)
    nat0 = None
    if g0 is not None:
        for row in g0.iter_rows(min_row=1, max_row=20, max_col=4):
            if isinstance(row[0].value, int) and isinstance(row[1].value, (int, float)):
                nat0 = nat0 or {}; nat0[row[0].value] = float(row[1].value)
    prov = None
    ws = next((w for w in wb.worksheets if w.title.strip().lower().startswith("cuadro 20")), None)
    if ws is None: raise ValueError("falta el Cuadro 20 (precios provinciales)")
    ttl = " ".join(str(c.value) for row in ws.iter_rows(min_row=1, max_row=3) for c in row if c.value is not None)
    ym = re.search(r"(20\d\d)", ttl)
    if not ym: raise ValueError("Cuadro 20: no se encontro el año en el titulo %r" % ttl)
    hdr = None
    for row in ws.iter_rows(min_row=1, max_row=8):
        if any(isinstance(c.value, str) and c.value.strip() == "Provincia" for c in row): hdr = row[0].row; break
    if hdr is None: raise ValueError("Cuadro 20: sin cabecera Provincia")
    cols = {}
    for c in ws[hdr]:
        if isinstance(c.value, str) and c.value.strip() in TYPES and c.value.strip() != "PT": cols[c.column] = c.value.strip()
    if len(cols) != len(TYPES) - 1: raise ValueError("Cuadro 20: columnas de tipos %d, se esperaban %d" % (len(cols), len(TYPES) - 1))
    values = {}; names = {}
    # desde 2025 el libro trae CCAA | Código | Provincia | tipos...; antes Código | Provincia | tipos...: las columnas se buscan por su cabecera
    ccol = next((c.column for c in ws[hdr] if isinstance(c.value, str) and norm(c.value) == "codigo"), 1)
    ncol = next((c.column for c in ws[hdr] if isinstance(c.value, str) and norm(c.value) == "provincia"), 2)
    for r in range(hdr + 1, hdr + 1 + 60):
        code = ws.cell(r, ccol).value
        if not isinstance(code, int): continue
        pid = "%02d" % code; names[pid] = str(ws.cell(r, ncol).value).strip()
        row = {}
        for c, t in cols.items():
            v = num(ws.cell(r, c).value)
            if v is not None: row[t] = v
        values[pid] = row
    if len(values) != 50 or sorted(values) != ["%02d" % i for i in range(1, 51)]: raise ValueError("Cuadro 20: %d provincias (codigos %s...), se esperaban las 50" % (len(values), sorted(values)[:3]))
    return price, weights, nat0, {"year": int(ym.group(1)), "values": values, "names": names}
def parse_rent(wb):
    rent = {}; sh = sheets_by_key(wb, r"^Cuadro[ _](\d+)(?:_.*)?$")
    for n, t in RENT_SHEETS.items():
        if n not in sh: raise ValueError("falta la hoja de canones %s (%s)" % (n, t))
        for code, ys in table(sh[n], "Canon").items(): rent.setdefault(code, {})[t] = ys
    c0 = sh.get("0"); nat0 = None
    if c0 is not None:
        for row in c0.iter_rows(min_row=1, max_row=20, min_col=2, max_col=4):
            if isinstance(row[0].value, int) and isinstance(row[1].value, (int, float)): nat0 = nat0 or {}; nat0[row[0].value] = float(row[1].value)
    other = None
    if "8" in sh:
        ws = sh["8"]; hr = None
        for row in ws.iter_rows(min_row=1, max_row=8):
            if any(isinstance(c.value, str) and c.value.strip() == "CCAA" for c in row): hr = row[0].row; break
        ym = re.search(r"(20\d\d)", " ".join(str(c.value) for row in ws.iter_rows(min_row=1, max_row=3) for c in row if c.value is not None))
        if hr is not None and ym:
            cols = {c.column: c.value.strip() for c in ws[hr + 1] if isinstance(c.value, str) and c.value.strip() in OTHER_CROPS}
            vals = {}
            for r in range(hr + 2, hr + 2 + 30):
                nm = ws.cell(r, 1).value
                if nm is None: continue
                code = "ES" if norm(nm) == "total nacional" else CCAA.get(norm(nm))
                if code is None: raise ValueError("Cuadro 8 fila %d: comunidad desconocida %r" % (r, nm))
                row = {t: num(ws.cell(r, c).value) for c, t in cols.items()}; row = {t: v for t, v in row.items() if v is not None}
                if row: vals[code] = row
                if code == "ES": break
            other = {"year": int(ym.group(1)), "values": vals}
    return rent, nat0, other
def check_price(price, weights, nat0, prov):
    errs = []
    if set(price) - {"ES"} != set(CCAA.values()): errs.append("comunidades distintas de las 17: %s" % sorted(set(CCAA.values()) ^ (set(price) - {"ES"})))
    for code, ts in price.items():
        for t, ys in ts.items():
            for y, v in ys.items():
                if not (v > 0): errs.append("%s %s %s: valor no positivo" % (code, t, y))
    if nat0:
        for y, v in nat0.items():
            g = price.get("ES", {}).get("PT", {}).get(y)
            if g is not None and abs(g - v) > 0.01 * max(1, v) / 100: errs.append("Total Nacional %s: Cuadro 1 (%s) != Cuadro 0 (%s)" % (y, g, v))
    # el total nacional del precio general es la media de las comunidades ponderada por las superficies base 2020 del propio libro
    w = weights.get("PT") or {}; yrs = sorted(price["ES"]["PT"])
    for y in yrs:
        num_, den = 0.0, 0.0
        for code in CCAA.values():
            p = price.get(code, {}).get("PT", {}).get(y)
            if p is not None and w.get(code): num_ += p * w[code]; den += w[code]
        if den and abs(num_ / den - price["ES"]["PT"][y]) > 0.005 * price["ES"]["PT"][y]: errs.append("Total Nacional %s: %.1f no coincide con la media ponderada de las comunidades (%.1f)" % (y, price["ES"]["PT"][y], num_ / den))
    pv = prov["values"]
    for pid, row in pv.items():
        for t, v in row.items():
            if not (v > 0): errs.append("provincia %s %s: valor no positivo" % (pid, t))
    if sum(len(r) for r in pv.values()) < 200: errs.append("pocas celdas provinciales (%d)" % sum(len(r) for r in pv.values()))
    return errs
def check_rent(rent, nat0):
    errs = []
    if set(rent) - {"ES"} != set(CCAA.values()): errs.append("canones: comunidades distintas de las 17")
    for code, ts in rent.items():
        for t, ys in ts.items():
            for y, v in ys.items():
                if not (0 < v < 20000): errs.append("canon %s %s %s fuera de rango: %s" % (code, t, y, v))
    if nat0:
        for y, v in nat0.items():
            g = rent["ES"]["GEN"].get(y)
            if g is not None and abs(g - v) > 0.06: errs.append("canon Total Nacional %s: Cuadro 1 (%s) != Cuadro 0 (%s)" % (y, g, v))
    return errs
def merge_years(old, new, years_new, what):
    """old/new: {codigo: {tipo: {año: valor}}}. Gana el libro nuevo en los años que trae; se conservan los anteriores. Devuelve (resultado, revisiones)."""
    out = json.loads(json.dumps(old)); rev = 0
    for code, ts in new.items():
        for t, ys in ts.items():
            d = out.setdefault(code, {}).setdefault(t, {})
            for y, v in ys.items():
                k = str(y)
                if k in d and abs(d[k] - v) > 0.5: rev += 1
                d[k] = v
    if rev: log("revision del MAPA: %d cifras de %s cambian respecto a lo guardado" % (rev, what))
    return out
def r2(v): return None if v is None else round(v, 2)
def to_arrays(d, years):
    return {code: {t: [r2(ys.get(str(y))) for y in years] for t, ys in ts.items()} for code, ts in d.items()}
def build(price_wb, rent_wb, prev, urls, stamp, price_year_def):
    price, weights, pnat0, prov = parse_price(price_wb)
    errs = check_price(price, weights, pnat0, prov)
    rent, rnat0, other = parse_rent(rent_wb); errs += check_rent(rent, rnat0)
    if errs: raise ValueError("; ".join(errs[:8]) + (" (+%d)" % (len(errs) - 8) if len(errs) > 8 else ""))
    oldp = {c: {t: {str(y): v for y, v in zip(prev["years"], arr) if v is not None} for t, arr in ts.items()} for c, ts in prev.get("price", {}).get("ccaa", {}).items()} if prev else {}
    oldr = {c: {t: {str(y): v for y, v in zip(prev["years"], arr) if v is not None} for t, arr in ts.items()} for c, ts in prev.get("rent", {}).get("ccaa", {}).items()} if prev else {}
    newp = {c: {t: {str(y): v for y, v in ys.items()} for t, ys in ts.items()} for c, ts in price.items()}
    newr = {c: {t: {str(y): v for y, v in ys.items()} for t, ys in ts.items()} for c, ts in rent.items()}
    mp = merge_years(oldp, newp, None, "precios de la tierra"); mr = merge_years(oldr, newr, None, "canones")
    years = sorted({int(y) for d in (mp, mr) for ts in d.values() for ys in ts.values() for y in ys})
    status = dict((prev or {}).get("status", {}))
    for y in sorted({y for ts in price.values() for ys in ts.values() for y in ys}): status[str(y)] = "definitive" if (y < prov["year"] or price_year_def) else "provisional"
    # provincial: un solo año; si el libro nuevo trae otro año, el provincial se sustituye (solo se publica el ultimo)
    doc = {"schemaVersion": 1, "generatedAt": stamp, "years": years, "status": status,
           "source": {"name": "MAPA, Estadística agraria: precios medios anuales de las tierras de uso agrario y cánones de arrendamiento rústico", "url": PAGE_PRICE, "licence": "MAPA aviso legal: reutilización citando la fuente"},
           "files": urls,
           "types": TYPES, "rentTypes": {"GEN": "General", "TAS": TYPES["TAS"], "TAR": TYPES["TAR"], "OLS": TYPES["OLS"], "VIS": TYPES["VIS"], "PRD": TYPES["PRD"], "PST": TYPES["PST"]},
           "price": {"unit": "EUR/ha", "ccaa": to_arrays(mp, years), "prov": {"year": prov["year"], "names": prov["names"], "values": {p: {t: r2(v) for t, v in row.items()} for p, row in prov["values"].items()}}},
           "rent": {"unit": "EUR/ha/year", "ccaa": to_arrays(mr, years), "other": ({"year": other["year"], "values": {c: {t: r2(v) for t, v in row.items()} for c, row in other["values"].items()}} if other else None)},
           "log": LOG[-40:]}
    return doc
def load_wb(src):
    import openpyxl
    return openpyxl.load_workbook(io.BytesIO(src) if isinstance(src, bytes) else src, data_only=True)
def main():
    args = sys.argv[1:]; outdir = ROOT / "data" / "spain-land"
    if "--out" in args: outdir = Path(args[args.index("--out") + 1])
    out = outdir / "land.json"; prev = None
    if out.exists():
        try: prev = json.loads(out.read_text("utf-8"))
        except Exception: prev = None
    pf = args[args.index("--price-file") + 1] if "--price-file" in args else None
    rf = args[args.index("--rent-file") + 1] if "--rent-file" in args else None
    try:
        urls = {}
        if pf and rf:
            pb, rb = load_wb(pf), load_wb(rf); urls = {"price": {"name": Path(pf).name}, "rent": {"name": Path(rf).name}}; pdef = True
        else:
            pu, py, pdef, pt = pick_price(http(PAGE_PRICE).decode("utf-8", "replace")); ru, ry, rt = pick_rent(http(PAGE_RENT).decode("utf-8", "replace"))
            log("precios de la tierra:", pt, pu); log("canones:", rt, ru)
            if prev and prev.get("files", {}).get("price", {}).get("url") == pu and prev.get("files", {}).get("rent", {}).get("url") == ru:
                log("libros sin cambios; no se escribe nada"); return 0
            pb, rb = load_wb(http(pu, timeout=300)), load_wb(http(ru, timeout=300))
            urls = {"price": {"url": pu, "name": pt}, "rent": {"url": ru, "name": rt}}
        stamp = datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
        doc = build(pb, rb, prev, urls, stamp, pdef)
    except Exception as e:
        log("ERROR:", e); return 1
    outdir.mkdir(parents=True, exist_ok=True)
    tmp = outdir / "land.json.tmp"; tmp.write_text(json.dumps(doc, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8"); os.replace(tmp, out)
    n = sum(len(a) for ts in doc["price"]["ccaa"].values() for a in ts.values())
    log("OK: años", doc["years"], "| precio: %d series CCAA/tipo" % sum(len(ts) for ts in doc["price"]["ccaa"].values()), "| provincias", len(doc["price"]["prov"]["values"]), "(año %d)" % doc["price"]["prov"]["year"],
        "| canones: %d series" % sum(len(ts) for ts in doc["rent"]["ccaa"].values()))
    return 0
if __name__ == "__main__":
    rc = main()
    try: (ROOT / "data" / "spain-land-log.txt").write_text("\n".join(LOG) + "\n", "utf-8")
    except Exception: pass
    sys.exit(rc)
