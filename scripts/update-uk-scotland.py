#!/usr/bin/env python3
"""Escocia: censo agrario de junio (Scottish Government, RESAS, "Results from the Scottish Agricultural Census", Main Tables; Open Government Licence v3.0)
-> data/uk-scotland-stats.json. Series anuales desde 2016 de superficie de cultivos (tabla 1), vacuno (3), ovino (4), porcino (5), aves (6), otros animales (7)
y mano de obra (8), tal como las publica el Gobierno de Escocia. La publicacion de cada ano se busca en la coleccion; si falla, no se toca el fichero.
Valores 'x' (confidenciales) o vacios se omiten; no se rellenan."""
import datetime, io, re, sys, html
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
from lib_ukstats import get, series, write
import openpyxl
ROOT = Path(__file__).resolve().parents[1]
DOCS = 'https://www.gov.scot/publications/results-from-the-scottish-agricultural-census-june-%d/documents/'
TABLES = {'Table_1': ('crops', 'ha', 'Area'), 'Table_3': ('livestock', 'head', 'Cattle'), 'Table_4': ('livestock', 'head', 'Sheep'), 'Table_5': ('livestock', 'head', 'Pigs'),
          'Table_6': ('livestock', 'head', 'Poultry'), 'Table_7': ('livestock', 'head', 'Other livestock'), 'Table_8': ('production', 'people', 'Workforce')}
LOG = []
def log(*a): s = ' '.join(str(x) for x in a); LOG.append(s); print(s, flush=True)
def main():
    url = None
    for y in range(datetime.date.today().year, datetime.date.today().year - 3, -1):
        try: t = html.unescape(get(DOCS % y).decode('utf-8', 'replace'))
        except SystemExit: continue
        c = [u for u in re.findall(r'href="([^"]+)"', t) if 'main-tables/' in u and '---ods' not in u]
        if c: url = c[0] if c[0].startswith('http') else 'https://www.gov.scot' + c[0]; log('ano', y, url); break
    if not url: raise SystemExit('no se encontraron las Main Tables del censo de Escocia')
    wb = openpyxl.load_workbook(io.BytesIO(get(url)), read_only=True, data_only=True)
    out = []
    for name, (grp, unit, pre) in TABLES.items():
        if name not in wb.sheetnames: log('falta', name); continue
        rows = [list(r) for r in wb[name].iter_rows(values_only=True)]
        hi = next((i for i, r in enumerate(rows[:12]) if r and r[1] and re.match(r'^\d{4}\b', str(r[1]))), None)
        if hi is None: log('sin cabecera', name); continue
        cols = [(j, int(str(c)[:4])) for j, c in enumerate(rows[hi]) if j and c and re.match(r'^\d{4} ', str(c)) and not re.search(r'(?i)average|change', str(c))]
        for r in rows[hi + 1:]:
            if not r or not r[0] or not str(r[0]).strip(): break
            lab = re.sub(r'\s*\[[^\]]*\]', '', str(r[0])).strip()
            if re.search(r'(?i)percentage|migrant', lab): continue
            pts = []
            for j, y in cols:
                v = r[j] if j < len(r) else None
                try: pts.append([str(y), round(float(v), 1)])
                except (TypeError, ValueError): pass
            if len(pts) < 4: continue
            sid = 'uk-sco-' + re.sub(r'[^a-z0-9]+', '-', ('%s-%s' % (pre, lab)).lower()).strip('-')[:70]
            out.append(series(sid, grp, '%s: %s, Scotland (June census)' % (pre, lab), unit, 'annual', pts, 'Scottish Government'))
    if len(out) < 40: raise SystemExit('muy pocas series (%d)' % len(out))
    ids = [s['id'] for s in out]; dup = {i for i in ids if ids.count(i) > 1}
    if dup: out = [s for s in out if s['id'] not in dup] + []; log('ids duplicados omitidos', sorted(dup))
    log('series', len(out))
    write(ROOT / 'data/uk-scotland-stats.json', {'name': 'Scottish Government (RESAS), Results from the Scottish Agricultural Census (June)', 'url': 'https://www.gov.scot/collections/june-scottish-agricultural-census/', 'license': 'Open Government Licence v3.0'}, out, LOG)
    (ROOT / 'data/uk-scotland-log.txt').write_text('\n'.join(LOG) + '\n', encoding='utf-8')
if __name__ == '__main__': main()
