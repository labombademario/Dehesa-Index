#!/usr/bin/env python3
import json, re, urllib.request
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DATA_JS = ROOT / "js" / "data.js"
SNAP_DIR = ROOT / "data" / "snapshots"
LANDING = "https://www.worldbank.org/en/research/commodity-markets"

def fetch(url):
    req = urllib.request.Request(url, headers={"User-Agent":"Dehesa-Index-data-bot/1.0"})
    with urllib.request.urlopen(req, timeout=60) as r:
        return r.read()

html = fetch(LANDING).decode("utf-8", "ignore")
matches = re.findall(r'https://thedocs\.worldbank\.org/[^"\']+CMO-Historical-Data-Monthly\.xlsx', html)
if not matches:
    raise RuntimeError("World Bank monthly XLSX link not found on commodity-markets landing page")
xlsx_url = matches[0].replace("&amp;", "&")

pub = re.search(r'\((\d{2}/\d{2}/\d{4})\)', re.sub(r'<[^>]+>', ' ', html))
publication_date = datetime.strptime(pub.group(1), "%m/%d/%Y").date().isoformat() if pub else None

xlsx = fetch(xlsx_url)
tmp = ROOT / ".worldbank-monthly.xlsx"
tmp.write_bytes(xlsx)
wb = load_workbook(tmp, read_only=True, data_only=True)
ws = wb["Monthly Prices"]

rows = list(ws.iter_rows(values_only=True))
header_row = None
urea_col = None
for idx, row in enumerate(rows[:100]):
    for col, val in enumerate(row):
        normalized = re.sub(r"\\s+", " ", str(val or "").strip().lower())
        if normalized in ("urea", "urea, e. europe", "urea e. europe") or ("urea" in normalized and "europe" in normalized):
            header_row, urea_col = idx, col
            break
    if urea_col is not None:
        break
if urea_col is None:
    raise RuntimeError("World Bank Urea series not found")

points = []
for row in rows[header_row + 1:]:
    if not row:
        continue
    period = str(row[0] or "").strip()
    if not re.fullmatch(r"\d{4}M\d{2}", period, re.I):
        continue
    value = row[urea_col]
    if isinstance(value, (int, float)):
        points.append((period.upper(), float(value)))
points = points[-12:]
if not points:
    raise RuntimeError("No World Bank urea observations found")

latest_period, latest_value = points[-1]
year, month = int(latest_period[:4]), int(latest_period[5:])
obs_date = f"{year:04d}-{month:02d}-01"
change_pct = None
if len(points) >= 2 and points[-2][1]:
    change_pct = round((latest_value / points[-2][1] - 1) * 100, 4)

data = DATA_JS.read_text(encoding="utf-8")
pattern = re.compile(r"(\{ nameKey: 'urea',.*?\n\s*eu: \{)(.*?)(\},\n\s*uk: \{)", re.S)
m = pattern.search(data)
if not m:
    raise RuntimeError("Urea EU block not found in js/data.js")

eu_body = f" price: {latest_value:g}, changePct: {change_pct if change_pct is not None else 'null'}, history: " + json.dumps([v for _, v in points], separators=(",", ":")) + ", currency: 'USD', kgPerUnit: 1000"
data = data[:m.start(2)] + eu_body + data[m.end(2):]
block = re.compile(r"('fertilizantes-urea-eu': \{)(.*?)(\n    \},)", re.S)
bm = block.search(data)
if not bm:
    raise RuntimeError("Data Trust EU urea block not found")
verified_at = datetime.utcnow().isoformat(timespec="seconds") + "Z"
trust_body = bm.group(2)
trust_body = re.sub(r"methodology: '[^']*'", "methodology: 'World Bank Urea, E. Europe international commodity reference; USD/metric ton.'", trust_body)
trust_body = re.sub(r"comparability: '[^']*', observationDate: [^,]+, publicationDate: [^\n]+", f"comparability: 'not_comparable', observationDate: '{obs_date}', publicationDate: {json.dumps(publication_date)}, status: 'verified', verifiedAt: '{verified_at}'", trust_body)
data = data[:bm.start(2)] + trust_body + data[bm.end(2):]
DATA_JS.write_text(data, encoding="utf-8")

snapshot_date = datetime.utcnow().date().isoformat()
history = []
for period, value in points:
    y, mo = int(period[:4]), int(period[5:])
    history.append({"period": mo, "year": y, "value": value})
observation = {
    "id": "di_fertilizantes_urea_eu",
    "product": "urea",
    "region": "eu",
    "sourceId": "world_bank",
    "observationDate": obs_date,
    "publicationDate": publication_date,
    "value": latest_value,
    "currency": "USD",
    "unit": "tonelada",
    "frequency": "monthly",
    "status": "verified",
    "verifiedAt": datetime.utcnow().isoformat(timespec="seconds") + "Z",
    "comparability": "not_comparable",
    "history": history
}
SNAP_DIR.mkdir(parents=True, exist_ok=True)
snapshot = SNAP_DIR / f"{snapshot_date}.json"
if snapshot.exists():
    old = json.loads(snapshot.read_text(encoding="utf-8"))
    obs = [o for o in old.get("observations", []) if not (o.get("product") == "urea" and o.get("region") == "eu")]
    old["observations"] = obs + [observation]
    snapshot.write_text(json.dumps(old, indent=2) + "\n", encoding="utf-8")
else:
    snapshot.write_text(json.dumps({"schemaVersion":"1.0","generatedAt":datetime.utcnow().isoformat(timespec="seconds")+"Z","observations":[observation]}, indent=2) + "\n", encoding="utf-8")
tmp.unlink(missing_ok=True)
print(f"World Bank urea EU updated: {latest_period} = {latest_value} USD/t; publication={publication_date}")
