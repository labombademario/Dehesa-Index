#!/usr/bin/env python3
"""Francia: gasoleo no de carretera (GNR, el gasoleo agricola), fioul domestico y gasoleo de carretera -> data/fr-gnr.json
1) Precios medios nacionales semanales de la DGEC (Ministere de la Transition ecologique, pagina «Prix des produits petroliers»): Excel
   «Prix HTT et TTC depuis janvier 2020», hojas «Carburants non routiers» (GNR), «Carburants routiers» (gazole) y «Combustibles» (fioul
   domestique). Contenido del sitio bajo Licence Ouverte Etalab 2.0 (mentions legales de ecologie.gouv.fr).
2) Precio del gazole en cada gasolinera (flujo instantaneo «Prix des carburants en France», data.economie.gouv.fr, Licence Ouverte 2.0):
   Dehesa Index calcula la mediana por departamento con los precios actualizados en los ultimos 8 dias. En las gasolineras no se vende GNR.
Si una de las dos partes falla se conserva la del fichero anterior. Uso: python3 scripts/update-fr-gnr.py"""
import datetime, html, io, json, re, statistics, sys, time, urllib.parse, urllib.request
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "fr-gnr.json"
PAGE = "https://www.ecologie.gouv.fr/politiques-publiques/prix-produits-petroliers"
FLUX = "https://data.economie.gouv.fr/api/explore/v2.1/catalog/datasets/prix-des-carburants-en-france-flux-instantane-v2/exports/json?select=code_departement,gazole_prix,gazole_maj"
H = {"User-Agent": "Dehesa-Index-data-bot/1.0 (+https://dehesaindex.com)"}
TODAY = datetime.date.today()
def get(u, t=240):
    last = None
    for i in range(4):
        try:
            with urllib.request.urlopen(urllib.request.Request(u, headers=H), timeout=t) as r: return r.read()
        except Exception as e: last = e; time.sleep(6 * (i + 1))
    raise RuntimeError("%s -> %s" % (u[-90:], last))
def num(v):
    try: x = float(v)
    except Exception: return None
    return x
def sheet(wb, pat):
    for ws in wb.worksheets:
        if re.search(pat, ws.title, re.I): return ws
    raise ValueError("hoja %s no encontrada: %s" % (pat, [w.title for w in wb.worksheets]))
def weekly(ws, cols):
    """filas con fecha en la columna A -> {fecha: [valores de cols]}"""
    out = {}
    for row in ws.iter_rows(values_only=True):
        d = row[0]
        if isinstance(d, datetime.datetime): d = d.date()
        if not isinstance(d, datetime.date): continue
        out[d.isoformat()] = [num(row[c]) if c < len(row) else None for c in cols]
    return out
def header(ws):
    rows = list(ws.iter_rows(min_row=1, max_row=2, values_only=True))
    return [" ".join(str(r[i] or "") for r in rows) for i in range(max(len(r) for r in rows))]
def national():
    import openpyxl
    page = get(PAGE).decode("utf-8", "replace")
    u = next(urllib.parse.urljoin(PAGE, html.unescape(h)) for h in re.findall(r'href="([^"]+\.xlsx)"', page) if re.search(r"Prix%20HTT%20et%20TTC|Prix HTT et TTC", h))
    wb = openpyxl.load_workbook(io.BytesIO(get(u)), read_only=True, data_only=True)
    g = sheet(wb, r"non routier"); hg = header(g)
    # columnas: la primera «HTT» y la primera «TTC» que digan «non routier»
    ih = next(i for i, h in enumerate(hg) if i and "non routier" in h.lower() and "h.t.t" in h.lower())
    it = next(i for i, h in enumerate(hg) if i and "non routier" in h.lower() and "t.t.c" in h.lower())
    r = sheet(wb, r"^Carburants routiers"); hr = header(r)
    ra = next(i for i in range(1, len(hr)) if "gazole" in hr[i].lower() and i >= next(j for j in range(1, len(hr)) if "t.t.c" in hr[j].lower()))
    c = sheet(wb, r"Combustibles"); hc = header(c)
    fc = next(i for i in range(1, len(hc)) if "fioul domestique" in hc[i].lower() and i >= next(j for j in range(1, len(hc)) if "t.t.c" in hc[j].lower()))
    G = weekly(g, [ih, it]); R = weekly(r, [ra]); F = weekly(c, [fc])
    dates = sorted(d for d in G if G[d][1] is not None)
    # la hoja mezcla semanas con medias mensuales (fechadas el dia 1): solo se queda el dia de la semana de la serie semanal
    wd = __import__("collections").Counter(datetime.date.fromisoformat(x).weekday() for x in dates).most_common(1)[0][0]
    dates = [x for x in dates if datetime.date.fromisoformat(x).weekday() == wd]
    if len(dates) < 100: raise ValueError("serie GNR corta: %d semanas" % len(dates))
    rows = [[d, G[d][1], G[d][0], (R.get(d) or [None])[0], (F.get(d) or [None])[0]] for d in dates]
    for x in rows:
        if not (0.4 < x[1] < 3.0): raise ValueError("GNR fuera de rango en %s: %s" % (x[0], x[1]))
        if x[3] is not None and x[1] >= x[3]: raise ValueError("GNR mas caro que el gazole en %s" % x[0])
    meta = {"file": u.rsplit("/", 1)[-1], "headers": {"gnr": hg[it].strip(), "gazole": hr[ra].strip(), "fod": hc[fc].strip()}}
    return rows, meta
def departments():
    L = json.loads(get(FLUX, 300)); by = {}; cut = (TODAY - datetime.timedelta(days=8)).isoformat()
    for e in L:
        v = num(e.get("gazole_prix")); d = str(e.get("code_departement") or "").strip().upper(); m = str(e.get("gazole_maj") or "")[:10]
        if v is None or not (0.8 < v < 3.5) or not d or m < cut: continue
        by.setdefault(d.zfill(2) if d.isdigit() else d, []).append(v)
    allv = [v for vs in by.values() for v in vs]
    if len(allv) < 5000: raise ValueError("solo %d gasolineras con gazole reciente" % len(allv))
    return {"date": TODAY.isoformat(), "median": round(statistics.median(allv), 3), "n": len(allv),
            "departments": {k: [round(statistics.median(v), 3), len(v)] for k, v in sorted(by.items())}}
def main():
    try: old = json.loads(OUT.read_text())
    except Exception: old = {}
    doc = {"schemaVersion": 1, "generatedAt": datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"), "unit": "EUR/l",
           "source": {"name": "Ministère de la Transition écologique (DGEC), prix des produits pétroliers; prix des carburants en France (flux instantané)", "url": PAGE,
                      "stationsUrl": "https://www.data.gouv.fr/datasets/prix-des-carburants-en-france-flux-instantane-v2-amelioree", "license": "Licence Ouverte 2.0"},
           "fields": ["date", "gnrTTC", "gnrHTT", "gazoleTTC", "fodTTC"]}
    ok = False
    try:
        rows, meta = national(); doc["weekly"] = rows; doc["source"].update(meta); ok = True
    except Exception as e:
        print("AVISO serie nacional:", e)
        if old.get("weekly"): doc["weekly"] = old["weekly"]; doc["source"].update({k: old["source"][k] for k in ("file", "headers") if k in old.get("source", {})})
    try: doc["stations"] = departments(); ok = True
    except Exception as e:
        print("AVISO gasolineras:", e)
        if old.get("stations"): doc["stations"] = old["stations"]
    if "weekly" not in doc or "stations" not in doc:
        print("FALLO: sin serie nacional o sin gasolineras"); return 0 if OUT.exists() else 1
    doc["asOf"] = doc["weekly"][-1][0]
    OUT.write_text(json.dumps(doc, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    w = doc["weekly"][-1]; s = doc["stations"]
    print("fr-gnr.json: GNR %s %.4f EUR/l TTC (gazole %.4f, fioul %s), %d semanas desde %s; gasolineras %d, mediana gazole %.3f en %d departamentos%s" % (w[0], w[1], w[3] or 0, w[4], len(doc["weekly"]), doc["weekly"][0][0], s["n"], s["median"], len(s["departments"]), "" if ok else " (sin cambios)"))
    return 0
if __name__ == "__main__": sys.exit(main())
