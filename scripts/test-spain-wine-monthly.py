#!/usr/bin/env python3
"""Vino mes a mes de España (MAPA, INFOVI informe mensual): lector, validacion, mes por el titulo, libro sin cuadro 4.6 y acumulacion con los libros reales de julio de 2026 y enero de 2021 (scripts/fixtures/spain-wine, sin red).
Uso: python3 scripts/test-spain-wine-monthly.py"""
import importlib.util, copy, sys
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
sp = importlib.util.spec_from_file_location("uswm", ROOT / "scripts" / "update-spain-wine-monthly.py"); M = importlib.util.module_from_spec(sp); sp.loader.exec_module(M)
import openpyxl
fail = 0
def ok(n, c):
    global fail
    if not c: fail += 1; print("FALLA", n)
F = ROOT / "scripts" / "fixtures" / "spain-wine"
p26, d26 = M.parse(openpyxl.load_workbook(F / "mensual_2026-07.xlsx", data_only=True))
p21, d21 = M.parse(openpyxl.load_workbook(F / "mensual_2021-01.xlsx", data_only=True))
ok("meses por el titulo del cuadro 5", p26 == "2026-07" and p21 == "2021-01")
T = lambda f: d26[f]["TOTAL"]
ok("existencias iniciales 30.838.312 y finales 28.368.423 hl (jul-2026)", T("stockStart") == 30838312 and T("stockEnd") == 28368423 and T("mustEnd") == 2065537)
ok("producción del mes 8.900 hl y uva 0 t (julio no es vendimia)", T("production") == 8900 and T("grape") == 0)
ok("entradas de España 1.836.800 y de otros países 31.660 hl", T("inSpain") == 1836800 and T("inAbroad") == 31660)
ok("salidas: interiores 2.682.534 + exteriores = 4.230.678; UE 1.039.198 y terceros 508.946", T("exitsDomestic") == 2682534 and T("exitsEU") == 1039198 and T("exitsThird") == 508946 and T("exitsTotal") == 4230678)
ok("operaciones propias 118.056 hl", T("ownOps") == 118056)
ok("Castilla-La Mancha: existencias finales 7.735.008 y salidas 1.983.288", d26["stockEnd"]["clm"] == 7735008 and d26["exitsTotal"]["clm"] == 1983288)
ok("17 comunidades", all(len([k for k in d26[f] if k != "TOTAL"]) == 17 for f in M.F))
ok("el libro de enero de 2021 no trae el cuadro 4.6: operaciones propias sin dato (None), no cero", all(v is None for v in d21["ownOps"].values()) and d21["stockEnd"]["TOTAL"] > 0)
for p, d in (("2026-07", d26), ("2021-01", d21)):
    e, w = M.validate(p, d); ok("el libro real de %s cuadra" % p, not e)
bad = copy.deepcopy(d26); bad["stockEnd"]["galicia"] += 90000; ok("comunidades que no suman el total se rechaza", M.validate("2026-07", bad)[0])
bad = copy.deepcopy(d26); bad["exitsTotal"]["TOTAL"] += 500000; ok("salidas distintas de interiores + exteriores se rechaza", M.validate("2026-07", bad)[0])
bad = copy.deepcopy(d26); bad["_exitsExt"]["TOTAL"] += 500000; ok("exteriores distintas de UE + terceros se rechaza", M.validate("2026-07", bad)[0])
bad = copy.deepcopy(d26); bad["inSpain"]["aragon"] = -5; ok("negativo se rechaza", M.validate("2026-07", bad)[0])
bad = copy.deepcopy(d26)
for k in bad["stockEnd"]: bad["stockEnd"][k] = (bad["stockEnd"][k] or 0) * 2
ok("identidad de existencias rota en mas de un 10 % se rechaza", M.validate("2026-07", bad)[0])
wb = openpyxl.load_workbook(F / "mensual_2026-07.xlsx", data_only=True); del wb["5. EXISTENCIAS FINALES"]
try: M.parse(wb); ok("sin cuadro 5 debe fallar", False)
except ValueError: pass
html = '<a href="/dam/x/infovi/ano-2023/infoviagosto2023.xlsx">a</a><a href="/dam/x/infovi/ano-2023/informeinfoviagosto2023.pdf">p</a><a href="/dam/x/infovi/ano-2023/informaciondeclaracionampliadajulio_2023.xlsx">amp</a><a href="/dam/x/infovi/ano-2023/231123declaracionesexistencias_julio23.xlsx">ex</a><a href="/dam/x/infovi/ano-2022/datosinfovijunio2022.xlsx">otro año</a>'
L = M.find_links(html, 2023); ok("solo el mensual de ese año (ni pdf, ni ampliado, ni existencias, ni otro año)", len(L) == 1 and L[0].endswith("infoviagosto2023.xlsx"))
ms, N, CC, rv = M.merge(None, {"2021-01": d21, "2026-07": d26}); ok("dos meses ordenados, sin revisiones", ms == ["2021-01", "2026-07"] and not rv and N["stockEnd"][1] == 28368423.0)
old = {"months": ms, "national": N, "ccaa": CC}; d26b = copy.deepcopy(d26); d26b["stockEnd"]["TOTAL"] += 100
ms2, N2, CC2, rv = M.merge(old, {"2026-07": d26b}); ok("la revisión del MAPA se registra y gana el libro nuevo", any(r[0] == "2026-07" and r[1] == "stockEnd" for r in rv) and N2["stockEnd"][1] == 28368523.0 and N2["stockEnd"][0] == N["stockEnd"][0])
w = M.chain(["2026-06", "2026-07"], {"stockEnd": [100.0, 0.0], "stockStart": [0.0, 90.0]}); ok("existencias que no encadenan entre meses seguidos avisan", len(w) == 1)
w = M.chain(["2026-01", "2026-07"], {"stockEnd": [100.0, 0.0], "stockStart": [0.0, 90.0]}); ok("meses no seguidos no se comparan", not w)
print("test-spain-wine-monthly:", "%d fallos" % fail if fail else "OK"); sys.exit(1 if fail else 0)
