#!/usr/bin/env python3
"""Francia (FranceAgriMer, a mano): cotizaciones mensuales de vino a granel, precios pagados por cereales ecológicos y calidad sanitaria nacional (micotoxinas)
-> data/france-extras-stats.json (formato country-stats, para el catálogo de la ficha de país).
  python3 scripts/build-france-extras.py <COT-VIN-HISTORIQUE_COTATION_UE_DEPUIS2009-*.xls> [<xls anterior>] <SCR-GRC-prixbio-*.xlsx> <SCR-CER-QUALITE_BTBDQS-*.xlsx> <SCR-CER-QUALITE_MaisQS-*.xlsx>
(las rutas se distinguen por el nombre del fichero). Los valores no numéricos (n.d., n.s., «<222») no se publican."""
import datetime, json, re, sys, unicodedata
from pathlib import Path
import openpyxl, xlrd
ROOT = Path(__file__).resolve().parents[1]
MONTHS = {m: i + 1 for i, m in enumerate("janvier fevrier mars avril mai juin juillet aout septembre octobre novembre decembre".split())}
def norm(x): return "".join(c for c in unicodedata.normalize("NFD", str(x).lower()) if unicodedata.category(c) != "Mn").strip()
def pct(a, b): return round((a - b) / abs(b) * 100, 2) if b else None
def mk(id_, group, label, unit, freq, pts, src):
    pts = [[p, round(v, 4)] for p, v in pts]
    return {"id": id_, "group": group, "label": label, "unit": unit, "frequency": freq, "latestPeriod": pts[-1][0], "latest": pts[-1][1], "changePct": pct(pts[-1][1], pts[-2][1]) if len(pts) > 1 else None, "points": pts, "sourceGroup": src}

def wine_file(f):
    s = xlrd.open_workbook(f).sheet_by_index(0); out = {}; camp = None
    for r in range(4, s.nrows):
        a, m = s.cell_value(r, 0), s.cell_value(r, 1)
        if isinstance(a, str) and re.match(r"^\d{4}-\d{2,4}$", a.strip()): camp = int(a.strip()[:4])
        if not isinstance(m, str) or camp is None: continue
        k = norm(m).replace("(*)", "").replace("*", "").strip()
        if k not in MONTHS: continue
        mo = MONTHS[k]; y = camp if mo >= 8 else camp + 1
        out["%04d-%02d" % (y, mo)] = [s.cell_value(r, c) for c in range(2, 10)]
    return out

WINE = [("red-aop", "Red and rosé wine, AOP"), ("red-igp", "Red and rosé wine, IGP"), ("red-vsig-var", "Red and rosé wine without GI, with grape variety"), ("red-vsig", "Red and rosé wine without GI, without grape variety"),
        ("white-aop", "White wine, AOP"), ("white-igp", "White wine, IGP"), ("white-vsig-var", "White wine without GI, with grape variety"), ("white-vsig", "White wine without GI, without grape variety")]

def wine(paths):
    data = {}
    for f in paths: data.update(wine_file(f))  # el fichero más reciente (último) manda: la cotización de julio se completa al mes siguiente
    series = []
    for i, (k, lab) in enumerate(WINE):
        pts = [(p, v[i]) for p, v in sorted(data.items()) if isinstance(v[i], (int, float)) and not isinstance(v[i], bool)]
        assert all(5 < v < 600 for _p, v in pts), "precio de vino fuera de rango en " + k
        if len(pts) > 12: series.append(mk("fr-fam-wine-" + k, "quotes", "%s: bulk price" % lab, "€/hl", "monthly", pts, "FranceAgriMer – cotations des vins (historique depuis 2009)"))
    return series

def organic(f):
    ws = openpyxl.load_workbook(f, data_only=True).worksheets[0]
    crops = [ws.cell(1, c).value for c in range(2, ws.max_column + 1)]
    names = {"Blé tendre": ("soft-wheat", "Soft wheat"), "Maïs": ("maize", "Maize"), "Triticale": ("triticale", "Triticale"), "Tournesol": ("sunflower", "Sunflower"), "Féveroles": ("faba-beans", "Faba beans")}
    series = []
    for j, c in enumerate(crops):
        pts = []
        for r in range(3, ws.max_row + 1):
            lab, v = ws.cell(r, 1).value, ws.cell(r, 2 + j).value
            if isinstance(lab, str) and re.match(r"^\d{4}/\d{2}$", lab.strip()) and isinstance(v, (int, float)): pts.append((lab.strip()[:4], float(v)))
        if pts and c in names:
            assert all(50 < v < 900 for _p, v in pts)
            series.append(mk("fr-fam-organic-" + names[c][0], "organic", "%s, organic: producer price (campaign starting that year)" % names[c][1], "€/t", "annual", pts, "FranceAgriMer – enquête prix payés aux producteurs (bio)"))
    return series

QUAL = [("Blé tendre - DON", "fr-fam-q-soft-wheat-don", "Soft wheat, deoxynivalenol (DON)"), ("Blé tendre - ergot", "fr-fam-q-soft-wheat-ergot", "Soft wheat, ergot sclerotia"),
        ("Blé dur - DON", "fr-fam-q-durum-don", "Durum wheat, deoxynivalenol (DON)"), ("Blé dur - ergot", "fr-fam-q-durum-ergot", "Durum wheat, ergot sclerotia")]
QUAL_MAIZE = [("DON", "fr-fam-q-maize-don", "Maize, deoxynivalenol (DON)"), ("Fumonisines", "fr-fam-q-maize-fumonisins", "Maize, fumonisins"),
              ("Zéaralénone", "fr-fam-q-maize-zearalenone", "Maize, zearalenone"), ("Aflatoxines totales", "fr-fam-q-maize-aflatoxins", "Maize, total aflatoxins")]
SRCQ = "FranceAgriMer / ARVALIS – enquête qualité sanitaire à l'entrée des silos"
def quality(f, sheets):
    wb = openpyxl.load_workbook(f, data_only=True); out = []
    for sh, sid, lab in sheets:
        ws = wb[sh]
        hdr = next(r for r in range(1, 40) if isinstance(ws.cell(r, 1).value, str) and norm(ws.cell(r, 1).value).startswith("annee de recolte"))
        heads = {c: str(ws.cell(hdr, c).value) for c in range(1, ws.max_column + 1) if ws.cell(hdr, c).value}
        hum = next(c for c, h in heads.items() if norm(h).replace("\xa0", " ").startswith("% des volumes collectes < seuil reglementaire"))
        ani = next((c for c, h in heads.items() if norm(h).replace("\xa0", " ").startswith("% des volumes collectes < seuil recommande")), None)
        for col, kind in ((hum, "human food limit"), (ani, "animal feed guidance level")):
            if col is None: continue
            pts = []
            for r in range(hdr + 1, ws.max_row + 1):
                y, v = ws.cell(r, 1).value, ws.cell(r, col).value
                if isinstance(y, (int, float)) and isinstance(v, (int, float)) and 0 <= v <= 1.0001: pts.append((str(int(y)), round(v * 100, 2)))
            if len(pts) >= 3: out.append(mk(sid + ("-human" if kind.startswith("human") else "-feed"), "crops", "%s: share of volume below the %s" % (lab, kind), "%", "annual", pts, SRCQ))
    return out

def main(args):
    by = lambda tag: [a for a in args if tag in Path(a).name]
    wines = sorted(by("COT-VIN"), key=lambda p: re.search(r"C(\d{2})-(\d{2})", p).group(0))
    series = wine(wines) + organic(by("prixbio")[0]) + quality(by("BTBDQS")[0], QUAL) + quality(by("MaisQS")[0], QUAL_MAIZE)
    doc = {"schemaVersion": 1, "generatedAt": datetime.datetime.utcnow().strftime("%Y-%m-%dT%H:%M:%SZ"),
           "countries": {"FR": {"name": "France", "source": {"name": "FranceAgriMer (VISIONet)", "url": "https://visionet.franceagrimer.fr/", "license": "Licence Ouverte 2.0"}, "series": series}}, "log": []}
    (ROOT / "data" / "france-extras-stats.json").write_text(json.dumps(doc, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    import collections; print(len(series), "series", dict(collections.Counter(s["group"] for s in series)))
if __name__ == "__main__":
    if len(sys.argv) < 5: raise SystemExit(__doc__)
    main(sys.argv[1:])
