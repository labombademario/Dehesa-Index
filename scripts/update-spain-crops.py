#!/usr/bin/env python3
"""España: superficies, producciones y rendimientos de cultivos por provincia (MAPA, «Superficies y producciones anuales de cultivos», Reglamento (CE) 543/2009)
-> data/spain-crops/crops-<grupo>.json + data/spain-crops/index.json.

Fuente: libros Excel del MAPA (pagina «superficies-producciones-anuales-cultivos»). Cada libro trae, ademas de sus tablas con formulas, una hoja de base de datos larga
(BD_*: provincia, grupo, cultivo, variable, valor) que es de donde sale todo. La campana de cada libro NO se deduce: se lee del titulo del enlace de la propia pagina
(«... DATOS PROVISIONALES CAMPAÑA 2025») y si falta, el libro se ignora. Los libros de la campana mas reciente (solo cultivos leñosos y hortalizas) llevan su propia BD;
los de la campana anterior comparten una BD con todos los grupos, asi que se lee un solo libro.
Nada se estima: una variable ausente es null, y los datos se guardan tal cual (ha, t) salvo la suma de regadio al aire libre + protegido en hortalizas.
Modo --verify: compara las sumas provinciales con las tablas nacionales que publica el propio MAPA en cada libro (solo diagnostico; descarga todos los libros del grupo)."""
import argparse, datetime, io, json, os, re, sys, time, urllib.request
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
PAGE = "https://www.mapa.gob.es/es/estadistica/temas/estadisticas-agrarias/agricultura/superficies-producciones-anuales-cultivos/"
SITE = "https://www.mapa.gob.es"
UA = {"User-Agent": "Mozilla/5.0 (compatible; DehesaIndex/1.0; +https://dehesaindex.com)"}
MEASURES = ["areaTotal", "areaSecano", "areaRegadio", "areaHarvested", "production"]
# grupo del MAPA (columna CSP_PST_TXTCORTO de la BD) -> clave, nombre y codigos de variable de cada medida (se suman si hay varios)
GROUPS = {
    "Cereales":             {"key": "cereales",      "total": ["SCUT"], "sec": ["SCUS"], "reg": ["SCUR"], "harv": ["SCOT"], "prod": ["PCOG"]},
    "Leguminosas":          {"key": "leguminosas",   "total": ["SCUT"], "sec": ["SCUS"], "reg": ["SCUR"], "harv": ["SCOT"], "prod": ["PCOG"]},
    "Tubérculos":           {"key": "tuberculos",    "total": ["SCUT"], "sec": ["SCUS"], "reg": ["SCUR"], "harv": ["SCOT"], "prod": ["PRCO"]},
    "Industriales":         {"key": "industriales",  "total": ["SCUT"], "sec": ["SCUS"], "reg": ["SCUR"], "harv": ["SCOT"], "prod": ["PRCO"]},
    "Hortalizas":           {"key": "hortalizas",    "total": ["SCUT"], "sec": ["SCUS"], "reg": ["SCRA", "SCRP"], "harv": ["SCOT"], "prod": ["PCOH"]},
    "Cítricos":             {"key": "citricos",      "total": ["STTT"], "sec": ["STST"], "reg": ["STRT"], "harv": ["SUAT"], "prod": ["PRTR"]},
    "Frutales No Cítricos": {"key": "frutales",      "total": ["STTT"], "sec": ["STST"], "reg": ["STRT"], "harv": ["SUAT"], "prod": ["PRTR"]},
    "Olivar":               {"key": "olivar",        "total": ["SRTT"], "sec": ["SRTS"], "reg": ["SRTR"], "harv": ["SRRT"], "prod": ["PTRE"]},
    "Viñedo":               {"key": "vinedo",        "total": ["SRTT"], "sec": ["SRTS"], "reg": ["SRTR"], "harv": ["SRRT"], "prod": ["PTRE"]},
    "Otros Leñosos":        {"key": "otros_lenosos", "total": ["SRTT"], "sec": ["SRTS"], "reg": ["SRTR"], "harv": ["SRRT"], "prod": ["PRTR"]},
}
SLOT = {"total": 0, "sec": 1, "reg": 2, "harv": 3, "prod": 4}
CROP_RE = re.compile(r"^[A-Z]{2}\d{4}$")
LOG = []
def log(*a):
    s = " ".join(str(x) for x in a); LOG.append(s); print(s, flush=True)

def http(url, method="GET", tries=3, timeout=300):
    last = None
    for i in range(tries):
        try:
            req = urllib.request.Request(url, headers=UA, method=method)
            with urllib.request.urlopen(req, timeout=timeout) as r:
                return r.read() if method == "GET" else dict(r.headers)
        except Exception as e:
            last = e; time.sleep(4 * (i + 1))
    raise RuntimeError("%s %s -> %s" % (method, url, str(last)[:160]))

def page_files(html):
    """[(url, nombre de fichero, etiqueta visible, campaña)] de los enlaces .xlsx de la pagina. Sin «CAMPAÑA aaaa» en la etiqueta no hay campana y se descarta."""
    out = []
    for h, t in re.findall(r'<a[^>]+href="([^"]+reglamento-n--543-2009/[^"]+\.xlsx)"[^>]*>(.*?)</a>', html, re.S):
        label = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", t)).strip()
        m = re.search(r"CAMPA[ÑN]A\s+(\d{4})", label, re.I)
        if not m: log("enlace sin campaña, se ignora:", h.split("/")[-1], "|", label[:80]); continue
        url = h if h.startswith("http") else SITE + h
        out.append((url, h.split("/")[-1], label, int(m.group(1))))
    return out

def pick_files(items):
    """campana -> lista de ficheros a leer. Campana con BD compartida (muchos grupos): se lee un solo libro (preferencia: cereales). Campana con BD propia: todos."""
    by = {}
    for it in items: by.setdefault(it[3], []).append(it)
    plan = {}
    for c, lst in by.items():
        big = [x for x in lst if "cereales" in x[1]]
        plan[c] = {"shared": big[:1], "own": [] if big else lst}
    return plan

def clean_name(s):
    s = re.sub(r"\s+", " ", str(s or "")).strip()
    lvl = 0
    while s.startswith("-"): lvl += 1; s = s[1:].lstrip()
    return s, lvl

def parse_bd(rows, stats):
    """rows: iterable de tuplas de la hoja BD (primera fila = cabecera). Devuelve {clave_grupo: {codigo_cultivo: {nombre, nivel, prov: {id: {var: valor}}}}}, {id: nombre}.
    Una misma (provincia, cultivo, variable) repetida con valores distintos es un error: no se elige."""
    out = {}; provs = {}; need = {}
    for g, cfg in GROUPS.items():
        need[g] = set(sum((cfg[k] for k in ("total", "sec", "reg", "harv", "prod")), []))
    first = True
    for r in rows:
        if first: first = False; continue
        if r is None or len(r) < 10: continue
        pid, pname, _gid, grp, _ver, crop, cname, var, _vdesc, val = r[:10]
        stats["rows"] += 1
        if grp not in GROUPS or var not in need[grp] or not crop: continue
        if val is None: stats["null"] += 1; continue
        if isinstance(val, bool) or not isinstance(val, (int, float)): stats["nonnumeric"] += 1; continue
        if val < 0: stats["negative"] += 1; continue
        pid = "%02d" % int(pid); provs[pid] = str(pname).strip()
        cfg = GROUPS[grp]; g = out.setdefault(cfg["key"], {}); nm, lvl = clean_name(cname)
        c = g.setdefault(str(crop), {"name": nm, "level": lvl, "prov": {}})
        d = c["prov"].setdefault(pid, {})
        if var in d and d[var] != val:
            raise ValueError("valor repetido y distinto: %s %s %s %r != %r" % (pid, crop, var, d[var], val))
        d[var] = val; stats["used"] += 1
    return out, provs

def build_group(gname, crops, provs):
    cfg = GROUPS[gname]; out = []
    for code in sorted(crops):
        c = crops[code]; v = {}
        for pid in sorted(c["prov"]):
            d = c["prov"][pid]; arr = [None] * 5
            for k, slot in SLOT.items():
                have = [d[x] for x in cfg[k] if x in d]
                if have: arr[slot] = sum(have)
            if any(x for x in arr if x): v[pid] = [clean(x) for x in arr]
        if not v: continue
        tot = [None] * 5
        for arr in v.values():
            for i, x in enumerate(arr):
                if x is not None: tot[i] = (tot[i] or 0) + x
        out.append({"c": code, "n": c["name"], "l": c["level"], "t": [clean(x) for x in tot], "v": v})
    return out
def clean(x):
    if x is None: return None
    return int(x) if float(x) == int(x) else round(float(x), 3)

def read_bd(data):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    bd = [w for w in wb.worksheets if w.title.upper().startswith("BD_")]
    if len(bd) != 1: raise RuntimeError("se esperaba una hoja BD_, hay: %s" % [w.title for w in wb.worksheets])
    stats = {"rows": 0, "used": 0, "null": 0, "nonnumeric": 0, "negative": 0}
    g, p = parse_bd(bd[0].iter_rows(values_only=True), stats)
    return g, p, stats

def national_tables(data):
    """Hoja 'TABLA GLOBAL NACIONAL*' -> {codigo_cultivo: {variable: valor}} tal cual la publica el MAPA (celdas con formulas, valores en cache)."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    sh = [w for w in wb.worksheets if w.title.upper().startswith("TABLA GLOBAL NACIONAL")]
    if not sh: return {}
    rows = list(sh[0].iter_rows(values_only=True)); codes = None; res = {}
    for r in rows:
        toks = [i for i, x in enumerate(r) if isinstance(x, str) and re.fullmatch(r"[A-Z][A-Z0-9]{3}", x.strip())]
        if codes is None and len(toks) >= 6 and not any(isinstance(x, str) and CROP_RE.match(x.strip()) for x in r[:3] if x): codes = {i: r[i].strip() for i in toks}; continue
        if codes is None: continue
        code = next((x.strip() for x in r[:3] if isinstance(x, str) and CROP_RE.match(x.strip())), None)
        if not code: continue
        d = res.setdefault(code, {})
        for i, v in codes.items():
            if i < len(r) and isinstance(r[i], (int, float)) and not isinstance(r[i], bool): d[v] = r[i]
    return res

def verify(allit, report):
    """Para cada libro de grupos relevantes: suma provincial (BD) de cada cultivo y variable usada frente a la tabla nacional del MAPA.
    Campana con BD compartida (hay libro de cereales): se lee una sola BD; campana con BD propia: la de cada libro."""
    skip = re.compile(r"(flores|ornament|plantones|semillas|viveros|energ|forrajeros|aprovech)")
    its = [it for it in allit if "formato_tabla_global" in it[1] and not skip.search(it[1])]
    shared = {}
    for it in allit:
        if "cereales" in it[1]: shared[it[3]] = read_bd(http(it[0]))
    tot_ok = tot_bad = tot_missing = 0; shown = 0; nobd = 0
    for it in its:
        url, fname, _lab, camp = it
        data = http(url); nat = national_tables(data)
        g, _p, _s = shared[camp] if camp in shared else read_bd(data)
        ok = bad = miss = 0
        for gname, cfg in GROUPS.items():
            crops = g.get(cfg["key"]) or {}
            for code, c in crops.items():
                if code not in nat: continue
                for var in set(sum((cfg[k] for k in ("total", "sec", "reg", "harv", "prod")), [])):
                    if var not in nat[code]: miss += 1; continue
                    have = [d[var] for d in c["prov"].values() if var in d]
                    if not have:   # fila agregada (p. ej. HO0000): la BD no trae esa variable, solo la tabla nacional; no hay nada que comparar
                        nobd += 1; report.append("sin BD %s %s %s %s MAPA=%s (no comparable)" % (camp, fname, code, var, nat[code][var])); continue
                    sm = sum(have)
                    if abs(sm - nat[code][var]) <= max(1.5, abs(nat[code][var]) * 0.0005): ok += 1
                    else:
                        bad += 1
                        if shown < 40: shown += 1; report.append("DIFF %s %s %s %s BD=%s MAPA=%s" % (camp, fname, code, var, sm, nat[code][var]))
        report.append("verificado %s %s: nacional=%d cultivos, ok=%d diferencias=%d sin columna=%d" % (camp, fname, len(nat), ok, bad, miss))
        tot_ok += ok; tot_bad += bad; tot_missing += miss
    report.append("variables sin datos en la BD (filas agregadas, no comparables): %d" % nobd)
    return tot_ok, tot_bad, tot_missing

def write_json(path, obj):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(ROOT / "data" / "spain-crops"))
    ap.add_argument("--local", action="append", default=[], help="campana:ruta.xlsx (sin red; la campana se da a mano, solo para pruebas)")
    ap.add_argument("--verify", action="store_true")
    ap.add_argument("--force", action="store_true")
    a = ap.parse_args(); out = Path(a.out); now = datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    prev = None
    if (out / "index.json").exists():
        try: prev = json.loads((out / "index.json").read_text(encoding="utf-8"))
        except Exception: prev = None
    files = []   # (campana, url, fichero, etiqueta, modo, meta)
    if a.local:
        for spec in a.local:
            c, p = spec.split(":", 1); files.append((int(c), p, Path(p).name, "local", "shared", {}))
    else:
        html = http(PAGE).decode("utf-8", "replace"); items = page_files(html)
        if not items: raise SystemExit("la pagina no tiene enlaces .xlsx con campaña: no se toca nada")
        plan = pick_files(items)
        for c in sorted(plan):
            for mode in ("shared", "own"):
                for it in plan[c][mode]:
                    try: h = http(it[0], "HEAD", tries=2, timeout=60)
                    except Exception: h = {}
                    files.append((c, it[0], it[1], it[2], mode, {"lastModified": h.get("Last-Modified"), "contentLength": h.get("Content-Length")}))
    if a.verify:
        rep = []; allit = page_files(http(PAGE).decode("utf-8", "replace"))
        ok, bad, miss = verify(allit, rep)
        Path("research_out").mkdir(exist_ok=True); Path("research_out/verify.txt").write_text("\n".join(rep) + "\nTOTAL ok=%d diferencias=%d sin columna=%d\n" % (ok, bad, miss), encoding="utf-8")
        print("\n".join(rep[-40:])); return 0 if bad == 0 else 1
    # ¿cambio algo? mismo fichero, misma fecha de modificacion y mismo tamano -> no se vuelve a bajar
    if prev and not a.force and not a.local:
        pf = {(x["campaign"], x["file"]): (x.get("lastModified"), x.get("contentLength")) for x in prev.get("files", [])}
        same = all(pf.get((f[0], f[2])) == (f[5].get("lastModified"), f[5].get("contentLength")) and f[5].get("lastModified") for f in files)
        if same and len(pf) == len(files):
            prev["checkedAt"] = now; write_json(out / "index.json", prev); log("sin cambios en el MAPA; se conserva lo publicado"); return 0
    camps = {}; provs_all = {}; meta_files = []
    for c, url, fname, label, mode, meta in files:
        data = Path(url).read_bytes() if a.local else http(url); log("leido", fname, len(data), "bytes, campaña", c, mode)
        g, p, st = read_bd(data); log("  BD:", st, "grupos:", sorted(g))
        provs_all.update(p)
        for gname, cfg in GROUPS.items():
            if cfg["key"] in g:
                if cfg["key"] in camps.get(c, {}): log("  grupo repetido en la campaña", c, cfg["key"], "-> se conserva el primero"); continue
                camps.setdefault(c, {})[cfg["key"]] = (gname, build_group(gname, g[cfg["key"]], p))
        meta_files.append({"campaign": c, "file": fname, "label": label, "mode": mode, "lastModified": meta.get("lastModified"), "contentLength": meta.get("contentLength"),
                           "groups": sorted(k for k in g if k in {x["key"] for x in GROUPS.values()}), "bdRows": st["rows"], "bdUsed": st["used"], "bdNegativeIgnored": st["negative"]})
    if not camps: raise SystemExit("no se leyó ningún grupo: no se toca nada")
    idx_groups = {}
    for gname, cfg in GROUPS.items():
        k = cfg["key"]; byc = {c: camps[c][k] for c in sorted(camps) if k in camps[c]}
        if not byc: continue
        doc = {"schemaVersion": 1, "generatedAt": now, "group": k, "groupName": gname, "source": {"id": "mapa_es", "name": "MAPA, Superficies y producciones anuales de cultivos", "url": PAGE},
               "measures": MEASURES, "units": {"area": "ha", "production": "t"}, "provinces": {pid: provs_all[pid] for pid in sorted(provs_all)},
               "campaigns": {str(c): {"status": "provisional", "crops": v[1]} for c, v in byc.items()}}
        write_json(out / ("crops-" + k + ".json"), doc)
        idx_groups[k] = {"name": gname, "file": "crops-" + k + ".json", "campaigns": [int(c) for c in byc], "crops": {str(c): len(v[1]) for c, v in byc.items()}}
        log("grupo", k, {str(c): len(v[1]) for c, v in byc.items()})
    write_json(out / "index.json", {"schemaVersion": 1, "generatedAt": now, "checkedAt": now, "source": {"id": "mapa_es", "url": PAGE}, "files": meta_files, "groups": idx_groups,
                                   "note": "Datos provisionales del MAPA tal como los publica; la campana se toma del titulo de cada enlace de la pagina."})
    return 0

if __name__ == "__main__":
    try: rc = main()
    finally:
        try:
            lp = ROOT / "data" / "spain-crops-log.txt"
            if "--local" not in sys.argv and "--out" not in sys.argv: lp.write_text("\n".join(LOG[-200:]) + "\n", encoding="utf-8")
        except Exception: pass
    sys.exit(rc)
