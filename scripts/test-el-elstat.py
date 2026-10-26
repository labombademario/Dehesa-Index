#!/usr/bin/env python3
"""Pruebas del lector de ELSTAT (scripts/update-el-elstat.py) con tablas sinteticas que imitan la estructura real de los ficheros:
meses vacios o 0,0 = null, bloques por año (mas recientes primero), promedios anuales, tipografia corregida y rechazo de cuentas regionales que no suman."""
import importlib.util, io, sys
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("el", ROOT / "scripts" / "update-el-elstat.py"); E = importlib.util.module_from_spec(spec); spec.loader.exec_module(E)
fails = []
def ok(c, m):
    print(("OK    " if c else "FALLO ") + m)
    if not c: fails.append(m)
# --- serie 04: dos bloques anuales, el reciente con meses futuros vacios ---
H = ["Code ", "Groups, sub-groups", "", "Weights"] + "Jan Feb Mar Apr May June July Aug Sep Oct Nov Dec".split() + ["Annual Average"]
def row(code, name, w, vals, avg=""): return [code, name, "", w] + vals + [""] * (12 - len(vals)) + [avg]
rows = [[""] * 17, ["Output Price Index in Agriculture - Livestock, 2026"] + [""] * 16, ["Base Year: 2020=100,0"] + [""] * 16, [""] * 17, H,
        row("140000", "AGRICULTURAL GOODS OUTPUT", 100000.0, [150.0, 151.5, 152.0]), row(209000.0, "Other goods and services", 5000.0, [120.0, 0.0, 121.0]),
        ["Note: The price indices are published rounded up to one decimal figure"] + [""] * 16, [""] * 17,
        ["Output Price Index in Agriculture - Livestock, 2025"] + [""] * 16, ["Base Year: 2020=100,0"] + [""] * 16, [""] * 17, H,
        row("140000", "AGRICULTURAL GOODS OUTPUT", 100000.0, [140.0] * 12, 140.0), row(209000.0, "Other goods and services", 5000.0, [110.0] * 12, 110.0),
        ["Note: The price indices are published rounded up to one decimal figure"] + [""] * 16]
S = E.parse_monthly(rows, "Output")
ok(set(S) == {"140000", "209000"}, "codigos normalizados (el numerico 209000.0 pasa a '209000')")
ok(S["140000"]["m"][(2026, 3)] == 152.0 and (2026, 4) not in S["140000"]["m"], "meses futuros vacios no generan valor")
ok((2026, 2) not in S["209000"]["m"], "0,0 en un indice mensual = sin precio (null)")
ok(S["140000"]["avg"] == {2025: 140.0}, "promedio anual leido de su columna")
try: E.parse_monthly(rows, "Input"); ok(False, "una hoja Output no se acepta como Input")
except RuntimeError: ok(True, "una hoja Output no se acepta como Input")
# --- serie 01: anual + mensual del año en curso, 0,0 = sin precio, tipografia corregida ---
P = [["Output price indices"] + [""] * 3, ["API codes", "", "041940", ""], ["Year and Month", "Wheat. soft", "Cuccumbers", "Olive oil"],
     [2019.0, 90.0, 95.0, 125.0], [2020.0, 100.0, 100.0, 100.0], [2021.0, 120.0, 110.0, 123.7], ["", "", "", ""], [2026.0, "", "", ""],
     ["Jan", 0.0, 180.2, 209.8], ["Feb", 0.0, 205.0, 208.5], ["Mar", "", 0.0, 208.4], ["Apr", "", "", ""], ["May", "", "", ""], ["June", "", "", ""], ["July", "", "", ""], ["Aug", "", "", ""], ["Sep", "", "", ""], ["Oct", "", "", ""], ["Nov", "", "", ""], ["Dec", "", "", ""]]
I = E.parse_products(P)
ok(set(I) == {"soft-wheat", "cucumbers", "olive-oil"}, "nombres corregidos (Cuccumbers -> Cucumbers) y ids estables")
ok(I["olive-oil"]["from"] == 2019 and I["olive-oil"]["annual"] == [125.0, 100.0, 123.7], "serie anual desde el primer año")
ok(I["cucumbers"]["code"] == "041940" and I["olive-oil"]["code"] is None, "codigo API cuando existe")
ok(I["cucumbers"]["monthly"]["v"][:4] == [180.2, 205.0, None, None], "0,0 mensual = null, sin estimar")
ok(I["soft-wheat"]["monthly"] is None, "producto sin ningun precio mensual no inventa meses")
# --- cuenta regional ---
import openpyxl
names = list(E.REG)
def book(mutate=None):
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append([None, "December  2024"]); ws.append(["EL.STAT."]); ws.append(["Economic Accounts for Agriculture by Nuts I,II"]); ws.append(["Year 2023"]); ws.append(["In million euro."])
    ws.append(["CATEGORIES", "DESCRIPTION"] + names)
    lines = {"10": 100.0, "13": 50.0, "14": 150.0, "15": 10.0, "16": 160.0, "17": 5.0, "18": 165.0, "19": 80.0, "20": 85.0, "27": 70.0}
    n2 = [k for k, nm in enumerate(names) if E.REG[nm][1] == 2]
    for c, tot in lines.items():
        v = [0.0] * len(names); v[0] = tot
        for k in n2: v[k] = round(tot / len(n2), 6)
        for n1, kids in (("Nisia Aigaiou, Kriti", ("Voreio Aigaio", "Notio Aigaio", "Kriti")), ("Voreia Ellada", ("Anatoliki Makedonia & Thraki", "Kentriki Makedonia", "Dytiki Makedonia", "Ipeiros")), ("Kentriki Ellada", ("Thessalia", "Ionia Nisia", "Dytiki Ellada", "Sterea Ellada", "Peloponnisos"))):
            v[names.index(n1)] = sum(v[names.index(x)] for x in kids)
        if mutate: mutate(c, v)
        ws.append([c, "LINE " + c] + v)
    b = io.BytesIO(); wb.save(b); return b.getvalue()
d = E.parse_accounts(book())
ok(d["year"] == 2023 and len(d["regions"]) == 17 and d["regions"][1] == {"id": "EL30", "level": 2, "name": "Attiki"}, "año, 17 columnas y codigos NUTS")
ok(d["rows"][0]["c"] == "10" and d["rows"][0]["v"][0] == 100.0, "lineas y total nacional")
def bad_sum(c, v):
    if c == "14": v[names.index("Kriti")] = 5.0
try: E.parse_accounts(book(bad_sum)); ok(False, "regiones que no suman el total se rechazan")
except RuntimeError as e: ok("suman" in str(e), "regiones que no suman el total se rechazan")
wb = openpyxl.Workbook(); ws = wb.active; ws.append(["Year 2023"]); ws.append(["CATEGORIES", "DESCRIPTION", "Ellada", "Region inventada"]); b = io.BytesIO(); wb.save(b)
try: E.parse_accounts(b.getvalue()); ok(False, "region desconocida se rechaza")
except RuntimeError as e: ok("desconocida" in str(e), "region desconocida se rechaza")
# --- cultivos y ganaderia por region y unidad regional ---
def rbook(tot=100.0, regs=((60.0, (30.0, 30.0)), (40.0, (25.0, 15.0)))):
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["Table 2a. Cereals for grain. Areas and production, by Region and Regional Unities, 2024"]); ws.append([]); ws.append([None, "Wheat", None, "Regions"]); ws.append([None, 1, 2])
    ws.append(["Σύνολο Ελλάδας", tot, tot * 3, "Greece Total"])
    for i, (rt, us) in enumerate(regs):
        ws.append(["Περιφέρεια Χ%d" % i, rt, rt * 3, "Region of X%d" % i])
        for j, u in enumerate(us): ws.append(["  Νομός%d_%d." % (i, j), u, u * 3, "  Unit%d_%d" % (i, j)])
    ws.merge_cells("B3:C3"); b = io.BytesIO(); wb.save(b); return b.getvalue()
t = E.rtable(rbook())
ok(t["year"] == 2024 and [r[2] for r in t["rows"]] == [0, 1, 2, 2, 1, 2, 2], "niveles: Grecia, region y unidad (por sangria) y año del titulo")
ok(t["rows"][2][0] == "Νομός0_0", "el punto final del nombre griego se quita")
ok(E.find_col(t["heads"], "Wheat | 1") == 2 and E.find_col(t["heads"], "Wheat | 2") == 3, "cabeceras combinadas se resuelven por su ruta")
try: E.find_col(t["heads"], "Barley | 1"); ok(False, "cabecera inexistente se rechaza")
except RuntimeError: ok(True, "cabecera inexistente se rechaza")
areas = E.aligned({"t2a": t}); v = E.series(t, areas, [2])
ok(len(areas) == 7 and areas[1]["parent"] == "EL" and areas[2]["parent"] == "R01" and v[0] == 100.0, "areas con padre y serie alineada por nombre griego")
ok(E.check_sums("x", areas, v) == [], "regiones suman Grecia y unidades suman su region")
t2 = E.rtable(rbook(tot=100.0, regs=((60.0, (30.0, 10.0)), (40.0, (25.0, 15.0)))))
ok(len(E.check_sums("x", areas, E.series(t2, areas, [2]))) == 1, "unidades que no suman su region se senalan")
t3 = E.rtable(rbook(tot=150.0)); ok(len(E.check_sums("x", areas, E.series(t3, areas, [2]))) == 1, "regiones que no suman Grecia se senalan")
html = '<a href="/x?documentID=7&amp;a=1">02a. Cereals</a> <a href="/x?documentID=8">02b. Pulses</a>'
ok(E.link_table(html, r"^02a\.").endswith("documentID=7&a=1"), "enlace de descarga por el texto del enlace")
try: E.link_table(html, r"^03"); ok(False, "enlace ausente se rechaza")
except RuntimeError: ok(True, "enlace ausente se rechaza")
print("\n%d fallos" % len(fails)); sys.exit(1 if fails else 0)
