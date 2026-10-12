#!/usr/bin/env python3
"""Austria: gasoleo y fioul de calefaccion semanales desde 2005, y coste aproximado del gasoleo para una explotacion con la devolucion
«Agrardieselrückvergütung» -> data/at-diesel.json
1) Precios semanales de Austria del Boletin Petrolero de la UE (Comision Europea, DG ENER): gasoleo con y sin impuestos y fioul de calefaccion
   con impuestos, en EUR por 1.000 l -> EUR/l. Reutilizacion: Decision 2011/833/UE.
2) Devolucion por litro y litros a tanto alzado por hectarea: scripts/ref/at-agrardiesel.json (se actualiza a mano; el tipo de 2026 es
   provisional, «voraussichtlich rund 15,5 Cent/Liter» segun la AMA, hasta que salga el reglamento de alivio del NEHG).
3) Coste aproximado (EUR/l), calculo de Dehesa Index:
   - explotacion en regimen general de IVA (regelbesteuert): precio / 1,20 - devolucion;
   - explotacion acogida al regimen a tanto alzado (pauschaliert, la mayoria en Austria; no recupera el IVA): precio - devolucion.
Uso: python3 scripts/update-at-diesel.py"""
import datetime, io, json, sys, time, urllib.request
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "at-diesel.json"
WOB = "https://energy.ec.europa.eu/document/download/906e60ca-8b6a-44e7-8589-652854d2fd3f_en?filename=Weekly_Oil_Bulletin_Prices_History_maticni_4web.xlsx"
H = {"User-Agent": "Dehesa-Index-data-bot/1.0 (+https://dehesaindex.com)"}
VAT = 0.20
def get(u, t=300):
    last = None
    for i in range(4):
        try:
            with urllib.request.urlopen(urllib.request.Request(u, headers=H), timeout=t) as r: return r.read()
        except Exception as e: last = e; time.sleep(6 * (i + 1))
    raise RuntimeError("%s -> %s" % (u[-90:], last))
def main():
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(get(WOB)), read_only=True, data_only=True)
        def col(ws, name):
            hdr = [str(x or "").strip() for x in next(ws.iter_rows(values_only=True, max_row=1))]; i = hdr.index(name); out = {}
            for r in ws.iter_rows(values_only=True, min_row=2):
                if not isinstance(r[0], datetime.datetime) or i >= len(r): continue
                try: out[r[0].date().isoformat()] = float(r[i]) / 1000
                except (TypeError, ValueError): pass
            return out
        wt = wb["Prices with taxes"]; wo = wb["Prices wo taxes"]
        dt = col(wt, "AT_price_with_tax_diesel"); ht = col(wt, "AT_price_with_tax_heating_oil"); dn = col(wo, "AT_price_wo_tax_diesel")
        rows = [[d, round(dt[d], 4), round(dn[d], 4) if d in dn else None, round(ht[d], 4) if d in ht else None] for d in sorted(dt) if 0.5 < dt[d] < 3.5]
        if len(rows) < 500: raise ValueError("serie corta: %d semanas" % len(rows))
        for a, b in zip(rows, rows[1:]):
            if abs(b[1] - a[1]) > 0.6: raise ValueError("salto raro entre %s y %s" % (a[0], b[0]))
    except Exception as e:
        print("FALLO boletin petrolero:", e); return 0 if OUT.exists() else 1
    ref = json.loads((ROOT / "scripts" / "ref" / "at-agrardiesel.json").read_text(encoding="utf-8"))
    y = max(ref["years"]); R = dict(ref["years"][y], year=int(y)); last = rows[-1]
    doc = {"schemaVersion": 1, "generatedAt": datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"), "asOf": last[0], "unit": "EUR/l", "vat": VAT,
           "source": {"name": "Comisión Europea, Weekly Oil Bulletin (Austria)", "url": "https://energy.ec.europa.eu/data-and-analysis/weekly-oil-bulletin_en", "license": "Reutilización de documentos de la Comisión (Decisión 2011/833/UE)"},
           "fields": ["date", "dieselTax", "dieselNoTax", "heatingOilTax"], "weekly": rows, "refund": R,
           "farmCostNet": round(last[1] / (1 + VAT) - R["eurPerLitre"], 4), "farmCostFlat": round(last[1] - R["eurPerLitre"], 4)}
    OUT.write_text(json.dumps(doc, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print("at-diesel.json: %d semanas %s..%s; gasoleo %.3f EUR/l; devolucion %s (%s%s); coste %.3f sin IVA / %.3f a tanto alzado" % (
        len(rows), rows[0][0], last[0], last[1], R["eurPerLitre"], y, ", provisional" if R.get("provisional") else "", doc["farmCostNet"], doc["farmCostFlat"]))
    return 0
if __name__ == "__main__": sys.exit(main())
