#!/usr/bin/env python3
"""Pruebas del adaptador de Uruguay (scripts/update-uruguay.py) con originales sinteticos en modo local: sin red, sin cifras reales.
Cubre los casos de calidad que obligan al adaptador a NO publicar: licencia distinta de odc-uy, HTTP 200 con HTML, fila URSEA Inicio>Fin, semestre 999 de tierras,
meses INIA incompletos, PD_AD frente a PD_PF y suma de bovinos con bovinos de leche."""
import csv, io, json, subprocess, sys, tempfile
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
fails = []
LAST = ['']
DEPS = [None]
def ok(name, cond):
    if not cond:
        fails.append(name); print('FALLA', name)
        print('::error title=test-uruguay::FALLA %s -- ultimo log: %s' % (name, LAST[0]))   # visible como anotacion en GitHub Actions

def pkg(name, urls, lic='odc-uy'):
    return {'name': name, 'license_id': lic, 'resources': [{'id': 'r%d' % i, 'url': u} for i, u in enumerate(urls)]}

def run(files, pkgs, extra=()):
    d = Path(tempfile.mkdtemp()); out = Path(tempfile.mkdtemp())
    (d / 'ckan-test.json').write_text(json.dumps({'result': {'results': pkgs}}))
    for n, b in files.items(): (d / n).write_bytes(b)
    p = subprocess.run([sys.executable, str(ROOT / 'scripts/update-uruguay.py'), '--local', str(d), '--out', str(out)], capture_output=True, text=True)
    st = out / 'uruguay-stats.json'
    series = {s['id']: s for s in json.loads(st.read_text())['countries']['UY']['series']} if st.exists() else {}
    lg = (out / 'uruguay-log.txt').read_text() if (out / 'uruguay-log.txt').exists() else ''
    DEPS[0] = json.loads((out / 'uruguay-departments.json').read_text()) if (out / 'uruguay-departments.json').exists() else None
    LAST[0] = (lg + p.stdout + p.stderr)[-500:].replace(chr(10), ' | ')
    return p, series, lg + p.stdout + p.stderr

def inia_csv(year, skip_days_in_feb=0):
    import calendar
    b = io.StringIO(); w = csv.writer(b); w.writerow(['fecha', 'temperaturaMaxima', 'temperaturaMinima', 'pluviometro'])
    for m in range(1, 13):
        for d in range(1, calendar.monthrange(year, m)[1] + 1):
            if m == 2 and d <= skip_days_in_feb: continue
            w.writerow(['%d-%02d-%02d' % (year, m, d), 20 + m, 10 + m, 1.0])
    return b.getvalue().encode()

# ---- 1. INIA: meses completos si; febrero con 5 dias menos (<90 %) no
files, pk = {}, []
for code in ('le',):
    urls = []
    for y in (2025, 2026):
        u = 'https://x/ppt_tmax_tmin_%s_%d.csv' % (code, y); urls.append(u)
    pk.append(pkg('inia-precipitacion-temps-extremas-' + code, urls))
    files['x--r0.csv'] = inia_csv(2025, skip_days_in_feb=0)
    files['x--r1.csv'] = inia_csv(2026, skip_days_in_feb=5)
p, s, log = run(files, pk)
ser = s.get('uy-inia-le-tmax')
ok('INIA genera serie', ser is not None)
if ser:
    per = [x[0] for x in ser['points']]
    ok('INIA: mes completo entra', '2026-01' in per)
    ok('INIA: febrero incompleto (23/28 dias) no entra', '2026-02' not in per)
    rain = dict(s['uy-inia-le-rain']['points'])
    ok('INIA: lluvia = suma mensual', abs(rain['2025-01'] - 31.0) < 1e-6)
    ok('INIA: maxima = media mensual', abs(dict(ser['points'])['2025-03'] - 23.0) < 1e-6)

# ---- 2. licencia distinta -> no se publica nada de ese conjunto
p, s, log = run(files, [pkg('inia-precipitacion-temps-extremas-le', ['https://x/ppt_tmax_tmin_le_2025.csv'], lic='cc-by')])
ok('licencia distinta de odc-uy: sin series', 'uy-inia-le-tmax' not in s)
ok('licencia distinta: queda en el log', 'licencia inesperada' in log)

# ---- 3. HTTP 200 con HTML no es dato
p, s, log = run({'x--r0.csv': b'<!DOCTYPE html><html>error</html>'}, [pkg('inia-precipitacion-temps-extremas-le', ['https://x/ppt_tmax_tmin_le_2025.csv'])])
ok('HTML con 200 no genera serie', 'uy-inia-le-tmax' not in s)

# ---- 4. URSEA: Inicio>Fin en cuarentena y sin corregir
rows = ['Inicio,Fin,Gasoil50S_ExP,Gasoil50S_ExP_imp']
import datetime
d0 = datetime.date(2025, 1, 26)
for k in range(8):
    a = datetime.date(2025 + (k + 0) // 12, (k % 12) + 1, 26); b = datetime.date(2025 + (k + 1) // 12, ((k + 1) % 12) + 1, 25)
    rows.append('%s,%s,%d.5,%d.5' % (a, b, 40 + k, 50 + k))
rows.append('2022-12-26,2022-12-25,99.0,99.0')
p, s, log = run({'x--r0.csv': '\n'.join(rows).encode()}, [pkg('ursea-ppi_vs_pe_v2', ['https://x/ppi_da_nuevo.csv'])])
u = s.get('uy-ursea-gasoil50s-ex')
ok('URSEA genera serie', u is not None)
if u:
    ok('URSEA: fila Inicio>Fin descartada', all(x[1] != 99.0 for x in u['points']))
    ok('URSEA: periodo = mes en que termina la ventana', u['points'][0][0] == '2025-02')
    ok('URSEA: descarte anotado en el log', 'posterior a Fin' in log)

# ---- 5. MGAP tierras: Sem 999 fuera, solo departamentos con los dos semestres, USD/ha ponderado
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Datos'
ws.append(['idDepto', 'ano', 'Sem', 'montous_sum', 'suptran_sum', 'cuenta_sum'])
for y in (2022, 2023, 2024):
    ws.append([1, y, 1, 1000.0, 100.0, 2]); ws.append([1, y, 2, 3000.0, 100.0, 2])   # dpto 1: (4000/200 = 20 USD/ha)
    ws.append([2, y, 1, 9000.0, 100.0, 1])                                          # dpto 2 solo semestre 1: fuera
ws.append([3, 2024, 999, 5e9, 1.0, 1])                                              # Sem 999: no se asigna ni contamina
ws.append([1, 2021, 999, 7.0, 1.0, 1]); ws.append([1, 2021, 1, 1000.0, 100.0, 2]); ws.append([1, 2021, 2, 3000.0, 100.0, 2])   # dpto-anio con fila 999: se excluye entero
bio = io.BytesIO(); wb.save(bio)
p, s, log = run({'x--r0.xlsx': bio.getvalue()}, [pkg('mgap-serie-precio-de-la-tierra-compraventas', ['https://x/compraventas-por-semestre-ao-y-departamento.xlsx'])])
l = s.get('uy-mgap-land-sale-usd-ha')
ok('Tierra genera serie', l is not None)
if l:
    ok('Tierra: USD/ha ponderado = suma montos / suma ha (20)', abs(l['latest'] - 20.0) < 1e-6)
    ok('Tierra: Sem 999 no contamina', l['latest'] < 1000)
    ok('Tierra: operaciones sumadas (4)', s['uy-mgap-land-sale-count']['latest'] == 4)
    ok('Tierra: dpto-anio con fila 999 excluido (2021 ausente)', '2021' not in dict(l['points']))
    ok('Tierra: por departamento (id ISO, USD/ha 20)', DEPS[0]['land']['sale']['2023'].get('AR', {}).get('usdHa') == 20.0)

# ---- 5b. Arrendamientos: otra cabecera (DEPCOD, montot_sum), unidad USD/ha/anio
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Datos'
ws.append(['DEPCOD', 'ano', 'Sem', 'montot_sum', 'suptran_sum', 'cuenta_sum'])
for y in (2023, 2024, 2025):
    ws.append([1, y, 1, 5000.0, 100.0, 2]); ws.append([1, y, 2, 5000.0, 100.0, 3])
bio = io.BytesIO(); wb.save(bio)
p, s, log = run({'x--r0.xlsx': bio.getvalue()}, [pkg('mgap-serie-precio-de-la-tierra-arrendamientos', ['https://x/arrendamientos-por-departamento.xlsx'])])
r = s.get('uy-mgap-land-rent-usd-ha')
ok('Alquiler genera serie', r is not None)
if r:
    ok('Alquiler: 10000/200 = 50 USD/ha/anio', abs(r['latest'] - 50.0) < 1e-6 and r['unit'] == 'USD/ha/year')
    ok('Alquiler: contratos sumados (5)', s['uy-mgap-land-rent-count']['latest'] == 5)

# ---- 6. DICOSE: PD_AD, no PD_PF; bovinos y bovinos de leche separados; especies sin control de calidad fuera
def dic(year):
    b = io.StringIO(); w = csv.writer(b, delimiter=';')
    w.writerow(['ns1:Ejercicio', 'ns1:DepartamentoCodigo', 'ns1:EspecieCodigo', 'ns1:AnimalesPorEspecieConsumoPD_AD', 'ns1:AnimalesPorEspecieNacimientosPD_AD', 'ns1:AnimalesPorEspecieMortandadPD_AD', 'ns1:CantidadTenedoresPorEspecie', 'ns1:AnimalesPorEspeciePD_AD', 'ns1:AnimalesPorEspeciePD_PF'])
    for i in range(1200):
        dpt = 1 + (i % 2)   # departamentos 1 (AR) y 2 (CA)
        w.writerow([year, dpt, 1, 0, 2, 0, 1, 10, 12]); w.writerow([year, dpt, 11, 0, 0, 0, 1, 1, 2]); w.writerow([year, dpt, 5, 0, 0, 0, 1, 7, 7])
    return b.getvalue().encode('latin-1')
pk = []; files = {}
for i, y in enumerate((2023, 2024, 2025)):
    n = ['mgap-datos-actualizados-de-la-declaracion-jurada-de-existencias-dicose-snig-2023', 'mgap-datos-preliminares-basados-en-la-declaracion-jurada-de-existencias-dicose-snig-2024', 'mgap-datos-preliminares-basados-en-la-declaracion-jurada-de-existencias-dicose-snig-2025'][i]
    pk.append({'name': n, 'license_id': 'odc-uy', 'resources': [{'id': 'd%d' % i, 'url': 'https://x/datosanimales.csv'}]}); files['x--d%d.csv' % i] = dic(y)
# 2016-2020: los CSV son solo metadatos (diccionario de datos) y los datos estan en el XML de «Datos Animales»
def dic_xml(year):
    parts = ['<?xml version="1.0" encoding="UTF-8"?><root xmlns:ns1="urn:x">']
    for i in range(1200):
        dpt = 1 + (i % 2)
        for sp, ad, nac in ((1, 10, 2), (11, 1, 0), (5, 7, 0)):
            parts.append('<ns1:fila><Ejercicio>%d</Ejercicio><DepartamentoCodigo>%d</DepartamentoCodigo><EspecieCodigo>%d</EspecieCodigo><AnimalesPorEspeciePD_AD>%d</AnimalesPorEspeciePD_AD><AnimalesPorEspecieNacimientosPD_AD>%d</AnimalesPorEspecieNacimientosPD_AD></ns1:fila>' % (year, dpt, sp, ad, nac))
    parts.append('</root>'); return ''.join(parts).encode('utf-8')
old = 'mgap-datos-generados-en-base-a-la-declaracion-jurada-de-existencias-dicose-snig-2016'
pk.insert(0, {'name': old, 'license_id': 'odc-uy', 'resources': [{'id': 'm0', 'url': 'https://x/datosanimales.csv'}, {'id': 'x0', 'url': 'https://x/datosanimales.xml'}]})
files['x--m0.csv'] = 'ATRIBUTO;XML;DESCRIPCION ;ORIGEN;LIMITANTES CONOCIDAS;NOTAS\nEjercicio;ns1:Ejercicio;Ano;DICOSE;;\n'.encode('latin-1'); files['x--x0.xml'] = dic_xml(2016)
p, s, log = run(files, pk)
c = s.get('uy-dicose-cattle-head')
ok('DICOSE genera bovinos', c is not None)
if c:
    ok('DICOSE: bovinos = PD_AD (12000), no PD_PF (14400)', c['latest'] == 12000.0)
    ok('DICOSE: bovinos de leche separado de bovinos', s['uy-dicose-dairy-cattle-head']['latest'] == 1200.0)
    ok('DICOSE: no hay serie de porcinos (sin control de calidad)', not any('pig' in k or 'swine' in k for k in s))
    ok('DICOSE: no hay total bovinos+leche', not any('total' in k for k in s))
    dd = DEPS[0]
    ok('Departamentos: fichero generado con 19 departamentos', bool(dd) and len(dd['departments']) == 19)
    if dd:
        cat = dd['livestock']['cattle']['2025']
        ok('Departamentos: ids = sufijo ISO (AR, CA), no codigo numerico', sorted(cat) == ['AR', 'CA'])
        ok('Departamentos: la suma de los departamentos = total nacional (PD_AD)', sum(cat.values()) == c['latest'])
    ok('DICOSE 2016-2020: sin CSV de datos se lee el XML (2016 entra, bovinos PD_AD = 12000)', dict(c['points']).get('2016') == 12000.0)
    ok('DICOSE 2016: el CSV de metadatos no cuenta como datos', 'ATRIBUTO' not in log or '2016' in str(dict(c['points'])))
    ok('DICOSE: nota avisa de no sumar bovinos y leche', 'No se suma con bovinos' in s['uy-dicose-dairy-cattle-head'].get('note', ''))

# ---- 7. Rejillas INIA: solo columna mensual (sufijo 0), media de celdas, mes con celda vacia fuera, ANR identico a PAD descartado
def grid(year, months, scale=1.0, hole=None):
    cols = ['id', 'Longitud', 'Latitud']
    for m in range(1, months + 1): cols += ['%d_%02d1' % (year, m), '%d_%02d0' % (year, m)]
    b = io.StringIO(); w = csv.writer(b); w.writerow(cols)
    for i in range(160):
        row = [i, -56.0, -33.0]
        for m in range(1, months + 1):
            mv = '' if hole == (m, i) else (10.0 * m + (i % 2)) * scale   # media mensual de las celdas = 10*m + 0.5
            row += [9999.0, mv]   # la decada (sufijo 1) trae un valor que NO debe usarse
        w.writerow(row)
    return b.getvalue().encode()
def gpk(name, ids): return {'name': name, 'license_id': 'odc-uy', 'resources': [{'id': i, 'url': 'https://x/inia_gras_%s%s.csv' % (name.split('-')[1], y)} for i, y in ids]}
files = {'x--p25.csv': grid(2025, 12), 'x--p26.csv': grid(2026, 4, hole=(2, 3)), 'x--a25.csv': grid(2025, 12), 'x--a26.csv': grid(2026, 4, hole=(2, 3))}
pk = [gpk('inia-pad-por-grilla', [('p25', 2025), ('p26', 2026)]), gpk('inia-anr-por-grilla', [('a25', 2025), ('a26', 2026)])]
p, s, log = run(files, pk)
g = s.get('uy-inia-grid-pad-mean')
ok('Rejilla PAD genera serie', g is not None)
if g:
    pts = dict(g['points'])
    ok('Rejilla: media mensual de las celdas (10*m+0.5), no la decada', abs(pts['2025-03'] - 30.5) < 1e-6)
    ok('Rejilla: mes con celda vacia no entra (no se rellena con 0)', '2026-02' not in pts and '2026-03' in pts)
ok('Rejilla ANR identica a PAD: no se publica', 'uy-inia-grid-anr-mean' not in s)
ok('Rejilla ANR identica: anotado en el log', 'identica a PAD' in log)

# ---- 8. Mapa de departamentos: 19 contornos con id ISO, los mismos que usa el adaptador
mp = json.loads((ROOT / 'vendor/uy-departments.js').read_text(encoding='utf-8').split('window.DEHESA_UY_DEPTS = ', 1)[1].rstrip().rstrip(';'))
import importlib.util
sp = importlib.util.spec_from_file_location('uy', ROOT / 'scripts/update-uruguay.py'); mod = importlib.util.module_from_spec(sp); sp.loader.exec_module(mod)
ok('Mapa: 19 departamentos', len(mp['states']) == 19)
ok('Mapa: ids del mapa = sufijos ISO del adaptador', sorted(s['id'] for s in mp['states']) == sorted(i for _, _, i in mod.UY_DEPTS))
ok('Mapa: nombres del mapa = nombres del adaptador', sorted(s['name'] for s in mp['states']) == sorted(n for _, n, _ in mod.UY_DEPTS))

# ---- 9. Inventario de GEI: agricultura sin la categoria Tierra, solo gases directos, factor GWP comprobado
def ghg_xlsx(years, bad_year=None):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Datos'
    ws.append(['Sector', 'Categoría', 'Subcategoría', 'Año', 'Tipo de GEI', 'GEI', 'Gg Gas', 'Gg CO2 - eq (GWP)', 'Gg CO2 - eq (GTP)', 'Notation Key'])
    A3 = '3 - Agricultura, Silvicultura y Otros Usos de la Tierra'
    for y in years:
        k = 30.0 if y == bad_year else 28.0
        ws.append([A3, '3.A - Ganado', '3.A.1 - Fermentación entérica', y, 'Directo', 'CH4', 10.0, str(10.0 * k).replace('.', ','), '1', None])
        ws.append([A3, '3.A - Ganado', '3.A.2 - Gestión del estiércol', y, 'Directo', 'CH4', 1.0, '28,0', '1', None])
        ws.append([A3, '3.C - Fuentes agregadas', '3.C.4 - Emisiones directas de N2O provenientes de suelos manejados', y, 'Directo', 'N2O', 2.0, '530,0', '1', None])
        ws.append([A3, '3.B - Tierra', '3.B.3 - Pastizales', y, 'Directo', 'CO2', -100.0, '-100,0', '1', None])
        ws.append(['1 - Energía', '1.A - Actividades de quema de combustible', '1.A.3 - Transporte', y, 'Directo', 'CO2', 50.0, '50,0', '50,0', None])
        ws.append(['1 - Energía', '1.A - Actividades de quema de combustible', '1.A.3 - Transporte', y, 'Indirecto', 'NOx', 9.0, None, None, None])
    b = io.BytesIO(); wb.save(b); return b.getvalue()
mod.LOG.clear(); d = mod.ghg_parse(ghg_xlsx(range(2010, 2022)))
ok('GEI: genera datos', d is not None)
if d:
    ok('GEI: agricultura = 3.A + 3.C sin 3.B (280+28+530)', abs(d['agriculture'][0] - 838.0) < 1e-6)
    ok('GEI: total nacional = agricultura + energia, sin la Tierra (838+50)', abs(d['national'][0] - 888.0) < 1e-6)
    ok('GEI: la Tierra va aparte y negativa', abs(d['landUseNet'][0] + 100.0) < 1e-6)
    ok('GEI: categorias enteric/manure/soils', d['byCategory']['enteric'][0] == 280.0 and d['byCategory']['manure'][0] == 28.0 and d['byCategory']['soils'][0] == 530.0)
    ok('GEI: gas indirecto (NOx) no suma', 'NOx' not in d['byGas'])
mod.LOG.clear(); d = mod.ghg_parse(ghg_xlsx(range(2008, 2022), bad_year=2015))
ok('GEI: anio con factor de CH4 distinto (30 en vez de 28) se descarta', d is not None and 2015 not in d['years'] and 2014 in d['years'])
ok('GEI: el anio descartado queda anotado', any('factor' in l for l in mod.LOG))
mod.LOG.clear(); ok('GEI: menos de 10 anios no se publica', mod.ghg_parse(ghg_xlsx(range(2018, 2022))) is None)

print('test-uruguay:', 'OK' if not fails else '%d fallos' % len(fails)); sys.exit(1 if fails else 0)
