#!/usr/bin/env python3
"""Proyecciones agrarias a 10 anos de EE. UU. (USDA ERS, Agricultural Baseline Database; Interagency Agricultural Projections Committee).
Busca en la pagina del producto el ZIP mas reciente (la edicion temprana de noviembre sustituye a la de febrero sin tocar codigo), lee el CSV maquina-legible
(year_type,year,value,unit,attribute,commodity,commodity_type,year_of_projection) y escribe data/us-markets/baseline.json: una serie por commodity+atributo,
con los 12 anos que publica (1 de historia, el actual y 10 de proyeccion). Valores y unidades tal como los publica ERS; no se calcula nada. Si no puede
validar lo descargado, NO toca el JSON anterior."""
import csv, io, json, re, sys, urllib.request, zipfile
from datetime import datetime, timezone
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
PAGE = "https://ers.usda.gov/data-products/agricultural-baseline-database"
KNOWN = ["https://ers.usda.gov/media/5014/usda-agricultural-baseline-projections-to-2035.zip?v=91724"]
OUT = ROOT / "data" / "us-markets" / "baseline.json"
LOGF = ROOT / "data" / "ers-baseline-log.txt"
UA = {"User-Agent": "dehesaindex.com data pipeline (+https://dehesaindex.com)"}
LOG = []
def log(*a):
    s = " ".join(str(x) for x in a); LOG.append(s); print(s, flush=True)
def get(u, t=180):
    with urllib.request.urlopen(urllib.request.Request(u, headers=UA), timeout=t) as r: return r.read()
def candidates():
    out = []
    try:
        h = get(PAGE, 60).decode("utf-8", "replace")
        for m in re.finditer(r'href="([^"]+\.zip[^"]*)"', h):
            u = m.group(1).replace("&amp;", "&")
            if u.startswith("/"): u = "https://ers.usda.gov" + u
            if re.search(r"baseline|projection", u, re.I) and u not in out: out.append(u)
        log("pagina:", len(out), "enlaces zip candidatos")
    except Exception as e: log("pagina no legible:", repr(e)[:120])
    for k in KNOWN:
        if k not in out: out.append(k)
    return out
def read_zip(raw):
    z = zipfile.ZipFile(io.BytesIO(raw)); best = None
    for i in z.infolist():
        if not i.filename.lower().endswith(".csv"): continue
        t = z.read(i).decode("utf-8-sig", "replace").splitlines()
        if t and t[0].strip().lower().startswith("year_type,year,value,unit,attribute,commodity,commodity_type,year_of_projection"):
            best = (i.filename, t)
    readme = ""
    for i in z.infolist():
        if i.filename.lower().endswith(".txt"):
            m = re.search(r"Updated\s+([0-9/]+)", z.read(i).decode("utf-8", "replace"))
            if m: readme = m.group(1)
    return best, readme
def parse(t):
    ser = {}; order = []; ed = set()
    for r in csv.DictReader(t):
        try: v = float(r["value"])
        except Exception: continue
        k = (r["commodity_type"], r["commodity"], r["attribute"], r["unit"], r["year_type"]); ed.add(r["year_of_projection"])
        if k not in ser: ser[k] = {}; order.append(k)
        ser[k][r["year"]] = v
    return ser, order, ed

SKIP_XLSX = ("feed grains", "wheat", "soybeans", "rice", "upland cotton")  # ya estan en el CSV
YR = re.compile(r"^(\d{4})(?:\s*\d/)?$|^(\d{4})/(\d{2})(?:\s*\d/)?$|^(\d{4})\s*[\u2013-]\s*(\d{2})$")
def ylab(c):
    if c is None: return None
    if isinstance(c, (int, float)) and not isinstance(c, bool): c = str(int(c)) if float(c).is_integer() else None
    if c is None: return None
    m = YR.match(str(c).strip())
    if not m: return None
    if m.group(1): y = m.group(1); return y if 1990 <= int(y) <= 2045 else None
    if m.group(2): return "%s/%s" % (m.group(2), m.group(3)) if 1990 <= int(m.group(2)) <= 2045 else None
    return "%s\u2013%s" % (m.group(4), m.group(5)) if 1990 <= int(m.group(4)) <= 2045 else None
def isnum(c): return isinstance(c, (int, float)) and not isinstance(c, bool)
def tables(raw, ed_y, url, readme):
    """Lee los libros .xlsx del ZIP que NO estan en el CSV (macro, PIB, poblacion, cambio, superficie, comercio, renta agraria, frutas y hortalizas, azucar, lacteos, ganaderia)
    y los vuelca tal cual (filas, secciones, unidades y años) en data/us-markets/baseline-tables.json. No se calcula nada."""
    import openpyxl
    z = zipfile.ZipFile(io.BytesIO(raw)); out = []
    for i in z.infolist():
        n = i.filename
        if not n.lower().endswith(".xlsx") or n.startswith("__MACOSX"): continue
        if any(k in n.lower() for k in SKIP_XLSX): continue
        wb = openpyxl.load_workbook(io.BytesIO(z.read(i)), read_only=True, data_only=True)
        for sh in wb.sheetnames:
            cur = None; title = ""; sec = ""; unit = ""
            for row in wb[sh].iter_rows(values_only=True):
                row = list(row)
                c0 = next((str(c).strip() for c in row[:3] if c is not None and str(c).strip()), "")
                if re.match(r"^Table\s+\d+", c0): title = c0; cur = None; sec = ""; unit = ""; continue
                lab = [ylab(c) for c in row]
                runs = []; k = 0
                while k < len(row):
                    if lab[k]:
                        e = k
                        while e + 1 < len(row) and lab[e + 1]: e += 1
                        runs.append((k, e)); k = e + 1
                    else: k += 1
                best = max(runs, key=lambda r: r[1] - r[0]) if runs else None
                if best and best[1] - best[0] >= 4 and not any(isnum(c) and not lab[j] for j, c in enumerate(row) if j >= best[0] and j <= best[1]):
                    ys = [lab[j] for j in range(best[0], best[1] + 1)]
                    omit = [str(c).strip() for j, c in enumerate(row) if c is not None and str(c).strip() and not (best[0] <= j <= best[1]) and j > 0 and str(c).strip().lower() not in ("units", "unit", "item")]
                    cur = {"id": "t%d" % (len(out) + 1), "file": n, "sheet": sh, "title": title or sh, "years": ys, "col0": best[0], "rows": [], "notes": [], "omitted": omit[:8], "unit": ""}
                    out.append(cur); sec = ""; unit = ""; continue
                if cur is None: continue
                vals = row[cur["col0"]:cur["col0"] + len(cur["years"])]
                nums = [round(float(c), 4) if isnum(c) else None for c in vals]
                head = [(j, str(c).strip()) for j, c in enumerate(row[:cur["col0"]]) if c is not None and not isnum(c) and str(c).strip()]
                if any(v is not None for v in nums):
                    if not head: continue
                    label = head[0][1]; ind = head[0][0] + (len(str(row[head[0][0]])) - len(str(row[head[0][0]]).lstrip())) // 3
                    u = head[1][1] if len(head) > 1 else unit
                    cur["rows"].append({"section": sec, "label": label, "indent": ind, "unit": u, "values": nums})
                else:
                    txt = [str(c).strip() for c in row if c is not None and str(c).strip()]
                    if not txt: continue
                    inyr = [str(c).strip() for j, c in enumerate(row) if c is not None and str(c).strip() and cur["col0"] <= j < cur["col0"] + len(cur["years"])]
                    inhd = [t for (j, t) in head]
                    if re.match(r"^(\d+/|Note|Source|Totals|Components|Figures|n\.a|NA\b)", txt[0]) or len(txt[0]) > 90: cur["notes"].append(" ".join(txt)[:500])
                    elif inyr: unit = inyr[0]; cur["unit"] = cur["unit"] or unit
                    elif inhd and len(inhd) == 1 and not inyr and row.index(next(c for c in row if c is not None and str(c).strip())) < cur["col0"]: sec = inhd[0]
                    elif len(txt) == 1: unit = txt[0]
    out = [t for t in out if len(t["rows"]) >= 2]
    for n_, t in enumerate(out): t["id"] = "t%d" % (n_ + 1)
    for t in out:
        for k in ("col0",): t.pop(k, None)
        t["notes"] = t["notes"][:12]
    log("TABLAS", len(out))
    for t in out: log("  ", t["id"], t["file"][:50], "|", t["title"][:80], "|", t["years"][0], "a", t["years"][-1], len(t["years"]), "cols", len(t["rows"]), "filas; unidad:", t["unit"][:40], "; omitidas:", t["omitted"][:4], "; notas", len(t["notes"]))
    if len(out) < 6: log("tablas: demasiado pocas, no se escribe"); return
    doc = {"schemaVersion": 1, "generatedAt": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"), "sourceId": "usda_ers", "edition": ed_y, "readmeUpdated": readme, "zipUrl": url.split("?")[0], "pageUrl": PAGE, "tables": out}
    (ROOT / "data" / "us-markets" / "baseline-tables.json").write_text(json.dumps(doc, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
    log("escrito baseline-tables.json", len(out), "tablas")
def main():
    ok = None
    for u in candidates():
        try:
            raw = get(u); best, readme = read_zip(raw)
            if not best: log("sin CSV valido en", u[:100]); continue
            ser, order, ed = parse(best[1])
            if len(ser) < 150: log("pocas series", len(ser), u[:100]); continue
            ed_y = max(ed); log("candidato", u[:110], "edicion", ed_y, "series", len(ser), "readme", readme)
            if ok is None or (ed_y, readme[-2:], len(raw)) > (ok[0], ok[1], ok[2]): ok = (ed_y, readme[-2:], len(raw), u, best[0], ser, order, readme, raw)
        except Exception as e: log("error", u[:100], repr(e)[:150])
    if not ok: log("sin edicion valida: se conserva el JSON anterior"); return 1
    ed_y, _, _, url, fn, ser, order, readme, raw = ok
    yrs = sorted({y for s in ser.values() for y in s})
    out = {"schemaVersion": 1, "generatedAt": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"), "sourceId": "usda_ers", "edition": ed_y,
           "readmeUpdated": readme, "file": fn, "zipUrl": url.split("?")[0], "pageUrl": PAGE,
           "series": [{"type": k[0], "commodity": k[1], "attribute": k[2], "unit": k[3], "yearType": k[4], "years": sorted(ser[k]), "values": [ser[k][y] for y in sorted(ser[k])]} for k in order]}
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(out, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
    log("escrito", OUT.name, len(out["series"]), "series; edicion", ed_y, "; anos", yrs[0], "a", yrs[-1], "; origen", url[:110])
    try: tables(raw, ed_y, url, readme)
    except Exception as e: log("tablas xlsx: error", repr(e)[:200])
    return 0
if __name__ == "__main__":
    try: rc = main()
    except Exception as e: log("error", repr(e)[:300]); rc = 1
    LOGF.write_text("\n".join(LOG) + "\n", encoding="utf-8"); sys.exit(rc)
