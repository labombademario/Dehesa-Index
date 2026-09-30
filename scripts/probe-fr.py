import time,urllib.request,urllib.parse,re,http.cookiejar,json,io
cj=http.cookiejar.CookieJar()
op=urllib.request.build_opener(urllib.request.HTTPCookieProcessor(cj))
def get(u,n=8000000):
    last=''
    for i in range(3):
        try:
            r=urllib.request.Request(u,headers={'User-Agent':'Mozilla/5.0 (X11; Linux x86_64) DehesaIndex','Accept':'*/*','Accept-Language':'fr,en'})
            with op.open(r,timeout=120) as x:
                b=x.read(n); ct=x.headers.get('content-type')
            return b,ct,x.geturl()
        except Exception as e: last='ERR %s'%e; time.sleep(5)
    return last.encode(),None,None
out=[]
def P(*a): out.append(' '.join(str(x) for x in a))
# 1) el xlsx de precios pagados a productores
g=json.loads(get('https://www.data.gouv.fr/api/1/datasets/historique-des-prix-moyens-mensuels-et-trimestriels-payes-aux-producteurs-depuis-2005-cereales-et-oleoproteagineux/')[0])
for r in g['resources']: P('RES',r['title'][:80],r['url'])
url=g['resources'][0]['url']
b,ct,fu=get(url); P('DL',ct,len(b),fu,b[:8])
try:
    import openpyxl
    wb=openpyxl.load_workbook(io.BytesIO(b),data_only=True)
    for ws in wb.worksheets:
        P('SHEET',ws.title,ws.max_row,ws.max_column)
        for row in list(ws.iter_rows(values_only=True))[:12]: P('  ',[c for c in row[:10]])
except Exception as e: P('XLSX ERR',e)
# 2) pagina menu: buscar fileurl/xlsx en html
for u in ['https://visionet.franceagrimer.fr/Pages/SeriesChronologiques.aspx?menuurl=SeriesChronologiques/productions%20vegetales/grandes%20cultures/cotations']:
    b,ct,fu=get(u); s=b.decode('utf-8','replace'); P('PAGE',len(s),fu)
    for m in list(dict.fromkeys(re.findall(r'[^"\'<>\s]*(?:fileurl|\.xlsx|\.csv)[^"\'<>\s]*',s,flags=re.I)))[:40]: P(' F',urllib.parse.unquote(m)[:250])
    P(' FORMS',re.findall(r'__VIEWSTATE[^>]{0,80}',s)[:1],re.findall(r'__doPostBack\([^)]*\)',s)[:15])
    i=s.find('cotations'); P(' CTX',re.sub(r'\s+',' ',s[i-200:i+600]))
open('data/probe/fr4.txt','w').write('\n'.join(out))
