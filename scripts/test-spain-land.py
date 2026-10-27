#!/usr/bin/env python3
"""Prueba sin red de update-spain-land.py con extractos reales de los libros del MAPA (scripts/fixtures/spain-land: precios de la tierra 2024 sin las hojas de ponderaciones y sin imagenes,
canones 2024, y las listas reales de enlaces de las dos paginas). Comprueba cifras concretas, la acumulacion de años, y que se rechaza lo corrupto."""
import copy, importlib.util, json, sys
from pathlib import Path
import openpyxl
ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("sl", ROOT / "scripts" / "update-spain-land.py"); sl = importlib.util.module_from_spec(spec); spec.loader.exec_module(sl)
FX = ROOT / "scripts" / "fixtures" / "spain-land"
fails = []
def check(c, m):
    if not c: fails.append(m)
def wbs(): return openpyxl.load_workbook(FX / "precios-tierra-2024.xlsx", data_only=True), openpyxl.load_workbook(FX / "canones-2024.xlsx", data_only=True)
def build(pw, rw, prev=None, defin=True): return sl.build(pw, rw, prev, {"price": {"name": "p"}, "rent": {"name": "r"}}, "2026-10-13T00:00:00Z", defin)
def rejects(label, mutate, expect):
    pw, rw = wbs()
    try: mutate(pw, rw); build(pw, rw); check(False, "debia rechazar: " + label)
    except ValueError as e: check(expect in str(e), "%s: mensaje %r no contiene %r" % (label, str(e)[:160], expect))
pw, rw = wbs(); d = build(pw, rw)
P, R = d["price"]["ccaa"], d["rent"]["ccaa"]
check(d["years"] == [2020, 2021, 2022, 2023, 2024], "años")
check(P["ES"]["PT"] == [9007.34, 9258.24, 9588.3, 9966.7, 10247.82], "precio general nacional %s" % P["ES"]["PT"])
check(P["CB"]["PT"][-1] == 5843.71 and P["CN"]["PT"][-1] == 148415.16, "Cantabria / Canarias")
check(P["ES"]["TAS"][0] == 6117.28 and P["ES"]["TAR"][-1] == 17212.2 and P["ES"]["PST"][-1] == 3279.29, "tipos nacionales")
check("ARZ" in P["EX"] and "ARZ" not in P["GA"], "una comunidad sin esa tierra no tiene serie (hueco, no cero)")
check(len(P) == 18 and len(P["ES"]) == 19, "comunidades %d, tipos nacionales %d" % (len(P), len(P["ES"])))
pv = d["price"]["prov"]; check(pv["year"] == 2024 and len(pv["values"]) == 50 and pv["names"]["02"] == "Albacete", "provincias")
check(pv["values"]["02"]["TAR"] == 14081.38 and "ARZ" not in pv["values"]["02"], "Albacete TAR; sin arrozal")
check(pv["values"]["03"].get("SBT") is None and "SBT" not in pv["values"]["03"], "«..» de Alicante es hueco")
check(R["ES"]["GEN"] == [151.26, 154.79, 158.87, 159.71, 166.75] and R["ES"]["TAR"][-1] == 587.79, "canones nacionales %s" % R["ES"]["GEN"])
check(d["rent"]["other"]["year"] == 2024 and d["rent"]["other"]["values"]["CT"]["ARZ"] == 1200.0, "otros cultivos")
check(set(d["rentTypes"]) == set(sl.RENT_TYPES) and d["status"]["2024"] == "definitive", "tipos de canon y estado")
check(len(json.dumps(d)) < 40000, "tamaño del fichero %d" % len(json.dumps(d)))
# acumulacion: un libro anterior con 2019 conserva 2019; si el MAPA revisa un año, gana el nuevo y se registra
prev = copy.deepcopy(d); prev["years"] = [2019] + prev["years"]
for ts in prev["price"]["ccaa"].values():
    for t, a in ts.items(): ts[t] = [a[0] * 0.97 if a[0] else None] + a
for ts in prev["rent"]["ccaa"].values():
    for t, a in ts.items(): ts[t] = [None] + a
prev["price"]["ccaa"]["AN"]["PT"][3] = 1.0   # 2022: valor antiguo distinto
sl.LOG.clear(); pw, rw = wbs(); d2 = build(pw, rw, prev)
check(d2["years"] == [2019, 2020, 2021, 2022, 2023, 2024], "acumula 2019: %s" % d2["years"])
check(d2["price"]["ccaa"]["ES"]["PT"][0] == round(9007.34 * 0.97, 2), "2019 conservado")
check(d2["price"]["ccaa"]["AN"]["PT"][3] == 14999.0 or d2["price"]["ccaa"]["AN"]["PT"][3] > 1000, "gana el libro nuevo en 2022")
check(any("revision del MAPA" in l for l in sl.LOG), "revision registrada")
pw, rw = wbs(); d3 = build(pw, rw, d2)
check(d3["price"]["ccaa"] == d2["price"]["ccaa"] and d3["rent"]["ccaa"] == d2["rent"]["ccaa"], "reejecutar con el mismo libro no cambia nada")
# enlaces reales de las paginas
L = json.loads((FX / "pages-links.json").read_text("utf-8"))
html = lambda k: "".join('<a href="%s">%s</a>' % (u, t) for u, t in L[k])
u, y, df, t = sl.pick_price(html("tierra-precios")); check(y == 2024 and df and u.endswith("resultados-2024_definitivos.xlsx"), "enlace de precios: %s %s" % (y, u))
u, y, t = sl.pick_rent(html("canones-anuales-arrendamientos-rusticos")); check(y == 2024 and u.endswith("2024-base-2020-2.xlsx"), "enlace de canones: %s %s" % (y, u))
h2 = html("tierra-precios") + '<a href="/x/resultados-2025.xlsx">Precios medios de las tierras de uso agrario. Resultados provisionales 2025</a>'
check(sl.pick_price(h2)[1:3] == (2025, False), "un libro provisional de 2025 gana a uno definitivo de 2024")
try: sl.pick_rent(html("canones-anuales-arrendamientos-rusticos") + '<a href="/y.xlsx">Tablas Cánones de arrendamiento 2025 Base 2025</a>'); check(False, "debia parar si cambia la base")
except RuntimeError as e: check("base 2025" in str(e), "mensaje de cambio de base")
try: sl.pick_price("<p>sin enlaces</p>"); check(False, "debia fallar sin enlaces")
except RuntimeError: pass
# rechazos
def rename_ccaa(pw, rw): pw["Cuadro_1_PT"]["B6"] = "Atlantis"
rejects("comunidad desconocida", rename_ccaa, "comunidad desconocida")
def bad_total(pw, rw): pw["Cuadro_1_PT"]["I23"] = pw["Cuadro_1_PT"]["I23"].value * 1.1
rejects("total nacional que no cuadra", bad_total, "Total Nacional")
def neg(pw, rw): pw["Cuadro_1_PT"]["I21"] = -5
rejects("precio negativo", neg, "no positivo")
def text_cell(pw, rw): pw["Cuadro_1_PT"]["H10"] = "n/d"
rejects("texto en una celda numerica", text_cell, "celda no numerica")
def drop_prov(pw, rw): pw["Cuadro 20_PT_2024_provinciales"]["A20"] = None
rejects("provincia que falta", drop_prov, "50")
def no_header(pw, rw): pw["Cuadro_3_CBP"]["B4"] = "Comunidades"
rejects("cabecera sin CCAA", no_header, "cabecera")
def drop_sheet(pw, rw): del pw["Cuadro_9_HUE"]
rejects("hoja de tipo que falta", drop_sheet, "HUE")
def rent_total(pw, rw): rw["Cuadro 1 "]["H22"] = 999
rejects("canon general que no cuadra", rent_total, "canon Total Nacional")
def rent_missing(pw, rw): rw["Cuadro 1 "]["A5"] = "Narnia"
rejects("canon: comunidad desconocida", rent_missing, "comunidad desconocida")
# libro de 2025: una columna «CCAA» delante (CCAA | Código | Provincia | tipos); las columnas se buscan por cabecera
pw, rw = wbs(); ws = pw["Cuadro 20_PT_2024_provinciales"]; ws.insert_cols(1)
for row in ws.iter_rows(min_row=1, max_row=8):
    if any(isinstance(c.value, str) and c.value.strip() == "Provincia" for c in row): ws.cell(row[0].row, 1).value = "CCAA"
d2 = build(pw, rw); check(d2["price"]["prov"]["names"]["02"] == "Albacete" and len(d2["price"]["prov"]["values"]) == 50 and d2["price"]["prov"]["values"]["02"]["TAR"] == 14081.38, "provincias con la columna CCAA delante")
print("FALLOS:" if fails else "test-spain-land OK", *fails, sep="\n" if fails else " ")
sys.exit(1 if fails else 0)
