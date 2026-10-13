#!/usr/bin/env python3
"""Portugal: gasoleo semanal (Boletin Petrolero de la UE) y precio de compra del gasoleo por las explotaciones agrarias (Eurostat) -> data/pt-diesel.json
Portugal tiene gasoleo agricola con impuesto reducido («gasoleo colorido e marcado») pero su precio no se publica en abierto con licencia que permita reutilizarlo:
la DGEG (Direcao-Geral de Energia e Geologia) lo tiene en su portal de precios, cuyos terminos de uso prohiben el uso comercial, asi que Dehesa Index no lo usa
ni lo estima. Hay dos series oficiales:
1) Boletin Petrolero de la UE (Comision Europea, DG ENER): gasoleo de automocion con y sin impuestos y gasoleo de calefaccion de Portugal, semanal, EUR/1000 l -> EUR/l.
   Reutilizacion: Decision 2011/833/UE.
2) Precio de compra del gasoleo por las explotaciones agrarias, EUR por 100 l, anual: Eurostat apri_ap_ina (producto 20232000, 2000-ultimo ano), sin IVA
   (metodologia de Eurostat para los precios de insumos). Licencia CC BY 4.0.
Si una parte falla se conserva la anterior. Uso: python3 scripts/update-pt-diesel.py"""
import datetime, io, json, sys, time, urllib.request
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "pt-diesel.json"
WOB = "https://energy.ec.europa.eu/document/download/906e60ca-8b6a-44e7-8589-652854d2fd3f_en?filename=Weekly_Oil_Bulletin_Prices_History_maticni_4web.xlsx"
ES = "https://ec.europa.eu/eurostat/api/dissemination/statistics/1.0/data/apri_ap_ina?format=JSON&lang=EN&geo=PT&currency=EUR&prod_inp=20232000"
H = {"User-Agent": "Dehesa-Index-data-bot/1.0 (+https://dehesaindex.com)", "Accept": "application/json, */*"}
def get(u, t=120, tries=4):
    last = None
    for i in range(tries):
        try:
            with urllib.request.urlopen(urllib.request.Request(u, headers=H), timeout=t) as r: return r.read()
        except Exception as e: last = e; time.sleep(4 * (i + 1))
    raise RuntimeError("%s -> %s" % (u[-90:], last))
def weekly():
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(get(WOB, 300)), read_only=True, data_only=True)
    def col(ws, name):
        hdr = [str(x or "").strip() for x in next(ws.iter_rows(values_only=True, max_row=1))]; i = hdr.index(name); out = {}
        for r in ws.iter_rows(values_only=True, min_row=2):
            if not isinstance(r[0], datetime.datetime) or i >= len(r): continue
            try: out[r[0].date().isoformat()] = float(r[i]) / 1000
            except (TypeError, ValueError): pass
        return out
    wt = wb["Prices with taxes"]; wo = wb["Prices wo taxes"]
    dt = col(wt, "PT_price_with_tax_diesel"); dn = col(wo, "PT_price_wo_tax_diesel")
    try: ht = col(wt, "PT_price_with_tax_heating_oil")
    except ValueError: ht = {}
    out = [[d, round(dt[d], 4), round(dn[d], 4) if d in dn else None, round(ht[d], 4) if d in ht else None] for d in sorted(dt) if 0.3 < dt[d] < 3.5]
    if len(out) < 500: raise ValueError("serie semanal corta: %d" % len(out))
    for a, b in zip(out, out[1:]):
        if abs(b[1] - a[1]) > 0.6: raise ValueError("salto raro entre %s y %s" % (a[0], b[0]))
    return out
def farm():
    d = json.loads(get(ES)); T = d["dimension"]["time"]["category"]["index"]; v = d["value"]; out = {}
    if isinstance(T, list): T = {k: i for i, k in enumerate(T)}
    for y, i in T.items():
        if str(i) in v and v[str(i)] is not None: out[int(y)] = round(v[str(i)] / 100, 4)
    rows = [[y, out[y]] for y in sorted(out)]
    if len(rows) < 20 or not all(0.2 < x[1] < 2.5 for x in rows): raise ValueError("precio agrario anual raro: %s" % rows[-3:])
    return rows
def main():
    try: old = json.loads(OUT.read_text())
    except Exception: old = {}
    doc = {"schemaVersion": 1, "generatedAt": datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"), "unit": "EUR/l",
           "source": {"name": "Comisión Europea, Weekly Oil Bulletin; Eurostat apri_ap_ina", "url": "https://energy.ec.europa.eu/data-and-analysis/weekly-oil-bulletin_en", "wobUrl": "https://energy.ec.europa.eu/data-and-analysis/weekly-oil-bulletin_en",
                      "farmUrl": "https://ec.europa.eu/eurostat/databrowser/view/apri_ap_ina/default/table", "license": "CC BY 4.0 (Eurostat); Decisión 2011/833/UE (Comisión)"},
           "weeklyFields": ["date", "dieselTax", "dieselNoTax", "heatingOilTax"], "farmFields": ["year", "dieselNoVat"], "vat": 0.23}
    for key, fn in (("weekly", weekly), ("farm", farm)):
        try: doc[key] = fn()
        except Exception as e:
            print("AVISO %s:" % key, e)
            if old.get(key): doc[key] = old[key]
    if not all(k in doc for k in ("weekly", "farm")): print("FALLO: falta una parte"); return 0 if OUT.exists() else 1
    doc["asOf"] = doc["weekly"][-1][0]
    OUT.write_text(json.dumps(doc, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print("pt-diesel.json: %d semanas (ultima %s, %.3f EUR/l); gasoleo agrario %d anos, ultimo %s: %.3f EUR/l" % (len(doc["weekly"]), doc["weekly"][-1][0], doc["weekly"][-1][1], len(doc["farm"]), doc["farm"][-1][0], doc["farm"][-1][1]))
    return 0
if __name__ == "__main__": sys.exit(main())
