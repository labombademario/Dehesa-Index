#!/usr/bin/env python3
"""Precios oficiales de la madera por surtido (€/m3 o moneda local) en seis paises -> data/forest-timber-prices.json
  SI  Eslovenia   SURS (SiStat, tabla 1656401S): compra de madera en rollo, precio medio por surtido. Mensual desde 2006.
  CZ  Chequia     CSU (conjunto abierto CEN0205C): precio medio de la madera en rollo por clase de calidad y especie, vendedor propietario (CZK/m3). Trimestral desde 2010.
  (BG Bulgaria    NSI, Forestry1.1_EUR.xlsx: el adaptador esta escrito pero DESACTIVADO hasta que se decida su licencia, ver mas abajo.)
  LV  Letonia     CSP (MEI020): precios medios de compra de madera sin IVA por surtido. Semestral desde 2006.
  FI  Finlandia   Luke (metryv/0100): precios de la madera en pie en las ventas (semanal desde 2020), pais y 8 regiones de precios.
  PL  Polonia     GUS (comunicados en Monitor Polski): precio medio de venta de la madera de los Bosques Estatales, base del impuesto forestal (anual, PLN/m3).
Son precios distintos entre paises: los surtidos, las clases de calidad y el punto de la cadena (madera en pie, en cargadero, comprada por la industria) no
coinciden; la bandera `bench` marca la serie que Dehesa Index usa para comparar un mismo tipo de madera, pero NO es una cifra armonizada.
Si un pais falla se conserva su serie anterior; si fallan mas de dos, el fichero no se toca y el script sale con error.
Licencias (data/license-registry.json): SURS libre con cita de la fuente, CSU sin derechos de autor, CSP y Luke CC BY 4.0, GUS comunicado oficial.
Bulgaria: la licencia del NSI permite el uso comercial citando la fuente pero prohibe «difundir obras derivadas o colectivas»; esta PENDIENTE y no se publica.
Uso: python3 scripts/update-forest-timber-prices.py"""
import csv, datetime, html, io, json, re, sys, time, urllib.request
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "forest-timber-prices.json"
UA = {"User-Agent": "Dehesa-Index-data-bot/1.0 (+https://dehesaindex.com)"}
MAX_FAIL = 2
def http(url, data=None, headers=None, tries=4, timeout=120):
    last = None
    h = dict(UA); h.update(headers or {})
    for i in range(tries):
        try:
            with urllib.request.urlopen(urllib.request.Request(url, data=data, headers=h), timeout=timeout) as r: return r.read()
        except Exception as e: last = e; time.sleep(4 * (i + 1))
    raise RuntimeError("%s: %s" % (url[:90], repr(last)[:160]))
def jpost(url, query):
    return json.loads(http(url, json.dumps(query).encode(), {"Content-Type": "application/json"}).decode("utf-8-sig"))
def jget(url): return json.loads(http(url).decode("utf-8-sig"))
def jsonstat_rows(d, dims):
    """json-stat2 con 'dims' como (unicas) dimensiones de interes -> {tupla de codigos: valor} (el resto de dimensiones debe tener tamaño 1)."""
    ids = d["id"]; size = d["size"]; vals = d["value"]
    codes = {i: list(d["dimension"][i]["category"]["index"]) if isinstance(d["dimension"][i]["category"]["index"], list) else sorted(d["dimension"][i]["category"]["index"], key=d["dimension"][i]["category"]["index"].get) for i in ids}
    out = {}; n = len(vals)
    strides = []; s = 1
    for z in reversed(size): strides.insert(0, s); s *= z
    for flat in range(n):
        v = vals[flat]
        if v is None: continue
        key = tuple(codes[i][(flat // strides[k]) % size[k]] for k, i in enumerate(ids) if i in dims)
        out[key] = v
    return out, {i: d["dimension"][i]["category"].get("label", {}) for i in ids}
# ------------------------------------------------------------------ Eslovenia
SI_SERIES = [  # codigo SURS, id, grupo, bench
    ("01", "all", "all", False), ("02", "sawlog", "sawlog", False), ("03", "sawlog-conifer", "sawlog_conifer", True), ("04", "sawlog-oak", "sawlog_broadleaf", False),
    ("05", "sawlog-beech", "sawlog_broadleaf", True), ("06", "sawlog-other-broadleaf", "sawlog_broadleaf", False), ("07", "pulpwood", "pulpwood", False),
    ("08", "pulpwood-conifer", "pulpwood_conifer", True), ("09", "pulpwood-broadleaf", "pulpwood_broadleaf", True), ("10", "other-industrial", "other", False), ("13", "fuelwood", "fuelwood", True)]
def adapter_si():
    url = "https://pxweb.stat.si/SiStatData/api/v1/en/Data/1656401S.px"
    m = jget(url); meas = next(v for v in m["variables"] if v["code"] == "MERITVE")
    price = next(c for c, t in zip(meas["values"], meas["valueTexts"]) if t.lower().startswith("price"))
    d = jpost(url, {"query": [{"code": "LES", "selection": {"filter": "all", "values": ["*"]}}, {"code": "MESEC", "selection": {"filter": "all", "values": ["*"]}},
                              {"code": "MERITVE", "selection": {"filter": "item", "values": [price]}}], "response": {"format": "json-stat2"}})
    rows, labels = jsonstat_rows(d, ("LES", "MESEC"))
    months = list(d["dimension"]["MESEC"]["category"]["index"])
    periods = ["%s-%s" % (x[:4], x[5:7]) for x in months]
    series = []
    for code, sid, grp, bench in SI_SERIES:
        vals = [rows.get((code, mm)) for mm in months]
        if sum(v is not None for v in vals) < 24: continue
        series.append({"id": "si-" + sid, "label": labels["LES"][code], "group": grp, "bench": bench, "values": [None if v is None else round(float(v), 2) for v in vals]})
    return {"freq": "M", "currency": "EUR", "unit": "EUR/m3", "vat": "n/d", "priceType": "purchase",
            "periods": periods, "series": series,
            "source": {"name": "Statistical Office of the Republic of Slovenia (SURS), SiStat 1656401S: Purchase of roundwood", "url": "https://pxweb.stat.si/SiStatData/pxweb/en/Data/-/1656401S.px", "license": "SURS: reuse with source credit"},
            "note": "Precio medio de compra de madera en rollo por las empresas (valor / cantidad). Los dos ultimos meses son provisionales."}
# ------------------------------------------------------------------ Chequia
CZ_SERIES = [  # codigo de reprezentant, id, grupo, bench
    ("11021", "spruce-sawlog-II", "sawlog_conifer", False), ("11031", "spruce-sawlog-IIIAB", "sawlog_conifer", True), ("11032", "pine-sawlog-IIIAB", "sawlog_conifer", False),
    ("12031", "oak-sawlog-IIIAB", "sawlog_broadleaf", False), ("12032", "beech-sawlog-IIIAB", "sawlog_broadleaf", True),
    ("11071", "spruce-pulpwood", "pulpwood_conifer", True), ("12062", "beech-pulpwood", "pulpwood_broadleaf", True),
    ("11080", "fuelwood-conifer", "fuelwood", True), ("12070", "fuelwood-broadleaf", "fuelwood", False)]
CZ_LABEL = {"11021": "Sawlogs class II, spruce", "11031": "Sawlogs class III.A/B, spruce", "11032": "Sawlogs class III.A/B, pine", "12031": "Sawlogs class III.A/B, oak",
            "12032": "Sawlogs class III.A/B, beech", "11071": "Pulpwood (class V), spruce", "12062": "Pulpwood (class V), beech", "11080": "Fuelwood (class VI), coniferous", "12070": "Fuelwood (class VI), broadleaf"}
def adapter_cz():
    raw = http("https://data.csu.gov.cz/opendata/sady/CEN0205C/distribuce/csv").decode("utf-8-sig")
    got = {}; per = set()
    for r in csv.DictReader(io.StringIO(raw)):
        if r["PRODEJCEDREV"] != "1": continue                       # 1 = vendedor propietario del bosque
        p = r["CASQKQR"]
        if not re.fullmatch(r"\d{4}-Q[1-4]", p): continue          # solo trimestres sueltos (sin acumulados 'K' ni años)
        code = r["REPRCENLES"]
        try: got[(code, p)] = float(r["Hodnota"].replace(",", "."))
        except ValueError: continue
        per.add(p)
    periods = sorted(per)
    series = []
    for code, sid, grp, bench in CZ_SERIES:
        vals = [got.get((code, p)) for p in periods]
        if sum(v is not None for v in vals) < 12: continue
        series.append({"id": "cz-" + sid, "label": CZ_LABEL[code], "group": grp, "bench": bench, "values": [None if v is None else round(v, 2) for v in vals]})
    return {"freq": "Q", "currency": "CZK", "unit": "CZK/m3", "vat": "n/d", "priceType": "sale-by-owner",
            "periods": periods, "series": series,
            "source": {"name": "Czech Statistical Office (CSU), open dataset CEN0205C: Average prices of domestic roundwood", "url": "https://data.csu.gov.cz/opendata/sady/CEN0205C", "license": "CSU open data (data.gov.cz): no copyright, no database right"},
            "note": "Precio medio de venta de madera en rollo cuando el vendedor es el propietario del bosque, por clase de calidad y especie."}
# ------------------------------------------------------------------ Bulgaria (DESACTIVADO: licencia pendiente, ver la cabecera)
BG_ROWS = [  # fila, id, etiqueta, grupo, bench
    (4, "conifer-sawlog-Ia", "Conifer sawlogs, class Ia (>50 cm)", "sawlog_conifer", False), (5, "conifer-sawlog-I", "Conifer sawlogs, class I (30-49 cm)", "sawlog_conifer", False),
    (6, "conifer-sawlog-II", "Conifer sawlogs, class II (18-29 cm)", "sawlog_conifer", True), (7, "conifer-industrial", "Conifer industrial wood (large)", "pulpwood_conifer", True),
    (19, "conifer-firewood", "Conifer firewood and brash", "fuelwood", True), (20, "conifer-stumpage-individuals", "Conifer standing timber sold to individuals", "stumpage", False),
    (24, "broadleaf-sawlog-Ia", "Broadleaf sawlogs, class Ia (>50 cm)", "sawlog_broadleaf", False), (26, "broadleaf-sawlog-II", "Broadleaf sawlogs, class II (18-29 cm)", "sawlog_broadleaf", True),
    (28, "broadleaf-veneer", "Broadleaf veneer logs", "other", False), (30, "broadleaf-industrial", "Broadleaf industrial wood (large)", "pulpwood_broadleaf", True),
    (44, "broadleaf-hard-firewood", "Hardwood firewood", "fuelwood", False), (47, "broadleaf-stumpage-individuals", "Broadleaf standing timber sold to individuals", "stumpage", False),
    (49, "marked-standing-total", "Marked standing timber (all species)", "stumpage", False)]
def adapter_bg():
    import openpyxl
    raw = http("https://www.nsi.bg/sites/default/files/files/data/timeseries/Forestry1.1_EUR.xlsx")
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    got = {}; per = set()
    for sh in wb.sheetnames:
        if not re.fullmatch(r"20\d\d", sh): continue
        for r in wb[sh].iter_rows(values_only=True):
            if not isinstance(r[0], int) or not isinstance(r[4], (int, float, type(None), str)): continue
            for q in range(4):
                v = r[4 + q]
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    p = "%s-Q%d" % (sh, q + 1); got[(r[0], p)] = float(v); per.add(p)
    periods = sorted(per)
    series = []
    for row, sid, label, grp, bench in BG_ROWS:
        vals = [got.get((row, p)) for p in periods]
        if sum(v is not None for v in vals) < 6: continue
        series.append({"id": "bg-" + sid, "label": label, "group": grp, "bench": bench, "values": [None if v is None else round(v, 2) for v in vals]})
    return {"freq": "Q", "currency": "EUR", "unit": "EUR/m3", "vat": "n/d", "priceType": "producer",
            "periods": periods, "series": series,
            "source": {"name": "National Statistical Institute of Bulgaria (NSI): Producer prices in forestry", "url": "https://www.nsi.bg/statistical-data/6/12", "license": "NSI: reuse with source credit"},
            "note": "Precios de productor de la silvicultura en leva convertidos a euros por el INE bulgaro (series anteriores a 2026 en BGN a tipo fijo). Los surtidos del punto de venta (cargadero, en pie) siguen las clases de la ley forestal bulgara."}
# ------------------------------------------------------------------ Letonia
def adapter_lv():
    url = "https://data.stat.gov.lv/api/v1/en/OSP_PUB/NOZ/ME/MEI/MEI020"
    m = jget(url); ind = next(v for v in m["variables"] if v["code"] == "INDICATOR")
    d = jpost(url, {"query": [{"code": "INDICATOR", "selection": {"filter": "all", "values": ["*"]}}, {"code": "TIME", "selection": {"filter": "all", "values": ["*"]}}], "response": {"format": "json-stat2"}})
    rows, labels = jsonstat_rows(d, ("INDICATOR", "TIME"))
    times = list(d["dimension"]["TIME"]["category"]["index"])
    periods = [t.replace("H", "-H") for t in times]
    series = []
    bench = {"SPR_SAW_18-26": ("sawlog_conifer", True), "BIR_SAW_18-24": ("sawlog_broadleaf", True)}
    for code, text in zip(ind["values"], ind["valueTexts"]):
        vals = [rows.get((code, t)) for t in times]
        if sum(v is not None for v in vals) < 6: continue
        t = text.lower()
        grp = "sawlog_conifer" if ("pine" in t or "spruce" in t) and "saw" in t else "sawlog_broadleaf" if "saw" in t else "other"
        g, b = bench.get(code, (grp, False))
        series.append({"id": "lv-" + re.sub(r"[^a-z0-9]+", "-", code.lower()).strip("-"), "label": text, "group": g, "bench": b, "values": [None if v is None else round(float(v), 2) for v in vals]})
    return {"freq": "H", "currency": "EUR", "unit": "EUR/m3", "vat": "excl", "priceType": "purchase",
            "periods": periods, "series": series,
            "source": {"name": "Central Statistical Bureau of Latvia (CSP): Average purchase prices of wood, MEI020", "url": "https://data.stat.gov.lv/pxweb/en/OSP_PUB/START__NOZ__ME__MEI/MEI020/", "license": "CSP open data"},
            "note": "Precio medio de compra de madera (sin IVA) por surtido y diametro; semestral."}
# ------------------------------------------------------------------ Finlandia
FI_ASSORT = {"Pine logs": ("pine-logs", "sawlog_conifer", False), "Spruce logs": ("spruce-logs", "sawlog_conifer", True), "Birch logs": ("birch-logs", "sawlog_broadleaf", True),
             "Pine pulpwood": ("pine-pulpwood", "pulpwood_conifer", False), "Spruce pulpwood": ("spruce-pulpwood", "pulpwood_conifer", True), "Birch pulpwood": ("birch-pulpwood", "pulpwood_broadleaf", True)}
FI_REGION_ASSORT = ("pine-logs", "spruce-logs", "spruce-pulpwood")      # por region solo tres surtidos, para no inflar el fichero
def adapter_fi():
    url = "https://statdb.luke.fi/PxWeb/api/v1/en/LUKE/met/metryv/0100_metryv.px"
    m = jget(url); V = {v["code"]: dict(zip(v["values"], v["valueTexts"])) for v in m["variables"]}
    standing = next(c for c, t in V["HAKT"].items() if t.strip() == "Standing sales")
    d = jpost(url, {"query": [{"code": "W", "selection": {"filter": "all", "values": ["*"]}}, {"code": "MPKH", "selection": {"filter": "all", "values": ["*"]}},
                              {"code": "HAKT", "selection": {"filter": "item", "values": [standing]}}, {"code": "PTL", "selection": {"filter": "all", "values": ["*"]}}], "response": {"format": "json-stat2"}})
    rows, _ = jsonstat_rows(d, ("W", "MPKH", "PTL"))
    weeks = list(d["dimension"]["W"]["category"]["index"])
    periods = ["%s-%s" % (w[:4], w[4:]) for w in weeks]
    series = []
    for pc, ptxt in V["PTL"].items():
        if ptxt not in FI_ASSORT: continue
        sid, grp, bench = FI_ASSORT[ptxt]
        for rc, rtxt in V["MPKH"].items():
            whole = rtxt.upper().startswith("WHOLE")
            if not whole and sid not in FI_REGION_ASSORT: continue
            vals = [rows.get((w, rc, pc)) for w in weeks]
            if sum(v is not None for v in vals) < 20: continue
            reg = None if whole else rtxt
            series.append({"id": "fi-" + sid + ("" if whole else "-" + re.sub(r"[^a-z0-9]+", "-", rtxt.lower()).strip("-")), "label": ptxt, "region": reg, "group": grp, "bench": bool(bench and whole),
                           "values": [None if v is None else round(float(v), 2) for v in vals]})
    return {"freq": "W", "currency": "EUR", "unit": "EUR/m3", "vat": "excl", "priceType": "stumpage",
            "periods": periods, "series": series,
            "source": {"name": "Natural Resources Institute Finland (Luke): Stumpage prices of roundwood, weekly (metryv/0100)", "url": "https://statdb.luke.fi/PxWeb/pxweb/en/LUKE/LUKE__04%20Metsa__04%20Metsavarat%20ja%20-markkinat/", "license": "CC BY 4.0"},
            "note": "Precio de la madera en pie en las ventas con comprador (todas las clases de corta), semanal; pais y regiones de precios."}
# ------------------------------------------------------------------ Polonia
PL_LIST = "https://stat.gov.pl/sygnalne/komunikaty-i-obwieszczenia/lista-komunikatow-i-obwieszczen/"
def adapter_pl():
    page = http(PL_LIST).decode("utf-8", "ignore")
    cur = re.search(r'href="([^"]*sredniej-ceny-sprzedazy-drewna[^"]*,268,\d+\.html)"', page)
    if not cur: raise RuntimeError("no se encuentra el comunicado del precio de la madera en la lista de GUS")
    def full(u): return u if u.startswith("http") else "https://stat.gov.pl" + u
    first = http(full(cur.group(1))).decode("utf-8", "ignore")
    links = list(dict.fromkeys([cur.group(1)] + re.findall(r'href="([^"]*sredniej-ceny-sprzedazy-drewna[^"]*)"', first)))
    got = {}
    for i, l in enumerate(links):
        raw = first if i == 0 else http(full(l)).decode("utf-8", "ignore")
        t = re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", " ", re.sub(r"<script.*?</script>|<style.*?</style>", "", raw, flags=re.S))))
        m = re.search(r"za pierwsze trzy kwarta\w+ (\d{4}) r\.,? wynios\w+ ([\d ]+,\d+) zł", t)
        if m: got[int(m.group(1))] = float(m.group(2).replace(" ", "").replace(",", "."))
    if len(got) < 8: raise RuntimeError("solo %d años de precio de la madera en GUS" % len(got))
    ys = sorted(got); periods = [str(y) for y in range(ys[0], ys[-1] + 1)]
    return {"freq": "A", "currency": "PLN", "unit": "PLN/m3", "vat": "n/d", "priceType": "sale-by-state-forests",
            "periods": periods, "series": [{"id": "pl-forest-tax-price", "label": "Average timber sale price of State Forests districts, Q1-Q3 (forest-tax base)", "group": "all", "bench": False,
                                           "values": [got.get(int(p)) for p in periods]}],
            "source": {"name": "Statistics Poland (GUS): Communique on the average timber sale price (forest tax base, Monitor Polski)", "url": PL_LIST, "license": "GUS: official communique, source credit"},
            "note": "Precio medio de venta de la madera obtenido por las unidades de los Bosques Estatales en los tres primeros trimestres del año (se usa para el impuesto forestal del año siguiente). Un dato al año."}
ADAPTERS = [("SI", adapter_si), ("CZ", adapter_cz), ("LV", adapter_lv), ("FI", adapter_fi), ("PL", adapter_pl)]   # sin ("BG", adapter_bg): licencia pendiente
def main():
    try: old = json.loads(OUT.read_text(encoding="utf-8")).get("countries", {})
    except Exception: old = {}
    out = {}; fails = []
    for cc, fn in ADAPTERS:
        try:
            c = fn()
            c["series"] = [x for x in c["series"] if any(v is not None for v in x["values"][-3:])]   # series descontinuadas
            if not c["series"] or not c["periods"]: raise RuntimeError("sin series")
            out[cc] = c; print(cc, "OK", len(c["series"]), "series,", len(c["periods"]), "periodos, ultimo", c["periods"][-1])
        except Exception as e:
            fails.append("%s (%s)" % (cc, str(e)[:140])); print("FALLO", cc, str(e)[:200])
            if cc in old: out[cc] = old[cc]
    if len(fails) > MAX_FAIL or len(out) < len(ADAPTERS) - MAX_FAIL:
        print("ERROR: fallan %d paises; no se toca el fichero" % len(fails)); return 1
    doc = {"schemaVersion": 1, "generatedAt": datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
           "groups": ["sawlog_conifer", "sawlog_broadleaf", "pulpwood_conifer", "pulpwood_broadleaf", "fuelwood", "stumpage", "sawlog", "pulpwood", "other", "all"],
           "countries": out}
    OUT.write_text(json.dumps(doc, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print("forest-timber-prices.json", OUT.stat().st_size, "bytes;", len(out), "paises; fallos:", len(fails)); return 0
if __name__ == "__main__": sys.exit(main())
