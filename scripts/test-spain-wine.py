#!/usr/bin/env python3
"""Vino de España por campaña y comunidad (MAPA, INFOVI informe ampliado de julio): lector, validacion y acumulacion con los informes reales de julio de 2021 y julio de 2026 (scripts/fixtures/spain-wine, sin red).
Uso: python3 scripts/test-spain-wine.py"""
import importlib.util, copy, sys
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
sp = importlib.util.spec_from_file_location("usw", ROOT / "scripts" / "update-spain-wine.py"); M = importlib.util.module_from_spec(sp); sp.loader.exec_module(M)
import openpyxl
fail = 0
def ok(n, c):
    global fail
    if not c: fail += 1; print("FALLA", n)
F = ROOT / "scripts" / "fixtures" / "spain-wine"
d25, e25 = M.parse(openpyxl.load_workbook(F / "ampliada_julio_2026.xlsx", data_only=True))
d20, e20 = M.parse(openpyxl.load_workbook(F / "ampliada_julio_2021.xlsx", data_only=True))
ok("fechas de extraccion del libro (8 sep 2026 y 1 sep 2021)", e25 == "2026-09-08" and e20 == "2021-09-01")
ok("campaña 2025/26: 28.857.382 hl de vino, 4.409.698,657 t de uva", d25["wine"]["TOTAL"] == 28857382 and abs(d25["grape"]["TOTAL"] - 4409698.657) < 0.001)
ok("campaña 2020/21: 40.948.526 hl y 6.105.801,137 t", d20["wine"]["TOTAL"] == 40948526 and abs(d20["grape"]["TOTAL"] - 6105801.137) < 0.001)
ok("Castilla-La Mancha 2025/26 15.819.660 hl", d25["wine"]["clm"] == 15819660)
ok("Madrid se lee aunque la hoja diga C.MADRID", d25["wine"]["madrid"] == 48618)
ok("categorias 2025/26: DOP 10.856.963, IGP 2.832.633, varietales 9.459.148, sin IG 5.708.638", (d25["dop"]["TOTAL"], d25["igp"]["TOTAL"], d25["varietal"]["TOTAL"], d25["sinig"]["TOTAL"]) == (10856963, 2832633, 9459148, 5708638))
ok("salidas 48.118.975 = exportaciones 16.927.700 + interiores 31.191.275", d25["exits"]["TOTAL"] == 48118975 and d25["exports"]["TOTAL"] == 16927700 and d25["domestic"]["TOTAL"] == 31191275)
ok("existencias a 31 de julio de 2026: vino 30.068.742 y mosto 2.068.747 hl", d25["stockWine"]["TOTAL"] == 30068742 and d25["stockMust"]["TOTAL"] == 2068747)
ok("17 comunidades y TOTAL en cada serie", all(len([k for k in d25[f] if k != "TOTAL"]) == 17 for f in M.FIELDS))
for c, d in (("2025", d25), ("2020", d20)):
    e, w = M.validate(c, d); ok("el informe real de %s cuadra" % c, not e)
bad = copy.deepcopy(d25); bad["wine"]["galicia"] += 9000; ok("comunidades que no suman el total se rechaza", M.validate("2025", bad)[0])
bad = copy.deepcopy(d25); bad["red"]["TOTAL"] += 50000; ok("tinto+blanco distinto del total se rechaza", M.validate("2025", bad)[0])
bad = copy.deepcopy(d25); bad["exports"]["clm"] += 80000; ok("salidas distintas de exportaciones + interiores se rechaza", M.validate("2025", bad)[0])
bad = copy.deepcopy(d25); bad["stockMust"]["TOTAL"] += 90000; ok("existencias que no suman se rechaza", M.validate("2025", bad)[0])
bad = copy.deepcopy(d25); bad["wine"]["TOTAL"] *= 3; bad["red"]["TOTAL"] *= 3; bad["white"]["TOTAL"] *= 3; ok("rendimiento uva-vino imposible se rechaza", M.validate("2025", bad)[0])
bad = copy.deepcopy(d25); bad["dop"]["aragon"] = -4; ok("negativo se rechaza", M.validate("2025", bad)[0])
wb = openpyxl.load_workbook(F / "ampliada_julio_2026.xlsx", data_only=True); del wb["4.1 EXISTENCIAS VINO+MOSTO"]
try: M.parse(wb); ok("sin cuadro 4.1 debe fallar", False)
except ValueError: pass
html = '<a href="/dam/x/infovi/ano-2024/informaciondeclaracionesampliadasjulio2024.xlsx">a</a><a href="/dam/x/infovi/ano-2024/informaciondeclaracionesampliadasjulio2024.pdf">p</a><a href="/dam/x/infovi/ano-2024/informaciondeclaracionexistenciasjulio2024_corregida.xlsx">e</a><a href="/dam/x/infovi/ano-2024/informaciondeclaracionampliadanoviembre2024.xlsx">n</a><a href="/dam/x/infovi/ano-2023/informaciondeclaracionampliadajulio_2023.xlsx">otro año</a>'
L = M.find_links(html, 2024); ok("solo el ampliado de julio de ese año (ni pdf, ni existencias, ni noviembre, ni otro año)", len(L) == 1 and L[0].endswith("ampliadasjulio2024.xlsx") and L[0].startswith("https://www.mapa.gob.es"))
ok("paginas de cada año probadas en orden", M.year_pages(2025)[0].endswith("datos_infovi_anteriores/infovi_2025") and M.year_pages(2025)[1].endswith("vitivinicultura/infovi_2025"))
cs, N, CC, revs = M.merge(None, {"2020": d20, "2025": d25}, {}, {}); ok("dos campañas: ordenadas y sin revisiones", cs == ["2020", "2025"] and not revs and N["wine"] == [40948526.0, 28857382.0])
d25b = copy.deepcopy(d25); d25b["wine"]["TOTAL"] += 100
old = {"campaigns": cs, "national": N, "ccaa": CC}
cs2, N2, CC2, revs = M.merge(old, {"2025": d25b}, {}, {}); ok("la revision de una campaña ya publicada se registra y gana el informe nuevo", any(r[0] == "2025" and r[1] == "wine" for r in revs) and N2["wine"][1] == 28857482.0 and N2["wine"][0] == 40948526.0)
print("test-spain-wine:", "%d fallos" % fail if fail else "OK"); sys.exit(1 if fail else 0)
