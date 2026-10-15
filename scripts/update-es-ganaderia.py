#!/usr/bin/env python3
"""España: ganadería del MAPA, tres estadísticas -> data/spain-ganaderia/{eggs,dairy,forecasts}.json

1. Huevos para consumo humano (Estadística de lana, miel y huevos): nº de ponedoras y producción (miles de docenas) por provincia, un libro por año.
   Hasta 2024 el libro separa «jaula y suelo» de «campera y ecológica»; desde 2025 separa jaula, suelo, campera y ecológica. Para poder comparar años se guardan siempre
   gallinas totales, «jaula+suelo» y «campera+ecológica»; los cuatro sistemas solo cuando el libro los da. Celdas con «DC» (dato confidencial) = null.
2. Estadística láctea anual: Cuadro B (industrias: leche recogida y productos obtenidos, por código de producto), Cuadro C (explotaciones: producción y destinos),
   Cuadro I (leche de vaca por comunidad). Series por año; los años que el MAPA publica con otro diseño de libro se ignoran y se anotan en el log.
3. Previsiones de la producción de ganado (bovino, porcino, ovino y caprino): la página solo enlaza la edición vigente; las anteriores se conservan en el fichero.
Nada se estima: celda vacía, «*» o «DC» = null."""
import argparse, datetime, html, io, json, re, sys, urllib.parse
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
import importlib.util
spec = importlib.util.spec_from_file_location("esy", str(ROOT / "scripts" / "update-es-esyrce.py")); esy = importlib.util.module_from_spec(spec); spec.loader.exec_module(esy)
http, slug, clean = esy.http, esy.slug, esy.clean
E = "https://www.mapa.gob.es/es/estadistica/temas/estadisticas-agrarias/ganaderia/"
PG_EGGS = E + "lana-miel-huevos-consumo-humano0"; PG_DAIRY = E + "estadistica-lacteas/industrias-anual"; PG_FC = E + "previsiones-produccion-ganado"
ALT = {"p-de-asturias": "principado-de-asturias", "asturias": "principado-de-asturias", "navarra": "comunidad-foral-de-navarra", "c-valenciana": "comunitat-valenciana", "comunidad-valenciana": "comunitat-valenciana",
       "murcia": "region-de-murcia", "madrid": "comunidad-de-madrid", "baleares": "illes-balears", "r-de-murcia": "region-de-murcia", "castilla-leon": "castilla-y-leon", "castilla-la-mancha": "castilla-la-mancha",
       "c-valen-cian": "comunitat-valenciana", "extremadu-ra": "extremadura", "c-valen-ciana": "comunitat-valenciana", "pais-vasco": "pais-vasco", "aragon": "aragon"}
def log(*a): esy.log(*a)
def tid(n): s = slug(n); return ALT.get(s, s)
SMALL = {"de", "del", "la", "las", "los", "y", "el"}
def nice(n):
    n = clean(n)
    if not n.isupper(): return n
    return " ".join(w.lower() if (i and w.lower() in SMALL) else w.capitalize() if not re.match(r"^[A-Z]\.$", w) else w for i, w in enumerate(n.lower().split(" "))).replace("Castilla-La Mancha", "Castilla-La Mancha")
def kind(n):
    t = tid(n)
    return "country" if t == "espana" else "ccaa" if t in esy.CCAA else "province"
def num(v):
    return None if isinstance(v, bool) or not isinstance(v, (int, float)) else round(float(v), 3)
def links(page, base):
    out = []; seen = set()
    for h, t in re.findall(r'href="([^"#]+\.(?:xlsx?|xlsm))"[^>]*>(.*?)</a>', page, re.S | re.I):
        u = urllib.parse.urljoin(base, html.unescape(h))
        if u not in seen: seen.add(u); out.append((u, urllib.parse.unquote(u.split("/")[-1]), clean(re.sub(r"<[^>]+>", " ", t))))
    return out
def year_of(u, name):
    m = re.search(r"/(20\d\d)/", u) or re.search(r"(?<!\d)(20\d\d)(?!\d)", name) or re.search(r"[-_](\d\d)(?:[-_.]|web)", name)
    if not m: return None
    y = int(m.group(1)); return y + 2000 if y < 100 else y

# ---------------- huevos ----------------
EGG_COLS = ("cage", "floor", "free", "organic", "cage_floor", "free_organic", "hens", "turkeys", "ducks", "geese", "all")
def parse_eggs_sheet(rows):
    hr = next((i for i, r in enumerate(rows[:10]) if r and isinstance(r[0], str) and r[0].strip().lower().startswith("provincias")), None)
    if hr is None: return None
    ncol = max(len(r) for r in rows[hr:hr + 3])
    raw1 = [clean(rows[hr][c]).lower() if c < len(rows[hr]) and isinstance(rows[hr][c], str) else "" for c in range(ncol)]
    h1 = []; cur = ""
    for c in range(ncol):   # cabecera superior con relleno hacia la derecha solo para «Gallinas» (celda combinada sobre sus subcolumnas)
        if raw1[c]: cur = raw1[c]
        h1.append(cur if c > 0 else "")
    h2 = [clean(rows[hr + 1][c]).lower() if c < len(rows[hr + 1]) and isinstance(rows[hr + 1][c], str) else "" for c in range(ncol)]
    if not any(h2): h2 = [clean(rows[hr + 2][c]).lower() if c < len(rows[hr + 2]) and isinstance(rows[hr + 2][c], str) else "" for c in range(ncol)]   # a veces la segunda fila de cabecera está vacía
    colmap = {}
    for c in range(1, ncol):
        a, b = h1[c], h2[c]
        if a.startswith("gallinas") and b:
            if "ecol" in b and "campera" not in b: colmap[c] = "organic"
            elif "campera" in b and "ecol" in b: colmap[c] = "free_organic"
            elif "campera" in b: colmap[c] = "free"
            elif "jaula" in b and ("suelo" in b or "producci" in b): colmap[c] = "cage_floor"
            elif "jaula" in b: colmap[c] = "cage"
            elif "suelta" in b or "gallinero" in b: colmap[c] = "floor"
            elif "total" in b: colmap[c] = "hens"
        elif raw1[c].startswith("pavas"): colmap[c] = "turkeys"
        elif raw1[c].startswith("patas"): colmap[c] = "ducks"
        elif raw1[c].startswith("ocas"): colmap[c] = "geese"
        elif raw1[c].startswith("total"): colmap[c] = "all"
    if "all" not in colmap.values() or not any(v in colmap.values() for v in ("hens", "cage_floor", "cage")): return None
    d0 = hr + 1
    while d0 < len(rows) and not (rows[d0] and isinstance(rows[d0][0], str) and clean(rows[d0][0]) and tid(rows[d0][0]) not in ("",) and any(isinstance(x, (int, float)) for x in rows[d0][1:ncol])): d0 += 1
    names = []; vals = {k: [] for k in set(colmap.values())}
    for r in rows[d0:]:
        if not r or not isinstance(r[0], str) or not clean(r[0]): continue   # filas separadoras (a veces con un 0 en la última columna)
        nm = clean(r[0])
        if nm.upper().startswith("DC"): break
        names.append(nm)
        for c, k in colmap.items(): vals[k].append(num(r[c]) if c < len(r) else None)
        if tid(nm) == "espana": break
    if len(names) < 40 or tid(names[-1]) != "espana": return None
    return names, vals
def build_eggs(a, out, now):
    sets = {}
    if a.local_eggs:
        for s in a.local_eggs: y, p = s.split(":", 1); sets[int(y)] = (p, Path(p).name)
    else:
        for u, n, t in links(http(PG_EGGS).decode("utf-8", "replace"), PG_EGGS):
            if "huevos" not in n.lower(): continue
            y = year_of(u, n)
            if y and y not in sets: sets[y] = (u, n)
    if not sets: log("huevos: sin libros"); return
    import openpyxl
    terr = {}; terr_order = []; years = {}
    for y in sorted(sets, reverse=True):
        u, n = sets[y]
        try:
            data = Path(u).read_bytes() if a.local_eggs else http(u)
            if n.lower().endswith(".xls"):
                import xlrd
                book = xlrd.open_workbook(file_contents=data); shs = {s.name: [s.row_values(i) for i in range(s.nrows)] for s in book.sheets()}
            else:
                wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True); shs = {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb}
        except Exception as e: log("huevos", y, "ERROR", n[:50], str(e)[:100]); continue
        got = {}
        for title, rows in shs.items():
            t = title.lower(); key = "hens" if "ponedora" in t else "dozens" if "huevo" in t else None
            if not key: continue
            res = parse_eggs_sheet(rows)
            if res: got[key] = res
            else: log("huevos", y, "hoja", title, "ignorada (estructura distinta)")
        if set(got) != {"hens", "dozens"}: log("huevos", y, "incompleto, se omite"); continue
        years[y] = {}
        for key, (names, vals) in got.items():
            ids = [tid(x) for x in names]
            for nm, i in zip(names, ids):
                if i not in terr: terr[i] = {"id": i, "name": nice(nm), "kind": kind(nm)}; terr_order.append(i)
            years[y][key] = {"ids": ids, "vals": vals}
        log("huevos", y, n[:45], len(got["hens"][0]), "territorios")
    if not years: log("huevos: ningún año legible, no se toca nada"); return
    ids = terr_order
    doc = {"schemaVersion": 1, "generatedAt": now, "source": {"id": "mapa_es", "name": "MAPA, Estadística de huevos para consumo humano", "url": PG_EGGS},
           "units": {"hens": "número de ponedoras", "dozens": "miles de docenas"}, "territories": [terr[i] for i in ids], "years": {}}
    for y in sorted(years):
        o = {}
        for key in ("hens", "dozens"):
            ent = years[y][key]; pos = {i: j for j, i in enumerate(ent["ids"])}; o[key] = {}
            for k, v in ent["vals"].items():
                arr = [v[pos[i]] if i in pos else None for i in ids]
                if any(x is not None for x in arr): o[key][k] = arr
        # «jaula+suelo» y «campera+ecológica» también para los años con los cuatro sistemas (suma solo si ambos existen)
        for key in ("hens", "dozens"):
            m = o[key]
            if "cage_floor" not in m and "cage" in m and "floor" in m: m["cage_floor"] = [round(a_ + b_, 3) if a_ is not None and b_ is not None else None for a_, b_ in zip(m["cage"], m["floor"])]
            if "hens" not in m and all(k in m for k in ("cage", "floor", "free", "organic")): m["hens"] = [round(sum(x), 3) if all(v is not None for v in x) else None for x in zip(m["cage"], m["floor"], m["free"], m["organic"])]   # gallinas totales = suma de los cuatro sistemas publicados
            if "free_organic" not in m and "free" in m and "organic" in m: m["free_organic"] = [round(a_ + b_, 3) if a_ is not None and b_ is not None else None for a_, b_ in zip(m["free"], m["organic"])]
        doc["years"][str(y)] = o
    doc["note"] = "Estadística anual del MAPA. Hasta 2024 solo hay dos grupos de sistema (jaula y suelo / campera y ecológica); desde 2025 hay cuatro. Dato confidencial = vacío. El último año puede ser provisional."
    (out / "eggs.json").write_text(json.dumps(doc, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")

# ---------------- lácteas ----------------
CODE = re.compile(r"^\d+(\.\d+)*\.?$")
def cells(r): return [(j, v) for j, v in enumerate(r) if v not in (None, "")]
def parse_B(rows):
    """-> (disponibilidades {etiqueta: [1000 t, grasa t, proteína t]}, productos {código: (etiqueta, 1000 t)})"""
    avail = {}; prods = {}
    for r in rows:
        cs = cells(r)
        i = 0
        while i < len(cs):
            j, v = cs[i]
            if isinstance(v, str) and CODE.match(v.strip()) and i + 1 < len(cs) and isinstance(cs[i + 1][1], str) and not CODE.match(cs[i + 1][1].strip()):
                lab = clean(cs[i + 1][1]); val = None
                if i + 2 < len(cs):
                    nv = cs[i + 2][1]
                    if isinstance(nv, (int, float)) and not isinstance(nv, bool): val = round(float(nv), 3)
                if lab and not lab.startswith("("): prods.setdefault(v.strip(), (lab, val))
                i += 2; continue
            i += 1
        if cs and isinstance(cs[0][1], str) and cs[0][0] <= 2 and not CODE.match(cs[0][1].strip()):   # panel izquierdo: disponibilidades
            lab = clean(cs[0][1]); nums = []
            for j, v in cs[1:]:
                if j > 8: break
                if isinstance(v, (int, float)) and not isinstance(v, bool): nums.append(round(float(v), 3))
                elif isinstance(v, str) and v.strip() in ("*", "**"): nums.append(None)
                else: break
            if nums and re.match(r"^(i{1,3}\.|\d[a-e]?\.)", lab.lower()): avail.setdefault(lab, nums)
    return avail, prods
def parse_C(rows):
    out = {}; sec = None
    for r in rows:
        cs = cells(r)
        if not cs: continue
        s0 = clean(cs[0][1]) if isinstance(cs[0][1], str) else ""
        m = re.match(r"^([ABC])\.\s*(DISPONIBILIDADES|DESTINOS|PRODUCTOS)", s0.upper())
        if m: sec = m.group(1); continue
        if sec is None: continue
        lab = None; val = None
        for j, v in cs:
            if j > 5: break
            if lab is None and isinstance(v, str) and not CODE.match(v.strip()) and clean(v): lab = clean(v); continue
            if lab is not None and isinstance(v, (int, float)) and not isinstance(v, bool): val = round(float(v), 3); break
        if lab and val is not None: out.setdefault((sec, lab), val)
    return out
def parse_I(rows):
    res = {}; started = False
    for r in rows:
        cs = cells(r)
        if len(cs) >= 2 and isinstance(cs[0][1], str) and clean(cs[0][1]).upper() == "CCAA": started = True; continue
        if not started or len(cs) < 2: continue
        if isinstance(cs[0][1], str) and isinstance(cs[1][1], (int, float)):
            res[tid(cs[0][1])] = (clean(cs[0][1]), round(float(cs[1][1]), 3))
            if tid(cs[0][1]) == "espana": break
    return res
def classify_dairy(n):
    m = n.lower()
    if re.search(r"cuadros?[_-]?b", m): return "BH"
    if re.search(r"cuadros?[_-]?c", m): return "CI"
    return None
def build_dairy(a, out, now):
    import openpyxl
    sets = {}
    if a.local_dairy:
        for s in a.local_dairy:
            y, m, p = s.split(":", 2); sets.setdefault(int(y), {})[m] = (p, Path(p).name)
    else:
        for u, n, t in links(http(PG_DAIRY).decode("utf-8", "replace"), PG_DAIRY):
            if not re.search(r"cuadro", n.lower()) or not n.lower().endswith((".xlsx", ".xls")): continue
            m = classify_dairy(n)
            if not m: continue
            y = year_of(u, n)
            if y and m not in sets.setdefault(y, {}): sets[y][m] = (u, n)
    avail = {}; prods = {}; farm = {}; ccaa = {}; ccaa_names = {}; years = []
    for y in sorted(sets):
        got = False
        for m, (u, n) in sets[y].items():
            try:
                data = Path(u).read_bytes() if a.local_dairy else http(u)
                wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
                shs = {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb}
            except Exception as e: log("lácteas", y, m, "ERROR", n[:50], str(e)[:90]); continue
            for title, rows in shs.items():
                t = title.lower()
                try:
                    if re.search(r"cuadro_?b", t):
                        av, pr = parse_B(rows)
                        if len(pr) < 20: log("lácteas", y, title, "ignorada: pocos productos leídos (%d)" % len(pr)); continue
                        for k, v in av.items(): avail.setdefault(k, {})[y] = v
                        for k, (lab, v) in pr.items(): prods.setdefault(k, {"label": lab, "v": {}})["v"][y] = v
                        got = True; log("lácteas", y, title, len(av), "disponibilidades", len(pr), "productos")
                    elif re.search(r"cuadro_?c", t):
                        r = parse_C(rows)
                        if len(r) < 8: log("lácteas", y, title, "ignorada: pocas filas (%d)" % len(r)); continue
                        for k, v in r.items(): farm.setdefault(k, {})[y] = v
                        got = True; log("lácteas", y, title, len(r), "filas")
                    elif re.search(r"cuadro_?i", t):
                        r = parse_I(rows)
                        if len(r) < 10: log("lácteas", y, title, "ignorada: pocas CCAA (%d)" % len(r)); continue
                        for k, (nm, v) in r.items(): ccaa.setdefault(k, {})[y] = v; ccaa_names[k] = nm
                        got = True; log("lácteas", y, title, len(r), "territorios")
                except Exception as e: log("lácteas", y, title, "ERROR", str(e)[:90])
        if got: years.append(y)
    if not years: log("lácteas: ningún año legible, no se toca nada"); return
    ys = sorted(years)
    def ser(d): return [d.get(y) for y in ys]
    doc = {"schemaVersion": 1, "generatedAt": now, "source": {"id": "mapa_es", "name": "MAPA, Estadística láctea anual", "url": PG_DAIRY}, "years": ys,
           "industry": {"collected": [{"label": k, "v": [(avail[k].get(y) or [None])[0] for y in ys], "fat": [(avail[k].get(y) or [None, None])[1] if len(avail[k].get(y) or []) > 1 else None for y in ys],
                                       "protein": [(avail[k].get(y) or [None, None, None])[2] if len(avail[k].get(y) or []) > 2 else None for y in ys]} for k in avail],
                        "products": [{"code": k, "label": prods[k]["label"], "v": ser(prods[k]["v"])} for k in sorted(prods, key=lambda c: [int(x) for x in c.strip(".").split(".")])]},
           "farm": [{"section": k[0], "label": k[1], "v": ser(v)} for k, v in farm.items()],
           "cowMilkByCcaa": {"territories": [{"id": k, "name": nice(ccaa_names[k]), "kind": kind(ccaa_names[k])} for k in ccaa], "v": {k: ser(v) for k, v in ccaa.items()}},
           "units": {"industry": "miles de toneladas de producto (grasa y proteína en toneladas)", "farm": "miles de toneladas", "cowMilkByCcaa": "miles de toneladas"},
           "note": "Cuadros B, C e I de la estadística láctea anual del MAPA. Algunos años se publican con otro diseño de libro y no se leen (ver el registro). El último año puede ser provisional."}
    (out / "dairy.json").write_text(json.dumps(doc, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")

# ---------------- previsiones ----------------
def parse_forecast(rows):
    sp = {}; cur = None
    for r in rows:
        cs = cells(r)
        if not cs: continue
        c1 = r[1] if len(r) > 1 else None; c2 = r[2] if len(r) > 2 else None
        if isinstance(c2, str) and c1 in (None, ""):
            cur = {"species": clean(c2), "columns": [clean(x) for x in r[2:] if isinstance(x, str) and clean(x)], "periods": []}; sp[len(sp)] = cur; continue
        if cur is not None and isinstance(c1, str) and re.search(r"(semestre|trimestre)\s+20\d\d", c1.lower()):
            vals = [num(x) for x in r[2:2 + max(1, len(cur["columns"]))]]
            cur["periods"].append({"period": clean(c1), "v": vals})
    return [v for v in sp.values() if v["periods"]]
def build_forecasts(a, out, now):
    import openpyxl
    cand = None
    if a.local_fc: cand = (a.local_fc, Path(a.local_fc).name)
    else:
        for u, n, t in links(http(PG_FC).decode("utf-8", "replace"), PG_FC):
            if re.search(r"gip\d{4}s\d", n.lower()) and n.lower().endswith((".xlsx", ".xls")): cand = (u, n); break
    if not cand: log("previsiones: sin libro"); return
    u, n = cand; m = re.search(r"gip(\d{4})s(\d)", n.lower())
    if not m: log("previsiones: nombre sin edición", n); return
    ed = "%s-S%s" % (m.group(1), m.group(2))
    wb = openpyxl.load_workbook(io.BytesIO(Path(u).read_bytes() if a.local_fc else http(u)), data_only=True, read_only=True)
    rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    unit = next((clean(r[1]) for r in rows if len(r) > 1 and isinstance(r[1], str) and r[1].lower().startswith("unidad")), "")
    sp = parse_forecast(rows)
    if len(sp) < 3: log("previsiones: estructura distinta (%d especies), se omite" % len(sp)); return
    fp = out / "forecasts.json"
    old = {}
    if fp.exists():
        try: old = json.loads(fp.read_text(encoding="utf-8")).get("editions", {})
        except Exception as e: log("previsiones: fichero anterior ilegible", str(e)[:80])
    old[ed] = {"unit": unit.replace("unidad:", "").strip(), "firstSeen": old.get(ed, {}).get("firstSeen", now[:10]), "species": sp}
    doc = {"schemaVersion": 1, "generatedAt": now, "source": {"id": "mapa_es", "name": "MAPA, Previsiones de la producción de ganado", "url": PG_FC}, "editions": dict(sorted(old.items())),
           "note": "Previsión semestral del MAPA (bovino, ovino y caprino por semestre; porcino por trimestre). La página solo enlaza la edición vigente: las anteriores se guardan desde que Dehesa Index las lee."}
    fp.write_text(json.dumps(doc, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
    log("previsiones", ed, len(sp), "especies")

def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--out", default=str(ROOT / "data" / "spain-ganaderia"))
    ap.add_argument("--local-eggs", action="append", default=[], help="año:ruta"); ap.add_argument("--local-dairy", action="append", default=[], help="año:BH|CI:ruta"); ap.add_argument("--local-fc", default="")
    ap.add_argument("--only", default="eggs,dairy,forecasts"); ap.add_argument("--log", default=str(ROOT / "data" / "spain-ganaderia-log.txt"))
    a = ap.parse_args(); out = Path(a.out); out.mkdir(parents=True, exist_ok=True)
    now = datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    for k, fn in (("eggs", build_eggs), ("dairy", build_dairy), ("forecasts", build_forecasts)):
        if k not in a.only.split(","): continue
        try: fn(a, out, now)
        except Exception as e: log(k, "ERROR", str(e)[:140])
    if not any((out / f).exists() for f in ("eggs.json", "dairy.json", "forecasts.json")): raise SystemExit("no se generó ningún fichero")
    Path(a.log).write_text("\n".join(esy.LOG) + "\n", encoding="utf-8")
if __name__ == "__main__": main()
