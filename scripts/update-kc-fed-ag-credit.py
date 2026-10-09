#!/usr/bin/env python3
"""EE. UU.: Encuesta trimestral de credito agrario de la Fed de Kansas City (Tenth District) -> data/us-kcfed-stats.json (countries.US, extend).
Tipos de interes fijos y variables (operativos, maquinaria/intermedios, inmobiliarios) por estado (KS, MO, NE, OK), Estados de montana y total del distrito;
variacion anual del valor de la tierra (secano, regadio, pastos) e indices de difusion de condiciones de credito.
Los xlsx cambian de ID cada trimestre: se rastrean desde https://www.kansascityfed.org/agriculture/ag-credit-survey/.
Desde el 3T 2025 la muestra incluye bancos comerciales e instituciones del Farm Credit System (antes solo bancos comerciales): se anota en periodNote.
Un fallo no borra lo anterior."""
import datetime, io, json, re, sys, urllib.request
from pathlib import Path
import openpyxl
ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'data/us-kcfed-stats.json'; LOGF = ROOT / 'data/us-kcfed-log.txt'
BASE = 'https://www.kansascityfed.org'; PAGE = BASE + '/agriculture/ag-credit-survey/'
LOG = []
def log(*a):
    s = ' '.join(str(x) for x in a); LOG.append(s); print(s)
H = {'User-Agent': 'DehesaIndex/1.0 (+https://dehesaindex.com)'}
def get(u): return urllib.request.urlopen(urllib.request.Request(u, headers=H), timeout=90).read()
NOTE = 'Desde 2025-Q3 la muestra incluye bancos comerciales e instituciones del Farm Credit System; antes solo bancos comerciales.'
PLACES = {'Kansas': 'Kansas', 'Missouri': 'Missouri', 'Nebraska': 'Nebraska', 'Oklahoma': 'Oklahoma', 'Mountain States': 'Mountain States', 'Tenth District': 'Tenth District'}
KINDS = {'Operating': 'operating loans', 'Intermediate': 'machinery/intermediate loans', 'Real Estate': 'real estate loans'}
CREDIT = {'for Loans': 'Loan demand', 'Availability': 'Loan fund availability', 'Rates': 'Loan repayment rates', 'or Extensions': 'Loan renewals or extensions', 'Required': 'Collateral required',
          'Income': 'Farm income', 'Spending': None}
def slug(s): return re.sub(r'[^a-z0-9]+', '-', s.lower()).strip('-')
def num(v): return float(v) if isinstance(v, (int, float)) else None
def rows(ws):
    """Devuelve [(periodo, [valores...])] rellenando el anio (celda vacia) y descartando '---' y '.'."""
    out = []; year = None
    for r in ws.iter_rows(values_only=True):
        if isinstance(r[0], (int, float)) and 1970 < r[0] < 2100: year = int(r[0])
        elif r[0] is not None: continue
        q = r[1]
        if year is None or not isinstance(q, (int, float)) or not 1 <= q <= 4: continue
        out.append(('%d-Q%d' % (year, q), list(r[2:])))
    return out
def series(sid, label, unit, pts, group):
    pts = [[p, v] for p, v in pts if v is not None]
    if len(pts) < 8: log('pocos datos', label, len(pts)); return None
    last, prev = pts[-1], pts[-2]
    s = {'id': sid, 'group': group, 'label': label, 'unit': unit, 'frequency': 'quarterly', 'latestPeriod': last[0], 'latest': last[1],
         'changePct': None, 'points': pts, 'sourceGroup': 'Kansas City Fed Agricultural Credit Survey'}
    if any(p[0] >= '2025-Q3' for p in pts): s['periodNote'] = NOTE
    return s
def main():
    html = get(PAGE).decode('utf8', 'replace')
    links = sorted(set(re.findall(r'href="(/documents/\d+/[^"]+\.xlsx)"', html)))
    log('enlaces', links)
    files = {}
    for l in links:
        n = l.lower()
        for key in ('creditconditions', 'fixedinterestrates', 'variableinterestrates', 'landvalues'):
            if key in n: files[key] = openpyxl.load_workbook(io.BytesIO(get(BASE + l)), data_only=True)
    log('ficheros', sorted(files))
    out = []
    for key, mode in (('fixedinterestrates', 'fixed'), ('variableinterestrates', 'variable')):
        wb = files.get(key)
        if not wb: continue
        for sh, kind in KINDS.items():
            if sh not in wb.sheetnames: log('falta hoja', key, sh); continue
            ws = wb[sh]; hdr = None
            for r in ws.iter_rows(values_only=True):
                if r[0] == 'Year': hdr = r; break
            if not hdr: continue
            names = ['Mountain States' if (c or '').startswith('States') else c for c in hdr[2:]]
            names = [n for n in names if n]
            data = rows(ws)
            for i, col in enumerate(names):
                # el rotulo de 'Mountain / States**' ocupa dos filas; se normaliza arriba
                nm = col.replace('*', '').strip()
                nm = 'Mountain States' if nm.startswith('Mountain') or nm == 'States' else nm
                if nm not in PLACES and nm != 'District': continue
                nm = 'Tenth District' if nm == 'District' else nm
                pts = [(p, num(v[i]) if i < len(v) else None) for p, v in data]
                s = series('us-kcfed-%s-%s-%s' % (mode, slug(kind), slug(nm)), '%s: %s interest rate, %s (KC Fed)' % (nm, mode, kind), '%', pts, 'rates')
                if s: out.append(s)
    wb = files.get('landvalues')
    if wb:
        ws = wb.worksheets[0]; data = rows(ws)
        for i, nm in enumerate(('Non-irrigated cropland', 'Irrigated cropland', 'Ranchland')):
            pts = [(p, num(v[i])) for p, v in data]
            s = series('us-kcfed-land-' + slug(nm), 'Tenth District: farmland value change from a year ago, %s (KC Fed)' % nm.lower(), '%', pts, 'rates')
            if s: out.append(s)
    wb = files.get('creditconditions')
    if wb:
        ws = wb.worksheets[0]; data = rows(ws)
        names = ['Loan demand', 'Loan fund availability', 'Loan repayment rates', 'Loan renewals or extensions', 'Collateral required', 'Farm income', 'Household spending', 'Capital spending']
        for i, nm in enumerate(names):
            pts = [(p, num(v[i])) for p, v in data]
            s = series('us-kcfed-credit-' + slug(nm), 'Tenth District: %s diffusion index (KC Fed)' % nm.lower(), 'index', pts, 'rates')
            if s: out.append(s)
    log('series', len(out))
    LOGF.write_text('\n'.join(LOG[-60:]) + '\n', encoding='utf-8')
    if len(out) < 40: log('demasiado pocas series; no se escribe'); sys.exit(1)
    src = {'name': 'Federal Reserve Bank of Kansas City - Agricultural Credit Survey (Tenth District)', 'url': PAGE, 'license': 'Federal Reserve Bank of Kansas City; free use with attribution'}
    doc = {'schemaVersion': 1, 'generatedAt': datetime.datetime.now(datetime.timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ'), 'countries': {'US': {'name': 'United States', 'extend': True, 'source': src, 'series': out}}, 'log': LOG[-30:]}
    OUT.write_text(json.dumps(doc, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
main()
