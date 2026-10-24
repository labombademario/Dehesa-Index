#!/usr/bin/env python3
"""Pruebas sin red de scripts/update-pl-zsrir.py con libros xlsx sinteticos que imitan los boletines del ZSRIR: celdas sucias ('nld', '---', numeros como texto), semanas ISO
(incluida la que cruza de año), bloques de dos grupos de columnas (trigo y maiz), contraste del ultimo punto semanal con la tabla resumen del propio boletin, columnas con nombre
desconocido (no se publican), el año 2021 de la colza fuera, y que un bloque que falla conserve lo publicado. El contrato de los ficheros se comprueba con scripts/contract_tests.py."""
import datetime, importlib.util, json, sys, tempfile
from pathlib import Path
import openpyxl
HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sp = importlib.util.spec_from_file_location("pz", str(HERE / "update-pl-zsrir.py")); pz = importlib.util.module_from_spec(sp); sp.loader.exec_module(pz)
import contract_tests as CT


def check(c, m):
    if not c:
        print("FALLO:", m)
        sys.exit(1)


D = datetime.datetime
# --- num / wk
check(pz.num("12,5") == 12.5 and pz.num("nld") is None and pz.num("---") is None and pz.num(None) is None and pz.num(True) is None and pz.num(7) == 7.0 and pz.num(" 1330.5 ") == 1330.5, "num")
check(pz.wk(datetime.date(2023, 1, 1)) == "2022-W52" and pz.wk(datetime.date(2022, 1, 9)) == "2022-W01" and pz.wk(datetime.date(2021, 1, 3)) == "2020-W53", "semanas ISO")


# --- weekly: dos grupos de columnas, el de la izquierda (col 2..12) y el de la derecha (14..)
def wkbook(left, right):
    wb = openpyxl.Workbook(); ws = wb.active
    for blk, (c0, vals) in enumerate(((2, left), (14, right))):
        for j, (start, series) in enumerate(vals):
            for i, v in enumerate(series):
                ws.cell(row=26 + i, column=c0 + 2 * j + 1, value=D.combine(start + datetime.timedelta(days=7 * i), datetime.time()))
                ws.cell(row=26 + i, column=c0 + 2 * j + 2, value=v)
    return wb, ws


s22 = datetime.date(2022, 1, 9)
wb, ws = wkbook([(s22, [100.0 + i for i in range(50)] + ["nld", "---"])], [(s22, [200.0 + i for i in range(52)])])
a = pz.weekly(ws, 2, 13); b = pz.weekly(ws, 14, 25)
check(a["2022-W01"][1] == 100 and "2022-W51" not in a and len(a) == 50, "trigo semanal sin celdas sucias (%d)" % len(a))
check(b["2022-W01"][1] == 200 and b["2022-W52"][1] == 251, "maiz semanal")
check(not (set(a) & set(b) - set(a)), "grupos separados")
# año ISO < min_year se omite (la semana 52 de 2021 que acaba el 2022-01-02 no entra)
wb, ws = wkbook([(datetime.date(2022, 1, 2), [5857.07, 3242, 3300] + [3000 + i for i in range(45)])], [])
w = pz.weekly(ws, 2, 13)
check("2021-W52" not in w and w["2022-W01"][1] == 3242, "semana de 2021 fuera: %s" % sorted(w))
# futuro fuera
fut = datetime.date.today() + datetime.timedelta(days=30)
wb, ws = wkbook([(fut, [1.0 + i for i in range(45)])], [])
check(pz.weekly(ws, 2, 13) == {}, "fechas futuras fuera")

# --- summary_weekly: cabecera dd.mm.aaaa, rotulo repartido en dos columnas, primera aparicion manda
sw = openpyxl.Workbook().active
sw["D7"] = "11.10.2026"; sw["E7"] = "04.10.2026"; sw["I7"] = "11.10.2026"
sw["B9"] = "PSZENICA "; sw["C9"] = "KONSUMPCYJNA"; sw["D9"] = "947.79"; sw["E9"] = 936.13; sw["I9"] = 36.7
sw["C10"] = "PASZOWA"; sw["D10"] = 907.49
sw["B16"] = "KUKURYDZA"; sw["C16"] = "MOKRA"; sw["D16"] = 630.4
sw["C17"] = "SUCHA"; sw["D17"] = 971.55
sw["B22"] = "PSZENICA "; sw["C22"] = "KONSUMPCYJNA"; sw["D22"] = 928.7
r = pz.summary_weekly(sw, r"\bPszenica konsumpcyjna")
check(r == {"2026-10-11": 947.79, "2026-10-04": 936.13}, "resumen trigo: %s" % r)
check(pz.summary_weekly(sw, r"\bKukurydza sucha").get("2026-10-11") == 971.55, "resumen maiz (rotulo arrastrado)")
# cabecera dd.mm + año en la fila siguiente
sy = openpyxl.Workbook().active
sy["B2"] = "18.10"; sy["C2"] = "11.10"; sy["B3"] = "2026"; sy["C3"] = "2026"; sy["A5"] = "Pszenica konsumpcyjna"; sy["B5"] = 965.36; sy["C5"] = 947.79
check(pz.summary_weekly(sy, r"\bPszenica konsumpcyjna") == {"2026-10-18": 965.36, "2026-10-11": 947.79}, "cabecera dd.mm + año")

# --- check_last
pts = {"2026-W41": (datetime.date(2026, 10, 11), 947.79), "2026-W40": (datetime.date(2026, 10, 4), 936.13)}
check(pz.check_last(pts, {"2026-10-11": 947.8}, "t") == datetime.date(2026, 10, 11), "contraste ok")
for bad, why in (({"2026-10-11": 971.55}, "valor distinto"), ({"2026-10-04": 936.13}, "semana ausente del resumen"), ({}, "sin resumen")):
    try:
        pz.check_last(pts, bad, "t"); check(False, "debia fallar: " + why)
    except RuntimeError:
        pass
try:
    pz.check_last({}, {}, "t"); check(False, "sin puntos")
except RuntimeError:
    pass

# --- monthly_dates
mw = openpyxl.Workbook().active
mw["B3"] = "x"; mw["B4"] = "MC/ROK"
for j, nm in enumerate(["Pszenica konsumpcyjna", "Pszenica  paszowa", "Mystery"]):
    mw.cell(row=4, column=3 + j, value=nm)
rows = [(D(2026, 9, 1), 915.7, "874,5", 1), (D(2026, 8, 1), "---", 810.86, 2), (D(2026, 7, 1), 835.67, 808.83, 3)]
for i, r in enumerate(rows):
    mw.cell(row=5 + i, column=2, value=r[0])
    for j in range(3):
        mw.cell(row=5 + i, column=3 + j, value=r[1 + j])
mw["B9"] = "tekst po tabeli"
names, data = pz.monthly_dates(mw)
check(names == {2: "Pszenica konsumpcyjna", 3: "Pszenica paszowa", 4: "Mystery"} and len(data) == 3, "mensual: %s %d" % (names, len(data)))

# --- bloques completos con libros sinteticos
CER_NAMES = ["Pszenica konsumpcyjna", "Pszenica paszowa", "Zboże nieznane"]


def cereal_book(last_week_value=947.79, summary_wheat=947.79, summary_maize=971.55):
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    m = wb.create_sheet("ZAKUP WYKRESY MCE"); m["B4"] = "MC/ROK"
    for j, nm in enumerate(CER_NAMES):
        m.cell(row=4, column=3 + j, value=nm)
    for i in range(30):
        y, mo = divmod(8 * 12 + 9 - 1 - i, 12)
        m.cell(row=5 + i, column=2, value=D(2018 + y, mo + 1, 1))
        m.cell(row=5 + i, column=3, value=500 + i); m.cell(row=5 + i, column=4, value="---" if i == 3 else 450 + i); m.cell(row=5 + i, column=5, value=1)
    t = wb.create_sheet("ZAKUP WYKRESY TYG.")
    start = datetime.date(2022, 1, 9); n = 249
    for i in range(n):
        d = start + datetime.timedelta(days=7 * i)
        v = 1000 - i * 0.5
        if i == n - 1:
            v = last_week_value
        for c0, val in ((3, v), (15, 971.55 if i == n - 1 else 800 + i * 0.1)):
            t.cell(row=26 + i, column=c0, value=D.combine(d, datetime.time())); t.cell(row=26 + i, column=c0 + 1, value=round(val, 2))
    last = start + datetime.timedelta(days=7 * (n - 1))
    z = wb.create_sheet("ZAKUP ZIARNA")
    z["D7"] = last.strftime("%d.%m.%Y")
    z["B9"] = "PSZENICA "; z["C9"] = "KONSUMPCYJNA"; z["D9"] = summary_wheat
    z["B16"] = "KUKURYDZA"; z["C16"] = "MOKRA"; z["D16"] = 630.4; z["C17"] = "SUCHA"; z["D17"] = summary_maize
    return wb, last


wb, last = cereal_book()
check(last <= datetime.date.today(), 'el boletin sintetico acaba en el pasado')
pz.workbook = lambda ds, tries=4: (wb, "2026-10-15", "x")
ed, tabs = pz.cereales()
check(ed.startswith("semana ") and [t["id"] for t in tabs] == ["cereales-mensual", "cereales-semanal"], "cereales: %s" % ed)
ids = [s["id"] for s in tabs[0]["series"]]
check(ids == ["pl-cer-m-pszenica-kons", "pl-cer-m-pszenica-pasz"], "cabecera desconocida no se publica: %s" % ids)
check(all(len(s["points"]) >= 29 for s in tabs[0]["series"]) and tabs[0]["series"][1]["points"][0][0] < tabs[0]["series"][1]["points"][-1][0], "mensual ordenado")
check(["pl-cer-w-pszenica-kons", "pl-cer-w-kukurydza"] == [s["id"] for s in tabs[1]["series"]], "semanal")
# contrato
errs, warns, st = [], [], {}
CT.poland_zsrir(pz.doc(tabs, ed, "2026-10-23T00:00:00Z"), errs, warns, st)
check(not errs, "contrato: %s" % errs)
# el resumen no coincide -> no se publica
wb, last = cereal_book(summary_wheat=999.0)
try:
    pz.cereales(); check(False, "debia fallar con el resumen distinto")
except RuntimeError as e:
    check("no se publica" in str(e), "mensaje: %s" % e)
# el contrato rechaza periodos mezclados, semanas inexistentes y valores no positivos
bad = pz.doc([pz.mk_table("t", "g", ["a", "a", "a", "a"], "PLN/t", "", [pz.mk_series("s", ["a", "a", "a", "a"], {"2026-01": 1, "2026-W02": 2})])], "e", "2026-10-23T00:00:00Z")
errs = []; CT.poland_zsrir(bad, errs, [], {}); check(errs, "mezcla de periodos")
bad = pz.doc([pz.mk_table("t", "g", ["a", "a", "a", "a"], "PLN/t", "", [pz.mk_series("s", ["a", "a", "a", "a"], {"2025-W53": 1, "2026-W01": 2})])], "e", "2026-10-23T00:00:00Z")
errs = []; CT.poland_zsrir(bad, errs, [], {}); check(errs and "inexistente" in errs[0], "semana 53 de 2025 no existe: %s" % errs)
bad = pz.doc([pz.mk_table("t", "g", ["a", "a", "a", "a"], "PLN/t", "", [pz.mk_series("s", ["a", "a", "a", "a"], {"2026-01": 0, "2026-02": 2})])], "e", "2026-10-23T00:00:00Z")
errs = []; CT.poland_zsrir(bad, errs, [], {}); check(errs, "valor cero")

# --- clones: dos series que comparten >= 20 valores identicos se recortan a lo posterior
pa = {"2024-W%02d" % i: 100.0 + i for i in range(1, 31)}; pa.update({"2025-W%02d" % i: 200.0 + i for i in range(1, 6)})
pb = dict(pa); pb.update({"2025-W%02d" % i: 300.0 + i for i in range(1, 6)})
pb.pop("2024-W30"); pb["2024-W30"] = 5.0
pc = {"2024-W%02d" % i: 900.0 + i for i in range(1, 31)}
sa, sb, sc = (pz.mk_series(n, ["a", "a", "a", "a"], p_) for n, p_ in (("a", pa), ("b", pb), ("c", pc)))
out = pz.strip_clones([sa, sb, sc], "t")
check(out[2]["points"] == sc["points"], "la serie sin clon no se toca")
check(all(p[0] > "2024-W29" for p in out[0]["points"]) and len(out[0]["points"]) == 6 and len(out[1]["points"]) == 6, "clones recortados: %s" % out[0]["points"][:2])
few = pz.strip_clones([pz.mk_series("a", ["a"] * 4, {"2024-W01": 1, "2024-W02": 2}), pz.mk_series("b", ["a"] * 4, {"2024-W01": 1, "2024-W02": 2})], "t")
check(len(few[0]["points"]) == 2, "pocas coincidencias no recortan")


# --- mercados mayoristas: hoja de medias con otro nombre y columnas desplazadas
import datetime as _dt
hb = openpyxl.Workbook(); hs = hb.active; hs.title = "zmiany cen hurt"
hs.cell(6, 2, "Produkt"); hs.cell(6, 3, "Jedn."); hs.cell(6, 4, _dt.datetime(2026, 1, 8)); hs.cell(6, 6, _dt.datetime(2025, 12, 30))
def hrow(r, nm, unit=None, v=()):
    hs.cell(r, 2, nm)
    if unit: hs.cell(r, 3, unit)
    for j, x in enumerate(v): hs.cell(r, 4 + j, x)
hrow(9, "Warzywa krajowe"); hrow(10, "Marchew", "kg", (1.0, 2.0, 1.2, 1.8)); hrow(11, "Ziemniaki", "kg", (0.7, 1.1, None, None))
hrow(12, "Kapusta biała", "kg", (5.0, 40.0, 1.0, 1.4))
hrow(13, "Owoce krajowe"); hrow(14, "Gruszki", "kg", (4.0, 5.0, 4.2, 5.2)); hrow(15, "Jabłka:"); hrow(16, "Gala", "kg", (2.0, 3.0, 2.0, 3.0))
hrow(17, "Owoce importowe"); hrow(18, "Gruszki", "kg", (9.0, 9.5, 9.0, 9.5))
hm, hd = pz.hurt_avg(hb)
check(hd == [_dt.date(2026, 1, 8), _dt.date(2025, 12, 30)], "fechas de cabecera: %s" % hd)
check(hm["Marchew"] == {_dt.date(2026, 1, 8): (1.0, 2.0), _dt.date(2025, 12, 30): (1.2, 1.8)}, "mayorista: dos fechas: %s" % hm.get("Marchew"))
check(hm["Ziemniaki"] == {_dt.date(2026, 1, 8): (0.7, 1.1)}, "mayorista: fecha sin dato se omite")
check("Kapusta biała" in hm and _dt.date(2026, 1, 8) not in hm["Kapusta biała"], "mayorista: maximo > 3 veces el minimo se descarta")
check(hm["Gruszki"] == {_dt.date(2026, 1, 8): (4.0, 5.0), _dt.date(2025, 12, 30): (4.2, 5.2)}, "mayorista: pera nacional, no la importada: %s" % hm.get("Gruszki"))
check("Gala" not in hm, "mayorista: la manzana no se publica")
hs.cell(6, 4, _dt.datetime(2025, 12, 1))
try:
    pz.hurt_avg(hb); check(False, "fechas invertidas deben fallar")
except RuntimeError:
    check(True, "fechas invertidas rechazadas")
check(pz.iso_week(_dt.date(2025, 12, 30)) == "2026-W01" and pz.iso_week(_dt.date(2026, 10, 22)) == "2026-W43", "iso_week")

# --- year_rows, has_value, cell_value
yw = openpyxl.Workbook().active
yw["B7"] = "BYDŁO - M.P. UZUPEŁNIAJĄCE"
yw["B8"] = "2026"; yw["C8"] = 1454.32; yw["D8"] = "1431,08"; yw["E8"] = "nld"
yw["B9"] = 2025; yw["C9"] = 1535.47
yw["B18"] = "DRÓB - M.P. PEŁNOPORCJOWE"; yw["B19"] = "2026"; yw["C19"] = 1490.02
check(pz.year_rows(yw, "BYDŁO") == {"2026-01": 1454.32, "2026-02": 1431.08, "2025-01": 1535.47}, "year_rows seccion: %s" % pz.year_rows(yw, "BYDŁO"))
check(pz.year_rows(yw, "DRÓB") == {"2026-01": 1490.02}, "year_rows otra seccion")
check(pz.has_value(yw, 1454.32) and not pz.has_value(yw, 1454.4), "has_value")
cw = openpyxl.Workbook().active
cw["B10"] = "M.P. PEŁNOPORCJOWE"; cw["C10"] = "OGÓŁEM"; cw["D10"] = 1253.94
cw["B11"] = None; cw["C11"] = "PRESTARTER"; cw["D11"] = 1605.88
check(pz.cell_value(cw, r"PEŁNOPORCJOWE OGÓŁEM") == 1253.94 and pz.cell_value(cw, r"PRESTARTER") == 1605.88 and pz.cell_value(cw, r"NADA") is None, "cell_value")
# cabecera con fechas reales (datetime) en el resumen
dw = openpyxl.Workbook().active
dw["C6"] = D(2026, 10, 18); dw["D6"] = D(2026, 10, 11); dw["B7"] = "Masło Ekstra \nw blokach"; dw["C7"] = 1800.69; dw["D7"] = 1837.71
check(pz.summary_weekly(dw, r"^Masło Ekstra w blokach") == {"2026-10-18": 1800.69, "2026-10-11": 1837.71}, "cabecera datetime")
# una celda numerica 13.97 no es una fecha 13.97
nw = openpyxl.Workbook().active
nw["C8"] = "18.10.2026"; nw["B9"] = "POLSKA"; nw["B10"] = "Bydło ogółem"; nw["C10"] = 25540.75; nw["D10"] = 13.97
check(pz.summary_weekly(nw, r"\bBydło ogółem") == {"2026-10-18": 25540.75}, "numeros con forma de fecha no cuentan")

# --- fechas fuera de orden (año mal escrito en la hoja): la fila se omite
fw = openpyxl.Workbook().active
fw["B25"] = "TYDZIEŃ"; fw["C25"] = "Gala"; fw["D25"] = "Ligol"
seq = [(D(2026, 1, 25), 349, 1), (D(2025, 2, 2), 346.11, 2), (D(2026, 2, 8), 340, 3), (D(2026, 2, 15), "nld", 4), (D(2026, 2, 22), 324.33, 5)]
for i, (d, v, l) in enumerate(seq):
    fw.cell(row=26 + i, column=1, value=str(i)); fw.cell(row=26 + i, column=2, value=d); fw.cell(row=26 + i, column=3, value=v); fw.cell(row=26 + i, column=4, value=l)
wd = pz.wide_weekly(fw)
check(sorted(wd["Gala"]) == ["2026-W04", "2026-W06", "2026-W08"] and sorted(wd["Ligol"]) == ["2026-W04", "2026-W06", "2026-W07", "2026-W08"], "wide_weekly orden/sucios: %s" % {k: sorted(v) for k, v in wd.items()})
# accumulate: lo publicado se conserva y lo nuevo manda
pz.PREV["fruta.json"] = {"tables": [{"series": [{"id": "x", "points": [["2026-W01", 1], ["2026-W02", 2]]}]}]}
check(pz.accumulate("x", {"2026-W02": 22, "2026-W03": 3}) == {"2026-W01": 1, "2026-W02": 22, "2026-W03": 3} and pz.accumulate("y", {"a": 1}) == {"a": 1}, "accumulate")
pz.PREV.clear()

# --- main: un bloque que falla conserva lo publicado; sin cambios no reescribe
tmp = Path(tempfile.mkdtemp())
pz.OUT = tmp / "poland-zsrir"; pz.ROOT = tmp; pz.time.sleep = lambda s: None
(tmp / "data").mkdir()
wb, last = cereal_book()
pz.cereales.__globals__["workbook"] = lambda ds, tries=4: (wb, "2026-10-15", "x")


def colza_ok():
    return "semana 2026-10-18", [pz.mk_table("colza-mensual", "colza", ["c", "c", "c", "c"], "PLN/t", "", [pz.mk_series("pl-colza-m-semilla", ["c", "c", "c", "c"], {"2026-%02d" % m: 2000 + m for m in range(1, 10)})])]


def colza_ko():
    raise RuntimeError("sin red")


pz.colza = colza_ok
# main() usa los nombres de modulo: se sustituyen en el espacio de nombres
g = pz.main.__globals__
g["colza"] = colza_ok
for other in ("carne", "lacteos", "piensos", "fruta", "mayorista", "azucar"):
    g[other] = colza_ko
pz.main()
check((pz.OUT / "cereales.json").exists() and (pz.OUT / "colza.json").exists() and (pz.OUT / "index.json").exists(), "ficheros escritos")
before = (pz.OUT / "colza.json").read_text()
g["colza"] = colza_ko
log0 = len(pz.LOG)
pz.main()
check((pz.OUT / "colza.json").read_text() == before, "colza se conserva cuando falla")
check(any("conserva" in x for x in pz.LOG[log0:]), "el registro lo dice")
st = json.loads((tmp / "data" / "poland-zsrir-stats.json").read_text())
check(st["countries"]["PL"]["series"] and all(s["frequency"] == "monthly" for s in st["countries"]["PL"]["series"]), "estadisticas solo mensuales")
idx = json.loads((pz.OUT / "index.json").read_text())
errs = []; CT.poland_zsrir_index(idx, errs, [], {}); check(not errs, "indice: %s" % errs)
print("OK test-pl-zsrir")
