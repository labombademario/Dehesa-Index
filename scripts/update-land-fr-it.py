#!/usr/bin/env python3
"""Francia e Italia: precio de la tierra agrícola -> data/france-land/land.json y data/italy-land/land.json.

ITALIA (CREA, Indagine sul mercato fondiario in Italia):
  Página: https://www.crea.gov.it/en/web/politiche-e-bioeconomia/-/indagine-mercato-fondiario ; se toma el libro «MF Regione 1960-AAAA» más reciente.
  Hoja «Prezzi», Tab. 1: valor fondiario medio por región (€/ha corrientes; euro-liras hasta 2000), 1960-AAAA, más las cinco circoscrizioni y el total de Italia.
  Ojo: en el libro, desde 2000 las regiones vienen en MILES de €/ha y antes en €/ha (el total y las circoscrizioni, siempre en €/ha); aquí todo queda en €/ha y se comprueba la continuidad.
  Licencia: los informes del CREA dicen «È consentita la riproduzione citando la fonte» (registro de licencias, fuente crea_fondiario).
  Aviso del propio CREA: el rapporto 2024 revisó superficies y regiones agrarias y sus valores 2024 no son comparables con esta serie; esta serie llega hasta el último año del libro.
FRANCIA (SSP/Agreste con la FNSafer, «Valeur vénale des terres en AAAA», Chiffres et données):
  Conjunto data.gouv.fr «Agreste - Valeur vénale des terres» (Licence Ouverte). Los PDF de Agreste traen, por departamento y por pequeña región agrícola, el precio medio trienal de las
  terres et prés libres de más de 70 áreas y de los loués (€/ha, tres últimos años). La edición de julio 2025 (datos 2024) solo trae metodología (las tablas van en una hoja de cálculo
  que Agreste no enlaza de forma estable), así que se usa la última edición con tablas legibles. Se leen los departamentos (filas sin mínimo ni máximo).
Validaciones (si falla alguna no se escribe nada y queda en rojo): 20 regiones y 1960..último año completos, continuidad interanual, total de Italia en rango; en Francia al menos 85
departamentos de tierras libres, códigos reconocidos, variación impresa coherente con los precios, mediana en rango.
Uso: update-land-fr-it.py [--out DIR] [--it-file X.xlsx] [--fr-text T.txt]   (con ficheros no se toca la red; así lo usan los tests)"""
import datetime, io, json, os, re, subprocess, sys, tempfile, time, unicodedata, urllib.parse, urllib.request
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
UA = {"User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36"}
CREA = "https://www.crea.gov.it"
CREA_PAGE = CREA + "/en/web/politiche-e-bioeconomia/-/indagine-mercato-fondiario"
CREA_KNOWN = CREA + "/documents/68457/0/MF+Regione+1960-2023+%282024.09%29.xlsx/2baf48a0-0a57-c012-190c-32f91c4843f5?t=1728635105737"
AG = "https://agreste.agriculture.gouv.fr/agreste-web/download/publication/publie/"
AG_PUBLISHED = "https://www.data.gouv.fr/datasets/agreste-valeur-venale-des-terres/"
DRAAF = "https://draaf.pays-de-la-loire.agriculture.gouv.fr/IMG/pdf/"   # espejo de la DRAAF Pays de la Loire (solo para pruebas: desde los runners de GitHub da timeout)
AG_DETAIL = "https://www.agreste.agriculture.gouv.fr/agreste-web/disaron/%s/detail/"
DG_API = "https://www.data.gouv.fr/api/1/datasets/agreste-valeur-venale-des-terres/"
FR_EDITIONS = [  # (año de los precios, referencia Agreste, enlaces candidatos), de la edición más reciente a la más antigua; cada una se lee con la primera fuente que responde
    ("2025", "Chd2609", [AG + "Chd2609/cd2026-09_Valeur-venale-des-terres-en-2025.pdf", AG + "Chd2609/cd2026-09_VVT-2025.pdf"]),
    ("2024", "Chd2513", [AG + "Chd2513/cd2025-13_Valeur-venale-des-terres-en-2024.pdf"]),
    ("2023", "Chd2412", [AG + "Chd2412/cd2024-12_VVT-2023.pdf"]),
]
FR_PDFS = [(t, c[0]) for t, _, c in FR_EDITIONS]   # compatibilidad: primer enlace de cada edición
IT_REG = {"piemonte": "PIE", "valle d'aosta": "VDA", "lombardia": "LOM", "trentino a.a.": "TAA", "veneto": "VEN", "friuli v.g.": "FVG", "liguria": "LIG", "emilia romagna": "EMR",
          "toscana": "TOS", "umbria": "UMB", "marche": "MAR", "lazio": "LAZ", "abruzzo": "ABR", "molise": "MOL", "campania": "CAM", "puglia": "PUG", "basilicata": "BAS",
          "calabria": "CAL", "sicilia": "SIC", "sardegna": "SAR"}
IT_MACRO = {"nord-ovest": "NO", "nord-est": "NE", "centro": "CE", "meridione": "SU", "isole": "IS"}
LOG = []
def log(*a):
    s = " ".join(str(x) for x in a); LOG.append(s); print(s, flush=True)
def norm(s):
    s = unicodedata.normalize("NFD", str(s).lower()); s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", re.sub(r"[-'’\"“”–]", " ", s)).strip()
def http(url, tries=3, timeout=180):
    last = None
    for i in range(tries):
        try:
            with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=timeout) as r: return r.read()
        except Exception as e: last = e; time.sleep(5 * (i + 1))
    raise RuntimeError("%s -> %s" % (url, str(last)[:160]))

# ───────────────────────────── Italia ─────────────────────────────
def pick_crea(html):
    """Libro «MF Regione 1960-AAAA» con el año final más alto (y, a igualdad, el de marca de tiempo mayor)."""
    best = None
    for h in re.findall(r'href=["\']([^"\']+)["\']', html):
        d = urllib.parse.unquote(h.replace("+", " "))
        m = re.search(r"MF[ _]Regione[ _]1960-(\d{4}).*?\.xlsx", d, re.I)
        if not m: continue
        t = re.search(r"[?&]t=(\d+)", h); k = (int(m.group(1)), int(t.group(1)) if t else 0)
        if best is None or k > best[0]: best = (k, h if h.startswith("http") else CREA + h)
    if not best: raise RuntimeError("no se encontro el libro MF Regione 1960-AAAA en la pagina del CREA")
    return best[0][0], best[1]
def parse_italy(src):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(src) if isinstance(src, bytes) else src, data_only=True)
    if "Prezzi" not in wb.sheetnames: raise RuntimeError("el libro del CREA no tiene la hoja «Prezzi»")
    rows = list(wb["Prezzi"].iter_rows(values_only=True))
    hi = next((i for i, r in enumerate(rows) if r and r[0] and str(r[0]).strip().startswith("Regione/circoscrizione")), None)
    if hi is None: raise RuntimeError("no se encontro la cabecera de la Tab. 1")
    cols = [(j, int(h)) for j, h in enumerate(rows[hi]) if isinstance(h, (int, float)) and 1900 < h < 2100]
    years = [y for _, y in cols]
    if years != list(range(years[0], years[-1] + 1)) or years[0] != 1960: raise RuntimeError("anos de la Tab. 1 no consecutivos desde 1960: %s..%s" % (years[:2], years[-2:]))
    reg, macro, italy = {}, {}, None
    for r in rows[hi + 1:]:
        if not r or r[0] is None: continue
        nm = str(r[0]).strip()
        if nm.startswith("Fonte") or nm.startswith("Tab."): break
        key = nm.lower()
        vals = [r[j] if isinstance(r[j], (int, float)) else None for j, _ in cols]
        if key in IT_REG: reg[IT_REG[key]] = vals
        elif key in IT_MACRO: macro[IT_MACRO[key]] = vals
        elif key == "italia": italy = vals
    if set(reg) != set(IT_REG.values()): raise RuntimeError("regiones distintas de las 20 esperadas: faltan %s" % sorted(set(IT_REG.values()) - set(reg)))
    if italy is None or set(macro) != set(IT_MACRO.values()): raise RuntimeError("faltan el total de Italia o las circoscrizioni")
    # unidades: regiones desde 2000 en miles de €/ha; se pasan a €/ha y se exige continuidad
    for k, a in reg.items():
        for i, y in enumerate(years):
            if a[i] is None: raise RuntimeError("hueco en %s %d" % (k, y))
            if y >= 2000 and a[i] < 1000: a[i] = a[i] * 1000
        for i in range(1, len(a)):
            ch = a[i] / a[i - 1] - 1
            if not -0.40 < ch < 0.70: raise RuntimeError("salto interanual absurdo en %s %d (%+.0f %%): ¿cambio de unidades?" % (k, years[i], ch * 100))
    for i in range(1, len(italy)):
        if italy[i] is None or not -0.30 < italy[i] / italy[i - 1] - 1 < 0.40: raise RuntimeError("salto interanual absurdo en el total de Italia %d" % years[i])
    if not 8000 <= italy[-1] <= 60000: raise RuntimeError("total de Italia fuera de rango: %s" % italy[-1])
    for k, a in macro.items():
        for i in range(len(a)):
            if a[i] is not None and years[i] >= 1970 and a[i] < 1000: raise RuntimeError("circoscrizione %s %d con unidades dudosas" % (k, years[i]))
    r1 = lambda a: [None if v is None else round(v, 1) for v in a]
    return {"years": years, "regions": {k: r1(v) for k, v in sorted(reg.items())}, "macro": {k: r1(v) for k, v in sorted(macro.items())}, "italy": r1(italy)}

# ───────────────────────────── Francia ─────────────────────────────
NUM = re.compile(r"^(NS|ND|\d{1,3}(?: \d{3})+|\d+)$")
EVO = re.compile(r"^(NS|ND|[+-]?\s?\d+\s?%)$")
def pnum(t): return None if t in ("NS", "ND") else int(t.replace(" ", ""))
def pevo(t): return None if t in ("NS", "ND") else int(re.sub(r"[\s%+]", "", t))
def fr_rows(text):
    """[(seccion, nombre, [p1, p2, p3], variacion, [min, max])]; seccion = free | rent. Los departamentos son las filas sin mínimo ni máximo."""
    sec = None; pend = []; out = []
    for raw in text.split("\n"):
        ln = raw.rstrip()
        if not ln.strip() or ".." in ln: continue
        if re.search(r"Prix moyen des terres et prés libres", ln): sec = "free"; pend = []; continue
        if re.search(r"Prix moyen des terres et prés loués", ln): sec = "rent"; pend = []; continue
        if re.search(r"Prix moyen des terres agricoles dans les DOM|Prix moyen des vignes", ln): sec = None; pend = []; continue
        if sec is None: continue
        parts = [p.strip() for p in re.split(r"\s{2,}", ln.strip()) if p.strip()]
        first_is_val = bool(NUM.match(parts[0])) and len(parts) >= 4
        vals = parts if first_is_val else parts[1:]
        name = " ".join(pend) if first_is_val else parts[0]
        ok = len(vals) in (4, 6) and all(NUM.match(v) for v in vals[:3]) and EVO.match(vals[3]) and all(NUM.match(v) for v in vals[4:])
        if ok:
            if not first_is_val and pend and not re.match(r"^[A-ZÀ-Ý]", name): name = " ".join(pend) + " " + name
            out.append((sec, name, [pnum(v) for v in vals[:3]], pevo(vals[3]), [pnum(v) for v in vals[4:]])); pend = []
        elif len(parts) == 1 or not any(NUM.match(v) for v in parts[1:]): pend = (pend + [ln.strip()])[-3:]
        else: pend = []
    return out
def dept_codes():
    s = (ROOT / "vendor" / "fr-departements.js").read_text("utf-8")
    m = dict((norm(n), i) for i, n in re.findall(r'"id"\s*:\s*"([^"]+)"\s*,\s*"name"\s*:\s*"([^"]+)"', s))
    if len(m) < 90: raise RuntimeError("no se pudo leer la lista de departamentos de vendor/fr-departements.js")
    return m
def fr_year_header(text):
    m = re.search(r"\b(20\d\d)\s+(20\d\d)\s+(20\d\d)\b", text)
    if not m: raise RuntimeError("no se encontro la cabecera de anos de las tablas")
    y = [int(m.group(i)) for i in (1, 2, 3)]
    if y[1] != y[0] + 1 or y[2] != y[1] + 1: raise RuntimeError("anos de cabecera incoherentes %s" % y)
    return y
def parse_france(text):
    codes = dept_codes(); R = fr_rows(text)
    years = fr_year_header(text); res = {"free": {}, "rent": {}}; bad = []
    for sec, name, p, evo, mm in R:
        if mm: continue   # fila de pequeña región agrícola (trae mínimo y máximo)
        nm = norm(name.split(" - ")[-1]) if " - " in name else norm(name)
        code = codes.get(nm)
        if not code: bad.append(name); continue
        if p[1] and p[2] and evo is not None and abs(round((p[2] / p[1] - 1) * 100) - evo) > 1.01:
            raise RuntimeError("variacion impresa incoherente en %s %s: %s %s" % (sec, name, p, evo))
        if code in res[sec]: raise RuntimeError("departamento repetido en %s: %s" % (sec, name))
        res[sec][code] = p
    if bad: raise RuntimeError("departamentos no reconocidos: %s" % bad[:8])
    if len(res["free"]) < 85: raise RuntimeError("solo %d departamentos de tierras libres (se esperan >= 85)" % len(res["free"]))
    if len(res["rent"]) < 40: raise RuntimeError("solo %d departamentos de tierras alquiladas (se esperan >= 40)" % len(res["rent"]))
    v = sorted(p[2] for p in res["free"].values() if p[2]); med = v[len(v) // 2]
    if not 3000 <= med <= 15000: raise RuntimeError("mediana departamental fuera de rango: %s" % med)
    return years, res
def pdf_text(data):
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "x.pdf"; p.write_bytes(data)
        return subprocess.run(["pdftotext", "-layout", str(p), "-"], capture_output=True, text=True, check=True).stdout
def merge_fr(prev, years, res):
    """Acumula años: el PDF nuevo manda sobre el anterior (revisiones al log)."""
    old = (prev or {}).get("years", []); allY = sorted(set(old) | set(years))
    out = {}
    for sec in ("free", "rent"):
        out[sec] = {}
        for code in set(res[sec]) | set((prev or {}).get(sec, {})):
            arr = [None] * len(allY)
            for i, y in enumerate(old):
                if code in (prev or {}).get(sec, {}): arr[allY.index(y)] = prev[sec][code][i]
            if code in res[sec]:
                for k, y in enumerate(years):
                    j = allY.index(y)
                    if arr[j] is not None and res[sec][code][k] is not None and arr[j] != res[sec][code][k]: log("revision %s %s %d: %s -> %s" % (sec, code, y, arr[j], res[sec][code][k]))
                    if res[sec][code][k] is not None or arr[j] is None: arr[j] = res[sec][code][k]
            out[sec][code] = arr
    return allY, out

def probe_links(url, tag):
    """Diagnostico: enlaces de una pagina de Agreste que parecen datos (xls, xlsx, csv, ods, zip, pdf, download) para saber de donde salen las tablas nuevas."""
    try:
        html = http(url, tries=1, timeout=45).decode("utf-8", "replace")
        hs = [h for h in re.findall(r'(?:href|src)=["\']([^"\']+)["\']', html, re.I) if re.search(r"\.(xlsx?|csv|ods|zip|pdf)\b|download|export|telecharg", h, re.I)]
        log("sondeo %s: %d bytes, %d enlaces de datos" % (tag, len(html), len(hs)))
        for h in list(dict.fromkeys(hs))[:12]: log("  " + urllib.parse.urljoin(url, h)[:170])
        t = re.search(r"<title>(.*?)</title>", html, re.S | re.I)
        if t: log("  titulo: " + re.sub(r"\s+", " ", t.group(1))[:100])
    except Exception as e: log("sondeo %s: no legible (%s)" % (tag, str(e)[-100:]))
def probe_xlsx(url, tag, nrows=8):
    """Diagnostico: descarga un xlsx de Agreste y registra hojas, dimensiones y primeras filas, para escribir el lector sobre el formato real."""
    try:
        import io, openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(http(url, tries=2, timeout=90)), read_only=True, data_only=True)
        log("sondeo xlsx %s: %d hojas %s" % (tag, len(wb.sheetnames), wb.sheetnames[:12]))
        for n in wb.sheetnames[:4]:
            ws = wb[n]; log("  hoja %r %sx%s" % (n, ws.max_row, ws.max_column))
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= nrows: break
                log("    " + " | ".join("" if v is None else str(v)[:22] for v in row[:9]))
    except Exception as e: log("sondeo xlsx %s: fallo (%s)" % (tag, str(e)[-120:]))
def licence_ok(src):
    """Puerta de licencia: solo se publica lo que el registro marca VERIFIED y en uso."""
    try: v = json.loads((ROOT / "data" / "license-registry.json").read_text("utf-8"))["sources"][src]
    except Exception: return False
    return v.get("status") == "VERIFIED" and bool(v.get("used"))
def write(path, doc):
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".json.tmp"); tmp.write_text(json.dumps(doc, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8"); os.replace(tmp, path)
def main():
    args = sys.argv[1:]; outdir = ROOT / "data"
    if "--out" in args: outdir = Path(args[args.index("--out") + 1])
    itf = args[args.index("--it-file") + 1] if "--it-file" in args else None
    frt = args[args.index("--fr-text") + 1] if "--fr-text" in args else None
    fru = args[args.index("--fr-url") + 1] if "--fr-url" in args else None   # solo para sembrar con ficheros locales: enlace oficial del que proceden
    itu = args[args.index("--it-url") + 1] if "--it-url" in args else None
    stamp = datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"); rc = 0
    # Italia
    try:
        if itf: src, url, fname = itf, itu, Path(itf).name
        else:
            try: ly, url = pick_crea(http(CREA_PAGE).decode("utf-8", "replace"))
            except Exception as e: log("aviso: pagina del CREA no legible (%s); se usa el enlace conocido" % str(e)[:100]); url = CREA_KNOWN
            src = http(url, timeout=300); fname = urllib.parse.unquote(url.split("/")[-2] if "/" in url else url)
        it = parse_italy(src)
        doc = {"schemaVersion": 1, "generatedAt": stamp, "unit": "EUR/ha", "years": it["years"], "regions": it["regions"], "macro": it["macro"], "italy": it["italy"],
               "source": {"name": "CREA, Indagine sul mercato fondiario in Italia (Banca Dati dei Valori Fondiari)", "url": CREA_PAGE, "file": url or fname,
                          "licence": "CREA: «È consentita la riproduzione citando la fonte»",
                          "note": "Valor fondiario medio por región (€/ha corrientes; euro-liras hasta 2000). Encuesta a testigos cualificados, no un registro de compraventas. Serie con el método anterior a 2024: el CREA revisó superficies y regiones agrarias en el rapporto 2024 y sus valores 2024 no son comparables."}}
        if licence_ok("crea_fondiario") or "--force-italy" in args:
            write(outdir / "italy-land" / "land.json", doc)
            log("OK Italia: %d-%d, 20 regiones, total %s €/ha (%d)" % (it["years"][0], it["years"][-1], it["italy"][-1], it["years"][-1]))
        else: log("Italia leida y validada (%d-%d) pero NO publicada: la licencia de crea_fondiario no esta VERIFIED en data/license-registry.json" % (it["years"][0], it["years"][-1]))
    except Exception as e: log("ERROR Italia:", e); rc = 1
    # Francia
    try:
        prev = None; pf = outdir / "france-land" / "land.json"
        if pf.exists():
            try: prev = json.loads(pf.read_text("utf-8"))
            except Exception: prev = None
        parsed = None; editions = []
        if frt: parsed = (parse_france(Path(frt).read_text("utf-8")), "texto " + Path(frt).name, fru)
        else:
            dg = []   # PDF de la ficha de data.gouv.fr, por si los enlaces directos cambian
            try:
                for r in json.loads(http(DG_API, tries=2, timeout=60)).get("resources", []): log("data.gouv recurso: %s %s" % (r.get("title"), r.get("url")))
            except Exception as e: log("data.gouv: ficha no legible (%s)" % str(e)[:120])
            for tag, ref, cands in FR_EDITIONS:
                found = []
                try:   # enlaces del PDF en la ficha de Agreste de esa edición
                    html = http(AG_DETAIL % ref, tries=1, timeout=45).decode("utf-8", "replace")
                    for h in re.findall(r'href=["\']([^"\']+\.pdf)["\']', html, re.I): found.append(urllib.parse.urljoin(AG_DETAIL % ref, h))
                    log("Francia %s: ficha de Agreste con %d PDF" % (tag, len(found)))
                except Exception as e: log("Francia %s: ficha de Agreste no legible (%s)" % (tag, str(e)[:120]))
                got = None
                for u in found + cands:
                    try: data = http(u, tries=2, timeout=240)
                    except Exception as e: log("Francia %s: no se pudo descargar %s (%s)" % (tag, u.rsplit("/", 1)[-1], str(e)[-110:])); continue
                    if data[:4] != b"%PDF": log("Francia %s: %s no es un PDF" % (tag, u.rsplit("/", 1)[-1])); continue
                    try: got = (parse_france(pdf_text(data)), "edicion " + tag, u)
                    except Exception as e: log("Francia %s: %s sin tablas utilizables (%s)" % (tag, u.rsplit("/", 1)[-1], str(e)[:100])); continue
                    break
                if got: editions.append(got); log("Francia %s: leida desde %s" % (tag, got[2]))
                else: probe_links(AG_DETAIL % ref, "ficha " + ref)
            B = "https://www.agreste.agriculture.gouv.fr/agreste-web/download/publication/publie/Chd2609/DonneesAssociees-cd2026-9_VVT-2025_Publication-"
            probe_xlsx(B + "LIBRE.xlsx", "2025 LIBRE"); probe_xlsx(B + "LOUE.xlsx", "2025 LOUE")
            if not editions or editions[0][1] != "edicion " + FR_EDITIONS[0][0]: probe_links("https://agreste.agriculture.gouv.fr/agreste-web/disaron/venale/745fcc0d-4e89-4a99-b151-835d2a9be2df/search/", "disar venale")
            if editions: parsed = editions[0]
        if frt: editions = [parsed]
        if not parsed: raise RuntimeError("ninguna edicion de Agreste con tablas utilizables")
        (years, res), label, url = parsed
        allY, merged = None, None
        for (yy, rr), _l, _u in reversed(editions):   # de la más antigua a la más reciente: la nueva manda y las revisiones quedan en el log
            allY, merged = merge_fr(prev if merged is None else {"years": allY, "free": merged["free"], "rent": merged["rent"]}, yy, rr)
        doc = {"schemaVersion": 1, "generatedAt": stamp, "unit": "EUR/ha", "years": allY, "free": merged["free"], "rent": merged["rent"],
               "source": {"name": "SSP (Agreste) con la FNSafer, Valeur vénale des terres", "url": AG_PUBLISHED, "file": url or label, "licence": "Licence Ouverte (data.gouv.fr)",
                          "note": "Precio medio trienal de las terres et prés libres (más de 70 áreas) y de los loués por departamento, en euros por hectárea del año central; observatorio de los Safer sobre las notificaciones de venta. Departamentos sin dato fiable (NS/ND) quedan en blanco."}}
        write(outdir / "france-land" / "land.json", doc)
        log("OK Francia: %s, años %s, %d departamentos libres, %d alquilados" % (label, allY, len(merged["free"]), len(merged["rent"])))
    except Exception as e: log("ERROR Francia:", e); rc = 1
    return rc
if __name__ == "__main__":
    rc = main()
    if "--out" not in sys.argv:
        try: (ROOT / "data" / "land-fr-it-log.txt").write_text("\n".join(LOG) + "\n", "utf-8")
        except Exception: pass
    sys.exit(rc)
