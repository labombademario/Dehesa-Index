#!/usr/bin/env python3
"""Polonia: gasoleo y fioul de calefaccion, y coste aproximado del gasoleo para una explotacion con la devolucion del impuesto especial
(zwrot podatku akcyzowego) -> data/pl-diesel.json
1) Precios semanales de Polonia del Boletin Petrolero de la UE (Comision Europea, DG ENER, historico desde 2005): gasoleo con y sin impuestos
   y fioul de calefaccion con impuestos, en EUR por 1.000 l, y el tipo de cambio EUR/PLN de esa semana (para expresarlos en zloty).
   Reutilizacion: Decision 2011/833/UE.
2) Devolucion por litro (zl/l) y limites por hectarea y por cabeza de ganado: reglamento anual del Consejo de Ministros (scripts/ref/pl-akcyza.json,
   se actualiza a mano cada enero; es un dato legal).
3) Coste aproximado para la explotacion (zl/l) = precio con impuestos / (1 + IVA) - devolucion (los agricultores en regimen general recuperan el IVA).
   El IVA normal es del 23 %, pero en 2026 Polonia lo ha bajado al 8 % en varios tramos (paquete «CPN»; desde el 3-10-2026 hasta el 31-12-2026).
   El tipo de cada semana se deduce del propio boletin: con IVA del 23 %, precio/1,23 - precio sin impuestos (impuesto especial y recargo de
   combustible) ronda 0,39 EUR/l; si sale por debajo de 0,30 la semana tuvo IVA reducido del 8 %. Calculo de Dehesa Index. Uso: python3 scripts/update-pl-diesel.py"""
import datetime, io, json, sys, time, urllib.request
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "pl-diesel.json"
WOB = "https://energy.ec.europa.eu/document/download/906e60ca-8b6a-44e7-8589-652854d2fd3f_en?filename=Weekly_Oil_Bulletin_Prices_History_maticni_4web.xlsx"
H = {"User-Agent": "Dehesa-Index-data-bot/1.0 (+https://dehesaindex.com)"}
VAT = 0.23
def vat_of(r):
    """Tipo de IVA de la semana, deducido de la diferencia entre el precio con y sin impuestos (ver arriba)."""
    if r[2] is None: return VAT
    return 0.08 if r[1] / 1.23 - r[2] < 0.30 else VAT
def get(u, t=300):
    last = None
    for i in range(4):
        try:
            with urllib.request.urlopen(urllib.request.Request(u, headers=H), timeout=t) as r: return r.read()
        except Exception as e: last = e; time.sleep(6 * (i + 1))
    raise RuntimeError("%s -> %s" % (u[-90:], last))
def cols(ws, names):
    hdr = [str(x or "").strip() for x in next(ws.iter_rows(values_only=True, max_row=1))]
    return [hdr.index(n) for n in names]
def series(ws, ci, scale=1000.0):
    out = {}
    for r in ws.iter_rows(values_only=True, min_row=2):
        d = r[0]
        if not isinstance(d, datetime.datetime) or ci >= len(r): continue
        try: v = float(r[ci]) / scale
        except (TypeError, ValueError): continue
        out[d.date().isoformat()] = v
    return out
def main():
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(get(WOB)), read_only=True, data_only=True)
        wt = wb["Prices with taxes"]; wo = wb["Prices wo taxes"]
        cx, cd, ch = cols(wt, ["PL_exchange_rate", "PL_price_with_tax_diesel", "PL_price_with_tax_heating_oil"]); (cn,) = cols(wo, ["PL_price_wo_tax_diesel"])
        fx = series(wt, cx, 1.0); dt = series(wt, cd); ht = series(wt, ch); dn = series(wo, cn)
        rows = []
        for d in sorted(dt):
            x = fx.get(d)
            if not x or not 0.15 < x < 0.4 or not 0.4 < dt[d] < 3.0: continue   # el boletin da el cambio como euros por zloty
            rows.append([d, round(dt[d], 4), round(dn[d], 4) if d in dn else None, round(ht[d], 4) if d in ht else None, round(1 / x, 4)])
        if len(rows) < 500: raise ValueError("serie corta: %d semanas" % len(rows))
    except Exception as e:
        print("FALLO boletin petrolero:", e); return 0 if OUT.exists() else 1
    ref = json.loads((ROOT / "scripts" / "ref" / "pl-akcyza.json").read_text(encoding="utf-8"))
    y = max(ref["years"]); R = ref["years"][y]
    last = rows[-1]; pln = last[1] * last[4]; v = vat_of(last)
    doc = {"schemaVersion": 1, "generatedAt": datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"), "asOf": last[0], "unit": "EUR/l",
           "source": {"name": "Comisión Europea, Weekly Oil Bulletin (Polonia)", "url": "https://energy.ec.europa.eu/data-and-analysis/weekly-oil-bulletin_en", "license": "Reutilización de documentos de la Comisión (Decisión 2011/833/UE)"},
           "fields": ["date", "dieselTax", "dieselNoTax", "heatingOilTax", "plnPerEur"], "weekly": rows, "vat": v, "vatNormal": VAT, "reducedVatWeeks": [r[0] for r in rows if r[0] >= "2026-01-01" and vat_of(r) != VAT],
           "refund": dict(R, year=int(y)), "farmCostPLN": round(pln / (1 + v) - R["zlPerLitre"], 3), "dieselPLN": round(pln, 3)}
    OUT.write_text(json.dumps(doc, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print("pl-diesel.json: %d semanas %s..%s; %s gasoleo %.3f EUR/l = %.2f zl/l; devolucion %s zl/l (%s); coste aprox. %.2f zl/l (IVA %d %%)" % (len(rows), rows[0][0], last[0], last[0], last[1], pln, R["zlPerLitre"], y, doc["farmCostPLN"], round(v * 100)))
    return 0
if __name__ == "__main__": sys.exit(main())
