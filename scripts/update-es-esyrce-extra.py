#!/usr/bin/env python3
"""España: ESYRCE, tablas complementarias por provincia y comunidad (MAPA)
-> data/spain-esyrce/extra-<año>.json + data/spain-esyrce/extra-index.json

Fuente: libros de la página ESYRCE del MAPA
  3.11 Estimaciones de rendimientos de los principales cultivos (.xls): nº de parcelas aforadas y rendimiento (kg/ha) por provincia, cultivo y régimen
  3.7  Tipificación de regadíos e invernaderos: ha por sistema de riego (gravedad, aspersión, automotriz, localizado) y por nivel de tecnificación de invernaderos
  3.8  Superficie de cultivos leñosos según el estado de la producción: ha por cultivo (1.er año, joven, en producción, abandonada, no comercial)
  3.6  Cubiertas del suelo en cultivos leñosos y barbechos (solo el año más reciente): ha por técnica de mantenimiento del suelo
El año se toma de la carpeta de la URL (.../esyrce/AAAA/...); los libros sin año en la URL se asignan al año siguiente al mayor con carpeta, como en update-es-esyrce.py.
Nada se estima: celda vacía o «-» = null; un rendimiento con 0 parcelas aforadas = null. Las hojas que no cumplen la estructura esperada se ignoran y se anotan en el log."""
import argparse, datetime, importlib.util, io, json, re, sys
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("esy", str(ROOT / "scripts" / "update-es-esyrce.py")); esy = importlib.util.module_from_spec(spec); spec.loader.exec_module(esy)
http, slug, clean, page_files, PAGE, CCAA = esy.http, esy.slug, esy.clean, esy.page_files, esy.PAGE, esy.CCAA
ALT = dict(esy.ALT); ALT.update({"p-de-asturias": "principado-de-asturias", "c-foral-de-navarra": "comunidad-foral-de-navarra", "c-valenciana": "comunitat-valenciana", "comunidad-valenciana": "comunitat-valenciana",
                                 "navarra": "comunidad-foral-de-navarra", "murcia": "region-de-murcia", "madrid": "comunidad-de-madrid", "baleares": "illes-balears", "asturias": "principado-de-asturias"})
def log(*a): esy.log(*a)
def tid(name): s = slug(name); return ALT.get(s, s)
def tkind(name):
    t = tid(name)
    return "country" if t == "espana" else "ccaa" if t in CCAA else "province"
def classify(name):
    n = name.lower().replace(" ", "_")
    if re.search(r"(^|[_/])3\.?11[_.]?(estimaciones|[a-z_]*rendimientos)", n) or re.search(r"(^|[^0-9.])311_?estimaciones", n): return "yields"
    if re.search(r"(^|[_/])3\.?7[_.]?tipificaci", n) or re.search(r"(^|[^0-9.])37tipificacion", n): return "irrigation"
    if re.search(r"(^|[_/])(3\.?)?8[_.]?superficie_de_cultivos_le", n) or re.search(r"(^|[^0-9.])38_?superficie[_a-z]*lenosos", n): return "woody"
    if re.search(r"(^|[_/])3\.6[_.]?cubiertas", n): return "soil"
    return None
REG = {"secano": "secano", "regadio": "regadio", "regadío": "regadio", "invernadero": "invernadero", "inverandero": "invernadero"}
def num(v, positive=False):
    if isinstance(v, bool) or not isinstance(v, (int, float)): return None
    v = float(v)
    if positive and v <= 0: return None
    return round(v, 1)
def is_terr_text(v): return isinstance(v, str) and v.strip() != "" and not v.strip().lower().startswith(("provincias", "cultivos", "tecnicas", "técnicas"))

# ---------------- 3.11 rendimientos (.xls con celdas combinadas) ----------------
def parse_yields(data):
    import xlrd
    book = xlrd.open_workbook(file_contents=data, formatting_info=True)
    out = []; terr_names = None
    for s in book.sheets():
        try: res = parse_yield_sheet(s)
        except Exception as e: log("  hoja", s.name, "error:", str(e)[:90]); continue
        if not res: log("  hoja", s.name, "ignorada (estructura distinta)"); continue
        names, crops = res
        if not crops: continue
        out.append((s.name, names, crops))
    return out
def drop_copies(sheets):
    """Si dos cultivos distintos traen casi las mismas cifras (≥90 % iguales en ≥10 provincias con dato) es un error de copia en el libro original: se conserva el primero y se descarta el resto."""
    kept = []; res = []
    for sname, names, items in sheets:
        keep = []
        for it in items:
            dup = None
            if sum(1 for v in it["y"] if v is not None) >= 10:
                for (nm, rg, y, nn) in kept:
                    if nn != len(names) or nm == it["name"]: continue
                    both = [(a, b) for a, b in zip(it["y"], y) if a is not None and b is not None]
                    if len(both) >= 10 and sum(1 for a, b in both if a == b) >= 0.9 * len(both): dup = nm; break
            if dup: log("  descartado", it["name"], it["regime"], "- cifras idénticas a", dup, "(error de copia en el libro original)"); continue
            keep.append(it)
            if sum(1 for v in it["y"] if v is not None) >= 10: kept.append((it["name"], it["regime"], it["y"], len(names)))
        res.append((sname, names, keep))
    return res
def parse_yield_sheet(s):
    nr, nc = s.nrows, s.ncols
    grid = [[s.cell_value(r, c) for c in range(nc)] for r in range(nr)]
    for rlo, rhi, clo, chi in s.merged_cells:   # el valor de la celda superior izquierda se propaga a todo el rango combinado
        v = grid[rlo][clo]
        for r in range(rlo, min(rhi, nr)):
            for c in range(clo, min(chi, nc)):
                if grid[r][c] == "": grid[r][c] = v
    hr = next((r for r in range(min(nr, 12)) if any(isinstance(x, str) and x.strip().lower().startswith("provincias") for x in grid[r])), None)
    if hr is None: return None
    tcols = [c for c, x in enumerate(grid[hr]) if isinstance(x, str) and x.strip().lower().startswith("provincias")]
    tcols = [c for k, c in enumerate(tcols) if k == 0 or c != tcols[k - 1] + 1]   # la celda «Provincias» puede estar combinada en dos columnas
    rr = None
    for r in range(hr + 1, min(nr, hr + 6)):   # fila de régimen: mayoría de celdas en {Secano, Regadío, Invernadero}
        cells = [clean(x).lower() for x in grid[r] if isinstance(x, str) and clean(x)]
        if cells and sum(1 for x in cells if x in REG) >= max(2, 0.6 * len(cells)): rr = r; break
    if rr is None: return None
    d0 = rr + 1
    panels = []
    for k, tc in enumerate(tcols):
        end = tcols[k + 1] if k + 1 < len(tcols) else nc
        names = []; r = d0
        while r < nr and is_terr_text(grid[r][tc]): names.append(clean(grid[r][tc])); r += 1
        if len(names) < 10: log("  ", s.name, "panel en la columna", tc, "sin territorios completos: se ignora"); continue   # p. ej. el segundo panel de hortícolas, desalineado en el original
        panels.append((tc, end, names))
    if not panels: return None
    base = panels[0][2]
    crops = []
    for tc, end, names in panels:
        if [tid(x) for x in names] != [tid(x) for x in base]: log("  ", s.name, "panel con territorios distintos: se ignora"); continue
        crop_row = None   # fila de nombres de cultivo: la más cercana por encima de la cabecera con texto en mayúsculas dentro del panel
        for r in range(hr - 1, -1, -1):
            cells = [x for x in grid[r][tc + 1:end] if isinstance(x, str) and clean(x)]
            if cells and not any("parcela" in x.lower() or "rendimiento" in x.lower() for x in cells) and any(x.strip().upper() == x.strip() for x in cells): crop_row = r; break
        if crop_row is None: return None
        cur = ""
        for c in range(tc + 1, end):
            kindtxt = clean(grid[hr][c]).lower(); reg = REG.get(clean(grid[rr][c]).lower())
            if c < len(grid[crop_row]) and isinstance(grid[crop_row][c], str) and clean(grid[crop_row][c]): cur = clean(grid[crop_row][c])
            if not reg or not kindtxt: continue
            kind = "n" if "parcela" in kindtxt else "y" if "rendimiento" in kindtxt else None
            if not kind or not cur: continue
            crops.append({"crop": cur, "regime": reg, "kind": kind, "col": [grid[r][c] for r in range(d0, d0 + len(names))]})
    res = {}; order = []
    for c in crops:   # parcelas y rendimiento de un mismo cultivo y régimen: se toma la primera columna de cada tipo
        key = (c["crop"], c["regime"])
        res.setdefault(key, {"n": None, "y": None})
        if key not in order: order.append(key)
        if res[key][c["kind"]] is None: res[key][c["kind"]] = c["col"]
    items = []
    for key in order:
        v = res[key]
        if v["y"] is None: continue
        n = [num(x) for x in v["n"]] if v["n"] else [None] * len(base)
        y = [num(x, True) if (not v["n"] or n[j] is None or n[j] > 0) else None for j, x in enumerate(v["y"])]
        if not any(x is not None for x in y): continue
        items.append({"name": key[0], "regime": key[1], "n": n, "y": y})
    return base, items

# ---------------- 3.7 / 3.8 / 3.6 (xlsx) ----------------
def sheet_rows(ws): return [list(r) for r in ws.iter_rows(values_only=True)]
def parse_irrigation(data):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    res = {}
    for ws in wb:
        rows = sheet_rows(ws)
        if not rows: continue
        head = [clean(x) for x in rows[0]]
        if "Provincias y CC.AA" not in head: continue
        tc = head.index("Provincias y CC.AA")
        i0 = next((i for i, h in enumerate(head) if re.match(r"total (regad|invern)", h.lower())), None)
        if i0 is None: continue
        cols = [(j, h) for j, h in enumerate(head) if j >= i0 and h]
        # las columnas «Superficie en invernadero» sueltas de la hoja de regadíos se ignoran: el nivel de tecnificación está en la hoja «invernaderos»
        terrs = []; vals = {h: [] for _, h in cols}
        for r in rows[1:]:
            if len(r) <= tc or not is_terr_text(r[tc]): continue
            terrs.append(clean(r[tc]))
            for j, h in cols: vals[h].append(num(r[j]) if j < len(r) else None)
        key = "irrigation" if "regad" in ws.title.lower() else "greenhouses" if "invern" in ws.title.lower() else None
        if key: res[key] = (terrs, {h: v for h, v in vals.items() if any(x is not None for x in v)})
    return res
def parse_woody(data):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    out = []; base = None
    for ws in wb:
        rows = sheet_rows(ws)
        if len(rows) < 4 or not rows[0] or clean(rows[0][0]) != "Provincias y CC.AA": continue
        crop = clean(rows[0][1]); states = [clean(x) for x in rows[1][1:7]]
        if not crop or len(states) != 6 or not states[0].lower().startswith("plantaci"): continue
        names = []; cols = [[] for _ in range(7)]
        for r in rows[2:]:
            if not r or not is_terr_text(r[0]):
                if names: break
                continue
            names.append(clean(r[0]))
            for k in range(7): cols[k].append(num(r[1 + k]) if len(r) > 1 + k else None)
        if base is None: base = names
        elif [tid(x) for x in names] != [tid(x) for x in base]: log("  3.8 hoja con territorios distintos:", ws.title); continue
        out.append({"crop": crop.title(), "states": states, "v": cols})
    return base or [], out
def parse_soil(data):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    res = {}
    for ws in wb:
        t = ws.title.lower()
        key = "fallow" if re.match(r"3\.6\.2\.1", t) else "woody" if re.match(r"3\.6\.2\.2", t) else None
        if not key: continue
        rows = sheet_rows(ws); hr = next((i for i, r in enumerate(rows[:6]) if any(isinstance(x, str) and x.strip().lower().startswith("provincias") for x in r)), None)
        if hr is None: continue
        tc = next(j for j, x in enumerate(rows[hr]) if isinstance(x, str) and x.strip().lower().startswith("provincias"))
        hdr = None
        for r in rows[hr + 1: hr + 3]:
            if sum(1 for x in r if isinstance(x, str) and x.strip()) >= 4: hdr = r; break
        if hdr is None: continue
        cols = [(j, clean(h)) for j, h in enumerate(hdr) if isinstance(h, str) and clean(h) and j != tc]
        terrs = []; vals = {h: [] for _, h in cols}
        for r in rows[hr + 1:]:
            if len(r) <= tc or not is_terr_text(r[tc]): continue
            terrs.append(clean(r[tc]))
            for j, h in cols: vals[h].append(num(r[j]) if j < len(r) else None)
        res[key] = (terrs, {h: v for h, v in vals.items() if any(x is not None for x in v)})
    return res

def doc_territories(names): return [{"id": tid(t), "name": t, "kind": tkind(t)} for t in names]
def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--out", default=str(ROOT / "data" / "spain-esyrce")); ap.add_argument("--local", action="append", default=[], help="año:tabla:ruta (tabla = yields|irrigation|woody|soil)")
    a = ap.parse_args(); out = Path(a.out); out.mkdir(parents=True, exist_ok=True)
    now = datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    sets = {}
    if a.local:
        for s in a.local:
            y, m, p = s.split(":", 2); sets.setdefault(int(y), {})[m] = (p, Path(p).name)
    else:
        items = page_files(http(PAGE).decode("utf-8", "replace")); folder = {}; nofolder = []
        for u, n, t in items:
            m = classify(n)
            if not m: continue
            fy = re.search(r"/esyrce/(20\d\d)/", u)
            (folder.setdefault(int(fy.group(1)), {}).setdefault(m, (u, n))) if fy else nofolder.append((m, u, n))
        for y, d in folder.items(): sets.setdefault(y, {}).update(d)
        if nofolder:
            y = (max(folder) + 1) if folder else datetime.date.today().year - 1
            for m, u, n in nofolder: sets.setdefault(y, {}).setdefault(m, (u, n))
            log("libros sin año en la URL asignados a", y)
    if not sets: raise SystemExit("no hay libros ESYRCE complementarios: no se toca nada")
    src = {"id": "mapa_es", "name": "MAPA, ESYRCE", "url": PAGE}
    for y in sorted(sets):
        doc = {"schemaVersion": 1, "generatedAt": now, "year": y, "source": src}
        terr_ref = None
        def ref(names):
            nonlocal terr_ref
            if terr_ref is None: terr_ref = names; return True
            return [tid(x) for x in names] == [tid(x) for x in terr_ref]
        for m in ("yields", "irrigation", "woody", "soil"):
            if m not in sets[y]: continue
            u, n = sets[y][m]
            try: data = Path(u).read_bytes() if a.local else http(u)
            except Exception as e: log("ERROR descarga", y, m, n[:60], str(e)[:100]); continue
            try:
                if m == "yields":
                    sheets = drop_copies(parse_yields(data))
                    crops = []; sets_ = []
                    for sname, names, items in sheets:
                        sig = [tid(x) for x in names]
                        if sig not in [x[0] for x in sets_]: sets_.append((sig, names))
                        k = [x[0] for x in sets_].index(sig)
                        for it in items: it["group"] = clean(sname).title(); it["set"] = k
                        crops += items
                    if crops:
                        doc["yieldSets"] = [doc_territories(nm) for _, nm in sets_]; doc["yields"] = crops
                        log("leído", y, m, n[:50], len(crops), "series cultivo×régimen", len(sets_), "conjuntos de territorios")
                elif m == "irrigation":
                    r = parse_irrigation(data)
                    for k, (terrs, vals) in r.items():
                        doc[k] = {"territories": doc_territories(terrs), "measures": vals}; log("leído", y, m, k, len(terrs), "territorios", len(vals), "columnas")
                elif m == "woody":
                    names, crops = parse_woody(data)
                    if crops: doc["woody"] = {"territories": doc_territories(names), "crops": crops}; log("leído", y, m, len(crops), "cultivos")
                elif m == "soil":
                    r = parse_soil(data)
                    for k, (terrs, vals) in r.items():
                        doc.setdefault("soil", {})[k] = {"territories": doc_territories(terrs), "measures": vals}; log("leído", y, m, k, len(terrs), "territorios", len(vals), "técnicas")
            except Exception as e: log("ERROR lectura", y, m, n[:60], str(e)[:120])
        keys = [k for k in ("yields", "irrigation", "greenhouses", "woody", "soil") if k in doc]
        if not keys: log("sin tablas para", y); continue
        doc["tables"] = keys
        (out / ("extra-%d.json" % y)).write_text(json.dumps(doc, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
    idx = {"years": {}}
    for fp in sorted(out.glob("extra-20*.json")):
        m = re.fullmatch(r"extra-(20\d\d)\.json", fp.name)
        try: o = json.loads(fp.read_text(encoding="utf-8")); idx["years"][m.group(1)] = {"file": fp.name, "tables": o["tables"]}
        except Exception as e: log("no se pudo reindexar", fp.name, str(e)[:80])
    if not idx["years"]: raise SystemExit("no se leyó ningún año: no se toca nada")
    idx.update({"schemaVersion": 1, "generatedAt": now, "source": src,
                "note": "Complemento de ESYRCE por provincia. Rendimientos: kg/ha estimados con las parcelas aforadas (no es la estadística oficial de producción). El último año es avance provisional del MAPA."})
    (out / "extra-index.json").write_text(json.dumps(idx, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
    (ROOT / "data" / "spain-esyrce-extra-log.txt").write_text("\n".join(esy.LOG) + "\n", encoding="utf-8")
if __name__ == "__main__": main()
