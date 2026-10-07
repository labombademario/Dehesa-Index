#!/usr/bin/env python3
"""Francia: estado semanal del maiz grano (Cere'Obs, FranceAgriMer) -> data/france-cereobs.json
Excel historico "depuis 2015" publicado en data.gouv.fr (la direccion se lee del conjunto de datos). Nacional todas las semanas desde 2015; regiones los ultimos 6 anos.
Por semana: % de superficie en cada etapa (siembra, emergencia, 6-8 hojas, floracion femenina, humedad del grano 50 %, cosecha) y % por estado (muy malo ... muy bueno). Vacios se omiten (null).
Licencia: data.gouv declara Licence Ouverte 2.0 pero el propio fichero dice 'tous droits reserves'; el registro la mantiene VERIFIED por decision del propietario (4 oct 2026, OWNER_ACCEPTED); la contradiccion sigue sin aclararse con FranceAgriMer.
Modo sin red: FR_FIXTURES=<carpeta con cereobs.xlsx>."""
import datetime, io, json, os, re, sys, time, urllib.request
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
DS = "https://www.data.gouv.fr/api/1/datasets/historique-de-letat-des-cultures-cereobs-mais/"
FIELDS = [("Semis", "sowing"), ("Levée", "emergence"), ("6/8 feuilles visibles", "leaves"), ("Floraison femelle", "flowering"), ("Humidité du grain 50%", "moisture50"), ("Récolte", "harvest"),
          ("Très mauvaises", "c1"), ("Mauvaises", "c2"), ("Assez bonnes", "c3"), ("Bonnes", "c4"), ("Très bonnes", "c5")]
LOG = []; FIX = os.environ.get("FR_FIXTURES")
def log(*a):
    s = " ".join(str(x) for x in a); LOG.append(s); print(s, flush=True)
def http(url, accept="*/*"):
    last = None
    for i in range(3):
        try:
            with urllib.request.urlopen(urllib.request.Request(url, headers={"User-Agent": "DehesaIndex/1.0 (+https://dehesaindex.com)", "Accept": accept}), timeout=180) as r: return r.read()
        except Exception as e: last = e; time.sleep(5 * (i + 1))
    raise RuntimeError("%s: %s" % (url[:80], str(last)[:160]))
def workbook():
    import openpyxl
    if FIX: return openpyxl.load_workbook(Path(FIX) / "cereobs.xlsx", data_only=True), "fixture"
    ds = json.loads(http(DS, "application/json").decode("utf-8"))
    res = [r for r in ds.get("resources", []) if (r.get("format") or "").lower() == "xlsx"]
    if not res: raise RuntimeError("data.gouv no lista ningun recurso xlsx")
    log("recurso", res[0]["title"], res[0].get("last_modified"))
    return openpyxl.load_workbook(io.BytesIO(http(res[0]["url"])), data_only=True), res[0].get("last_modified")
def num(v):
    return round(float(v), 2) if isinstance(v, (int, float)) and not isinstance(v, bool) else None
def table(ws, key):
    rows = list(ws.iter_rows(values_only=True)); h = next(i for i, r in enumerate(rows) if r and "Semaine" in r)
    head = [str(c).strip() if c is not None else "" for c in rows[h]]; wi = head.index("Semaine"); ri = head.index(key) if key else None
    idx = []
    for fr, k in FIELDS:
        if fr not in head: raise RuntimeError("columna ausente en la hoja %s: %s" % (ws.title, fr))
        idx.append(head.index(fr))
    out = {}
    for r in rows[h + 1:]:
        m = re.fullmatch(r"(\d{4})-S(\d{2})", str(r[wi]).strip()) if r and r[wi] else None
        if not m: continue
        vals = [num(r[j]) for j in idx]
        if all(v is None for v in vals): continue
        out.setdefault(r[ri] if ri is not None else "FR", []).append(["%s-W%s" % m.groups()] + vals)
    return out
def build():
    wb, mod = workbook(); y = datetime.date.today().year
    nat = table(wb["Données France"], None)["FR"]
    reg = table(wb["Données régions"], "Région")
    reg = {k: [r for r in v if int(r[0][:4]) >= y - 5] for k, v in reg.items()}
    stated = None
    pres = [c for r in wb["Présentation"].iter_rows(values_only=True) for c in r if isinstance(c, str) and "arrêtées au" in c]
    if pres:
        m = re.search(r"(\d{2})/(\d{2})/(\d{4})", pres[0]); stated = "%s-%s-%s" % (m.group(3), m.group(2), m.group(1)) if m else None
    log("nacional", len(nat), "semanas;", len(reg), "regiones; datos hasta", stated)
    return {"schemaVersion": 1, "generatedAt": datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"), "country": "FR",
            "source": {"id": "franceagrimer", "name": "FranceAgriMer, Cere'Obs (etat des cultures, mais grain)", "url": "https://cereobs.franceagrimer.fr/", "license": "Licence Ouverte 2.0 on data.gouv.fr; file states all rights reserved (under clarification)", "dataUntil": stated, "resourceModified": mod},
            "fields": [k for _, k in FIELDS], "units": "% of area (stages) / % of area by condition class (c1 very poor ... c5 very good)",
            "national": nat, "regions": reg}
def main():
    args = sys.argv[1:]; outdir = ROOT / "data"
    if "--outdir" in args: outdir = Path(args[args.index("--outdir") + 1])
    outdir.mkdir(parents=True, exist_ok=True)
    try:
        doc = build()
        if len(doc["national"]) < 100 or not doc["regions"]: raise RuntimeError("bloques esenciales vacios")
    except Exception as e:
        log("FALLO:", type(e).__name__, str(e)[:300])
        if os.environ.get("GITHUB_ACTIONS"): print("::error title=fr-cereobs::" + " | ".join(LOG[-6:])[:3000].replace("%", "%25").replace("\n", "%0A"), flush=True)
        return 1
    s = json.dumps(doc, ensure_ascii=False, separators=(",", ":")); (outdir / "france-cereobs.json").write_text(s, encoding="utf-8")
    log("escrito france-cereobs.json", len(s) // 1024, "KB"); return 0
if __name__ == "__main__": sys.exit(main())
