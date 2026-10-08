#!/usr/bin/env python3
"""Francia: superficie, rendimiento y producción por cultivo (SSP/Agreste, difundido por FranceAgriMer) -> data/france-crops-stats.json (serie nacional 2000-2026,
formato de country-stats para el catálogo de la ficha de país) y data/france-crops-regions.json (regiones y departamentos 2021-2026, para la ficha de región).
Entrada, a mano (los ficheros de FranceAgriMer cambian de URL cada mes):
  python3 scripts/build-france-crops.py <hist_nat_surface_prod_cult-*.xlsx> <hist_reg_surface_prod_cult_cer-*.csv> <hist_dep_surface_prod_cult_cer-*.csv>
Francia metropolitana (sin DOM). El último año es provisional (situación mensual de grandes cultivos). Se comprueba que las regiones suman el total nacional."""
import csv, datetime, json, re, sys
from pathlib import Path
import openpyxl
ROOT = Path(__file__).resolve().parents[1]
CROPS = [  # (hoja del Excel / nombre en CSV, id, etiqueta en inglés)
    ("Céréales", "cereals", "Cereals"), ("Blé tendre", "soft-wheat", "Soft wheat"), ("Blé dur", "durum-wheat", "Durum wheat"), ("Orge", "barley", "Barley"),
    ("Maïs grain", "grain-maize", "Grain maize"), ("Maïs (grain et semence)", "maize-grain-seed", "Maize (grain and seed)"), ("Avoine", "oats", "Oats"), ("Seigle", "rye", "Rye"), ("Triticale", "triticale", "Triticale"),
    ("Sorgho", "sorghum", "Sorghum"), ("Riz", "rice", "Rice"), ("Oléagineux", "oilseeds", "Oilseeds"), ("Colza", "rapeseed", "Rapeseed"),
    ("Tournesol", "sunflower", "Sunflower"), ("Soja", "soybeans", "Soybeans"), ("Lin oléagineux", "linseed", "Oilseed flax"),
    ("Pois protéagineux", "protein-peas", "Protein peas"), ("Féveroles", "faba-beans", "Faba beans"), ("Lupin doux", "sweet-lupin", "Sweet lupin"),
    ("Lentilles", "lentils", "Lentils"),
]
SRC = "FranceAgriMer / Agreste (SSP) – surfaces, rendements et productions (France métropolitaine; dernière année provisoire)"
def num(x):
    return float(str(x).strip().replace(",", ".")) if str(x).strip() not in ("", "-") else None
def pct(a, b): return round((a - b) / abs(b) * 100, 2) if b else None

def national(xlsx):
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    out = {}
    for sheet, cid, _l in CROPS:
        ws = wb[sheet]
        hdr = next(r for r in range(1, 12) if ws.cell(r, 1).value == "Année")
        assert [ws.cell(hdr, c).value for c in (2, 3, 4)] == ["Suf.(ha)", "Rend(qx/ha)", "Prod.(t)"], "columnas inesperadas en " + sheet
        pts = {}
        for r in range(hdr + 1, ws.max_row + 1):
            y = ws.cell(r, 1).value
            if not isinstance(y, int): continue
            a, q, p = (ws.cell(r, c).value for c in (2, 3, 4))
            if not all(isinstance(v, (int, float)) for v in (a, q, p)) or (a == 0 and p == 0): continue
            assert a >= 0 and q >= 0 and p >= 0
            if a and p and abs(p / a * 10 - q) > max(1.0, 0.03 * q): raise SystemExit("rendimiento no cuadra con producción/superficie: %s %s" % (sheet, y))
            pts[y] = (a, q, p)
        out[cid] = pts
    return out

def read_csv(path):
    rows = list(csv.reader(open(path, encoding="latin-1"), delimiter=";"))
    return rows[0], [[c.strip() for c in r] for r in rows[1:] if r]

def main(xl, regcsv, depcsv):
    nat = national(xl)
    series = []
    for sheet, cid, label in CROPS:
        pts = nat[cid]
        for key, idx, unit, lab in (("area", 0, "ha", "cultivated area"), ("yield", 1, "q/ha", "yield"), ("prod", 2, "t", "production")):
            p = [[str(y), round(v[idx], 2)] for y, v in sorted(pts.items())]
            if len(p) < 3: continue
            series.append({"id": "fr-ssp-%s-%s" % (cid, key), "group": "crops", "label": "%s: %s (FranceAgriMer/Agreste SSP)" % (label, lab), "unit": unit, "frequency": "annual",
                           "latestPeriod": p[-1][0], "latest": p[-1][1], "changePct": pct(p[-1][1], p[-2][1]), "points": p, "sourceGroup": SRC})
    doc = {"schemaVersion": 1, "generatedAt": datetime.datetime.utcnow().strftime("%Y-%m-%dT%H:%M:%SZ"),
           "countries": {"FR": {"name": "France", "source": {"name": "FranceAgriMer (VISIONet)", "url": "https://visionet.franceagrimer.fr/", "license": "Licence Ouverte 2.0"}, "series": series}}, "log": []}
    (ROOT / "data" / "france-crops-stats.json").write_text(json.dumps(doc, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    # regiones y departamentos
    h, reg = read_csv(regcsv); hd, dep = read_csv(depcsv)
    names = {c[1]: c[2] for c in CROPS}
    byname = {c[0]: c[1] for c in CROPS}
    R, D, tot = {}, {}, {}
    import collections
    cnt = collections.Counter((r[0], r[1], r[3]) for r in reg)
    bad = sorted({r[0] for r in reg if cnt[(r[0], r[1], r[3])] > 1})
    dcnt = collections.Counter((r[0], r[1], r[4]) for r in dep)
    bad = sorted(set(bad) | {r[0] for r in dep if dcnt[(r[0], r[1], r[4])] > 1})
    print("cultivos descartados en regiones/departamentos (el CSV repite filas de sus componentes en lugar del total):", bad)
    reg = [r for r in reg if r[0] not in bad]; dep = [r for r in dep if r[0] not in bad]
    for r in reg:
        cid = byname.get(r[0])
        if r[0] == "TOTAL FRANCE": continue
        if not cid or not r[1].isdigit() and r[1] != "0": continue
        a, q, p = num(r[4]), num(r[5]), num(r[6])
        if a is None or p is None: continue
        R.setdefault(r[1], {"name": r[2]}).setdefault("c", {}).setdefault(cid, {})[r[3]] = [round(a), round(q, 1) if q is not None else None, round(p)]
    for r in dep:
        cid = byname.get(r[0])
        if not cid: continue
        a, q, p = num(r[5]), num(r[6]), num(r[7])
        if a is None or p is None: continue
        D.setdefault(r[1], {"name": r[2], "reg": r[3]}).setdefault("c", {}).setdefault(cid, {})[r[4]] = [round(a), round(q, 1) if q is not None else None, round(p)]
    # control: suma de regiones ~ total nacional del Excel (el CSV regional no incluye algunos cultivos menores en todas las regiones)
    worst = 0
    for cid in ("soft-wheat", "barley", "maize-grain-seed", "rapeseed", "sunflower"):
        for y in ("2021", "2022", "2023", "2024", "2025", "2026"):
            n = nat[cid].get(int(y))
            if not n: continue
            s = sum(v["c"][cid][y][2] for v in R.values() if cid in v["c"] and y in v["c"][cid])
            worst = max(worst, abs(s - n[2]) / n[2])
    assert worst < 0.03, "las regiones no suman el total nacional (desvío máx. %.1f %%)" % (worst * 100)
    worst_d = 0
    for cid in ("soft-wheat", "barley", "maize-grain-seed", "rapeseed", "sunflower"):  # los departamentos suman su región
        for y in ("2021", "2024", "2026"):
            for code, rv in R.items():
                if cid not in rv["c"] or y not in rv["c"][cid]: continue
                s2 = sum(dv["c"][cid][y][2] for dv in D.values() if dv["reg"].strip().upper() == rv["name"].strip().upper() and cid in dv["c"] and y in dv["c"][cid])
                if rv["c"][cid][y][2] > 5000: worst_d = max(worst_d, abs(s2 - rv["c"][cid][y][2]) / rv["c"][cid][y][2])
    assert worst_d < 0.03, "los departamentos no suman su región (%.1f %%)" % (worst_d * 100)
    out = {"schemaVersion": 1, "generatedAt": doc["generatedAt"], "source": {"id": "franceagrimer", "name": "FranceAgriMer / Agreste (SSP)", "url": "https://visionet.franceagrimer.fr/"},
           "note": "Francia metropolitana (13 regiones). Cada cultivo: año -> [superficie ha, rendimiento q/ha, producción t]. El último año es provisional. Los departamentos se usan solo para comprobar que suman su región.",
           "crops": {cid: names[cid] for _s, cid, _l in CROPS}, "regions": R}
    (ROOT / "data" / "france-crops-regions.json").write_text(json.dumps(out, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print("series nacionales: %d; regiones: %d; departamentos: %d; desvío máx. regiones vs nacional: %.2f %%; departamentos vs región: %.2f %%" % (len(series), len(R), len(D), worst * 100, worst_d * 100))

if __name__ == "__main__":
    if len(sys.argv) != 4: raise SystemExit(__doc__)
    main(*sys.argv[1:])
