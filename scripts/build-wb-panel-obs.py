#!/usr/bin/env python3
"""Banco Mundial, Commodity Price Data «Pink Sheet» (mensual, USD) al panel de precios: aceites vegetales, cafe, cacao, te, platano, naranja, algodon,
caucho, cacahuete, harina de pescado, carne/pollo/cordero internacionales, trigo de EE. UU. de exportacion, arroz de Tailandia y Vietnam, azucar,
TSP y roca fosforica. Son precios internacionales de referencia (no lo que cobra el agricultor): comparabilidad «not_comparable», region 'eu'
(igual que la urea y el gas del Banco Mundial). Escribe observaciones en el snapshot de hoy y, despues, el motor de precios reconstruye todo.
Licencia: CC BY 4.0 (World Bank). Solo se usa lo que publica la fuente: una serie sin dato en los ultimos 3 meses o con menos de 12 puntos se omite.
Claves: las mismas que js/data.js (bloque WB_PANEL). Uso: python3 scripts/build-wb-panel-obs.py [ruta/al/CMO-Historical-Data-Monthly.xlsx]"""
import json, re, sys, urllib.request
from datetime import datetime, timezone
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SNAP_DIR = ROOT / "data" / "snapshots"
LANDING = "https://www.worldbank.org/en/research/commodity-markets"
FIRST_YEAR = 1990
MIN_POINTS = 12
MAX_AGE_MONTHS = 3

# clave: (categoria, encabezado exacto de la columna (sin ' **'), unidad del panel)
PRODUCTS = {
    "aceite_palma": ("aceite", "Palm oil", "tonelada"),
    "aceite_palmiste": ("aceite", "Palm kernel oil", "tonelada"),
    "aceite_coco": ("aceite", "Coconut oil", "tonelada"),
    "aceite_cacahuete": ("aceite", "Groundnut oil", "tonelada"),
    "aceite_soja": ("aceite", "Soybean oil", "tonelada"),
    "aceite_colza": ("aceite", "Rapeseed oil", "tonelada"),
    "aceite_girasol": ("aceite", "Sunflower oil", "tonelada"),
    "harina_pescado": ("pienso", "Fish meal", "tonelada"),
    "cacao": ("tropicales", "Cocoa", "kg"),
    "cafe_arabica": ("tropicales", "Coffee, Arabica", "kg"),
    "cafe_robusta": ("tropicales", "Coffee, Robusta", "kg"),
    "te_subastas": ("tropicales", "Tea, avg 3 auctions", "kg"),
    "banana_ue": ("tropicales", "Banana, Europe", "kg"),
    "banana_us": ("tropicales", "Banana, US", "kg"),
    "naranja": ("tropicales", "Orange", "kg"),
    "algodon_a": ("tropicales", "Cotton, A Index", "kg"),
    "caucho_rss3": ("tropicales", "Rubber, RSS3", "kg"),
    "caucho_tsr20": ("tropicales", "Rubber, TSR20", "kg"),
    "cacahuete": ("tropicales", "Groundnuts", "tonelada"),
    "gambas_mx": ("tropicales", "Shrimps, Mexican", "kg"),
    "tabaco_us": ("tropicales", "Tobacco, US import u.v.", "tonelada"),
    "ternera_mundo": ("ganado", "Beef", "kg"),
    "cordero_mundo": ("ovino", "Lamb", "kg"),
    "pollo_mundo": ("avicultura", "Chicken", "kg"),
    "tsp": ("fertilizantes", "TSP", "tonelada"),
    "roca_fosforica": ("fertilizantes", "Phosphate rock", "tonelada"),
    "trigo_srw_us": ("cereales", "Wheat, US SRW", "tonelada"),
    "trigo_hrw_us": ("cereales", "Wheat, US HRW", "tonelada"),
    "arroz_tai_5": ("cereales", "Rice, Thai 5%", "tonelada"),
    "arroz_tai_25": ("cereales", "Rice, Thai 25%", "tonelada"),
    "arroz_tai_a1": ("cereales", "Rice, Thai A.1", "tonelada"),
    "arroz_viet_5": ("cereales", "Rice, Viet Namese 5%", "tonelada"),
    "azucar_mundo": ("azucar", "Sugar, world", "kg"),
    "azucar_ue": ("azucar", "Sugar, EU", "kg"),
    "azucar_us": ("azucar", "Sugar, US", "kg"),
}


def fetch(url):
    req = urllib.request.Request(url, headers={"User-Agent": "Dehesa-Index-data-bot/1.0"})
    with urllib.request.urlopen(req, timeout=90) as r:
        return r.read()


def norm(h):
    return re.sub(r"\s+", " ", str(h or "").replace("**", "")).strip().lower()


def main():
    publication = None
    if len(sys.argv) > 1:
        raw = Path(sys.argv[1]).read_bytes()
    else:
        html = fetch(LANDING).decode("utf-8", "ignore")
        m = re.findall(r'https://thedocs\.worldbank\.org/[^"\']+CMO-Historical-Data-Monthly\.xlsx', html)
        if not m:
            sys.exit("No se encuentra el enlace del XLSX mensual del Banco Mundial")
        pub = re.search(r"\((\d{2}/\d{2}/\d{4})\)", re.sub(r"<[^>]+>", " ", html))
        if pub:
            publication = datetime.strptime(pub.group(1), "%m/%d/%Y").date().isoformat()
        raw = fetch(m[0].replace("&amp;", "&"))
    tmp = ROOT / ".worldbank-commodities.xlsx"
    tmp.write_bytes(raw)
    try:
        wb = load_workbook(tmp, read_only=True, data_only=True)
        rows = list(wb["Monthly Prices"].iter_rows(values_only=True))
    finally:
        tmp.unlink(missing_ok=True)
    hdr = next((i for i, r in enumerate(rows[:20]) if any(norm(v).startswith("crude oil") for v in r)), None)
    if hdr is None:
        sys.exit("Fila de encabezados no encontrada")
    cols = {norm(v): c for c, v in enumerate(rows[hdr]) if v}
    units = rows[hdr + 1]
    periods = [str(r[0] or "").strip() for r in rows[hdr + 2:]]
    last_period = max(p for p in periods if re.fullmatch(r"\d{4}M\d{2}", p))
    ly, lm = int(last_period[:4]), int(last_period[5:])
    today = datetime.now(timezone.utc)
    verified_at = today.isoformat(timespec="seconds").replace("+00:00", "Z")
    pubdate = publication or today.date().isoformat()
    obs = []
    for key, (cat, head, unit) in PRODUCTS.items():
        c = cols.get(norm(head))
        if c is None:
            print("SIN COLUMNA:", key, head); continue
        want_unit = "$/kg" if unit == "kg" else "$/mt"
        got = str(units[c] or "").replace(" ", "").lower()
        if want_unit not in got:
            print("UNIDAD INESPERADA:", key, head, units[c]); continue
        pts = []
        for r in rows[hdr + 2:]:
            p = str(r[0] or "").strip()
            if not re.fullmatch(r"\d{4}M\d{2}", p) or int(p[:4]) < FIRST_YEAR:
                continue
            v = r[c] if c < len(r) else None
            if isinstance(v, (int, float)) and v > 0:  # el Banco Mundial marca con 0 los meses sin dato
                pts.append(("%s-%s" % (p[:4], p[5:7]), round(float(v), 4)))
        if len(pts) < MIN_POINTS:
            print("POCOS PUNTOS:", key, len(pts)); continue
        pp, val = pts[-1]
        age = (ly * 12 + lm) - (int(pp[:4]) * 12 + int(pp[5:]))
        if age > MAX_AGE_MONTHS:
            print("SIN DATO RECIENTE:", key, pp); continue
        prev = pts[-2][1]
        change = round((val / prev - 1) * 100, 4) if prev else None
        obs.append({
            "id": "di_%s_%s_eu" % (cat, key), "product": key, "region": "eu", "sourceId": "world_bank",
            "observationDate": pp, "publicationDate": pubdate, "status": "verified", "verifiedAt": verified_at,
            "comparability": "not_comparable",
            "methodology": "Banco Mundial, Commodity Price Data (Pink Sheet): %s, precio internacional de referencia en dolares de EE. UU. (%s). No es el precio que cobra el agricultor. Serie mensual." % (head, "por kg" if unit == "kg" else "por tonelada metrica"),
            "value": val, "currency": "USD", "unit": unit, "frequency": "monthly", "changePct": change,
            "history": [{"period": q[5:7], "year": int(q[:4]), "value": v} for q, v in pts],
        })
        print("OK", key, pp, val, len(pts), "puntos")
    if not obs:
        sys.exit("Sin observaciones")
    SNAP_DIR.mkdir(parents=True, exist_ok=True)
    snap = SNAP_DIR / (today.date().isoformat() + ".json")
    doc = {"schemaVersion": "1.0", "generatedAt": verified_at, "observations": []}
    if snap.exists():
        doc = json.loads(snap.read_text(encoding="utf-8"))
    mine = {o["product"] for o in obs}
    doc["observations"] = [o for o in doc.get("observations", []) if not (o.get("region") == "eu" and o.get("product") in mine)] + obs
    snap.write_text(json.dumps(doc, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(len(obs), "observaciones del Banco Mundial para el panel")


if __name__ == "__main__":
    main()
