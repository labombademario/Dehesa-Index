#!/usr/bin/env python3
"""ABARES Australian Crop Report (xlsx oficial de datos) -> data/abares-stats.json (countries.AU, extend)
Fuente: ABARES, Australian Crop Report (No. 219, septiembre 2026), https://doi.org/10.25814/5s62-8a44, licencia CC BY 4.0.
Es una carga MANUAL trimestral: ABARES no es accesible desde GitHub Actions, asi que el Excel oficial se guarda en manual-sources/abares/
y este script lo convierte. Para actualizar: sustituir el xlsx por el de la nueva edicion y ejecutar el script (ver docs/WORKFLOWS.md).
Cargamos solo lo que es de ABARES o de ABS (valores unitarios de exportacion):
  - Tablas 11 y 12: produccion (kt) y superficie ('000 ha) por cultivo, Australia, anual, campana 1989-90 -> ultima estimacion.
  - Tablas 15 y 16: exportaciones por campana comercial (kt).
  - Tabla 17: columnas "Export" (valor unitario FOB de ABS) y "Domestic" (A$/t, trimestral).
Tambien las columnas "Domestic" (compiladas de Farm Weekly, The Land, etc.; decision del propietario, con advertencia). NO se cargan las "International" (precios de terceros: CME, USDA, IGC...);
contenido de terceros excluido de la CC BY 4.0 de ABARES). Las previsiones (f) no entran en los puntos: van en la nota de la serie.
Convencion de periodo: la campana 2025-26 se registra como 2025 (anyo de siembra)."""
import json, re, sys, datetime
from pathlib import Path
import openpyxl
ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "manual-sources" / "abares" / "AustCropRrt20260901_CropData_v1.1.0.xlsx"
OUT = ROOT / "data" / "abares-stats.json"
EDITION = {"report": "Australian Crop Report", "number": 219, "date": "2026-09-01", "doi": "https://doi.org/10.25814/5s62-8a44"}
REF = {"source": "ABARES", "dataset": "Australian Crop Report (Sep 2026, No. 219)", "date": EDITION["date"], "license": "CC BY 4.0", "original": EDITION["doi"]}
SG = "ABARES Australian Crop Report"

wb = openpyxl.load_workbook(XLSX, data_only=True)
def num(v):
    if isinstance(v, (int, float)): return float(v)
    return None
def crop_year(s):
    m = re.match(r"^(\d{4})[–-](\d{2})\s*([sf]?)", str(s or "").strip())
    return (m.group(1), m.group(3)) if m else (None, None)
def clean(s): return re.sub(r"\s+", " ", re.sub(r"\s[a-z]$", "", str(s or "").strip())).strip()

series = []
def add(sid, group, label, unit, freq, pts, note=None, sg=SG):
    pts = sorted({p: v for p, v in pts}.items())
    pts = [[p, round(v, 3)] for p, v in pts if v is not None]
    if len(pts) < 2: return
    prev = pts[-2][1]
    s = {"id": sid, "group": group, "label": label, "unit": unit, "frequency": freq, "latestPeriod": pts[-1][0], "latest": pts[-1][1],
         "changePct": round((pts[-1][1] - prev) / prev * 100, 1) if prev else None, "points": pts, "sourceGroup": sg, "reference": REF}
    if note: s["periodNote"] = note
    series.append(s)

CROPS11 = {"Wheat": "wheat", "Barley": "barley", "Canola": "canola", "Chickpeas": "chickpeas", "Faba beans": "faba-beans", "Field peas": "field-peas",
           "Lentils": "lentils", "Lupins": "lupins", "Oats": "oats", "Triticale": "triticale"}
CROPS12 = {"Grain sorghum": "sorghum", "Cottonseed": "cottonseed", "Cotton lint": "cotton-lint", "Rice": "rice", "Corn (maize)": "maize", "Soybeans": "soybeans", "Sunflower": "sunflower"}
LABEL = {"canola": "Canola (rapeseed)", "maize": "Maize (corn)", "sorghum": "Grain sorghum", "cottonseed": "Cottonseed", "cotton-lint": "Cotton lint", "faba-beans": "Faba beans", "field-peas": "Field peas",
         "soybeans": "Soybeans", "sunflower": "Sunflower seed", "wheat": "Wheat", "barley": "Barley", "chickpeas": "Chickpeas", "lentils": "Lentils", "lupins": "Lupins", "oats": "Oats", "triticale": "Triticale", "rice": "Rice"}

def crop_table(name, crops, season):
    ws = wb[name]; hdr = None
    for r in range(1, 20):
        if ws.cell(r, 2).value == "Year": hdr = r; break
    cols = {}
    for c in range(3, ws.max_column + 1):
        h = ws.cell(hdr, c).value
        if h: cols[c] = clean(h)
    forecast = {}
    for c, h in cols.items():
        key = crops.get(h)
        if not key: print("columna sin mapear en", name, h); sys.exit(1)
        area, prod = [], []
        for r in range(hdr + 3, ws.max_row + 1):
            y, fl = crop_year(ws.cell(r, 2).value)
            if not y: continue
            a, p = num(ws.cell(r, c).value), num(ws.cell(r, c + 1).value)
            if fl == "f":
                forecast[key] = (y, a, p); continue
            if a is not None: area.append((y, a))
            if p is not None: prod.append((y, p))
        est = [y for y in range(2000, 2100) if False]
        note = "Campana %s registrada con el anyo de siembra (2025 = 2025-26). Las campanas marcadas 's' por ABARES son estimaciones." % season
        fy, fa, fp = forecast.get(key, (None, None, None))
        if fy and fp is not None: note += " Prevision ABARES %s-%s: %s kt de produccion (no incluida en los puntos)." % (fy, str(int(fy) + 1)[2:], format(fp, ",.0f").replace(",", "."))
        base = LABEL[key]
        add("au-ab-%s-prod" % key, "crops", "%s: production (ABARES)" % base, "kt", "annual", prod, note)
        if key not in ("cottonseed",):
            add("au-ab-%s-area" % key, "crops", "%s: area sown (ABARES)" % base + (" [harvested]" if key == "cotton-lint" else ""), "'000 ha", "annual", area, note)
crop_table("Table 11", CROPS11, "de siembra abril-marzo")
crop_table("Table 12", CROPS12, "de siembra abril-marzo")

# Tablas 15 y 16: exportaciones por campana comercial (columnas = campanas)
EXP = [("Table 15", "Wheat", "Exports", "wheat", "Wheat", "octubre-septiembre"), ("Table 15", "Canola", "Exports", "canola", "Canola (rapeseed)", "noviembre-octubre"),
       ("Table 16", "Barley", "Exports", "barley", "Barley", "noviembre-octubre"), ("Table 16", "Oats", "Exports", "oats", "Oats", "noviembre-octubre"),
       ("Table 16", "Grain sorg", "Exports", "sorghum", "Grain sorghum", "marzo-febrero"), ("Table 16", "Corn (maiz", "Exports", "maize", "Maize (corn)", "marzo-febrero")]
def balance(name):
    ws = wb[name]; hdr = None
    for r in range(1, 20):
        if ws.cell(r, 2).value == "Crop": hdr = r; break
    years = {c: crop_year(ws.cell(hdr, c).value) for c in range(3, ws.max_column + 1)}
    cur, out = None, {}
    for r in range(hdr + 2, ws.max_row + 1):
        lab = str(ws.cell(r, 2).value or "").strip()
        if not lab: continue
        if not ws.cell(r, 3).value and not isinstance(ws.cell(r, 3).value, (int, float)): cur = lab; continue
        out[(cur, lab)] = {years[c][0]: num(ws.cell(r, c).value) for c in years if years[c][0] and years[c][1] != "f" and num(ws.cell(r, c).value) is not None}
    return out
B15, B16 = balance("Table 15"), balance("Table 16")
def find(b, crop_prefix, lab_prefix):
    for (c, l), v in b.items():
        if c and c.startswith(crop_prefix) and l.startswith(lab_prefix): return v
    return None
for t, cp, lp, key, nm, season in EXP:
    d = find(B15 if t == "Table 15" else B16, cp, lp)
    if not d: print("sin bloque", t, cp); sys.exit(1)
    add("au-ab-exp-%s" % key, "trade", "Exports (marketing year): %s (ABARES)" % nm, "kt", "annual", list(d.items()),
        "Campana comercial %s, no comparable con las exportaciones por anyo fiscal de ABS. Registrada con el anyo de inicio (2024 = 2024-25). Ultimos anyos: estimacion ABARES." % season)
# pulsos (Table 15, bloque "Pulses")
for key, nm in (("lupins", "Lupins"), ("field peas", "Field peas"), ("chickpeas", "Chickpeas")):
    cur = None; got = None
    for (c, l), v in B15.items():
        if c and c.startswith("Exports") and l.strip().lower() == key: got = v
    if got: add("au-ab-exp-%s" % key.replace(" ", "-"), "trade", "Exports (marketing year): %s (ABARES)" % nm, "kt", "annual", list(got.items()),
                "Campana comercial, no comparable con las exportaciones por anyo fiscal de ABS. Registrada con el anyo de inicio. Ultimos anyos: estimacion ABARES.")
for key, nm in (("feed barley", "Barley, feed"), ("malting barley", "Barley, malting")):
    for (c, l), v in B16.items():
        if c and c.startswith("Barley") and l.strip().lower().startswith(key):
            add("au-ab-exp-%s" % key.replace(" ", "-"), "trade", "Exports (marketing year): %s (ABARES)" % nm, "kt", "annual", list(v.items()),
                "Campana comercial noviembre-octubre. Registrada con el anyo de inicio. Ultimos anyos: estimacion ABARES.")

# Tabla 17: SOLO valores unitarios de exportacion (ABS, FOB), trimestral
ws = wb["Table 17"]
heads = {c: str(ws.cell(9, c).value or "").strip() for c in range(4, ws.max_column + 1)}
crop_row = {}
cur = None
for c in range(4, ws.max_column + 1):
    v = ws.cell(8, c).value
    if v: cur = str(v).strip()
    crop_row[c] = cur
EUV = {"Export: feed b": ("barley-feed", "Barley, feed"), "Export: malting b": ("barley-malting", "Barley, malting"), "Export b": None, "Export: chickpeas b": ("chickpeas", "Chickpeas"), "Export: field peas b": ("field-peas", "Field peas")}
for c, h in heads.items():
    if not h.startswith("Export"): continue
    k = re.sub(r"\s+", " ", h)
    crop = crop_row[c]
    if k == "Export b": key, nm = {"Grain sorghum": ("sorghum", "Grain sorghum")}.get(crop, (None, None))
    else: key, nm = EUV[k]
    if not key: print("Export sin mapear", crop, h); sys.exit(1)
    pts, yr = [], None
    for r in range(10, ws.max_row + 1):
        lab = str(ws.cell(r, 2).value or "").strip()
        m = re.match(r"^(\d{4}) (Q[1-4])$", lab)
        if m: yr, q = m.group(1), m.group(2)
        elif re.match(r"^Q[1-4]$", lab): q = lab
        else: continue
        v = num(ws.cell(r, c).value)
        if v is not None: pts.append(("%s-%s" % (yr, q), v))
    add("au-ab-euv-%s" % key, "prices", "Export unit value (FOB): %s (ABS via ABARES)" % nm, "A$/t", "quarterly", pts,
        "Valor unitario medio FOB de las exportaciones del trimestre (ABS, recogido por ABARES): no es un precio de mercado corriente; puede haber desfase entre la negociacion y el embarque.", sg="ABARES Australian Crop Report (ABS export unit values)")

# Tabla 17: precios DOMESTICOS (decision del propietario, 7 oct 2026). ABARES los compila de informes de terceros (Farm Weekly, The Land, The Weekly Times, Jumbuk AG):
# se publican con esa advertencia en la nota y en el registro de licencias. Siguen excluidas las columnas "International".
DOM = {"Wheat": "wheat", "Barley": "barley", "Grain sorghum": "sorghum", "Oats": "oats", "Corn (maize)": "maize", "Oilseeds": "canola"}
PULSE = {"lupins": "lupins", "chickpeas": "chickpeas", "field peas": "field-peas"}
for c, h in heads.items():
    if not h.startswith("Domestic"): continue
    h1 = re.sub(r"\s+", " ", h); crop = crop_row[c]
    if crop == "Pulses":
        m = re.match(r"Domestic: ([a-z ]+),", h1); key = PULSE.get(m.group(1)) if m else None
    else: key = DOM.get(crop)
    if not key: print("Domestic sin mapear", crop, h); sys.exit(1)
    pts, yr = [], None
    for r in range(10, ws.max_row + 1):
        lab = str(ws.cell(r, 2).value or "").strip()
        m = re.match(r"^(\d{4}) (Q[1-4])$", lab)
        if m: yr, q = m.group(1), m.group(2)
        elif re.match(r"^Q[1-4]$", lab): q = lab
        else: continue
        v = num(ws.cell(r, c).value)
        if v is not None: pts.append(("%s-%s" % (yr, q), v))
    add("au-ab-dom-%s" % key, "prices", "Domestic price (%s): %s (ABARES)" % (h1.replace("Domestic: ", ""), LABEL[key]), "A$/t", "quarterly", pts,
        "Media trimestral de precios domesticos que ABARES compila de informes de terceros (Farm Weekly, The Land, The Weekly Times, Jumbuk AG); sin GST. Publicado por decision del propietario del sitio, citando a ABARES como compilador.",
        sg="ABARES Australian Crop Report (domestic prices compiled from third-party reports)")

ids = [s["id"] for s in series]
assert len(ids) == len(set(ids)), "ids duplicados"
now = datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
doc = {"schemaVersion": 1, "generatedAt": now, "edition": EDITION,
       "countries": {"AU": {"name": "Australia", "extend": True, "source": {"name": "ABARES, Australian Crop Report (September 2026, No. 219)", "url": EDITION["doi"], "license": "CC BY 4.0"}, "series": series}},
       "log": ["ABARES Australian Crop Report %s: %d series (Tablas 11, 12, 15, 16 y columnas Export de la 17). Domesticos incluidos por decision del propietario; no se cargan los internacionales (terceros)." % (EDITION["date"], len(series))]}
OUT.write_text(json.dumps(doc, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
print("OK", len(series), "series ->", OUT.name)
