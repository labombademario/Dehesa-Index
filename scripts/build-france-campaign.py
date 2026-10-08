#!/usr/bin/env python3
"""Francia (FranceAgriMer, a mano): precios al productor trimestrales, superficies PAC declaradas, molineria, piensos (incorporacion y stocks) y stocks en deposito
-> data/france-campaign-stats.json (formato country-stats).
  python3 scripts/build-france-campaign.py <HISTOPRIXTRIMESTRIELS.xlsx> <surface_PAC_NAT.xlsx> <meunerie_negoce.xlsx> <REG_incorporation_FAB.xlsx> <REG_stock_FAB.xlsx> <stock_DEPOT_mensuel.xlsx>
(los ficheros se distinguen por el nombre). 'Chiffres provisoires': los meses sin publicar vienen como 0 y se descartan."""
import datetime, json, re, sys, unicodedata
from pathlib import Path
import openpyxl, openpyxl.reader.excel as E
E.ExcelReader.read_properties = lambda self: None   # algunos xlsx de FranceAgriMer traen docProps/core.xml invalido
ROOT = Path(__file__).resolve().parents[1]
MON = {m: i + 1 for i, m in enumerate("janvier fevrier mars avril mai juin juillet aout septembre octobre novembre decembre".split())}
def norm(x): return "".join(c for c in unicodedata.normalize("NFD", str(x).lower()) if unicodedata.category(c) != "Mn").replace("\xa0", " ").strip()
def pct(a, b): return round((a - b) / abs(b) * 100, 2) if b else None
def mk(id_, group, label, unit, freq, pts, src):
    pts = [[p, round(v, 4)] for p, v in pts]
    return {"id": id_, "group": group, "label": label, "unit": unit, "frequency": freq, "latestPeriod": pts[-1][0], "latest": pts[-1][1], "changePct": pct(pts[-1][1], pts[-2][1]) if len(pts) > 1 else None, "points": pts, "sourceGroup": src}
EN = {"ble tendre": ("soft-wheat", "Soft wheat"), "ble dur": ("durum-wheat", "Durum wheat"), "orges": ("barley", "Barley"), "orge": ("barley", "Barley"), "mais": ("maize", "Maize"), "seigle": ("rye", "Rye"), "avoine": ("oats", "Oats"),
      "sorgho": ("sorghum", "Sorghum"), "triticale": ("triticale", "Triticale"), "colza": ("rapeseed", "Rapeseed"), "pois": ("pea", "Peas"), "tournesol": ("sunflower", "Sunflower"), "feverole": ("faba", "Faba beans"), "soja": ("soy", "Soybeans"),
      "epeautre": ("spelt", "Spelt"), "lentilles": ("lentil", "Lentils"), "lin": ("flax", "Flax"), "lupin": ("lupin", "Lupin"), "riz": ("rice", "Rice"), "sarrasin": ("buckwheat", "Buckwheat")}
def en(name):
    k = norm(name); return EN.get(k)
def ymq(camp, q):
    y = int(camp[:4]); return ("%d-Q3" % y, "%d-Q4" % y, "%d-Q1" % (y + 1), "%d-Q2" % (y + 1))[q - 1]
DUP_PAID = {"soft-wheat", "durum-wheat", "barley", "maize", "rye", "oats", "sorghum", "triticale", "rapeseed", "pea", "sunflower", "faba", "soy"}
def farmgate(f):
    ws = openpyxl.load_workbook(f, data_only=True)["Publication nationale"]
    hdr = next(r for r in range(1, 30) if str(ws.cell(r, 1).value).startswith("Trimestre")); rows = []
    for r in range(hdr + 1, ws.max_row + 1):
        q, c = ws.cell(r, 1).value, ws.cell(r, 2).value
        if isinstance(q, (int, float)) and isinstance(c, str) and re.match(r"^\d{4}/\d{2}$", c): rows.append((ymq(c, int(q)), r))
    out = []
    for col in range(3, ws.max_column + 1):
        h = ws.cell(hdr, col).value
        if not h or not en(h): continue
        pts = [(p, float(ws.cell(r, col).value)) for p, r in rows if isinstance(ws.cell(r, col).value, (int, float))]
        pts.sort()
        if len(pts) < 8: continue
        assert all(40 < v < 1500 for _p, v in pts), h
        s, lab = en(h)
        if s in DUP_PAID: continue   # ya publicados como fr-paid-* en data/france-stats.json (mismo fichero de FranceAgriMer)
        out.append(mk("fr-fam-farmgate-" + s, "prices", "%s: farm-gate price, campaign average" % lab, "€/t", "quarterly", pts, "FranceAgriMer – prix payés aux producteurs (trimestriels, cumul de campagne)"))
    return out
PAC_EN = {"total cereales": "All cereals", "ble tendre d'hiver": "Winter soft wheat", "ble dur d'hiver": "Winter durum wheat", "avoine d'hiver": "Winter oats", "avoine de printemps": "Spring oats", "ble tendre de printemps": "Spring soft wheat", "mais doux": "Sweet maize", "lin non textile": "Linseed", "colza de printemps": "Spring rapeseed", "orge d'hiver": "Winter barley", "orge de printemps": "Spring barley",
          "mais grain et ensilage": "Grain and silage maize", "triticale": "Triticale", "seigle": "Rye", "sorgho": "Sorghum", "total oleagineux": "All oilseeds",
          "colza d'hiver": "Winter rapeseed", "tournesol": "Sunflower", "soja": "Soybeans", "total legumineuses et fourrageres": "All pulses and forage legumes",
          "pois": "Peas", "feverole": "Faba beans", "luzerne": "Alfalfa", "lentilles": "Lentils", "ble tendre": "Soft wheat", "ble dur": "Durum wheat", "orge": "Barley", "mais grain": "Grain maize", "avoine": "Oats", "triticale d'hiver": "Winter triticale",
          "colza": "Rapeseed", "feverole d'hiver": "Winter faba beans", "trefle": "Clover"}
def pac(f):
    ws = openpyxl.load_workbook(f, data_only=True)["dossiers-surface"]; data = {}; grp = {}
    for r in range(6, ws.max_row + 1):
        g, c, y, n, s = [ws.cell(r, k).value for k in range(1, 6)]
        if g in ("Céréales", "Oléagineux", "Légumineuses et fourragères") and isinstance(y, (int, float)) and isinstance(s, (int, float)) and c:
            d = data.setdefault(c, {}); d[int(y)] = max(float(s), d.get(int(y), 0)); grp[c] = g  # un año repetido (blé dur 2009): se queda la fila grande
    out = []; unmapped = []
    for c, d in data.items():
        k = norm(c)
        if k not in PAC_EN:
            if d.get(2026, 0) >= 80000: unmapped.append((c, d.get(2026)))
            continue
        pts = [(str(y), v) for y, v in sorted(d.items())]
        if len(pts) < 8: continue
        slug = re.sub(r"[^a-z0-9]+", "-", k).strip("-")
        out.append(mk("fr-fam-pac-" + slug, "crops", "%s: area declared under CAP" % PAC_EN[k], "ha", "annual", pts, "FranceAgriMer / ASP – surfaces déclarées PAC"))
    if unmapped: print("PAC sin traducir (>=80 kha):", unmapped)
    return out
def camp_month(camp, mname):
    y = int(camp[:4]); m = MON[norm(mname)]; return "%04d-%02d" % (y if m >= 7 else y + 1, m)
CUT = "2026-08"   # ultimo mes publicado de la campana 2026/27
def clean(pts): return sorted((p, v) for p, v in pts if p <= CUT)
def fab(f, tag, label, sid, group, kind):
    ws = openpyxl.load_workbook(f, data_only=True)["CEREALES"]; out = []
    for c in range(2, ws.max_column + 1):
        h = ws.cell(5, c).value
        if not h or not en(h): continue
        s, lab = en(h); pts = []
        for k in range(4):   # 3 campanas + N/N-1
            camp = ws.cell(6, c + k).value
            if not isinstance(camp, str) or not re.match(r"^\d{4}/\d{2}$", camp): continue
            for r in range(7, 19):
                m, v = ws.cell(r, 1).value, ws.cell(r, c + k).value
                if isinstance(m, str) and norm(m) in MON and isinstance(v, (int, float)): pts.append((camp_month(camp, m), float(v)))
        pts = clean(pts)
        if len(pts) < 20 or sum(v for _p, v in pts) / len(pts) < 1500: continue
        out.append(mk("fr-fam-%s-%s" % (sid, s), group, "%s: %s" % (lab, label), "t", "monthly", pts, "FranceAgriMer – %s" % tag))
    return out
def milling(f):
    ws = openpyxl.load_workbook(f, data_only=True)["Mensuel"]; rows = {norm(ws.cell(r, 1).value).split("\n")[0]: r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value}
    want = [("grains mis en oeuvre", "fr-fam-milling-soft-wheat-ground", "production", "Soft wheat ground by mills"),
            ("farine produite pure", "fr-fam-milling-flour-pure", "production", "Pure flour produced by mills"),
            ("stock fin grains", "fr-fam-milling-stock-grain", "stocks", "Soft wheat stocks at mills"),
            ("total panification", "fr-fam-milling-flour-bakery", "production", "Flour delivered to bakeries"),
            ("total exportations en farine", "fr-fam-milling-flour-exports", "production", "Flour exports by mills")]
    out = []
    for key, sid, grp, lab in want:
        r = rows[key]; pts = []
        for c in range(2, 41):
            m = ws.cell(5, c).value
            if m: cur = m
            camp = ws.cell(6, c).value; v = ws.cell(r, c).value
            if isinstance(camp, str) and re.match(r"^\d{4}/\d{2}$", camp) and norm(cur) in MON and isinstance(v, (int, float)): pts.append((camp_month(camp, cur), float(v)))
        pts = clean(pts)
        out.append(mk(sid, grp, lab, "t", "monthly", pts, "FranceAgriMer – Etat 8 meunerie et négoce en farine"))
    return out
def depot(f):
    wb = openpyxl.load_workbook(f, data_only=True); out = []
    for sh in wb.sheetnames:
        if not en(sh) and norm(sh) not in ("cereales autres",): 
            if sh != "Maïs" and sh != "Orge": continue
        if not en(sh): continue
        ws = wb[sh]; tot = next((r for r in range(1, ws.max_row + 1) if norm(ws.cell(r, 1).value or "").startswith("total france")), None)
        if tot is None: continue
        pts = []; cur = None
        for c in range(2, 62):
            m = ws.cell(6, c).value
            if m: cur = m
            camp = ws.cell(7, c).value; v = ws.cell(tot, c).value
            if cur and norm(cur) in MON and isinstance(camp, str) and re.match(r"^\d{4}/\d{2}$", camp) and isinstance(v, (int, float)): pts.append((camp_month(camp, cur), float(v)))
        pts = clean(pts); s, lab = en(sh)
        if len(pts) < 20: continue
        out.append(mk("fr-fam-depot-" + s, "stocks", "%s: depot stocks held for farmers" % lab, "t", "monthly", pts, "FranceAgriMer – Etat 2 stocks de dépôt"))
    return out
def main(args):
    by = lambda tag: next(a for a in args if tag in Path(a).name)
    series = farmgate(by("HISTOPRIX")) + pac(by("surface_PAC_NAT")) + milling(by("meunerie_negoce")) + fab(by("REG_incorporation_FAB"), "Etat 13 aliments du bétail", "grain used by feed makers", "feed-use", "production", "use") \
             + fab(by("REG_stock_FAB"), "Etat 13 aliments du bétail", "feed makers' stocks", "feed-stock", "stocks", "stock") + depot(by("stock_DEPOT_mensuel"))
    ids = [s["id"] for s in series]; assert len(ids) == len(set(ids))
    doc = {"schemaVersion": 1, "generatedAt": datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
           "countries": {"FR": {"name": "France", "source": {"name": "FranceAgriMer (VISIONet)", "url": "https://visionet.franceagrimer.fr/", "license": "Licence Ouverte 2.0"}, "series": series}}, "log": []}
    (ROOT / "data" / "france-campaign-stats.json").write_text(json.dumps(doc, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    import collections; print(len(series), "series", dict(collections.Counter(s["group"] for s in series)))
    for s in series: print(" ", s["id"], s["latestPeriod"], s["latest"], len(s["points"]))
if __name__ == "__main__":
    if len(sys.argv) < 7: raise SystemExit(__doc__)
    main(sys.argv[1:])
