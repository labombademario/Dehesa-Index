#!/usr/bin/env python3
"""Croacia: precios maximos oficiales del gasoleo, del «plavi dizel» (gasoleo azul agricola) y de la gasolina, y gasoleo semanal de la UE -> data/hr-diesel.json
1) Precios maximos de venta al publico (EUR/l, IVA incluido) que la Vlada (Gobierno de Croacia) fija con la «Uredba o utvrđivanju najviših maloprodajnih
   cijena naftnih derivata»: cada siete dias en 2026 (antes, cada dos semanas). El decreto solo trae la formula y las primas; los precios finales los publica
   la Vlada en la nota de prensa del mismo dia («Nove cijene iznosit će: ... EUR/l za dizelsko gorivo ... za plavi dizel»), junto con los precios que habria
   sin sus medidas sobre trosarina (impuesto especial) y primas. Dehesa Index lee esas notas con el buscador de vlada.gov.hr, mes a mes, y guarda cada una
   con su fecha de entrada en vigor. Si una semana no aparece en el buscador, no se inventa: queda un hueco. Los contenidos de vlada.gov.hr se pueden
   reutilizar sin permiso especial citando la fuente (condiciones de uso de la web).
   El «plavi dizel» (gasoleo azul) lo compran los agricultores, pescadores y acuicultores con autorizacion de la Carinska uprava (Aduana); tiene trosarina
   0,00 EUR/1000 l (Uredba o visini trošarine, NN 156/2022) frente a unos 330-390 EUR/1000 l del gasoleo normal en 2026.
2) Boletin Petrolero de la UE (Comision Europea, DG ENER): gasoleo de automocion con y sin impuestos y fioul de calefaccion de Croacia, semanal desde 2005, en
   EUR por 1.000 l -> EUR/l (Croacia usa el euro desde 2023; los datos anteriores estan convertidos por la Comision). Reutilizacion: Decision 2011/833/UE.
Si una parte falla se conserva la del fichero anterior. Uso: python3 scripts/update-hr-diesel.py"""
import calendar, datetime, html, io, json, re, sys, time, urllib.parse, urllib.request
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "hr-diesel.json"
WOB = "https://energy.ec.europa.eu/document/download/906e60ca-8b6a-44e7-8589-652854d2fd3f_en?filename=Weekly_Oil_Bulletin_Prices_History_maticni_4web.xlsx"
VL = "https://vlada.gov.hr"
SEARCH = VL + "/vijesti?pojam=%s&datumod=%s&datumdo=%s&tip=0&tip2=0&tema=0&profil=0&page=%d"
TERMS = ["goriva", "dizel", "naftnih derivata"]
FIRST = (2023, 1)          # primer mes de la carga inicial (Croacia usa euros desde el 1-1-2023)
H = {"User-Agent": "Dehesa-Index-data-bot/1.0 (+https://dehesaindex.com)"}
VAT = 0.25
MES = {"siječnja": 1, "veljače": 2, "ožujka": 3, "travnja": 4, "svibnja": 5, "lipnja": 6, "srpnja": 7, "kolovoza": 8, "rujna": 9, "listopada": 10, "studenoga": 11, "prosinca": 12}
def get(u, t=90, raw=False, tries=4):
    last = None
    for i in range(tries):
        try:
            with urllib.request.urlopen(urllib.request.Request(u, headers=H), timeout=t) as r:
                b = r.read(); return b if raw else b.decode("utf-8", "replace")
        except Exception as e: last = e; time.sleep(4 * (i + 1))
    raise RuntimeError("%s -> %s" % (u[-90:], last))
def text(h): return re.sub(r"\s+", " ", html.unescape(re.sub(r"<(script|style)[^>]*>.*?</\1>|<[^>]+>", " ", h, flags=re.S)))
def num(s): return round(float(s.replace(",", ".")), 2)
def parse(x):
    """Texto de una nota de la Vlada -> (fecha de entrada en vigor, [gasolina, gasoleo, gasoleo azul], [gasoleo sin medidas, azul sin medidas]) o None."""
    cut = x.find("Da nema mjera"); a, b = (x[:cut], x[cut:]) if cut > 0 else (x, "")
    def pr(s, k):
        m = re.search(r"(\d+,\d{2}) EUR/l za " + k, s); return num(m.group(1)) if m else None
    cur = [pr(a, "benzinsko gorivo"), pr(a, "dizelsko gorivo"), pr(a, "plavi dizel")]
    if cur[1] is None or cur[2] is None: return None
    sn = re.search(r"stupa na snagu (\d{1,2})\. (\w+) (20\d\d)", a)
    if sn and sn.group(2) in MES: d = datetime.date(int(sn.group(3)), MES[sn.group(2)], int(sn.group(1)))
    else:
        pub = re.search(r"(\d{1,2})\.(\d{1,2})\.(20\d\d)\.?\s*Ispi", a)
        if not pub: return None
        d = datetime.date(int(pub.group(3)), int(pub.group(2)), int(pub.group(1))) + datetime.timedelta(days=1)   # nota del lunes, precios desde el martes
    return d.isoformat(), cur, [pr(b, "dizelsko gorivo"), pr(b, "plavi dizel")] if b else [None, None]
def months(start, end):
    y, m = start
    while (y, m) <= end:
        yield y, m; m += 1
        if m > 12: y, m = y + 1, 1
def official(old):
    seen = set(old.get("seen", [])); rows = {r[0]: r for r in old.get("official", [])}
    today = datetime.date.today()
    if rows:
        d0 = today - datetime.timedelta(days=75); start = (d0.year, d0.month)
    else: start = FIRST
    found = {}
    for y, m in months(start, (today.year, today.month)):
        last = calendar.monthrange(y, m)[1]
        for term in TERMS:
            for pg in range(1, 8):
                h = get(SEARCH % (urllib.parse.quote_plus(term), "01.%02d.%d" % (m, y), "%d.%02d.%d" % (last, m, y), pg)); n0 = len(found)
                for mm in re.finditer(r'<a[^>]+href="(/[^"]+/\d{4,6})"[^>]*>(.*?)</a>', h, re.S):
                    title = text(mm.group(2)).strip()
                    if re.search(r"goriv|dizel|naftn", title, re.I): found.setdefault(mm.group(1), title)
                if len(found) == n0: break
                time.sleep(0.3)
    new = 0
    for u in sorted(found):
        if u in seen: continue
        try: x = text(get(VL + u))
        except Exception as e: print("AVISO", u, e); continue
        seen.add(u); time.sleep(0.3)
        p = parse(x)
        if not p: continue
        d, cur, nm = p
        for v in cur + nm:
            if v is not None and not 0.3 < v < 3.5: raise ValueError("precio fuera de rango en %s: %s %s" % (u, cur, nm))
        if cur[2] >= cur[1]: print("AVISO: azul >= gasoleo en", u, cur); continue
        row = [d, cur[0], cur[1], cur[2], nm[0], nm[1]]
        if d not in rows or rows[d][4] is None and row[4] is not None: rows[d] = row; new += 1
    out = [rows[k] for k in sorted(rows)]
    if len(out) < 3: raise ValueError("solo %d notas con precios" % len(out))
    print("notas de la Vlada: %d encontradas, %d nuevas con precios, %d en total" % (len(found), new, len(out)))
    return out, sorted(seen)
def weekly():
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(get(WOB, 300, True)), read_only=True, data_only=True)
    def col(ws, name):
        hdr = [str(x or "").strip() for x in next(ws.iter_rows(values_only=True, max_row=1))]; i = hdr.index(name); out = {}
        for r in ws.iter_rows(values_only=True, min_row=2):
            if not isinstance(r[0], datetime.datetime) or i >= len(r): continue
            try: out[r[0].date().isoformat()] = float(r[i]) / 1000
            except (TypeError, ValueError): pass
        return out
    wt = wb["Prices with taxes"]; wo = wb["Prices wo taxes"]
    dt = col(wt, "HR_price_with_tax_diesel"); ht = col(wt, "HR_price_with_tax_heating_oil"); dn = col(wo, "HR_price_wo_tax_diesel")
    out = [[d, round(dt[d], 4), round(dn[d], 4) if d in dn else None, round(ht[d], 4) if d in ht else None] for d in sorted(dt) if 0.3 < dt[d] < 3.5]
    if len(out) < 500: raise ValueError("serie semanal corta: %d" % len(out))
    for a, b in zip(out, out[1:]):
        if abs(b[1] - a[1]) > 0.6: raise ValueError("salto raro entre %s y %s" % (a[0], b[0]))
    return out
def main():
    try: old = json.loads(OUT.read_text())
    except Exception: old = {}
    doc = {"schemaVersion": 1, "generatedAt": datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"), "unit": "EUR/l", "vat": VAT,
           "source": {"name": "Vlada Republike Hrvatske: notas sobre los precios máximos de los derivados del petróleo (Uredba, NN); Comisión Europea, Weekly Oil Bulletin",
                      "url": "https://vlada.gov.hr/vijesti", "wobUrl": "https://energy.ec.europa.eu/data-and-analysis/weekly-oil-bulletin_en",
                      "license": "Contenidos de vlada.gov.hr reutilizables citando la fuente (condiciones de uso); Decisión 2011/833/UE (Comisión)"},
           "officialFields": ["date", "petrol", "diesel", "blueDiesel", "dieselNoMeasures", "blueDieselNoMeasures"], "weeklyFields": ["date", "dieselTax", "dieselNoTax", "heatingOilTax"],
           "blueDieselExcise": 0.0}
    try: doc["official"], doc["seen"] = official(old)
    except Exception as e:
        print("AVISO notas de la Vlada:", e)
        if old.get("official"): doc["official"], doc["seen"] = old["official"], old.get("seen", [])
    try: doc["weekly"] = weekly()
    except Exception as e:
        print("AVISO boletin petrolero:", e)
        if old.get("weekly"): doc["weekly"] = old["weekly"]
    if "weekly" not in doc or "official" not in doc: print("FALLO: falta una de las dos partes"); return 0 if OUT.exists() else 1
    last = doc["official"][-1]; w = doc["weekly"][-1]; doc["asOf"] = last[0]
    OUT.write_text(json.dumps(doc, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print("hr-diesel.json: %d precios oficiales (%s..%s), ultimo gasoleo %s, azul %s EUR/l; %d semanas del boletin, ultima %s gasoleo %.3f" % (
        len(doc["official"]), doc["official"][0][0], last[0], last[2], last[3], len(doc["weekly"]), w[0], w[1]))
    return 0
if __name__ == "__main__": sys.exit(main())
