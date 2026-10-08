#!/usr/bin/env python3
"""Francia: estado semanal de los cereales de paja (trigo blando, trigo duro, cebada de invierno y de primavera; Cere'Obs, FranceAgriMer) -> data/france-cereobs-cp.json
A mano, desde el Excel «SCR-GRC-CEREOBS_CP_depuis_2015» (hojas «Données France» y «Données régions»):
  python3 scripts/build-france-cereobs-cp.py <CEREOBS_CP_depuis_2015-*.xlsx>
Por semana: % de superficie en cada etapa (siembra, emergencia, inicio de ahijado, espiga a 1 cm, 2 nudos, espigado, cosecha) y % por clase de estado (muy mala ... muy buena).
Nacional: últimas 6 campañas; regiones: últimas 2 (para no engordar el fichero). Campaña = año de cosecha (la semana 36 en adelante ya es la campaña siguiente)."""
import datetime, json, re, sys
from pathlib import Path
import openpyxl
ROOT = Path(__file__).resolve().parents[1]
CROPS = [("Blé tendre", "soft-wheat"), ("Blé dur", "durum-wheat"), ("Orge d'hiver", "winter-barley"), ("Orge de printemps", "spring-barley")]
FIELDS = [("Semis", "sowing"), ("Levée", "emergence"), ("Début tallage", "tillering"), ("Épi 1 cm", "ear1cm"), ("2 noeuds", "nodes2"), ("Épiaison", "heading"), ("Récolte", "harvest"),
          ("Très mauvaises", "c1"), ("Mauvaises", "c2"), ("Assez bonnes", "c3"), ("Bonnes", "c4"), ("Très bonnes", "c5")]
def num(v): return round(float(v), 1) if isinstance(v, (int, float)) and not isinstance(v, bool) else None
def season(label):
    y, w = int(label[:4]), int(label[6:]); return y + 1 if w >= 36 else y
def sheet(ws, regional):
    rows = list(ws.iter_rows(values_only=True))
    h = next(i for i, r in enumerate(rows) if r and "Culture" in r)
    head = [str(c).strip() if c is not None else "" for c in rows[h]]
    ci, wi = head.index("Culture"), head.index("Semaine"); ri = head.index("Région") if regional else None
    idx = [head.index(fr) for fr, _k in FIELDS]  # falla si el Excel cambia de columnas
    out = {}
    for r in rows[h + 1:]:
        m = re.fullmatch(r"(\d{4})-S(\d{2})", str(r[wi]).strip()) if r and r[wi] else None
        if not m or r[ci] not in dict(CROPS): continue
        vals = [num(r[j]) for j in idx]
        if all(v is None for v in vals): continue
        key = (dict(CROPS)[r[ci]], r[ri] if regional else "FR")
        out.setdefault(key, []).append(["%s-W%s" % m.groups()] + vals)
    return out
def main(xlsx):
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    nat, reg = sheet(wb["Données France"], False), sheet(wb["Données régions"], True)
    last = max(season(r[0]) for v in nat.values() for r in v)
    crops = {}
    for _fr, cid in CROPS:
        n = [r for r in sorted(nat.get((cid, "FR"), []), key=lambda r: r[0]) if season(r[0]) > last - 6]
        regs = {k[1]: [r for r in sorted(v, key=lambda r: r[0]) if season(r[0]) > last - 2] for k, v in reg.items() if k[0] == cid}
        regs = {k: v for k, v in regs.items() if v}
        assert n and regs, "sin datos para " + cid
        crops[cid] = {"national": n, "regions": regs}
    until = None
    for r in wb["Présentation"].iter_rows(values_only=True):
        for c in r:
            if isinstance(c, str) and "arrêtées au" in c:
                m = re.search(r"(\d{2})/(\d{2})/(\d{4})", c); until = "%s-%s-%s" % (m.group(3), m.group(2), m.group(1)) if m else None
    doc = {"schemaVersion": 1, "generatedAt": datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"), "country": "FR",
           "source": {"id": "franceagrimer", "name": "FranceAgriMer, Cere'Obs (état des cultures, céréales à paille)", "url": "https://cereobs.franceagrimer.fr/", "license": "Licence Ouverte 2.0", "dataUntil": until},
           "fields": [k for _f, k in FIELDS], "units": "% of area (stages) / % of area by condition class (c1 very poor ... c5 very good)", "crops": crops}
    (ROOT / "data" / "france-cereobs-cp.json").write_text(json.dumps(doc, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print({c: (len(v["national"]), len(v["regions"]), sum(len(x) for x in v["regions"].values())) for c, v in crops.items()}, "hasta", until, (ROOT / "data" / "france-cereobs-cp.json").stat().st_size // 1024, "KB")
if __name__ == "__main__":
    if len(sys.argv) != 2: raise SystemExit(__doc__)
    main(sys.argv[1])
