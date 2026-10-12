#!/usr/bin/env python3
"""Irlanda: gasoleo de automocion semanal (Boletin Petrolero de la UE), gasoleo marcado (agricola) despachado y precio de los combustibles agrarios -> data/ie-diesel.json
Irlanda NO publica en abierto un precio del «marked gas oil» (el gasoleo marcado o «green diesel» que usan agricultores y otros usuarios autorizados, con un
tipo de impuesto especial reducido): el Boletin Petrolero de la UE solo trae el gasoleo de automocion, y los precios del gasoleo marcado que se ven en la prensa no
tienen una serie oficial abierta. Dehesa Index muestra, en su lugar, tres series oficiales que si existen y no se mezclan:
1) Boletin Petrolero de la UE (Comision Europea, DG ENER): gasoleo de automocion con y sin impuestos y gasoleo de calefaccion de Irlanda, semanal, EUR por 1.000 l
   -> EUR/l. Reutilizacion: Decision 2011/833/UE.
2) CSO (Oficina Central de Estadistica), tabla FPM01 «Fuel Excise Clearance»: litros de gasoleo de automocion, gasolina, gasoleo marcado y queroseno que salen de los
   depositos fiscales cada mes, en millones de litros (datos de la Hacienda irlandesa, Revenue). Es volumen, no precio.
3) CSO, tabla AHM05 «Agricultural Input and Output Price Indices» (base 2020=100): indice del precio que pagan los agricultores por la energia (022), el combustible
   de motores (0222) y todos los insumos (02).
Licencia de los datos del CSO: CC BY 4.0. Si una parte falla se conserva la del fichero anterior. Uso: python3 scripts/update-ie-diesel.py"""
import datetime, io, json, sys, time, urllib.request
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "ie-diesel.json"
WOB = "https://energy.ec.europa.eu/document/download/906e60ca-8b6a-44e7-8589-652854d2fd3f_en?filename=Weekly_Oil_Bulletin_Prices_History_maticni_4web.xlsx"
CSO = "https://ws.cso.ie/public/api.restful/PxStat.Data.Cube_API.ReadDataset/%s/JSON-stat/2.0/en"
H = {"User-Agent": "Dehesa-Index-data-bot/1.0 (+https://dehesaindex.com)"}
def get(u, t=120, tries=4):
    last = None
    for i in range(tries):
        try:
            with urllib.request.urlopen(urllib.request.Request(u, headers=H), timeout=t) as r: return r.read()
        except Exception as e: last = e; time.sleep(4 * (i + 1))
    raise RuntimeError("%s -> %s" % (u[-90:], last))
def jsonstat(code, raw=None):
    """Tabla del CSO (JSON-stat 2.0) -> (periodos, categorias de la ultima dimension, funcion v(i, j)). Dimensiones: STATISTIC, tiempo, categoria."""
    d = json.loads(raw if raw is not None else get(CSO % code)); ids = d["id"]
    if len(ids) != 3 or d["size"][0] != 1: raise ValueError("%s: dimensiones inesperadas %s" % (code, d["size"]))
    ix = lambda k: k if isinstance(k, dict) else {c: i for i, c in enumerate(k)}
    T = ix(d["dimension"][ids[1]]["category"]["index"]); C = ix(d["dimension"][ids[2]]["category"]["index"]); v = d["value"]
    ts = sorted(T, key=lambda k: T[k]); nc = len(C)
    return ts, {k: C[k] for k in C}, (lambda t, c: v[T[t] * nc + C[c]])
def weekly():
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(get(WOB, 300)), read_only=True, data_only=True)
    def col(ws, name):
        hdr = [str(x or "").strip() for x in next(ws.iter_rows(values_only=True, max_row=1))]; i = hdr.index(name); out = {}
        for r in ws.iter_rows(values_only=True, min_row=2):
            if not isinstance(r[0], datetime.datetime) or i >= len(r): continue
            try: out[r[0].date().isoformat()] = float(r[i]) / 1000
            except (TypeError, ValueError): pass
        return out
    wt = wb["Prices with taxes"]; wo = wb["Prices wo taxes"]
    dt = col(wt, "IE_price_with_tax_diesel"); ht = col(wt, "IE_price_with_tax_heating_oil"); dn = col(wo, "IE_price_wo_tax_diesel")
    out = [[d, round(dt[d], 4), round(dn[d], 4) if d in dn else None, round(ht[d], 4) if d in ht else None] for d in sorted(dt) if 0.3 < dt[d] < 3.5]
    if len(out) < 500: raise ValueError("serie semanal corta: %d" % len(out))
    for a, b in zip(out, out[1:]):
        if abs(b[1] - a[1]) > 0.6: raise ValueError("salto raro entre %s y %s" % (a[0], b[0]))
    return out
def clearance(raw=None):
    ts, C, v = jsonstat("FPM01", raw); out = []
    for t in ts:
        if t < "201501": continue
        r = [v(t, c) for c in ("1", "2", "3", "4")]
        if any(x is None for x in r): continue
        out.append(["%s-%s" % (t[:4], t[4:]), r[0], r[1], r[2], r[3]])
    if len(out) < 90: raise ValueError("FPM01 corta: %d meses" % len(out))
    for r in out[-12:]:
        if not all(0 < x < 1000 for x in r[1:]): raise ValueError("FPM01 fuera de rango: %s" % r)
    return out
def prices(raw=None):
    ts, C, v = jsonstat("AHM05", raw); out = []
    for t in ts:
        r = [v(t, c) for c in ("0222", "022", "02")]
        if any(x is None for x in r): continue
        out.append(["%s-%s" % (t[:4], t[4:]), round(r[0], 2), round(r[1], 2), round(r[2], 2)])
    if len(out) < 60: raise ValueError("AHM05 corta: %d meses" % len(out))
    if not all(30 < x < 300 for r in out for x in r[1:]): raise ValueError("AHM05 fuera de rango")
    return out
def main():
    try: old = json.loads(OUT.read_text())
    except Exception: old = {}
    doc = {"schemaVersion": 1, "generatedAt": datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"), "unit": "EUR/l",
           "source": {"name": "Comisión Europea, Weekly Oil Bulletin; CSO Irlanda, Fuel Excise Clearance (FPM01) y Agricultural Input and Output Price Indices (AHM05)",
                      "url": "https://www.cso.ie/en/statistics/energy/fuelexciseclearances/", "wobUrl": "https://energy.ec.europa.eu/data-and-analysis/weekly-oil-bulletin_en",
                      "csoPricesUrl": "https://www.cso.ie/en/statistics/agriculture/agriculturalpriceindices/", "license": "CC BY 4.0 (CSO); Decisión 2011/833/UE (Comisión)"},
           "weeklyFields": ["date", "dieselTax", "dieselNoTax", "heatingOilTax"], "clearanceFields": ["month", "autodiesel", "petrol", "markedGasOil", "kerosene"], "clearanceUnit": "million litres",
           "priceFields": ["month", "motorFuel", "energy", "allInputs"], "priceBase": "2020=100", "vat": 0.23}
    for key, fn in (("weekly", weekly), ("clearance", clearance), ("prices", prices)):
        try: doc[key] = fn()
        except Exception as e:
            print("AVISO %s:" % key, e)
            if old.get(key): doc[key] = old[key]
    if not all(k in doc for k in ("weekly", "clearance", "prices")): print("FALLO: falta una parte"); return 0 if OUT.exists() else 1
    doc["asOf"] = doc["weekly"][-1][0]
    OUT.write_text(json.dumps(doc, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    w = doc["weekly"][-1]; c = doc["clearance"][-1]; p = doc["prices"][-1]
    print("ie-diesel.json: %d semanas (ultima %s gasoleo %.3f EUR/l); gasoleo marcado %s: %s Ml (%d meses); indice combustible de motores %s: %s" % (len(doc["weekly"]), w[0], w[1], c[0], c[3], len(doc["clearance"]), p[0], p[1]))
    return 0
if __name__ == "__main__": sys.exit(main())
