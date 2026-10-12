#!/usr/bin/env python3
"""Belgica: gasoleo agricola (precio maximo oficial diario) y gasoleo/fioul semanales desde 2005 -> data/be-diesel.json
1) Tarifa oficial de los productos petroliferos (precios maximos de venta, IVA incluido), que fija cada dia laborable la Direccion General de
   Energia del SPF Economia con el contrato-programa. Incluye el «Gasolie Diesel Landbouw» (gasoleo para uso agricola, con impuesto especial
   reducido) en dos tramos: desde 2.000 l y menos de 2.000 l. Se lee la vista publica de be.STAT (Statbel, SPF Economia) con las ultimas 365
   modificaciones; Dehesa Index guarda cada dia en un historico propio que crece. Statbel publica sus datos bajo CC BY 4.0 y el SPF Economia
   permite reutilizar la informacion de su web incluso con fines comerciales, citando la fuente.
2) Boletin Petrolero de la UE (Comision Europea, DG ENER): gasoleo de automocion con y sin impuestos y fioul de calefaccion con impuestos de
   Belgica, semanal desde 2005, en EUR por 1.000 l -> EUR/l. Reutilizacion: Decision 2011/833/UE.
Precio sin IVA = precio / 1,21. Los agricultores en regimen general recuperan el IVA; los acogidos al regimen agricola a tanto alzado no.
Si una parte falla se conserva la del fichero anterior. Uso: python3 scripts/update-be-diesel.py"""
import csv, datetime, io, json, re, sys, time, urllib.request
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "be-diesel.json"
BESTAT = "https://bestat.statbel.fgov.be/bestat/api/views/f8d4c0ce-bec4-4f58-b2a2-4f6b8667ad24/result/CSV"
WOB = "https://energy.ec.europa.eu/document/download/906e60ca-8b6a-44e7-8589-652854d2fd3f_en?filename=Weekly_Oil_Bulletin_Prices_History_maticni_4web.xlsx"
H = {"User-Agent": "Dehesa-Index-data-bot/1.0 (+https://dehesaindex.com)"}
VAT = 0.21
MES = {"jan": 1, "feb": 2, "mrt": 3, "apr": 4, "mei": 5, "jun": 6, "jul": 7, "aug": 8, "sep": 9, "okt": 10, "nov": 11, "dec": 12}
# columnas del historico diario: producto oficial -> patron sobre el nombre de be.STAT
PROD = [("farmBig", r"^Gasolie Diesel Landbouw \(vanaf 2000"), ("farmSmall", r"^Gasolie Diesel Landbouw \(minder dan 2000"),
        ("roadB7", r"^Diesel B7\b"), ("heatBig", r"^Gasolie verwarming .*\(vanaf 2000")]
def get(u, t=300):
    last = None
    for i in range(4):
        try:
            with urllib.request.urlopen(urllib.request.Request(u, headers=H), timeout=t) as r: return r.read()
        except Exception as e: last = e; time.sleep(6 * (i + 1))
    raise RuntimeError("%s -> %s" % (u[-90:], last))
def day(s):
    m = re.match(r"^(\d{1,2})([a-z]{3})(\d{2})$", s.strip().lower())
    if not m or m.group(2) not in MES: return None
    return datetime.date(2000 + int(m.group(3)), MES[m.group(2)], int(m.group(1))).isoformat()
def official():
    rows = list(csv.DictReader(io.StringIO(get(BESTAT).decode("utf-8-sig"))))
    by = {}
    for r in rows:
        d = day(r.get("Dag") or ""); p = (r.get("Product") or "").strip()
        try: v = float(r.get("Prijs incl. BTW"))
        except (TypeError, ValueError): continue
        if not d: continue
        for k, pat in PROD:
            if re.search(pat, p): by.setdefault(d, {})[k] = round(v, 4)
    out = [[d] + [by[d].get(k) for k, _ in PROD] for d in sorted(by)]
    out = [x for x in out if x[1] is not None]
    if len(out) < 100: raise ValueError("solo %d dias con gasoleo agricola" % len(out))
    for x in out:
        for v in x[1:]:
            if v is not None and not 0.3 < v < 3.5: raise ValueError("precio fuera de rango %s: %s" % (x[0], x))
    return out
def weekly():
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(get(WOB)), read_only=True, data_only=True)
    def col(ws, name):
        hdr = [str(x or "").strip() for x in next(ws.iter_rows(values_only=True, max_row=1))]; i = hdr.index(name); out = {}
        for r in ws.iter_rows(values_only=True, min_row=2):
            if not isinstance(r[0], datetime.datetime) or i >= len(r): continue
            try: out[r[0].date().isoformat()] = float(r[i]) / 1000
            except (TypeError, ValueError): pass
        return out
    wt = wb["Prices with taxes"]; wo = wb["Prices wo taxes"]
    dt = col(wt, "BE_price_with_tax_diesel"); ht = col(wt, "BE_price_with_tax_heating_oil"); dn = col(wo, "BE_price_wo_tax_diesel")
    out = [[d, round(dt[d], 4), round(dn[d], 4) if d in dn else None, round(ht[d], 4) if d in ht else None] for d in sorted(dt) if 0.5 < dt[d] < 3.5]
    if len(out) < 500: raise ValueError("serie semanal corta: %d" % len(out))
    return out
def main():
    try: old = json.loads(OUT.read_text())
    except Exception: old = {}
    doc = {"schemaVersion": 1, "generatedAt": datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"), "unit": "EUR/l", "vat": VAT,
           "source": {"name": "SPF Economía, DG Energía: tarifa oficial de los productos petrolíferos (vía be.STAT, Statbel); Comisión Europea, Weekly Oil Bulletin",
                      "url": "https://economie.fgov.be/fr/themes/energie/prix-de-lenergie/prix-maximum-des-produits/tarif-officiel-des-produits", "bestat": BESTAT,
                      "wobUrl": "https://energy.ec.europa.eu/data-and-analysis/weekly-oil-bulletin_en", "license": "CC BY 4.0 (Statbel); reutilización libre citando la fuente (SPF Economía); Decisión 2011/833/UE (Comisión)"},
           "dailyFields": ["date"] + [k for k, _ in PROD], "weeklyFields": ["date", "dieselTax", "dieselNoTax", "heatingOilTax"]}
    try:
        new = official(); hist = {x[0]: x for x in old.get("daily", [])}
        for x in new: hist[x[0]] = x
        doc["daily"] = [hist[d] for d in sorted(hist)]
    except Exception as e:
        print("AVISO tarifa oficial:", e)
        if old.get("daily"): doc["daily"] = old["daily"]
    try: doc["weekly"] = weekly()
    except Exception as e:
        print("AVISO boletin petrolero:", e)
        if old.get("weekly"): doc["weekly"] = old["weekly"]
    if "daily" not in doc or "weekly" not in doc: print("FALLO: falta una de las dos partes"); return 0 if OUT.exists() else 1
    last = doc["daily"][-1]; doc["asOf"] = last[0]; doc["farmNoVat"] = round(last[1] / (1 + VAT), 4)
    OUT.write_text(json.dumps(doc, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    w = doc["weekly"][-1]
    print("be-diesel.json: %d dias (%s..%s), gasoleo agricola %s EUR/l (sin IVA %.3f), automocion B7 %s; %d semanas, ultima %s gasoleo %.3f" % (
        len(doc["daily"]), doc["daily"][0][0], last[0], last[1], doc["farmNoVat"], last[3], len(doc["weekly"]), w[0], w[1]))
    return 0
if __name__ == "__main__": sys.exit(main())
