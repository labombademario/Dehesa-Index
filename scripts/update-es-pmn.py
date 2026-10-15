#!/usr/bin/env python3
"""España: Precios Medios Nacionales semanales (MAPA, Subdirección General de Análisis, Coordinación y Estadística)
-> data/spain-prices/pmn.json
y «Observatorio de precios de los alimentos» (precios semanales origen y mayorista MERCASA, 2004-hoy) -> data/spain-prices/obs.json

Fuente: un libro Excel por año en la página «precios-medios-nacionales» del MAPA. Cada libro: filas = productos, columnas = «Semana NN».
La fecha de cada semana NO se lee de la etiqueta (tiene erratas) sino que se calcula como el lunes de la semana ISO (año, nº de semana), que es lo que
coincide con los rangos publicados (p. ej. 2026 semana 1 = 29/12-04/01). Un «-» es dato ausente (null). Nada se estima ni se rellena."""
import argparse, datetime, io, json, re, sys, time, urllib.parse, urllib.request, html
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
PAGE = "https://www.mapa.gob.es/es/estadistica/temas/estadisticas-agrarias/economia/precios-medios-nacionales"
UA = {"User-Agent": "Mozilla/5.0 (compatible; DehesaIndex/1.0; +https://dehesaindex.com)"}
OBS_PAGE = "https://www.mapa.gob.es/es/estadistica/temas/estadisticas-alimentacion/observatorio-precios"
OBS_ALIAS = {"gallos": "gallo", "naranja-tipo-navel": "naranja-navel"}
ALIAS = {"trigo blando panificable": "trigo blando", "trigo blando (desde s27-2025)": "trigo blando"}
LOG = []
ZEROS = []
def log(*a):
    s = " ".join(str(x) for x in a); LOG.append(s); print(s, flush=True)

def http(url, tries=3, timeout=240):
    last = None
    for i in range(tries):
        try:
            with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=timeout) as r: return r.read()
        except Exception as e:
            last = e; time.sleep(4 * (i + 1))
    raise RuntimeError("%s -> %s" % (url, str(last)[:160]))

def page_files(htm):
    out = []; seen = set()
    for h, t in re.findall(r'href="([^"#]+\.xlsx?)"[^>]*>(.*?)</a>', htm, re.S | re.I):
        u = urllib.parse.urljoin(PAGE, html.unescape(h)); name = urllib.parse.unquote(u.split("/")[-1])
        m = re.search(r"precios_medios_nacionales.*?(20\d\d)", name, re.I)
        if not m or u in seen: continue
        seen.add(u); out.append((int(m.group(1)), u, name))
    return out

def split_unit(label):
    s = re.sub(r"\s+", " ", str(label or "")).strip()
    m = re.search(r"\(([^()]*(?:€|unidad)[^()]*)\)\s*$", s)
    if m: return s[:m.start()].strip(" ,"), m.group(1).strip()
    m = re.search(r"\((€[^()]*)\)", s)
    if m: return (s[:m.start()] + s[m.end():]).strip(" ,"), m.group(1).strip()
    return s, None

def num(v):
    if isinstance(v, bool) or v is None: return None
    if isinstance(v, (int, float)):
        if v <= 0: ZEROS.append(v); return None   # un precio de 0 o negativo no existe: el libro lo usa donde no hubo operaciones
        return round(float(v), 3)
    return None

def parse(data, year):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True); ws = wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    hi = next((i for i, r in enumerate(rows) if r and str(r[0] or "").strip().upper() == "SECTOR"), None)
    if hi is None: raise ValueError("sin fila de cabecera SECTOR")
    cols = {}
    for j, c in enumerate(rows[hi]):
        m = re.match(r"\s*Semana\s+(\d+)\s*$", str(c or ""), re.I)
        if m: cols[j] = int(m.group(1))
    if not cols: raise ValueError("sin columnas Semana")
    sector = None; foot = None; prods = []
    for r in rows[hi + 1:]:
        if not r: continue
        if r[0] and str(r[0]).strip().lower().startswith("posici"): foot = re.sub(r"\s+", " ", str(r[0])).strip(); continue
        name = r[2] if len(r) > 2 else None
        if r[0] and str(r[0]).strip(): sector = re.sub(r"\s+", " ", str(r[0])).strip()
        if not name or not str(name).strip(): continue
        vals = {w: num(r[j]) for j, w in cols.items() if j < len(r)}
        nm, unit = split_unit(name)
        code = str(r[1]).strip() if len(r) > 1 and r[1] else None
        prods.append({"label": nm, "unit": unit, "sector": sector, "pos": code, "vals": vals})
    return prods, foot, max(cols.values())

def monday(y, w):
    try: return datetime.date.fromisocalendar(y, w, 1)
    except ValueError: return None

def slug(s):
    s = s.lower(); s = ALIAS.get(s, s)
    for a, b in (("á", "a"), ("é", "e"), ("í", "i"), ("ó", "o"), ("ú", "u"), ("ñ", "n"), ("ü", "u")): s = s.replace(a, b)
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")

def build(parsed):
    """parsed: {año: (prods, foot, maxw)} -> documento."""
    weeks = []; widx = {}
    for y in sorted(parsed):
        for w in range(1, parsed[y][2] + 1):
            d = monday(y, w)
            if d and d.isoformat() not in widx: widx[d.isoformat()] = len(weeks); weeks.append(d.isoformat())
    ser = {}
    for y in sorted(parsed):
        for p in parsed[y][0]:
            key = slug(p["label"]); canon = ALIAS.get(p["label"].lower(), p["label"])
            s = ser.setdefault(key, {"id": key, "name": canon, "unit": p["unit"], "sector": p["sector"], "pos": p["pos"], "v": [None] * len(weeks), "_y": y})
            s["name"] = canon if y >= s["_y"] else s["name"]; s["unit"] = p["unit"] or s["unit"]; s["sector"] = p["sector"] or s["sector"]; s["pos"] = p["pos"] or s["pos"]; s["_y"] = y
            for w, v in p["vals"].items():
                d = monday(y, w)
                if d is None or v is None: continue
                i = widx[d.isoformat()]
                if s["v"][i] is None: s["v"][i] = v
    out = []
    for s in ser.values():
        if not any(x is not None for x in s["v"]): continue
        s.pop("_y"); out.append(s)
    return weeks, out


def parse_obs(data):
    """Libro del observatorio: una hoja por año (nombre = año). Cada hoja tiene dos bloques: «PRECIOS EN ORIGEN» y «PRECIOS MERCASA» (mayorista en destino);
    filas = productos (col C), unidad (col D), semanas «S n» desde la col E. Devuelve {(año, bloque, producto): (unidad, {semana: valor})}."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True); out = {}; order = []
    for ws in wb:
        if not re.fullmatch(r"20\d\d", ws.title.strip()): continue
        y = int(ws.title.strip()); rows = [list(r) for r in ws.iter_rows(values_only=True)]; blk = None; cols = None
        for r in rows:
            c = r[2] if len(r) > 2 else None; u = r[3] if len(r) > 3 else None
            if isinstance(c, str) and c.strip().upper().startswith("PRECIOS"):
                t = c.strip().upper(); blk = "origen" if "ORIGEN" in t else "mercasa" if ("MERCASA" in t or "DESTINO" in t) else None; cols = None; continue
            if isinstance(c, str) and c.strip().lower() == "productos":
                cols = {j: int(m.group(1)) for j, h in enumerate(r) if j >= 4 and isinstance(h, str) and (m := re.match(r"\s*S\s*(\d+)\s*$", h))}; continue
            if not blk or not cols or not isinstance(c, str) or not c.strip() or c.strip().startswith("("): continue
            if not isinstance(u, str) or not u.strip(): continue
            vals = {w: num(r[j]) for j, w in cols.items() if j < len(r)}
            key = (y, blk, re.sub(r"\s*\(\d+\)\s*$", "", re.sub(r"\s+", " ", c)).strip())
            if key in out: continue
            out[key] = (u.strip(), vals); order.append(key)
    return out, order

def build_obs(parsed):
    years = sorted({k[0] for k in parsed}); weeks = []; widx = {}
    for y in years:
        mw = max((w for k, (u, v) in parsed.items() if k[0] == y for w, x in v.items() if x is not None), default=0)
        for w in range(1, mw + 1):
            d = monday(y, w)
            if d and d.isoformat() not in widx: widx[d.isoformat()] = len(weeks); weeks.append(d.isoformat())
    ser = {}
    for (y, blk, name), (unit, vals) in sorted(parsed.items(), key=lambda kv: kv[0][0]):
        sid = blk + "-" + OBS_ALIAS.get(slug(name), slug(name))
        okunit = unit.lower() in ("kg", "unidad", "docena", "litro", "l")
        s = ser.setdefault(sid, {"id": sid, "name": name, "unit": unit if okunit else "kg", "stage": blk, "v": [None] * len(weeks)})
        s["name"] = name
        if okunit: s["unit"] = unit
        for w, v in vals.items():
            d = monday(y, w)
            if d is None or v is None: continue
            i = widx[d.isoformat()]
            if s["v"][i] is None: s["v"][i] = v
    return weeks, [s for s in ser.values() if any(x is not None for x in s["v"])]

def obs_url(htm):
    for h in re.findall(r'href="([^"#]+cuadro-de-precios-semanales\.xlsx)"', htm, re.I): return urllib.parse.urljoin(OBS_PAGE, html.unescape(h))
    return None

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(ROOT / "data" / "spain-prices"))
    ap.add_argument("--local-obs", dest="local_obs", default=None)
    ap.add_argument("--local", action="append", default=[], help="año:ruta.xlsx (sin red; solo pruebas)")
    a = ap.parse_args(); out = Path(a.out); out.mkdir(parents=True, exist_ok=True)
    now = datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    items = []
    if a.local:
        for spec in a.local:
            y, p = spec.split(":", 1); items.append((int(y), p, Path(p).name))
    else:
        items = page_files(http(PAGE).decode("utf-8", "replace"))
        if not items: raise SystemExit("la página no tiene libros de precios medios: no se toca nada")
    by = {}
    for y, u, n in items: by.setdefault(y, []).append((u, n))
    parsed = {}; files = []; foots = {}
    for y in sorted(by):
        best = None
        for u, n in by[y]:
            try:
                data = Path(u).read_bytes() if a.local else http(u); pr, foot, mw = parse(data, y)
            except Exception as e:
                log("ERROR", n, str(e)[:120]); continue
            cnt = sum(1 for p in pr for v in p["vals"].values() if v is not None)
            log("leído", n, "productos", len(pr), "celdas", cnt, "semanas", mw)
            if not best or cnt > best[0]: best = (cnt, pr, foot, mw, n, u)
        if best:
            parsed[y] = (best[1], best[2], best[3]); files.append({"year": y, "file": best[4], "products": len(best[1]), "cells": best[0]}); foots[y] = best[2]
    if not parsed: raise SystemExit("no se leyó ningún libro: no se toca nada")
    weeks, series = build(parsed)
    last = max(parsed)
    legend = foots.get(last) or ""
    doc = {"schemaVersion": 1, "generatedAt": now, "source": {"id": "mapa_es", "name": "MAPA, Precios Medios Nacionales", "url": PAGE},
           "weeks": weeks, "weekDefinition": "lunes de la semana ISO; los precios cubren lunes a domingo", "series": sorted(series, key=lambda s: (s["sector"] or "", s["name"])),
           "legend": legend, "files": files,
           "note": "Precios en la posición comercial indicada por «pos» (ver leyenda). Un valor null es dato no publicado. Semanas del año en curso: avance provisional."}
    # --- observatorio (origen / MERCASA) ---
    try:
        if a.local_obs: od = Path(a.local_obs).read_bytes(); ou = a.local_obs
        else:
            ou = obs_url(http(OBS_PAGE).decode("utf-8", "replace"))
            od = http(ou) if ou else None
        if od:
            pobs, _ = parse_obs(od); w2, s2 = build_obs(pobs)
            if s2:
                odoc = {"schemaVersion": 1, "generatedAt": now, "source": {"id": "mapa_es", "name": "MAPA, Observatorio de precios de los alimentos", "url": OBS_PAGE},
                        "weeks": w2, "weekDefinition": "lunes de la semana ISO", "unit": "€/kg salvo indicación en «unit»", "series": sorted(s2, key=lambda x: (x["stage"], x["name"])),
                        "stages": {"origen": "Precio en origen (productor)", "mercasa": "Precio mayorista en destino (Mercas, MERCASA)"},
                        "note": "Un valor null es dato no publicado. Los nombres de producto pueden cambiar entre años; se unen por nombre normalizado."}
                (out / "obs.json").write_text(json.dumps(odoc, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
                log("observatorio: series", len(s2), "semanas", len(w2), w2[0], w2[-1])
        else: log("observatorio: sin enlace al libro, se conserva lo anterior")
    except Exception as e:
        log("ERROR observatorio:", str(e)[:200])
    (out / "pmn.json").write_text(json.dumps(doc, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
    log("series", len(series), "semanas", len(weeks), weeks[0], weeks[-1])
    if ZEROS: log("valores <= 0 tratados como sin dato:", len(ZEROS))
    (ROOT / "data" / "spain-prices-log.txt").write_text("\n".join(LOG) + "\n", encoding="utf-8")
    return 0

if __name__ == "__main__":
    sys.exit(main())
