#!/usr/bin/env python3
"""Paquete de datos descargable para empresas -> data/pack/<mercado>/<producto>.csv|xlsx + data/pack/index.json

Un fichero por producto y mercado (EE. UU., UE, Reino Unido, Canada) con todo el historico publicado en data/prices/history.
Cada fila lleva fuente, licencia y texto de atribucion. SOLO entran las fuentes cuyo veredicto comercial es «apto» o «condiciones»
(data/commercial-summary.json, que se calcula de data/license-registry.json): si una fuente pasa a «no apto», sus ficheros
desaparecen del paquete en la siguiente ejecucion. Nada se estima ni se rellena. Uso: python3 scripts/build-data-pack.py [--check]"""
import csv, datetime, io, json, shutil, sys
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, Alignment

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(Path(__file__).resolve().parent))
from price_names import index_name
OUT = ROOT / 'data' / 'pack'
REGIONS = ['us', 'eu', 'uk', 'ca']
MONTHS = {m: i + 1 for i, m in enumerate(['JAN', 'FEB', 'MAR', 'APR', 'MAY', 'JUN', 'JUL', 'AUG', 'SEP', 'OCT', 'NOV', 'DEC'])}
FIELDS = ['product_key', 'product', 'market', 'date', 'value', 'unit', 'currency', 'frequency', 'source', 'license', 'attribution']

def J(p): return json.loads((ROOT / p).read_text(encoding='utf-8'))

def period_iso(year, period):
    """Normaliza el periodo del historico a AAAA-MM, AAAA-MM-DD o AAAA-Qn. None si no se reconoce (la fila se omite, no se inventa)."""
    p = str(period).strip()
    if p.upper()[:3] in MONTHS and not p[:1].isdigit(): return '%04d-%02d' % (year, MONTHS[p.upper()[:3]])
    if p.upper().startswith('Q') and p[1:].isdigit(): return '%04d-%s' % (year, p.upper())
    if '-' in p:
        a, b = p.split('-', 1)
        if a.isdigit() and b.isdigit() and 1 <= int(a) <= 12 and 1 <= int(b) <= 31: return '%04d-%02d-%02d' % (year, int(a), int(b))
        return None
    if p.isdigit() and 1 <= int(p) <= 12: return '%04d-%02d' % (year, int(p))
    return None

def main():
    check = '--check' in sys.argv
    reg = J('data/license-registry.json')['sources']
    alias = {}
    for k, v in reg.items():
        alias[k] = k
        for a in v.get('aliases', []): alias[a] = k
    verdicts = J('data/commercial-summary.json')['sources']
    cite = J('data/views/license-cite.json')['sources']
    names = {}
    for x in J('data/app/v1/prices.json')['prices']: names[(x['product'], x['region'])] = x
    # Fecha de referencia = la del ultimo dato verificado (no la hora de ejecucion): sin datos nuevos, el paquete no cambia y no genera commits.
    snaps = []
    for region in REGIONS:
        d = J('data/prices/latest/%s.json' % region)
        for o in (d.get('observations') if isinstance(d, dict) else d): snaps.append(str(o.get('snapshotDate') or o.get('observationDate') or ''))
    now = max(snaps) if snaps else ''
    index, excluded, written = [], [], set()
    for region in REGIONS:
        d = J('data/prices/latest/%s.json' % region)
        obs = d.get('observations') if isinstance(d, dict) else d
        for o in obs:
            prod, sid = o['product'], alias.get(o.get('sourceId'), o.get('sourceId'))
            n = names.get((prod, region), {})
            nmap = n.get('name') or index_name(prod) or {}
            label = lambda lg: nmap.get(lg) or prod
            verdict = (verdicts.get(sid) or {}).get('verdict')
            hp = ROOT / ('data/prices/history/%s/%s.json' % (region, prod))
            if verdict not in ('apto', 'condiciones') or not hp.exists():
                excluded.append({'product': prod, 'region': region, 'source': sid, 'reason': 'licencia' if verdict not in ('apto', 'condiciones') else 'sin historico'})
                continue
            h = J('data/prices/history/%s/%s.json' % (region, prod))
            c = cite.get(sid) or {}
            r = reg.get(sid, {})
            src = (r.get('name') or c.get('name') or sid)
            lic = r.get('licenseName') or c.get('licenseName') or ''
            lic_id = r.get('licenseId') or c.get('licenseId') or ''
            attr = r.get('attributionText') or c.get('attributionText') or src
            rows = []
            for e in h['history']:
                dt = period_iso(e['year'], e['period'])
                if dt is None or e.get('value') is None: continue
                rows.append((dt, e['value']))
            if not rows: continue
            rows.sort(key=lambda t: t[0])
            unit = h.get('unit') or o.get('unit'); cur = h.get('currency') or o.get('currency'); freq = h.get('frequency') or o.get('frequency')
            market = region.upper()
            base = OUT / region
            if not check: base.mkdir(parents=True, exist_ok=True)
            cp, xp = base / (prod + '.csv'), base / (prod + '.xlsx')
            changed = False
            line = lambda dt, v: [prod, label('es'), market, dt, v, unit, cur, freq, src, lic_id, attr]
            if not check:
                buf = io.StringIO(); w = csv.writer(buf, lineterminator='\n'); w.writerow(FIELDS)
                for dt, v in rows: w.writerow(line(dt, v))
                txt = '﻿' + buf.getvalue()
                changed = not cp.exists() or cp.read_text(encoding='utf-8') != txt or not xp.exists()
                if changed: cp.write_text(txt, encoding='utf-8')
            written.update([cp, xp])
            if not check and changed:
                wb = openpyxl.Workbook(); wb.properties.created = wb.properties.modified = datetime.datetime(2026, 1, 1); ws = wb.active; ws.title = 'Datos'; ws.append(FIELDS)
                for c_ in ws[1]: c_.font = Font(bold=True)
                for dt, v in rows: ws.append(line(dt, v))
                for col, wdt in zip('ABCDEFGHIJK', (14, 18, 8, 12, 12, 12, 9, 11, 40, 16, 60)): ws.column_dimensions[col].width = wdt
                ws.freeze_panes = 'A2'
                s2 = wb.create_sheet('Fuente y licencia')
                info = [('Producto', label('es') + ' (' + prod + ')'), ('Mercado', market + (' · ' + (n.get('place') or {}).get('es', '') if n.get('place') else '')),
                        ('Descripcion del dato', (n.get('note') or {}).get('es', '')), ('Unidad', '%s (%s)' % (unit, cur)), ('Frecuencia', freq),
                        ('Primer dato', rows[0][0]), ('Ultimo dato', rows[-1][0]), ('Filas', len(rows)),
                        ('Fuente', src), ('URL de la fuente', r.get('url') or c.get('url') or ''), ('Licencia', lic), ('URL de la licencia', r.get('licenseUrl') or c.get('licenseUrl') or ''),
                        ('Texto de atribucion', attr), ('Veredicto de uso comercial', 'Apto' if verdict == 'apto' else 'Apto con condiciones'),
                        ('Condiciones', r.get('additionalRestrictions') or ''), ('Datos a fecha de', now),
                        ('Aviso', 'Dehesa Index no estima ni rellena valores. Cita la fuente al reutilizar. Esto no es asesoramiento juridico: la licencia de la fuente manda. Mas informacion: https://dehesaindex.com/condiciones')]
                for k, v in info: s2.append([k, v])
                for row in s2.iter_rows(min_col=1, max_col=1):
                    for c_ in row: c_.font = Font(bold=True)
                s2.column_dimensions['A'].width = 28; s2.column_dimensions['B'].width = 110
                for row in s2.iter_rows(min_col=2, max_col=2):
                    for c_ in row: c_.alignment = Alignment(wrap_text=True, vertical='top')
                wb.save(xp)
            index.append({'product': prod, 'name': {lg: label(lg) for lg in ('es', 'en', 'fr', 'it')}, 'region': region, 'place': n.get('place') or {}, 'source': sid, 'sourceName': src,
                          'license': lic_id, 'verdict': verdict, 'attribution': attr, 'unit': unit, 'currency': cur, 'frequency': freq, 'rows': len(rows), 'first': rows[0][0], 'last': rows[-1][0],
                          'csv': 'data/pack/%s/%s.csv' % (region, prod), 'xlsx': 'data/pack/%s/%s.xlsx' % (region, prod),
                          'csvBytes': cp.stat().st_size if cp.exists() else 0, 'xlsxBytes': xp.stat().st_size if xp.exists() else 0})
    index.sort(key=lambda x: (REGIONS.index(x['region']), x['name']['es']))
    doc = {'schemaVersion': 1, 'generatedAt': now, 'policy': 'Solo fuentes con veredicto comercial apto o con condiciones (data/commercial-summary.json). Sin estimaciones ni relleno.',
           'columns': FIELDS, 'summary': {'files': len(index), 'excluded': len(excluded), 'rows': sum(x['rows'] for x in index)}, 'files': index, 'excluded': excluded}
    if check:
        cur_doc = J('data/pack/index.json') if (OUT / 'index.json').exists() else {}
        same = [(x['product'], x['region']) for x in cur_doc.get('files', [])] == [(x['product'], x['region']) for x in index]
        print('paquete al dia' if same else 'paquete DESFASADO: ejecuta scripts/build-data-pack.py'); sys.exit(0 if same else 1)
    for f in OUT.rglob('*'):
        if f.is_file() and f.suffix in ('.csv', '.xlsx') and f not in written: f.unlink()  # fuente que ya no es apta o producto retirado: fuera del paquete
    (OUT / 'index.json').write_text(json.dumps(doc, ensure_ascii=False, indent=1) + '\n', encoding='utf-8')
    print('Paquete: %d ficheros, %d filas, %d excluidos (%s)' % (len(index), doc['summary']['rows'], len(excluded), ', '.join('%s/%s' % (e['region'], e['product']) for e in excluded[:6])))

main()
