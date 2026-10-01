#!/usr/bin/env python3
"""Aranceles de importación de México (SNICE/SE), Canadá (CBSA) y la UE (TARIC, espejo del DG TAXUD) para los capítulos 01-24 y 31.
Mismo esquema que data/us-tariffs.json. Cada país falla por separado y conserva el fichero anterior."""
import csv, datetime, io, json, re, subprocess, sys, urllib.request, zipfile
from pathlib import Path
ROOT = Path(__file__).resolve().parent.parent
CHAPTERS = ["%02d" % i for i in range(1, 25)] + ["31"]
UA = {"User-Agent": "Mozilla/5.0 (X11; Linux x86_64) Chrome/124 Safari/537.36 DehesaIndex/1.0"}
TODAY = datetime.date.today()
LOG = []
def log(*a):
    m = datetime.datetime.utcnow().strftime("%H:%M:%S ") + " ".join(str(x) for x in a); LOG.append(m); print(m)
def fetch(u, timeout=300):
    for i in range(3):
        try:
            with urllib.request.urlopen(urllib.request.Request(u, headers=UA), timeout=timeout) as r: return r.read()
        except Exception as e: log("  error", u[:90], e)
    raise RuntimeError("descarga fallida " + u)

def classify(g):
    s = (g or "").strip()
    if not s: return "other", None
    if re.match(r"^(free|ex\.?|0(\.0+)?\s*%)$", s, re.I): return "free", 0.0
    m = re.match(r"^(\d+(?:\.\d+)?)\s*%$", s)
    if m: return "adv", float(m.group(1))
    if "+" in s and "%" in s: return "comp", float(re.search(r"(\d+(?:\.\d+)?)\s*%", s).group(1))
    if re.search(r"[€$¢]|/t\b|/kg|/100|/hl|EUR|cent", s) and "%" not in s: return "spec", None
    return "other", None

def stats(out):
    chs = {}; hds = {}
    for l in out:
        l["r"], l["a"] = classify(l["g"])
        for key, tab in ((l["c"], chs), (l["hd"], hds)):
            x = tab.setdefault(key, {"n": 0, "free": 0, "adv": 0, "spec": 0, "comp": 0, "other": 0, "sumadv": 0.0, "maxadv": 0.0, "trq": 0, "fta": 0})
            x["n"] += 1; x[l["r"]] += 1
            if l["a"] is not None and l["r"] == "adv": x["sumadv"] += l["a"]; x["maxadv"] = max(x["maxadv"], l["a"])
            x["trq"] += bool(l["q"]); x["fta"] += bool(l["sp"])
    for t in (chs, hds):
        for x in t.values(): x["avgadv"] = round(x["sumadv"] / x["adv"], 2) if x["adv"] else None; x["sumadv"] = round(x["sumadv"], 1)
    return chs, hds

def write(name, out, names, source, release, note):
    if len(out) < 1200: raise RuntimeError("pocas lineas %d" % len(out))
    chs, hds = stats(out)
    doc = {"schemaVersion": 1, "generatedAt": datetime.datetime.utcnow().strftime("%Y-%m-%dT%H:%M:%SZ"), "release": release, "source": source, "note": note,
           "headings": {k: names.get(k, "") for k in hds}, "chapters": chs, "headingStats": hds, "lines": out}
    (ROOT / "data" / ("tariffs-%s.json" % name)).write_text(json.dumps(doc, ensure_ascii=False, separators=(",", ":")))
    log(name, "lineas", len(out), "partidas", len(hds), "capitulos", len(chs))

def code8(c): return "%s.%s.%s" % (c[:4], c[4:6], c[6:8])

def mexico():
    import openpyxl
    page = fetch("https://www.snice.gob.mx/cs/avi/snice/ligie.info22.html").decode("utf-8", "replace")
    hs = re.findall(r'href="([^"]*FRACCIONESARANCELARIAS[^"]*)"', page, re.I)
    if not hs: raise RuntimeError("no se encontro el enlace de fracciones")
    url = hs[0] if hs[0].startswith("http") else "https://www.snice.gob.mx" + hs[0]
    wb = openpyxl.load_workbook(io.BytesIO(fetch(url)), read_only=True); ws = wb["FA"] if "FA" in wb.sheetnames else wb.worksheets[0]
    out = []; names = {}; ctx = {}
    for row in ws.iter_rows(values_only=True):
        if len(row) < 7 or not row[2]: continue
        code = str(row[2]).strip(); digits = re.sub(r"\D", "", code); desc = re.sub(r"\s+", " ", str(row[3] or "")).strip(" .:")
        if not digits or digits[:2] not in CHAPTERS: continue
        imp = row[5]
        if imp is None or str(imp).strip() == "":
            ctx = {k: v for k, v in ctx.items() if k < len(digits)}; ctx[len(digits)] = desc
            if len(digits) == 4: names[digits] = desc
            continue
        if len(digits) != 8: continue
        g = ("%g %%" % imp) if isinstance(imp, (int, float)) else ("Free" if str(imp).strip().lower().startswith("ex") else str(imp).strip())
        path = [ctx[k] for k in sorted(ctx)] + [desc]
        out.append({"h": code, "c": digits[:2], "hd": digits[:4], "d": " › ".join(path[-3:]), "g": g, "sp": "", "o": "", "u": str(row[4] or "").strip(), "q": "", "ad": "", "f": []})
    write("mx", out, names, {"name": "Secretaría de Economía (SNICE) — Tarifa de la LIGIE", "url": "https://www.snice.gob.mx/cs/avi/snice/ligie.info22.html", "license": "Publicación oficial del Gobierno de México"},
          {"name": url.rsplit("/", 1)[-1], "title": "LIGIE"}, "Arancel general de importación (NMF). Las preferencias del T-MEC no figuran en este fichero.")

def canada():
    page = fetch("https://www.cbsa-asfc.gc.ca/trade-commerce/tariff-tarif/2026/menu-eng.html").decode("utf-8", "replace")
    zs = sorted(set(re.findall(r'href="(/trade-commerce/tariff-tarif/2026/01-99/01-99-2026-(\d+)-eng\.zip)"', page)), key=lambda x: int(x[1]))
    if not zs: raise RuntimeError("sin zip de la CBSA")
    z = zipfile.ZipFile(io.BytesIO(fetch("https://www.cbsa-asfc.gc.ca" + zs[-1][0]))); nm = z.namelist()[0]; z.extract(nm, "/tmp/ca")
    txt = subprocess.run(["mdb-export", "/tmp/ca/" + nm, "TPHS"], capture_output=True, text=True, check=True).stdout
    PREF = [("UST", "CUSMA-US"), ("MXT", "CUSMA-MX"), ("CEUT", "CETA (EU)"), ("CPTPT", "CPTPP"), ("UKT", "UK"), ("AUT", "Australia"), ("NZT", "N. Zealand"), ("KRT", "Korea"), ("CIAT", "Chile"), ("COLT", "Colombia"), ("PT", "Peru")]
    out = []; names = {}; ctx = {}; cur = None; seen = {}
    rows = list(csv.DictReader(io.StringIO(txt)))
    def dt(s):
        try: return datetime.datetime.strptime(s.split()[0], "%m/%d/%y").date()
        except Exception: return datetime.date(1970, 1, 1)
    for r in rows:
        code = (r["TARIFF"] or "").strip(); digits = re.sub(r"\D", "", code)
        if not digits or digits[:2] not in CHAPTERS or dt(r["EFF_DATE"]) > TODAY: continue
        desc = re.sub(r"\s+", " ", " ".join(x for x in (r["DESC1"], r["DESC2"], r["DESC3"]) if x)).strip(" :")
        if len(digits) < 8:
            ctx = {k: v for k, v in ctx.items() if k < len(digits)}; ctx[len(digits)] = desc
            if len(digits) == 4: names[digits] = desc
            continue
        if len(digits) == 8:
            ctx8 = {k: v for k, v in ctx.items() if k < 8}
            g = (r["MFN"] or "").strip()
            sp = "; ".join("%s: %s" % (lab, r[k].strip()) for k, lab in PREF if (r.get(k) or "").strip())
            path = [ctx8[k] for k in sorted(ctx8)] + [desc]
            cur = {"h": code, "c": digits[:2], "hd": digits[:4], "d": " › ".join(path[-3:]), "g": g, "sp": sp, "o": (r.get("General Tariff") or "").strip(), "u": "", "q": "", "ad": "", "f": []}
            if code in seen:
                if dt(r["EFF_DATE"]) < seen[code][0]: cur = None; continue
                out[seen[code][1]] = cur
            else: seen[code] = (dt(r["EFF_DATE"]), len(out)); out.append(cur)
        elif len(digits) == 10 and cur is not None and digits[:8] == re.sub(r"\D", "", cur["h"]):
            if not cur["u"] and (r["UOM"] or "").strip(): cur["u"] = r["UOM"].strip()
            if not cur["g"] and (r["MFN"] or "").strip(): cur["g"] = r["MFN"].strip()
    out = [l for l in out if l["g"]]
    write("ca", out, names, {"name": "Canada Border Services Agency — Customs Tariff 2026", "url": "https://www.cbsa-asfc.gc.ca/trade-commerce/tariff-tarif/2026/menu-eng.html", "license": "Government of Canada (Open Government Licence – Canada, por confirmar)"},
          {"name": zs[-1][0].rsplit("/", 1)[-1], "title": "Customs Tariff 2026"}, "Arancel NMF; las preferencias (CUSMA, CETA, CPTPP…) están en la columna de preferencias.")

def eu():
    import openpyxl
    rel = json.loads(fetch("https://api.github.com/repos/rousseauxy/taric-opendata/releases?per_page=100"))
    cand = [(r["published_at"], a["browser_download_url"], a["name"]) for r in rel if r["tag_name"].startswith("eu-") for a in r["assets"] if re.match(r"eu-taric-\d{4}-\d{2}\.zip$", a["name"])]
    if not cand: raise RuntimeError("sin release TARIC")
    cand.sort(); _, url, name = cand[-1]
    z = zipfile.ZipFile(io.BytesIO(fetch(url)))
    # nomenclatura
    nom = {}; names = {}; stack = {}
    wb = openpyxl.load_workbook(io.BytesIO(z.read("Nomenclature EN.xlsx")), read_only=True)
    for row in wb.worksheets[0].iter_rows(min_row=2, values_only=True):
        if not row[0]: continue
        g = str(row[0])[:10]; hp = row[4]; lvl = str(row[5] or "").count("-"); d = re.sub(r"\s+", " ", str(row[6] or "")).strip(" :")
        if row[2] and str(row[2]) < TODAY.strftime("%Y"): pass
        if hp in (2, 4): stack = {}; 
        stack = {k: v for k, v in stack.items() if k < lvl}; stack[lvl] = d
        if hp == 4: names[g[:4]] = d
        nom[g] = [stack[k] for k in sorted(stack)]
    def dd(s):
        try: return datetime.datetime.strptime(s, "%d-%m-%Y").date()
        except Exception: return None
    def fmt(e):
        e = (e or "").strip()
        if not e: return ""
        e = re.sub(r"(\d+)\.(\d+)\s*%", lambda m: ("%g" % float(m.group(0).replace("%", "").strip())) + " %", e)
        e = re.sub(r"(\d+\.\d+) EUR TNE", lambda m: "%g €/t" % float(m.group(1)), e); e = re.sub(r"(\d+\.\d+) EUR DTN", lambda m: "%g €/100 kg" % float(m.group(1)), e)
        e = re.sub(r"(\d+\.\d+) EUR (KGM|HLT)", lambda m: "%g €/%s" % (float(m.group(1)), "kg" if m.group(2) == "KGM" else "hl"), e)
        return e
    wb = openpyxl.load_workbook(io.BytesIO(z.read("Duties Import 01-99.xlsx")), read_only=True)
    third = {}; pref = {}; quota = set()
    PO = {"CA": "Canada", "MX": "Mexico", "GB": "UK", "JP": "Japan", "KR": "Korea", "UA": "Ukraine", "NZ": "N. Zealand", "CL": "Chile"}
    for row in wb.worksheets[0].iter_rows(min_row=2, values_only=True):
        gc = str(row[0] or "")
        if len(gc) != 10 or gc[:2] not in CHAPTERS: continue
        s, e = dd(str(row[3] or "")), dd(str(row[4] or ""))
        if (s and s > TODAY) or (e and e < TODAY): continue
        mt = str(row[11] or ""); org = str(row[10] or ""); k8 = gc[:8]
        if mt in ("122", "123", "143"): quota.add(k8)
        if mt == "103" and org == "1011":
            third.setdefault(k8, []).append((gc[8:] != "00", fmt(str(row[9] or ""))))
        elif mt == "142" and org in PO and not str(row[2] or ""):
            pref.setdefault(k8, {})[org] = fmt(str(row[9] or ""))
    out = []
    for k8 in sorted(third):
        rows = sorted(third[k8])  # primero la linea de 8 digitos
        g = rows[0][1]
        if g.startswith("Cond"):
            m = re.findall(r"[\d.]+ (?:%|€/t|€/100 kg)", g); amts = sorted(set(m), key=lambda x: -float(x.split()[0]))
            g = (amts[0] + " (cond.)") if amts else "Conditional"
        sp = "; ".join("%s: %s" % (PO[o], v) for o, v in sorted(pref.get(k8, {}).items(), key=lambda x: PO[x[0]]) if v and not v.startswith("Cond"))
        path = nom.get(k8 + "00", [])
        out.append({"h": code8(k8), "c": k8[:2], "hd": k8[:4], "d": " › ".join(path[-3:]) if path else "", "g": g, "sp": sp, "o": "", "u": "", "q": "TRQ" if k8 in quota else "", "ad": "", "f": []})
    write("eu", out, names, {"name": "European Commission — TARIC (DG TAXUD), copia de rousseauxy/taric-opendata", "url": "https://taxation-customs.ec.europa.eu/customs/customs-tariff/eu-customs-tariff-taric_en", "license": "Política de reutilización de la Comisión Europea (atribución)"},
          {"name": name, "title": "TARIC"}, "Arancel de tercer país (NMF, erga omnes). 'cond.' = el derecho depende de condiciones (p. ej. precio de entrada); 'TRQ' = existe contingente arancelario. Preferencias solo de Canadá, México, Reino Unido, Japón, Corea, Ucrania, Nueva Zelanda y Chile.")

ok = 0
for fn in (mexico, canada, eu):
    try: fn(); ok += 1
    except Exception as e: log("ERROR", fn.__name__, repr(e)[:300])
(ROOT / "data" / "partner-tariffs-log.txt").write_text("\n".join(LOG) + "\n")
if not ok: sys.exit(1)
