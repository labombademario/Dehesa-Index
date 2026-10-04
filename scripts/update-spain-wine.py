#!/usr/bin/env python3
"""España: vino por campaña y comunidad autonoma (MAPA, INFOVI «informe ampliado» de julio) -> data/spain-wine/infovi.json.

Cada año el MAPA publica, con los datos de julio, el informe ampliado de la campaña (1 ago - 31 jul) con las declaraciones obligatorias de los operadores (INFOVI): entrada de uva (cuadro 1), produccion de vino por categoria (2.2),
salidas de vino (3.0), exportaciones (3.1), salidas interiores (3.3) y existencias a 31 de julio de vino y mosto (4.1), todo por comunidad. Son datos declarados por las bodegas, NO un balance del sector: no hay consumo ni
comercio exterior oficial (las exportaciones son las declaradas por los operadores y el MAPA las da solo como orientativas). Unidades: uva en toneladas (la hoja viene en kg), vino en hectolitros.
- Las hojas se localizan por el titulo «CUADRO n.m» (los nombres de hoja cambian cada año). Cada hoja trae las 17 comunidades y una fila TOTAL; debajo hay otros bloques que no se leen.
- Se leen las paginas de INFOVI de cada año (2021 en adelante), se busca el informe ampliado de julio de cada una y se acumulan las campañas; un informe nuevo sustituye al anterior de su campaña y las revisiones quedan en el log.
- Se rechaza la escritura entera si algo no cuadra: TOTAL distinto de la suma de comunidades, tinto+blanco distinto del total, categorias que no suman el vino, salidas distintas de exportaciones + interiores, existencias que no suman, rendimiento uva-vino fuera de rango.
--file usa un libro local (pruebas, campaña indicada con --campaign)."""
import argparse, datetime, io, json, os, re, sys, time, unicodedata, urllib.request
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "spain-wine"
SITE = "https://www.mapa.gob.es"
V = "/es/agricultura/temas/producciones-agricolas/vitivinicultura/"
UA = {"User-Agent": "Mozilla/5.0 (compatible; DehesaIndex/1.0; +https://dehesaindex.com)"}
FIRST = 2021
CCAA = {"andalucia": "Andalucía", "aragon": "Aragón", "asturias": "Asturias", "baleares": "Illes Balears", "canarias": "Canarias", "cantabria": "Cantabria", "clm": "Castilla-La Mancha", "cyl": "Castilla y León", "cataluna": "Cataluña",
        "extremadura": "Extremadura", "galicia": "Galicia", "madrid": "Comunidad de Madrid", "murcia": "Región de Murcia", "navarra": "Navarra", "paisvasco": "País Vasco", "larioja": "La Rioja", "valenciana": "Comunitat Valenciana"}
NAMES = {"andalucia": "andalucia", "aragon": "aragon", "asturias": "asturias", "baleares": "baleares", "canarias": "canarias", "cantabria": "cantabria", "castilla la mancha": "clm", "castilla y leon": "cyl", "cataluna": "cataluna", "extremadura": "extremadura",
         "galicia": "galicia", "madrid": "madrid", "c madrid": "madrid", "murcia": "murcia", "navarra": "navarra", "pais vasco": "paisvasco", "la rioja": "larioja", "c valenciana": "valenciana", "valenciana": "valenciana"}
FIELDS = ["grape", "wine", "red", "white", "dop", "igp", "varietal", "sinig", "exits", "exports", "domestic", "stockWine", "stockMust"]
LOG = []
def log(*a):
    s = " ".join(str(x) for x in a); LOG.append(s); print(s, flush=True)
def http(url, method="GET", tries=3, timeout=300):
    last = None
    for i in range(tries):
        try:
            with urllib.request.urlopen(urllib.request.Request(url, headers=UA, method=method), timeout=timeout) as r: return r.read() if method == "GET" else dict(r.headers)
        except Exception as e: last = e; time.sleep(4 * (i + 1))
    raise RuntimeError("%s %s -> %s" % (method, url, str(last)[:160]))
def norm(s):
    s = unicodedata.normalize("NFD", str(s).lower()); s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9 ]", " ", s)).strip()
def num(x): return float(x) if isinstance(x, (int, float)) and not isinstance(x, bool) else None

def find_links(html, year):
    """Informe ampliado de julio del año (xlsx) entre los enlaces de la pagina anual."""
    c = []
    for h in dict.fromkeys(re.findall(r'href="([^"]+\.xlsx)"', html)):
        n = norm(h.rsplit("/", 1)[-1])
        if ("ano-%d/" % year) in h and "ampliad" in n and "julio" in n and "existencias" not in n: c.append(h)
    return [h if h.startswith("http") else SITE + h for h in c]
def year_pages(year):
    return [SITE + V + "datos_infovi_anteriores/infovi_%d" % year, SITE + V + "infovi_%d" % year, SITE + V + "datos_infovi"]

def sheets(wb):
    """{clave de cuadro: hoja}: «1», «2.2», «3.0», «3.0d» (desglose por categoria), «3.1», «3.3», «4.1»."""
    out = {}
    for ws in wb.worksheets:
        t = None
        for r in ws.iter_rows(min_row=1, max_row=4):
            for c in r:
                if isinstance(c.value, str) and c.value.strip().upper().startswith("CUADRO"): t = re.sub(r"\s+", " ", c.value.strip()); break
            if t: break
        if not t: continue
        m = re.match(r"(?i)cuadro\s+(\d+)(?:\s*[.,]\s*(\d+))?", t)
        if not m: continue
        k = m.group(1) + ("." + m.group(2) if m.group(2) is not None else "")
        if "CATEGOR" in t.upper() and k == "3.0": k = "3.0d"
        out.setdefault(k, ws)
    return out
def block(ws):
    """Filas de las 17 comunidades + TOTAL del primer bloque: {slug|TOTAL: [valores de las columnas siguientes a la etiqueta]}."""
    rows = {}; lc = None
    for r in ws.iter_rows(min_row=1, max_row=40):
        v = [c.value for c in r]
        for i, x in enumerate(v[:3]):
            if isinstance(x, str) and norm(x) == "andalucia": lc = i; break
        if lc is not None:
            if rows and not isinstance(v[lc], str): break
            lab = norm(v[lc]) if isinstance(v[lc], str) else ""
            if lab in NAMES: rows[NAMES[lab]] = [num(x) for x in v[lc + 1: lc + 17]]
            elif lab == "total":
                rows["TOTAL"] = [num(x) for x in v[lc + 1: lc + 17]]; break
    return rows
def cells(rows, col):
    return {k: (v[col] if len(v) > col else None) for k, v in rows.items()}

def parse(wb):
    S = sheets(wb); need = ["1", "2.2", "3.0", "3.1", "3.3", "4.1"]
    miss = [k for k in need if k not in S]
    if miss: raise ValueError("faltan cuadros del informe: " + ", ".join(miss))
    B = {k: block(S[k]) for k in need}
    for k in need:
        if "TOTAL" not in B[k] or len([x for x in B[k] if x != "TOTAL"]) != 17: raise ValueError("cuadro %s: no se leen las 17 comunidades y el total" % k)
    d = {}
    g = B["1"]; tin, bla, tot = cells(g, 0), cells(g, 1), cells(g, 2)
    p = B["2.2"]
    d["grape"] = {k: None if v is None else v / 1000.0 for k, v in tot.items()}
    d["_grapeRed"] = {k: v / 1000.0 if v is not None else None for k, v in tin.items()}; d["_grapeWhite"] = {k: v / 1000.0 if v is not None else None for k, v in bla.items()}
    d["red"], d["white"], d["wine"] = cells(p, 12), cells(p, 13), cells(p, 14)
    d["dop"], d["igp"], d["varietal"], d["sinig"] = cells(p, 2), cells(p, 5), cells(p, 8), cells(p, 11)
    d["exits"] = cells(B["3.0"], 4); d["exports"] = cells(B["3.1"], 4); d["domestic"] = cells(B["3.3"], 4)
    e = B["4.1"]; d["stockWine"], d["stockMust"], d["_stockTotal"] = cells(e, 6), cells(e, 7), cells(e, 8)
    d["_stockParts"] = {k: sum(x for x in v[:4] if x is not None) if all(x is not None for x in v[:4]) else None for k, v in e.items()}
    ext = None
    for ws in wb.worksheets:
        for r in ws.iter_rows(min_row=1, max_row=40):
            for c in r:
                if isinstance(c.value, str):
                    m = re.search(r"(?i)extracci[oó]n de (\d{1,2}) de (\w+) de (\d{4})", c.value)
                    if m:
                        mo = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"]
                        try: ext = "%s-%02d-%02d" % (m.group(3), mo.index(norm(m.group(2))) + 1, int(m.group(1)))
                        except ValueError: pass
    return d, ext

def validate(c, d):
    errs = []; warns = []
    keys = [k for k in d["wine"] if k != "TOTAL"]
    def close(a, b, tol=2.0): return a is None or b is None or abs(a - b) <= tol + 0.0001 * abs(b)
    for f in FIELDS:
        for k, v in d[f].items():
            if v is not None and v < 0: errs.append("%s %s %s: valor negativo" % (c, f, k))
    for f in FIELDS:
        for k in ("TOTAL",):
            if d[f].get(k) is None: errs.append("%s %s: falta el total" % (c, f))
        s = sum(d[f][k] for k in keys if d[f].get(k) is not None)
        if d[f].get("TOTAL") is not None and not close(s, d[f]["TOTAL"], 3.0 if f != "grape" else 3.0): errs.append("%s %s: las comunidades suman %s y el TOTAL del MAPA es %s" % (c, f, round(s, 1), round(d[f]["TOTAL"], 1)))
    for k in keys + ["TOTAL"]:
        if not close((d["red"][k] or 0) + (d["white"][k] or 0), d["wine"][k]): errs.append("%s %s: tinto+blanco no es el total de vino" % (c, k))
        cats = [d[x][k] for x in ("dop", "igp", "varietal", "sinig")]
        if all(x is not None for x in cats) and d["wine"][k] is not None and not close(sum(cats), d["wine"][k]): errs.append("%s %s: las categorias no suman el vino producido" % (c, k))
        if d["exits"][k] is not None and d["exports"][k] is not None and d["domestic"][k] is not None and not close(d["exports"][k] + d["domestic"][k], d["exits"][k]): errs.append("%s %s: salidas distintas de exportaciones + interiores" % (c, k))
        if d["stockWine"][k] is not None and d["stockMust"][k] is not None and not close(d["stockWine"][k] + d["stockMust"][k], d["_stockTotal"][k]): errs.append("%s %s: existencias de vino + mosto distintas del total" % (c, k))
        if d["_stockParts"][k] is not None and d["stockWine"][k] is not None and not close(d["_stockParts"][k], d["stockWine"][k]): errs.append("%s %s: las existencias de vino no suman sus cuatro columnas" % (c, k))
        if not close((d["_grapeRed"][k] or 0) + (d["_grapeWhite"][k] or 0), d["grape"][k], 0.01): errs.append("%s %s: uva tinta+blanca distinta del total" % (c, k))
    g, w = d["grape"]["TOTAL"], d["wine"]["TOTAL"]
    if g and w and not 5.0 <= w / g <= 8.0: errs.append("%s: rendimiento de %.2f hl por tonelada de uva fuera de rango (5-8)" % (c, w / g))
    return errs, warns

def now(): return datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
def write_json(path, obj):
    tmp = str(path) + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f: json.dump(obj, f, ensure_ascii=False, separators=(",", ":"))
    os.replace(tmp, path)
def r3(x): return None if x is None else round(x, 3)

def pack(parsed):
    """{campaña: d} -> estructura del JSON (listas alineadas con campañas ordenadas)."""
    cs = sorted(parsed)
    nat = {f: [r3(parsed[c][f].get("TOTAL")) for c in cs] for f in FIELDS}
    cc = {k: {f: [r3(parsed[c][f].get(k)) for c in cs] for f in FIELDS} for k in CCAA}
    return cs, nat, cc
def merge(old, new, extracted, links):
    """Campañas nuevas sobre las del archivo; el informe nuevo gana y se anotan las revisiones."""
    revs = []; allc = {}
    if old:
        for i, c in enumerate(old["campaigns"]):
            allc[c] = {"national": {f: old["national"][f][i] for f in FIELDS}, "ccaa": {k: {f: old["ccaa"][k][f][i] for f in FIELDS} for k in CCAA}}
    for c, d in new.items():
        n = {f: r3(d[f]["TOTAL"]) for f in FIELDS}; cc = {k: {f: r3(d[f].get(k)) for f in FIELDS} for k in CCAA}
        if c in allc:
            for f in FIELDS:
                o = allc[c]["national"][f]
                if o is not None and n[f] is not None and abs(o - n[f]) > 1e-6: revs.append((c, f, o, n[f]))
        allc[c] = {"national": n, "ccaa": cc}
    cs = sorted(allc)
    return cs, {f: [allc[c]["national"][f] for c in cs] for f in FIELDS}, {k: {f: [allc[c]["ccaa"][k][f] for c in cs] for f in FIELDS} for k in CCAA}, revs

def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--file"); ap.add_argument("--campaign"); ap.add_argument("--force", action="store_true"); ap.add_argument("--out"); a = ap.parse_args()
    out = Path(a.out) if a.out else OUT; out.mkdir(parents=True, exist_ok=True); path = out / "infovi.json"
    old = json.loads(path.read_text("utf-8")) if path.exists() else None
    import openpyxl
    parsed = {}; ext = {}; src = dict((old or {}).get("sources", {})); changed = False
    try:
        if a.file:
            if not a.campaign: raise ValueError("--file necesita --campaign")
            wb = openpyxl.load_workbook(a.file, data_only=True); d, e = parse(wb); parsed[a.campaign] = d; ext[a.campaign] = e; src[a.campaign] = {"url": None, "lastModified": None, "extracted": e}; changed = True
        else:
            this = datetime.date.today().year
            for y in range(FIRST, this + 1):
                links = []
                for pu in year_pages(y):
                    try: links = find_links(http(pu).decode("utf-8", "replace"), y)
                    except Exception as ex: log("pagina %s: %s" % (pu.rsplit("/", 1)[-1], str(ex)[:80])); continue
                    if links: break
                c = str(y - 1)
                if not links: log("campaña %s/%s: sin informe ampliado de julio publicado" % (c, str(y)[2:])); continue
                url = links[0]; lm = None
                try: lm = http(url, "HEAD").get("Last-Modified")
                except Exception as ex: log("HEAD fallo:", str(ex)[:80])
                if not a.force and old and c in old["campaigns"] and src.get(c, {}).get("url") == url and lm and src[c].get("lastModified") == lm: continue
                wb = openpyxl.load_workbook(io.BytesIO(http(url)), data_only=True); d, e = parse(wb); parsed[c] = d; ext[c] = e; src[c] = {"url": url, "lastModified": lm, "extracted": e}; changed = True
    except Exception as ex: log("ERROR al leer el informe:", ex); return 1
    if not changed and old:
        old["checkedAt"] = now(); write_json(path, old); log("informes sin cambios"); return 0
    errs = []
    for c, d in parsed.items():
        e, w = validate(c, d); errs += e
        for x in w: log("AVISO:", x)
    if errs:
        for e in errs[:20]: log("ERROR de validacion:", e)
        log("no se escribe nada (%d errores)" % len(errs)); return 1
    cs, N, CC, revs = merge(old, parsed, ext, src)
    for c, f, o, n in revs[:30]: log("revision del MAPA: campaña %s %s %s -> %s" % (c, f, o, n))
    if len(cs) >= 2 and any(int(cs[i + 1]) - int(cs[i]) != 1 for i in range(len(cs) - 1)): log("AVISO: campañas no consecutivas:", cs)
    t = now()
    doc = {"schemaVersion": 1, "generatedAt": t, "checkedAt": t, "source": {"id": "mapa_es", "url": SITE + V + "datos_infovi_anteriores"}, "campaigns": cs,
           "units": {"grape": "toneladas de uva", "wine": "hectolitros", "stocks": "hectolitros a 31 de julio"}, "fields": FIELDS, "national": N, "ccaa": CC, "ccaaNames": CCAA,
           "sources": {c: src.get(c) for c in cs if src.get(c)}, "revisions": len(revs)}
    write_json(path, doc); log("OK: campañas %s..%s (%d), %d revisiones" % (cs[0], cs[-1], len(cs), len(revs))); return 0

if __name__ == "__main__":
    rc = main()
    try: (ROOT / "data" / "spain-wine-log.txt").write_text("\n".join(LOG) + "\n", "utf-8")
    except Exception: pass
    sys.exit(rc)
