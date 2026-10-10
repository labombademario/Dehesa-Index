import json, urllib.request, urllib.parse, os, re, io
os.makedirs('probe',exist_ok=True); S=open('probe/summary_ch2.txt','w')
def log(*a): s=' '.join(str(x) for x in a); print(s); S.write(s+'\n'); S.flush()
UA={'User-Agent':'DehesaIndex-research/1.0 (+https://dehesaindex.com)'}
def get(u,t=150,data=None,hdr=None):
    try:
        h=dict(UA); h.update(hdr or {})
        with urllib.request.urlopen(urllib.request.Request(u,data=data,headers=h),timeout=t) as r: return r.status,r.read(),r.geturl()
    except urllib.error.HTTPError as e: return e.code,e.read()[:2000],u
    except Exception as e: return 0,repr(e).encode()[:300],u
import subprocess, sys; subprocess.run([sys.executable,'-m','pip','install','--break-system-packages','--user','-q','openpyxl','xlrd'],check=False); import site, importlib; sys.path.append(site.getusersitepackages()); importlib.invalidate_caches()
import openpyxl
try: import xlrd
except Exception: xlrd=None
def show_xls(name,b):
    open('probe/'+name,'wb').write(b)
    try:
        if b[:2]==b'PK':
            wb=openpyxl.load_workbook(io.BytesIO(b),read_only=True,data_only=True)
            for ws in wb.worksheets[:6]:
                log('    SHEET',ws.title,ws.max_row,ws.max_column)
                for i,row in enumerate(ws.iter_rows(values_only=True)):
                    if i>=14: break
                    log('      ',[c for c in row[:14]])
        elif xlrd:
            wb=xlrd.open_workbook(file_contents=b)
            for sh in wb.sheets()[:6]:
                log('    SHEET',sh.name,sh.nrows,sh.ncols)
                for i in range(min(14,sh.nrows)): log('      ',sh.row_values(i)[:14])
    except Exception as e: log('    xls parse',repr(e)[:200])
for name in ('erlospreisstatistik-mengen-und-preise-zur-produktion-der-landwirtschaft2','produzenten-und-importpreisindex1','gesamtproduktion-der-landwirtschaft-zu-laufenden-preisen4','volumen-der-landwirtschaftlichen-gesamtproduktion-indizes-19901004','landwirtschaftliche-gesamtrechnung-produktionswert-der-landwirtschaft-historische-zeitreihen4'):
    st,b,_=get('https://opendata.swiss/api/3/action/package_show?id='+name)
    if st!=200: log('PKG',name,st); continue
    p=json.loads(b)['result']; d=p.get('description') or {}
    log('PKG',name,'| issued',p.get('issued'),'| modified',p.get('modified'),'| accrual',p.get('accrual_periodicity'),'|',(d.get('en') or d.get('de') or '')[:300])
    for r in p['resources']:
        u=r.get('download_url') or r.get('url'); log('   RES',r.get('format'),r.get('rights'),u)
        if r.get('format') in ('XLS','XLSX','CSV') and u:
            st2,b2,fu=get(u,t=200); log('    get',st2,len(b2),fu)
            if st2==200: show_xls(name[:40]+'.'+('xlsx' if b2[:2]==b'PK' else 'xls'),b2)
# BFS PxWeb: buscar tablas de precios agricolas
for q in ('Produzentenpreis Landwirtschaft','landwirtschaftliche Produzentenpreise','Erlöspreis','Preisindex landwirtschaftlicher'):
    st,b,_=get('https://www.pxweb.bfs.admin.ch/api/v1/de/?query='+urllib.parse.quote(q))
    log('PXWEB',q,st,b[:600])
# Swiss PPI CPA 01 en PxWeb
st,b,_=get('https://www.pxweb.bfs.admin.ch/api/v1/de/px-x-0502010000_101/px-x-0502010000_101.px'); log('PXWEB ppi meta',st,b[:1500])
# opendata.swiss: comercio exterior
for q in ('Aussenhandel','Import Export Waren','Swiss-Impex','foreign trade'):
    st,b,_=get('https://opendata.swiss/api/3/action/package_search?rows=30&q='+urllib.parse.quote(q))
    if st==200:
        for p in json.loads(b)['result']['results']:
            ti=p.get('title') or {}; log('TRADE',q,'|',(p.get('organization') or {}).get('name'),'|',p.get('name'),'|',(ti.get('en') or ti.get('de') or '')[:100],'|',sorted(set((r.get('rights') or '') for r in p.get('resources',[]))))
# agroscope cube
def sparql(q):
    st,b,_=get('https://lindas.admin.ch/query',t=170,data=urllib.parse.urlencode({'query':q}).encode(),hdr={'Accept':'application/sparql-results+json'})
    return [{k:v['value'] for k,v in r.items()} for r in json.loads(b)['results']['bindings']] if st==200 else [{'err':str(st)+b[:200].decode('utf-8','replace')}]
for r in sparql('PREFIX schema: <http://schema.org/> SELECT ?p ?o WHERE { GRAPH <https://lindas.admin.ch/agroscope/cube> { <https://agriculture.ld.admin.ch/agroscope/BLW_t_01/1> ?p ?o } } LIMIT 60'): log('AGROSCOPE',r)
NS='https://agriculture.ld.admin.ch/foag/'
for pid in ('277','279','337','281','282','271'):
    for r in sparql('PREFIX schema: <http://schema.org/> SELECT ?n WHERE { GRAPH <https://lindas.admin.ch/foag/agricultural-market-data> { <%sproduct/%s> schema:name ?n } }'%(NS,pid)): log('FOAG label',pid,r)
S.close()
